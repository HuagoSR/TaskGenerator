from pathlib import Path

import pytest
from openpyxl import Workbook

from task_generator.evaluation.agent_rubric_grader import (
    AgentRubricGradeDraftV1,
    AgentRubricGradeDraftV2,
    EvidenceRefV1,
    canonicalize_evidence_ref,
    finalize_agent_rubric_grade,
    finalize_agent_rubric_grade_v2,
)
from task_generator.evaluation.independent_rubric_grader import FrozenRubricItemV1
from task_generator.evaluation.spreadsheet_evidence import extract_spreadsheet_evidence


HASH = "a" * 64


def _stage(tmp_path: Path):
    root = tmp_path / "stage"
    book = root / "anonymous_submission/result.xlsx"
    book.parent.mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Header"
    ws["B2"] = "Evidence"
    wb.save(book)
    reference = root / "reference_files/source.xlsx"
    reference.parent.mkdir()
    wb.save(reference)
    wb.close()
    evidence = {
        "anonymous_submission/result.xlsx": extract_spreadsheet_evidence(book, relative_path="anonymous_submission/result.xlsx"),
        "reference_files/source.xlsx": extract_spreadsheet_evidence(reference, relative_path="reference_files/source.xlsx"),
    }
    return root, evidence


def test_structured_ref_generates_stable_locator_and_citation(tmp_path):
    root, evidence = _stage(tmp_path)
    ref = EvidenceRefV1(relative_path="anonymous_submission/result.xlsx", sheet_name="Summary", cell_range="B2", evidence_role="support")
    left = canonicalize_evidence_ref(ref, staging_root=root, evidence=evidence, verification_scope="formula")
    right = canonicalize_evidence_ref(ref, staging_root=root, evidence=evidence, verification_scope="formula")
    assert left.citation_id == right.citation_id
    assert left.locator.endswith("#sheet=Summary&range=B2")
    with pytest.raises(ValueError, match="range_empty"):
        canonicalize_evidence_ref(
            ref.model_copy(update={"cell_range": "A2"}), staging_root=root, evidence=evidence, verification_scope="formula"
        )


def test_agent_grade_requires_both_roots_for_quantified_full_credit(tmp_path):
    root, evidence = _stage(tmp_path)
    rubric = [FrozenRubricItemV1(rubric_item_id="one", max_score=2, criterion="Each row is correct.")]
    draft = AgentRubricGradeDraftV1.model_validate({
        "material_status": "complete",
        "material_notes": "Readable workbook.",
        "assessments": [{
            "rubric_item_id": "one", "awarded": 2, "satisfaction": "met", "verification_scope": "all_rows",
            "evidence_refs": [{"relative_path": "anonymous_submission/result.xlsx", "sheet_name": "Summary", "cell_range": "A1:B2", "evidence_role": "support"}],
            "support": ["All submitted rows checked."], "defects": [], "rationale": "All submitted rows agree with evidence.",
        }],
    })
    with pytest.raises(ValueError, match="candidate_and_reference"):
        finalize_agent_rubric_grade(
            task_id="t", submission_id="s", judge_id="j", rubric_items=rubric, draft=draft,
            staging_root=root, spreadsheet_evidence=evidence, input_sha256=HASH, delivery_sha256=HASH,
        )


def test_agent_partial_one_point_is_zero_and_program_sums(tmp_path):
    root, evidence = _stage(tmp_path)
    rubric = [FrozenRubricItemV1(rubric_item_id="one", max_score=1, criterion="One field is correct.")]
    draft = AgentRubricGradeDraftV1.model_validate({
        "material_status": "complete", "material_notes": "Readable workbook.",
        "assessments": [{
            "rubric_item_id": "one", "awarded": 0, "satisfaction": "partial", "verification_scope": "sample",
            "evidence_refs": [
                {"relative_path": "anonymous_submission/result.xlsx", "sheet_name": "Summary", "cell_range": "A1", "evidence_role": "support"},
                {"relative_path": "anonymous_submission/result.xlsx", "sheet_name": "Summary", "cell_range": "B2", "evidence_role": "defect"},
            ],
            "support": ["Header is present."], "defects": ["Required detail is incorrect."], "rationale": "Only part of the requirement is supported.",
        }],
    })
    grade, citations = finalize_agent_rubric_grade(
        task_id="t", submission_id="s", judge_id="j", rubric_items=rubric, draft=draft,
        staging_root=root, spreadsheet_evidence=evidence, input_sha256=HASH, delivery_sha256=HASH,
    )
    assert grade.total_score == 0
    assert len(citations) == 2


def test_v2_all_rows_can_use_formula_and_recalculated_methods(tmp_path):
    root, evidence = _stage(tmp_path)
    rubric = [FrozenRubricItemV1(rubric_item_id="one", max_score=2, criterion="Each row is correct.")]
    draft = AgentRubricGradeDraftV2.model_validate({
        "material_status": "complete", "material_notes": "All rows were checked.",
        "assessments": [{
            "rubric_item_id": "one", "awarded": 2, "satisfaction": "met", "coverage": "all_rows",
            "verification_methods": ["formula", "recalculated_value"],
            "evidence_refs": [
                {"relative_path": "anonymous_submission/result.xlsx", "sheet_name": "Summary", "cell_range": "A1:B2", "evidence_role": "support"},
                {"relative_path": "reference_files/source.xlsx", "sheet_name": "Summary", "cell_range": "A1:B2", "evidence_role": "support"},
            ],
            "support": ["Every row was formula-checked and recalculated."], "defects": [],
            "rationale": "Complete coverage and formula methods independently recorded.",
        }],
    })
    grade, _ = finalize_agent_rubric_grade_v2(
        task_id="t", submission_id="s", judge_id="j", rubric_items=rubric, draft=draft,
        staging_root=root, spreadsheet_evidence=evidence, input_sha256=HASH, delivery_sha256=HASH,
    )
    assert grade.total_score == 2


def test_v2_formula_method_does_not_imply_all_rows(tmp_path):
    root, evidence = _stage(tmp_path)
    rubric = [FrozenRubricItemV1(rubric_item_id="one", max_score=2, criterion="Each row is correct.")]
    draft = AgentRubricGradeDraftV2.model_validate({
        "material_status": "complete", "material_notes": "Only a sample was checked.",
        "assessments": [{
            "rubric_item_id": "one", "awarded": 2, "satisfaction": "met", "coverage": "sample",
            "verification_methods": ["formula"],
            "evidence_refs": [{"relative_path": "anonymous_submission/result.xlsx", "sheet_name": "Summary", "cell_range": "B2", "evidence_role": "support"}],
            "support": ["One formula was checked."], "defects": [], "rationale": "Only sampled formula evidence exists.",
        }],
    })
    with pytest.raises(ValueError, match="quantified_full_requires_all_rows"):
        finalize_agent_rubric_grade_v2(
            task_id="t", submission_id="s", judge_id="j", rubric_items=rubric, draft=draft,
            staging_root=root, spreadsheet_evidence=evidence, input_sha256=HASH, delivery_sha256=HASH,
        )
