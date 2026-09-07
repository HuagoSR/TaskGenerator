from pathlib import Path

import pytest
from openpyxl import Workbook

from task_generator.evaluation.independent_rubric_grader import (
    FrozenRubricItemV1,
    IndependentRubricGradeDraftV2,
    finalize_independent_grade_v2,
)
from task_generator.evaluation.spreadsheet_evidence import (
    extract_spreadsheet_evidence,
    validate_spreadsheet_locator,
)

HASH = "a" * 64


def _book(path: Path, *, cycle: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Header"
    ws["B2"] = "=B2+1" if cycle else "=1+1"
    wb.save(path)


def test_spreadsheet_evidence_preserves_coordinates_and_detects_cycle(tmp_path):
    path = tmp_path / "candidate.xlsx"
    _book(path, cycle=True)
    evidence = extract_spreadsheet_evidence(path, relative_path="anonymous_submission/candidate.xlsx")
    assert evidence.sheets[0].cells[1].address == "B2"
    assert any(row.kind == "formula_cycle" for row in evidence.diagnostics)
    validate_spreadsheet_locator(
        "anonymous_submission/candidate.xlsx#sheet=Summary&range=B2",
        {evidence.relative_path: evidence},
    )
    with pytest.raises(ValueError, match="outside_used_range"):
        validate_spreadsheet_locator(
            "anonymous_submission/candidate.xlsx#sheet=Summary&range=Z99",
            {evidence.relative_path: evidence},
        )


def _draft(*, satisfaction: str, awarded: int, scope: str = "formula", defects=None):
    return IndependentRubricGradeDraftV2.model_validate({
        "material_status": "complete", "material_notes": "Workbook is readable.",
        "assessments": [{
            "rubric_item_id": "one", "awarded": awarded, "satisfaction": satisfaction,
            "verification_scope": scope,
            "evidence_paths": ["anonymous_submission/candidate.xlsx#sheet=Summary&range=B2"],
            "support": ["Formula at Summary!B2."],
            "defects": defects or [],
            "rationale": "Evidence-based conclusion for this item.",
        }],
    })


def test_v2_rejects_partial_one_point_full_and_quantified_sample(tmp_path):
    path = tmp_path / "candidate.xlsx"
    _book(path, cycle=False)
    evidence = extract_spreadsheet_evidence(path, relative_path="anonymous_submission/candidate.xlsx")
    staging = tmp_path / "stage"
    (staging / "anonymous_submission").mkdir(parents=True)
    (staging / "anonymous_submission/candidate.xlsx").write_bytes(path.read_bytes())
    rubric = [FrozenRubricItemV1(rubric_item_id="one", max_score=1, criterion="Each row is correct.")]
    locator = lambda value: validate_spreadsheet_locator(value, {evidence.relative_path: evidence})
    with pytest.raises(ValueError, match="one_point_partial_must_zero"):
        finalize_independent_grade_v2(
            task_id="t", submission_id="s", judge_id="j", rubric_items=rubric,
            draft=_draft(satisfaction="partial", awarded=1, defects=["Missing support."]),
            staging_root=staging, input_sha256=HASH, delivery_sha256=HASH, validate_locator=locator,
        )
    with pytest.raises(ValueError, match="quantified_full_requires_all_rows"):
        finalize_independent_grade_v2(
            task_id="t", submission_id="s", judge_id="j", rubric_items=rubric,
            draft=_draft(satisfaction="met", awarded=1, scope="sample"),
            staging_root=staging, input_sha256=HASH, delivery_sha256=HASH, validate_locator=locator,
        )
