from __future__ import annotations

import hashlib
import json
import os
import shutil
import signal
import subprocess
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Literal, Optional

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from task_generator.v3_matched_screening import (
    MatchedScreeningAnalyzer,
    RUBRIC_WEIGHTS,
    ScreeningCriterionGradeV1,
    ScreeningGraderReviewV1,
    ScreeningTaskObservationV1,
)
from task_generator.v3_semantic_review_executor import (
    SemanticReviewExecutionError,
    SemanticReviewExecutor,
)
from task_generator.v3_skill_extractor import ProviderConfig
from task_generator.v3_source_fingerprint import governed_source_fingerprint
from task_generator.v3_validity_utility import UtilityProfileV1, ValidityVectorV1


MODEL = "gpt-5.6-sol"
CLI_PACKAGE = "@openai/codex"
PUBLIC_PROBE_TIMEOUT_SECONDS = 300
TASK_TIMEOUT_SECONDS = 1800
ALLOWED_PROJECTION_FILES = {
    "dataset_row.json",
    "deliverable_contract.json",
}
SECRET_ENV_NAMES = {
    "AGENT_API_KEY",
    "ANTHROPIC_API_KEY",
    "DEEPSEEK_API_KEY",
    "E2B_API_KEY",
    "OPENAI_API_KEY",
    "TUZI_API_KEY",
}
INFRASTRUCTURE_MARKERS = (
    "authentication",
    "authorization",
    "credential",
    "login required",
    "not logged in",
    "http 401",
    "http 403",
    "service unavailable",
    "connection refused",
    "connection reset",
    "stream disconnected",
    "protocol error",
    "tls",
    "dns",
    "rate limit",
    "runtime missing",
    "executable not found",
)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha_file(path: Path) -> str:
    return _sha_bytes(path.read_bytes())


def _sha_json(value: object) -> str:
    return _sha_bytes(
        json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
    )


def _tree(root: Path) -> str:
    rows = [
        f"{item.relative_to(root).as_posix()}:{_sha_file(item)}"
        for item in sorted(path for path in root.rglob("*") if path.is_file())
    ]
    return _sha_bytes("\n".join(rows).encode("utf-8"))


def _atomic_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _projected_files(package: Path) -> List[Path]:
    result = [
        package / "dataset_row.json",
        package / "deliverable_contract.json",
    ]
    reference_root = package / "reference_files"
    result.extend(
        sorted(path for path in reference_root.rglob("*") if path.is_file())
    )
    if not all(path.is_file() for path in result[:2]):
        raise ValueError("codex_local_candidate_contract_missing")
    if not reference_root.is_dir() or not result[2:]:
        raise ValueError("codex_local_reference_files_missing")
    if any(path.suffix.lower() != ".xlsx" for path in result[2:]):
        raise ValueError("codex_local_reference_type_not_allowed")
    return result


def _projection_tree(package: Path) -> str:
    rows = [
        f"{path.relative_to(package).as_posix()}:{_sha_file(path)}"
        for path in _projected_files(package)
    ]
    return _sha_bytes("\n".join(rows).encode("utf-8"))


def _copy_projection(package: Path, target: Path) -> None:
    if target.exists():
        raise FileExistsError("codex_local_workspace_already_exists")
    target.mkdir(parents=True)
    for source in _projected_files(package):
        relative = source.relative_to(package)
        destination = target / relative
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
    (target / "deliverable_files").mkdir()


def _expected_deliverable(package: Path) -> str:
    payload = json.loads(
        (package / "deliverable_contract.json").read_text(encoding="utf-8")
    )
    paths = [
        str(item.get("relative_path", ""))
        for item in payload.get("deliverables", [])
        if isinstance(item, dict)
    ]
    if len(paths) != 1:
        raise ValueError("codex_local_exactly_one_deliverable_required")
    expected = Path(paths[0])
    if (
        expected.is_absolute()
        or ".." in expected.parts
        or expected.parts[:1] != ("deliverable_files",)
        or expected.suffix.lower() != ".xlsx"
    ):
        raise ValueError("codex_local_deliverable_path_invalid")
    return expected.as_posix()


class CodexLocalDeliveryInspectionV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    relative_path: str
    exists: bool
    nonempty: bool
    xlsx_openable: bool
    has_nonempty_sheet: bool
    differs_from_inputs: bool
    sha256: Optional[str] = None
    size_bytes: int = 0
    workbook_sheets: List[str] = Field(default_factory=list)
    valid: bool
    failure: Optional[str] = None


def inspect_delivery(
    workspace: Path, expected: str, input_hashes: set[str]
) -> CodexLocalDeliveryInspectionV1:
    path = workspace / Path(expected)
    exists = path.is_file()
    size = path.stat().st_size if exists else 0
    digest = _sha_file(path) if exists else None
    openable = False
    nonempty_sheet = False
    sheets: List[str] = []
    failure = None
    if not exists:
        failure = "expected_delivery_missing"
    elif size == 0:
        failure = "expected_delivery_empty"
    else:
        try:
            workbook = load_workbook(path, data_only=False, read_only=True)
            sheets = list(workbook.sheetnames)
            for sheet in workbook.worksheets:
                if any(
                    cell.value not in (None, "")
                    for row in sheet.iter_rows()
                    for cell in row
                ):
                    nonempty_sheet = True
                    break
            workbook.close()
            openable = True
        except Exception:
            failure = "expected_delivery_not_real_xlsx"
    differs = bool(digest and digest not in input_hashes)
    if openable and not nonempty_sheet:
        failure = "expected_delivery_has_no_content"
    if openable and nonempty_sheet and not differs:
        failure = "expected_delivery_copies_input"
    valid = bool(exists and size and openable and nonempty_sheet and differs)
    return CodexLocalDeliveryInspectionV1(
        relative_path=expected,
        exists=exists,
        nonempty=size > 0,
        xlsx_openable=openable,
        has_nonempty_sheet=nonempty_sheet,
        differs_from_inputs=differs,
        sha256=digest,
        size_bytes=size,
        workbook_sheets=sheets,
        valid=valid,
        failure=None if valid else failure,
    )


class CodexLocalCLIIdentityV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    package_name: Literal["@openai/codex"] = CLI_PACKAGE
    version: str = Field(min_length=1)
    package_lock_sha256: str = Field(min_length=64, max_length=64)
    executable_sha256: str = Field(min_length=64, max_length=64)
    executable_path: str
    installation_tree_sha256: str = Field(min_length=64, max_length=64)


class CodexLocalTaskBindingV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blind_task_id: str
    brief_id: str
    route_id: Literal["skill_guided_llm", "llm_led_hybrid"]
    motif: Literal[
        "fan_in_reconciliation", "cross_check_validation", "policy_application"
    ]
    replicate_id: Literal["a", "b"]
    domain: Optional[Literal["audit_compliance", "procurement_operations"]] = None
    package_fingerprint: str = Field(min_length=64, max_length=64)
    candidate_tree_sha256: str = Field(min_length=64, max_length=64)
    candidate_projection_sha256: str = Field(min_length=64, max_length=64)
    reality_evidence_sha256: str = Field(min_length=64, max_length=64)
    expected_deliverable: str


class CodexLocalScreeningScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal[
        "v3.codex_local_screening_scope.1",
        "v3.codex_local_screening_scope.2",
    ] = (
        "v3.codex_local_screening_scope.1"
    )
    cohort_kind: Literal[
        "matched_screening",
        "representative_production_pilot",
    ] = "matched_screening"
    campaign_id: str
    campaign_manifest_sha256: str = Field(min_length=64, max_length=64)
    parity_report_sha256: str = Field(min_length=64, max_length=64)
    source_fingerprint: str = Field(min_length=64, max_length=64)
    cli: CodexLocalCLIIdentityV1
    model: Literal["gpt-5.6-sol"] = MODEL
    authentication: Literal["existing_chatgpt_codex_home"] = (
        "existing_chatgpt_codex_home"
    )
    public_probe_timeout_seconds: Literal[300] = PUBLIC_PROBE_TIMEOUT_SECONDS
    task_timeout_seconds: Literal[1800] = TASK_TIMEOUT_SECONDS
    process_attempts_per_task: Literal[1] = 1
    runner_retry_count: Literal[0] = 0
    builtin_transport_retry_control: Literal[
        "codex_chatgpt_provider_not_user_configurable"
    ] = "codex_chatgpt_provider_not_user_configurable"
    execution_isolation: Literal[
        "outer_managed_workspace_codex_danger_full_access"
    ] = "outer_managed_workspace_codex_danger_full_access"
    task_bindings: List[CodexLocalTaskBindingV1] = Field(
        min_length=12, max_length=24
    )
    private_package_upload_authorized: Literal[True] = True
    deepseek_grader_authorized: Literal[True] = True
    expert_review_authorized: Literal[False] = False
    training_authorized: Literal[False] = False
    registry_mutation_authorized: Literal[False] = False
    release_authorized: Literal[False] = False
    promotion_authorized: Literal[False] = False

    @model_validator(mode="after")
    def validate_cohort(self) -> "CodexLocalScreeningScopeV1":
        expected = (
            24
            if self.cohort_kind == "representative_production_pilot"
            else 12
        )
        if len(self.task_bindings) != expected or len(
            {item.blind_task_id for item in self.task_bindings}
        ) != expected:
            raise ValueError("codex_local_task_ids_not_unique")
        for field in (
            "package_fingerprint",
            "candidate_tree_sha256",
            "candidate_projection_sha256",
        ):
            if (
                len({getattr(item, field) for item in self.task_bindings})
                != expected
            ):
                raise ValueError(f"codex_local_{field}_not_unique")
        if self.cohort_kind == "representative_production_pilot":
            if self.scope_version != "v3.codex_local_screening_scope.2":
                raise ValueError("codex_local_representative_requires_v2")
            cells = {
                (
                    item.route_id,
                    item.motif,
                    item.replicate_id,
                    item.domain,
                )
                for item in self.task_bindings
            }
            if len(cells) != 24 or None in {
                item.domain for item in self.task_bindings
            }:
                raise ValueError(
                    "codex_local_representative_cells_incomplete"
                )
        else:
            if self.scope_version != "v3.codex_local_screening_scope.1":
                raise ValueError("codex_local_matched_requires_v1")
            cells = {
                (item.route_id, item.motif, item.replicate_id)
                for item in self.task_bindings
            }
            if len(cells) != 12:
                raise ValueError("codex_local_matched_cells_incomplete")
        return self


class CodexLocalScreeningReceiptV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    receipt_version: Literal["v3.codex_local_screening_receipt.1"] = (
        "v3.codex_local_screening_receipt.1"
    )
    scope_sha256: str = Field(min_length=64, max_length=64)
    authorized_by_user: Literal[True] = True
    authorization_basis: Literal["project_standing_external_model_authorization"] = (
        "project_standing_external_model_authorization"
    )
    issued_at: str


class CodexJSONLDiagnosticsV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    line_count: int = 0
    invalid_line_count: int = 0
    event_types: List[str] = Field(default_factory=list)
    turn_seen: bool = False
    command_event_seen: bool = False
    error_event_seen: bool = False
    usage: Dict[str, int] = Field(default_factory=dict)


def parse_codex_jsonl(text: str) -> CodexJSONLDiagnosticsV1:
    event_types: List[str] = []
    usage: Dict[str, int] = {}
    invalid = 0
    command = False
    turn = False
    error = False
    lines = [line for line in text.splitlines() if line.strip()]
    for line in lines:
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            invalid += 1
            continue
        if not isinstance(event, dict):
            invalid += 1
            continue
        kind = str(event.get("type", "unknown"))
        event_types.append(kind)
        if kind.startswith("turn."):
            turn = True
        if kind == "error":
            error = True
        item = event.get("item")
        item_type = item.get("type") if isinstance(item, dict) else None
        if item_type in {
            "command_execution",
            "mcp_tool_call",
            "file_change",
        }:
            command = True
        raw_usage = event.get("usage")
        if isinstance(raw_usage, dict):
            for key, value in raw_usage.items():
                if isinstance(value, int) and value >= 0:
                    usage[key] = usage.get(key, 0) + value
    return CodexJSONLDiagnosticsV1(
        line_count=len(lines),
        invalid_line_count=invalid,
        event_types=event_types,
        turn_seen=turn,
        command_event_seen=command,
        error_event_seen=error,
        usage=usage,
    )


class CodexLocalProcessOutcomeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    outcome_version: Literal["v3.codex_local_process_outcome.1"] = (
        "v3.codex_local_process_outcome.1"
    )
    kind: Literal["public_probe", "private_task"]
    blind_task_id: str
    status: Literal["pass", "task_failed", "infrastructure_failed"]
    started_at: str
    completed_at: str
    duration_seconds: float
    returncode: Optional[int]
    timed_out: bool
    command: List[str]
    workspace_path: str
    workspace_input_tree_sha256: str
    jsonl_path: str
    jsonl_sha256: str
    stderr_path: str
    stderr_sha256: str
    diagnostics: CodexJSONLDiagnosticsV1
    delivery: CodexLocalDeliveryInspectionV1
    output_tree_sha256: str
    first_failure: Optional[str] = None


class CodexLocalExecutionStateV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    state_version: Literal["v3.codex_local_execution_state.1"] = (
        "v3.codex_local_execution_state.1"
    )
    scope_sha256: str
    receipt_sha256: str
    public_probe_status: Literal[
        "not_started", "running", "completed", "interrupted"
    ] = "not_started"
    task_status: Dict[
        str, Literal["not_started", "running", "completed", "interrupted"]
    ]
    public_probe_outcome_path: Optional[str] = None
    task_outcome_paths: Dict[str, str] = Field(default_factory=dict)
    created_at: str
    updated_at: str


class CodexLocalCampaignManifestV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    manifest_version: Literal["v3.codex_local_campaign_manifest.1"] = (
        "v3.codex_local_campaign_manifest.1"
    )
    scope_sha256: str
    receipt_sha256: str
    status: Literal[
        "running", "solver_completed", "completed", "incomplete"
    ]
    public_probe_outcome_path: Optional[str] = None
    task_outcome_paths: Dict[str, str] = Field(default_factory=dict)
    grader_records: Dict[str, str] = Field(default_factory=dict)
    result_path: Optional[str] = None
    result_sha256: Optional[str] = None
    first_failure: Optional[str] = None
    created_at: str
    updated_at: str


class CodexLocalGraderDraftV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blind_task_id: str
    criteria: List[ScreeningCriterionGradeV1] = Field(
        min_length=7, max_length=7
    )
    major_defect: bool
    major_defect_reason: Optional[str] = None
    professional_plausibility: Literal["pass", "fail"]
    effective_rubric_dimensions: int = Field(ge=0, le=7)
    route_identity_seen: Literal[False] = False

    @model_validator(mode="after")
    def validate_draft(self) -> "CodexLocalGraderDraftV1":
        identities = [item.criterion_id for item in self.criteria]
        if len(set(identities)) != 7 or set(identities) != set(RUBRIC_WEIGHTS):
            raise ValueError("codex_local_grader_criteria_mismatch")
        if self.major_defect and not self.major_defect_reason:
            raise ValueError("codex_local_grader_major_defect_reason_missing")
        return self


class CodexLocalGraderRecordV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blind_task_id: str
    status: Literal["not_eligible", "completed", "infrastructure_failed"]
    attempt_count: int = Field(ge=0, le=2)
    review_path: Optional[str] = None
    review_sha256: Optional[str] = None
    first_failure_path: Optional[str] = None
    attempt_artifacts: Dict[str, str] = Field(default_factory=dict)


class CodexLocalTooling:
    def __init__(self, install_root: str | Path) -> None:
        self.root = Path(install_root).resolve()

    @property
    def executable(self) -> Path:
        suffix = ".cmd" if os.name == "nt" else ""
        return self.root / "node_modules" / ".bin" / f"codex{suffix}"

    def freeze_identity(self) -> CodexLocalCLIIdentityV1:
        package_json = self.root / "node_modules" / "@openai" / "codex" / "package.json"
        lock = self.root / "package-lock.json"
        if not package_json.is_file() or not lock.is_file() or not self.executable.is_file():
            raise FileNotFoundError("codex_local_cli_not_installed")
        package = json.loads(package_json.read_text(encoding="utf-8"))
        return CodexLocalCLIIdentityV1(
            version=str(package["version"]),
            package_lock_sha256=_sha_file(lock),
            executable_sha256=_sha_file(self.executable),
            executable_path=str(self.executable),
            installation_tree_sha256=_tree(self.root / "node_modules"),
        )

    def install_exact(
        self,
        *,
        version: str,
        npm_executable: str | Path,
        command_runner=subprocess.run,
    ) -> CodexLocalCLIIdentityV1:
        if self.root.exists() and any(self.root.iterdir()):
            return self.freeze_identity()
        self.root.mkdir(parents=True, exist_ok=True)
        command = [
            str(npm_executable),
            "install",
            "--prefix",
            str(self.root),
            "--save-exact",
            f"{CLI_PACKAGE}@{version}",
        ]
        completed = command_runner(
            command,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            check=False,
            timeout=600,
        )
        if completed.returncode != 0:
            raise RuntimeError("codex_local_cli_install_failed")
        return self.freeze_identity()


def compile_scope(
    *,
    campaign_root: str | Path,
    parity_report_path: str | Path,
    cli_identity: CodexLocalCLIIdentityV1,
    repository_root: str | Path,
    output_root: str | Path,
) -> tuple[CodexLocalScreeningScopeV1, Path, str]:
    campaign_dir = Path(campaign_root).resolve()
    campaign_path = campaign_dir / "matched_screening_campaign.json"
    campaign = json.loads(campaign_path.read_text(encoding="utf-8"))
    parity_path = Path(parity_report_path).resolve()
    parity = json.loads(parity_path.read_text(encoding="utf-8"))
    source_fingerprint = governed_source_fingerprint(repository_root)
    if not (
        parity.get("passed")
        and parity.get("network_mode") == "none"
        and parity.get("read_only_root") is True
        and parity.get("provider_credentials_mounted") is False
        and parity.get("source_fingerprint") == source_fingerprint
    ):
        raise ValueError("codex_local_parity_not_eligible")
    candidate_root = campaign_dir / "blind_staging" / "candidate_packages"
    bindings: List[CodexLocalTaskBindingV1] = []
    for assignment in campaign["assignments"]:
        package = candidate_root / assignment["blind_task_id"]
        bindings.append(
            CodexLocalTaskBindingV1(
                blind_task_id=assignment["blind_task_id"],
                brief_id=assignment["brief_id"],
                route_id=assignment["route_id"],
                motif=assignment["motif"],
                replicate_id=assignment["replicate_id"],
                package_fingerprint=assignment["package_fingerprint"],
                candidate_tree_sha256=_tree(package),
                candidate_projection_sha256=_projection_tree(package),
                reality_evidence_sha256=assignment["reality_evidence_sha256"],
                expected_deliverable=_expected_deliverable(package),
            )
        )
    scope = CodexLocalScreeningScopeV1(
        campaign_id=campaign["campaign_id"],
        campaign_manifest_sha256=_sha_file(campaign_path),
        parity_report_sha256=_sha_file(parity_path),
        source_fingerprint=source_fingerprint,
        cli=cli_identity,
        task_bindings=bindings,
    )
    encoded = json.dumps(
        scope.model_dump(mode="json"),
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    persisted = encoded + b"\n"
    digest = _sha_bytes(persisted)
    scope_path = Path(output_root).resolve() / "governance" / "scopes" / f"{digest}.json"
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    if scope_path.exists() and scope_path.read_bytes() != persisted:
        raise ValueError("codex_local_scope_immutable_collision")
    scope_path.write_bytes(persisted)
    return scope, scope_path, digest


def compile_standing_receipt(
    *, scope_path: str | Path, output_root: str | Path
) -> tuple[CodexLocalScreeningReceiptV1, Path]:
    path = Path(scope_path).resolve()
    digest = _sha_file(path)
    if path.name != f"{digest}.json":
        raise ValueError("codex_local_scope_not_content_addressed")
    receipt = CodexLocalScreeningReceiptV1(
        scope_sha256=digest,
        issued_at=_now(),
    )
    receipt_path = (
        Path(output_root).resolve()
        / "governance"
        / "receipts"
        / f"{digest}.json"
    )
    if receipt_path.exists():
        raise FileExistsError("codex_local_receipt_exists")
    _atomic_json(receipt_path, receipt.model_dump(mode="json"))
    return receipt, receipt_path


def _sanitized_environment(codex_home: Path) -> Dict[str, str]:
    result = {
        key: value
        for key, value in os.environ.items()
        if key.upper() not in SECRET_ENV_NAMES
        and not key.upper().startswith("CODEX_")
    }
    result["CODEX_HOME"] = str(codex_home)
    return result


def _build_command(executable: Path, workspace: Path) -> List[str]:
    return [
        str(executable),
        "--ask-for-approval",
        "never",
        "--model",
        MODEL,
        "exec",
        "-c",
        "project_doc_max_bytes=0",
        "-c",
        "agents.enabled=false",
        "--disable",
        "plugins",
        "--disable",
        "apps",
        "--disable",
        "multi_agent",
        "--disable",
        "skill_search",
        "--json",
        "--ephemeral",
        "--ignore-user-config",
        "--ignore-rules",
        "--sandbox",
        "danger-full-access",
        "--skip-git-repo-check",
        "-C",
        str(workspace),
        "-",
    ]


def _terminate_process_tree(process: subprocess.Popen) -> None:
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"],
            text=True,
            capture_output=True,
            check=False,
            timeout=30,
        )
    else:
        try:
            os.killpg(process.pid, signal.SIGTERM)
            process.wait(timeout=5)
        except (ProcessLookupError, subprocess.TimeoutExpired):
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass


def _prompt(expected: str, python_executable: Path) -> str:
    return f"""You are completing a route-blind spreadsheet work task.
Read dataset_row.json, deliverable_contract.json, and every workbook under reference_files/.
Use shell commands and Python with openpyxl as needed. The available Python executable is:
{python_executable}

Produce exactly this final deliverable path inside the current working directory:
{expected}

The output must be a genuine, non-empty XLSX that addresses the task. Do not merely copy an
input workbook. Do not write the final deliverable anywhere else. Verify it with openpyxl
before finishing.
"""


def _write_public_fixture(workspace: Path) -> str:
    workspace.mkdir(parents=True, exist_ok=True)
    (workspace / "reference_files").mkdir()
    (workspace / "deliverable_files").mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet.append(["Item", "Amount"])
    sheet.append(["A", 12])
    sheet.append(["B", 30])
    workbook.save(workspace / "reference_files" / "public_source.xlsx")
    row = {
        "task_id": "public_codex_local_probe",
        "prompt": (
            "Create an XLSX summary with the two source rows and a total amount. "
            "Add a Notes sheet stating that this is a public synthetic fixture."
        ),
    }
    contract = {
        "deliverables": [
            {"relative_path": "deliverable_files/codex_local_probe.xlsx"}
        ]
    }
    _atomic_json(workspace / "dataset_row.json", row)
    _atomic_json(workspace / "deliverable_contract.json", contract)
    return "deliverable_files/codex_local_probe.xlsx"


class CodexLocalRunner:
    def __init__(
        self,
        *,
        scope_path: str | Path,
        receipt_path: str | Path,
        campaign_root: str | Path,
        output_root: str | Path,
        codex_home: str | Path,
        python_executable: str | Path,
        popen_factory=subprocess.Popen,
    ) -> None:
        self.scope_path = Path(scope_path).resolve()
        self.receipt_path = Path(receipt_path).resolve()
        self.campaign_root = Path(campaign_root).resolve()
        self.output = Path(output_root).resolve()
        self.codex_home = Path(codex_home).resolve()
        self.python = Path(python_executable).resolve()
        self.popen_factory = popen_factory
        self.scope = CodexLocalScreeningScopeV1.model_validate_json(
            self.scope_path.read_text(encoding="utf-8")
        )
        self.receipt = CodexLocalScreeningReceiptV1.model_validate_json(
            self.receipt_path.read_text(encoding="utf-8")
        )
        self._validate_static()

    def _validate_static(self) -> None:
        if _sha_file(self.scope_path) != self.receipt.scope_sha256:
            raise PermissionError("codex_local_receipt_scope_mismatch")
        repository_root = Path(__file__).resolve().parents[2]
        if (
            governed_source_fingerprint(repository_root)
            != self.scope.source_fingerprint
        ):
            raise PermissionError("codex_local_source_fingerprint_drift")
        if not self.python.is_file():
            raise FileNotFoundError("codex_local_python_missing")
        executable = Path(self.scope.cli.executable_path)
        if not executable.is_file() or _sha_file(executable) != self.scope.cli.executable_sha256:
            raise PermissionError("codex_local_cli_identity_drift")
        if CodexLocalTooling(executable.parents[2]).freeze_identity() != self.scope.cli:
            raise PermissionError("codex_local_cli_installation_drift")
        campaign_path = self.campaign_root / (
            "representative_codex_campaign.json"
            if self.scope.cohort_kind == "representative_production_pilot"
            else "matched_screening_campaign.json"
        )
        if _sha_file(campaign_path) != self.scope.campaign_manifest_sha256:
            raise PermissionError("codex_local_campaign_manifest_drift")
        campaign = json.loads(campaign_path.read_text(encoding="utf-8"))
        assignments = {
            item["blind_task_id"]: item for item in campaign["assignments"]
        }
        if set(assignments) != {
            item.blind_task_id for item in self.scope.task_bindings
        }:
            raise PermissionError("codex_local_campaign_task_identity_drift")
        packages = (
            self.campaign_root
            / (
                "reality/blind_staging/candidate_packages"
                if self.scope.cohort_kind
                == "representative_production_pilot"
                else "blind_staging/candidate_packages"
            )
        )
        for binding in self.scope.task_bindings:
            assignment = assignments[binding.blind_task_id]
            for field in (
                "brief_id",
                "route_id",
                "motif",
                "replicate_id",
                "package_fingerprint",
                "reality_evidence_sha256",
            ):
                if getattr(binding, field) != assignment[field]:
                    raise PermissionError(
                        f"codex_local_campaign_binding_drift:{field}"
                    )
            if binding.domain != assignment.get("domain"):
                raise PermissionError(
                    "codex_local_campaign_binding_drift:domain"
                )
            package = packages / binding.blind_task_id
            if _tree(package) != binding.candidate_tree_sha256:
                raise PermissionError("codex_local_candidate_tree_drift")
            if _projection_tree(package) != binding.candidate_projection_sha256:
                raise PermissionError("codex_local_candidate_projection_drift")
            if _expected_deliverable(package) != binding.expected_deliverable:
                raise PermissionError("codex_local_deliverable_contract_drift")

    def _initial_state(self) -> CodexLocalExecutionStateV1:
        now = _now()
        return CodexLocalExecutionStateV1(
            scope_sha256=self.receipt.scope_sha256,
            receipt_sha256=_sha_file(self.receipt_path),
            task_status={
                item.blind_task_id: "not_started"
                for item in self.scope.task_bindings
            },
            created_at=now,
            updated_at=now,
        )

    def run_solver(self) -> tuple[CodexLocalCampaignManifestV1, Path]:
        self.output.mkdir(parents=True, exist_ok=True)
        state_path = self.output / "execution_state.json"
        manifest_path = self.output / "codex_local_campaign_manifest.json"
        if state_path.exists():
            state = CodexLocalExecutionStateV1.model_validate_json(
                state_path.read_text(encoding="utf-8")
            )
            if state.scope_sha256 != self.receipt.scope_sha256:
                raise PermissionError("codex_local_resume_scope_mismatch")
            if state.public_probe_status == "running":
                state.public_probe_status = "interrupted"
            for task_id, status in list(state.task_status.items()):
                if status == "running":
                    state.task_status[task_id] = "interrupted"
            state.updated_at = _now()
            _atomic_json(state_path, state.model_dump(mode="json"))
        else:
            state = self._initial_state()
            _atomic_json(state_path, state.model_dump(mode="json"))
        now = _now()
        if manifest_path.exists():
            manifest = CodexLocalCampaignManifestV1.model_validate_json(
                manifest_path.read_text(encoding="utf-8")
            )
        else:
            manifest = CodexLocalCampaignManifestV1(
                scope_sha256=self.receipt.scope_sha256,
                receipt_sha256=_sha_file(self.receipt_path),
                status="running",
                created_at=now,
                updated_at=now,
            )
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))

        if state.public_probe_status == "not_started":
            state.public_probe_status = "running"
            state.updated_at = _now()
            _atomic_json(state_path, state.model_dump(mode="json"))
            workspace = self.output / "workspaces" / "public_probe"
            expected = _write_public_fixture(workspace)
            outcome, outcome_path = self._run_one(
                kind="public_probe",
                task_id="public_codex_local_probe",
                workspace=workspace,
                expected=expected,
                timeout=PUBLIC_PROBE_TIMEOUT_SECONDS,
            )
            state.public_probe_status = "completed"
            state.public_probe_outcome_path = str(outcome_path)
            manifest.public_probe_outcome_path = str(outcome_path)
            state.updated_at = manifest.updated_at = _now()
            _atomic_json(state_path, state.model_dump(mode="json"))
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            if outcome.status != "pass":
                manifest.status = "incomplete"
                manifest.first_failure = outcome.first_failure
                manifest.updated_at = _now()
                _atomic_json(manifest_path, manifest.model_dump(mode="json"))
                return manifest, manifest_path
        elif state.public_probe_status != "completed":
            manifest.status = "incomplete"
            manifest.first_failure = "public_probe_interrupted_no_retry"
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            return manifest, manifest_path

        probe = CodexLocalProcessOutcomeV1.model_validate_json(
            Path(state.public_probe_outcome_path).read_text(encoding="utf-8")
        )
        if probe.status != "pass":
            manifest.status = "incomplete"
            manifest.first_failure = probe.first_failure
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            return manifest, manifest_path

        packages = (
            self.campaign_root
            / (
                "reality/blind_staging/candidate_packages"
                if self.scope.cohort_kind
                == "representative_production_pilot"
                else "blind_staging/candidate_packages"
            )
        )
        for binding in sorted(
            self.scope.task_bindings, key=lambda item: item.blind_task_id
        ):
            task_id = binding.blind_task_id
            if state.task_status[task_id] != "not_started":
                continue
            state.task_status[task_id] = "running"
            state.updated_at = _now()
            _atomic_json(state_path, state.model_dump(mode="json"))
            workspace = self.output / "workspaces" / task_id
            _copy_projection(packages / task_id, workspace)
            if _projection_tree(workspace) != binding.candidate_projection_sha256:
                raise PermissionError("codex_local_staged_projection_drift")
            outcome, outcome_path = self._run_one(
                kind="private_task",
                task_id=task_id,
                workspace=workspace,
                expected=binding.expected_deliverable,
                timeout=TASK_TIMEOUT_SECONDS,
            )
            state.task_status[task_id] = "completed"
            state.task_outcome_paths[task_id] = str(outcome_path)
            manifest.task_outcome_paths[task_id] = str(outcome_path)
            state.updated_at = manifest.updated_at = _now()
            _atomic_json(state_path, state.model_dump(mode="json"))
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            if outcome.status == "infrastructure_failed":
                manifest.status = "incomplete"
                manifest.first_failure = (
                    f"{task_id}:{outcome.first_failure or 'infrastructure_failed'}"
                )
                manifest.updated_at = _now()
                _atomic_json(manifest_path, manifest.model_dump(mode="json"))
                return manifest, manifest_path
        interrupted = any(
            status in {"running", "interrupted"}
            for status in state.task_status.values()
        )
        manifest.status = "incomplete" if interrupted else "solver_completed"
        manifest.first_failure = (
            "one_or_more_tasks_interrupted_no_retry" if interrupted else None
        )
        manifest.updated_at = _now()
        _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        return manifest, manifest_path

    def _run_one(
        self,
        *,
        kind: Literal["public_probe", "private_task"],
        task_id: str,
        workspace: Path,
        expected: str,
        timeout: int,
    ) -> tuple[CodexLocalProcessOutcomeV1, Path]:
        command = _build_command(Path(self.scope.cli.executable_path), workspace)
        env = _sanitized_environment(self.codex_home)
        input_hashes = {
            _sha_file(path)
            for path in (workspace / "reference_files").rglob("*.xlsx")
        }
        started_at = _now()
        started = time.monotonic()
        creation_flags = (
            subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0
        )
        process = self.popen_factory(
            command,
            cwd=str(workspace),
            env=env,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdin=subprocess.PIPE,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            creationflags=creation_flags,
            start_new_session=os.name != "nt",
        )
        timed_out = False
        try:
            stdout, stderr = process.communicate(
                input=_prompt(expected, self.python), timeout=timeout
            )
        except subprocess.TimeoutExpired as exc:
            timed_out = True
            _terminate_process_tree(process)
            tail_out, tail_err = process.communicate()
            stdout = (exc.stdout or "") + (tail_out or "")
            stderr = (exc.stderr or "") + (tail_err or "")
        duration = time.monotonic() - started
        log_root = self.output / "process_logs"
        log_root.mkdir(parents=True, exist_ok=True)
        jsonl_path = log_root / f"{task_id}.jsonl"
        stderr_path = log_root / f"{task_id}.stderr.txt"
        jsonl_path.write_text(stdout or "", encoding="utf-8")
        stderr_path.write_text(stderr or "", encoding="utf-8")
        diagnostics = parse_codex_jsonl(stdout or "")
        delivery = inspect_delivery(workspace, expected, input_hashes)
        combined = f"{stdout}\n{stderr}".lower()
        infrastructure = any(marker in combined for marker in INFRASTRUCTURE_MARKERS)
        if timed_out:
            status = (
                "task_failed"
                if diagnostics.command_event_seen
                else "infrastructure_failed"
            )
            failure = (
                "solver_timeout_with_tool_activity"
                if diagnostics.command_event_seen
                else "solver_timeout_without_tool_activity"
            )
        elif process.returncode == 0 and delivery.valid:
            status, failure = "pass", None
        elif infrastructure or (
            process.returncode not in (0, None)
            and not diagnostics.turn_seen
        ):
            status, failure = "infrastructure_failed", "codex_cli_infrastructure_failure"
        else:
            status = "task_failed"
            failure = delivery.failure or "codex_cli_task_failed"
        outcome = CodexLocalProcessOutcomeV1(
            kind=kind,
            blind_task_id=task_id,
            status=status,
            started_at=started_at,
            completed_at=_now(),
            duration_seconds=round(duration, 3),
            returncode=process.returncode,
            timed_out=timed_out,
            command=command,
            workspace_path=str(workspace),
            workspace_input_tree_sha256=_projection_tree(workspace),
            jsonl_path=str(jsonl_path),
            jsonl_sha256=_sha_file(jsonl_path),
            stderr_path=str(stderr_path),
            stderr_sha256=_sha_file(stderr_path),
            diagnostics=diagnostics,
            delivery=delivery,
            output_tree_sha256=_tree(workspace),
            first_failure=failure,
        )
        outcome_path = self.output / "outcomes" / f"{task_id}.json"
        _atomic_json(outcome_path, outcome.model_dump(mode="json"))
        return outcome, outcome_path


class CodexLocalScreeningFinalizer:
    """Grades valid local-Codex deliveries and applies the frozen analyzer."""

    def __init__(
        self,
        *,
        campaign_root: str | Path,
        execution_root: str | Path,
        provider_config: ProviderConfig,
        grader_call=None,
    ) -> None:
        self.campaign_root = Path(campaign_root).resolve()
        self.output = Path(execution_root).resolve()
        self.config = provider_config
        self.grader_call = grader_call
        if (
            provider_config.provider_name != "deepseek"
            or provider_config.model != "deepseek-v4-pro"
        ):
            raise PermissionError("codex_local_grader_provider_mismatch")

    def finalize(self) -> tuple[CodexLocalCampaignManifestV1, Path]:
        manifest_path = self.output / "codex_local_campaign_manifest.json"
        manifest = CodexLocalCampaignManifestV1.model_validate_json(
            manifest_path.read_text(encoding="utf-8")
        )
        if manifest.status != "solver_completed":
            raise PermissionError("codex_local_solver_not_complete")
        campaign = json.loads(
            (self.campaign_root / "matched_screening_campaign.json").read_text(
                encoding="utf-8"
            )
        )
        assignments = {
            item["blind_task_id"]: item for item in campaign["assignments"]
        }
        records: Dict[str, CodexLocalGraderRecordV1] = {}
        provider_calls = 0
        for task_id in sorted(assignments):
            outcome = CodexLocalProcessOutcomeV1.model_validate_json(
                Path(manifest.task_outcome_paths[task_id]).read_text(
                    encoding="utf-8"
                )
            )
            record = CodexLocalGraderRecordV1(
                blind_task_id=task_id,
                status="not_eligible",
                attempt_count=0,
            )
            if outcome.status == "pass" and outcome.delivery.valid:
                record, calls = self._grade_one(
                    task_id=task_id,
                    delivery=Path(outcome.workspace_path)
                    / outcome.delivery.relative_path,
                    candidate_package=self.campaign_root
                    / "blind_staging"
                    / "candidate_packages"
                    / task_id,
                    teacher_package=Path(assignments[task_id]["package_root"]),
                )
                provider_calls += calls
            record_path = self.output / "grader" / task_id / "record.json"
            _atomic_json(record_path, record.model_dump(mode="json"))
            records[task_id] = record
            manifest.grader_records[task_id] = str(record_path)
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        observations = self._observations(campaign, assignments, records)
        result = MatchedScreeningAnalyzer().analyze(
            campaign["campaign_id"], observations
        )
        result_path = self.output / "matched_screening_result.json"
        _atomic_json(result_path, result.model_dump(mode="json"))
        manifest.result_path = str(result_path)
        manifest.result_sha256 = _sha_file(result_path)
        manifest.status = (
            "incomplete" if result.decision == "incomplete" else "completed"
        )
        manifest.first_failure = (
            f"grader_infrastructure_failure:{provider_calls}_calls"
            if result.decision == "incomplete"
            and any(
                record.status == "infrastructure_failed"
                for record in records.values()
            )
            else None
        )
        manifest.updated_at = _now()
        _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        return manifest, manifest_path

    def _grade_one(
        self,
        *,
        task_id: str,
        delivery: Path,
        candidate_package: Path,
        teacher_package: Path,
    ) -> tuple[CodexLocalGraderRecordV1, int]:
        executor = SemanticReviewExecutor(
            self.config,
            max_tokens=4000,
            input_token_hard_limit=20000,
            max_retries=0,
        )
        payload = {
            "blind_task_id": task_id,
            "candidate_prompt": json.loads(
                (candidate_package / "dataset_row.json").read_text(
                    encoding="utf-8"
                )
            )["prompt"],
            "rubric": json.loads(
                (
                    teacher_package / "teacher" / "rubric_plan_v2.json"
                ).read_text(encoding="utf-8")
            ),
            "fact_anchors": json.loads(
                (
                    teacher_package
                    / "teacher"
                    / "deterministic_fact_anchors.json"
                ).read_text(encoding="utf-8")
            ),
            "deliverable": executor._xlsx_content(delivery),
        }
        system = (
            "Grade the route-blind professional XLSX using exactly the seven "
            "supplied rubric criteria. Return 0, 0.5, or 1 per criterion with "
            "evidence locators. Do not infer route identity. Do not return a "
            "total score."
        )
        record = CodexLocalGraderRecordV1(
            blind_task_id=task_id,
            status="not_eligible",
            attempt_count=0,
        )
        first_feedback = None
        provider_calls = 0
        for attempt in (1, 2):
            record.attempt_count = attempt
            attempt_root = self.output / "grader" / task_id
            raw_path = attempt_root / f"attempt_{attempt}.raw.txt"
            diagnostics_path = (
                attempt_root / f"attempt_{attempt}.diagnostics.json"
            )
            try:
                if self.grader_call:
                    draft = self.grader_call(
                        payload, system, first_feedback
                    )
                    if not isinstance(draft, CodexLocalGraderDraftV1):
                        draft = CodexLocalGraderDraftV1.model_validate(draft)
                    raw = draft.model_dump_json()
                    diagnostics = {"mock": True}
                else:
                    draft = executor._call(
                        system,
                        payload,
                        CodexLocalGraderDraftV1,
                        {"blind_task_id": task_id},
                        format_feedback=first_feedback,
                    )
                    raw = executor.last_raw_response_content
                    diagnostics = executor.last_diagnostics
                provider_calls += 1
                raw_path.parent.mkdir(parents=True, exist_ok=True)
                raw_path.write_text(raw, encoding="utf-8")
                _atomic_json(diagnostics_path, diagnostics)
                record.attempt_artifacts[str(raw_path)] = _sha_file(raw_path)
                record.attempt_artifacts[str(diagnostics_path)] = _sha_file(
                    diagnostics_path
                )
                scores = {
                    item.criterion_id: item.score for item in draft.criteria
                }
                weighted = round(
                    sum(
                        scores[key] * RUBRIC_WEIGHTS[key]
                        for key in RUBRIC_WEIGHTS
                    ),
                    6,
                )
                review = ScreeningGraderReviewV1(
                    **draft.model_dump(mode="json"),
                    weighted_score=weighted,
                )
                review_path = attempt_root / "review.json"
                _atomic_json(review_path, review.model_dump(mode="json"))
                record.status = "completed"
                record.review_path = str(review_path)
                record.review_sha256 = _sha_file(review_path)
                return record, provider_calls
            except Exception as exc:
                raw_path.parent.mkdir(parents=True, exist_ok=True)
                if not raw_path.exists():
                    raw_path.write_text(
                        getattr(executor, "last_raw_response_content", ""),
                        encoding="utf-8",
                    )
                diagnostics = {
                    "error_type": type(exc).__name__,
                    "failure_code": getattr(
                        exc, "failure_code", "schema_failure"
                    ),
                    **getattr(executor, "last_diagnostics", {}),
                }
                _atomic_json(diagnostics_path, diagnostics)
                record.attempt_artifacts[str(raw_path)] = _sha_file(raw_path)
                record.attempt_artifacts[str(diagnostics_path)] = _sha_file(
                    diagnostics_path
                )
                if getattr(exc, "failure_code", None) != "input_token_ceiling":
                    provider_calls += 1
                if attempt == 1:
                    record.first_failure_path = str(diagnostics_path)
                eligible = isinstance(
                    exc, (SemanticReviewExecutionError, ValueError)
                ) and (
                    not isinstance(exc, SemanticReviewExecutionError)
                    or exc.retry_eligible
                    or exc.failure_code
                    in {
                        "invalid_json",
                        "schema_failure",
                        "empty_response",
                        "truncated_response",
                    }
                )
                if attempt == 1 and eligible:
                    first_feedback = (
                        "Return valid JSON matching the required schema; "
                        "preserve the same substantive judgment."
                    )
                    continue
                record.status = "infrastructure_failed"
                return record, provider_calls
        raise AssertionError("unreachable")

    def _observations(
        self,
        campaign: dict,
        assignments: Dict[str, dict],
        grader_records: Dict[str, CodexLocalGraderRecordV1],
    ) -> List[ScreeningTaskObservationV1]:
        rows: List[ScreeningTaskObservationV1] = []
        for assignment in campaign["assignments"]:
            task_id = assignment["blind_task_id"]
            outcome = CodexLocalProcessOutcomeV1.model_validate_json(
                Path(
                    CodexLocalCampaignManifestV1.model_validate_json(
                        (
                            self.output / "codex_local_campaign_manifest.json"
                        ).read_text(encoding="utf-8")
                    ).task_outcome_paths[task_id]
                ).read_text(encoding="utf-8")
            )
            teacher_root = Path(assignments[task_id]["package_root"])
            validity = ValidityVectorV1.model_validate_json(
                (
                    teacher_root / "governance" / "validity_vector.json"
                ).read_text(encoding="utf-8")
            )
            utility = UtilityProfileV1.model_validate_json(
                (
                    teacher_root / "governance" / "utility_profile.json"
                ).read_text(encoding="utf-8")
            )
            grader = grader_records[task_id]
            review = (
                ScreeningGraderReviewV1.model_validate_json(
                    Path(grader.review_path).read_text(encoding="utf-8")
                )
                if grader.review_path
                else None
            )
            rows.append(
                ScreeningTaskObservationV1(
                    blind_task_id=task_id,
                    route_id=assignment["route_id"],
                    motif=assignment["motif"],
                    replicate_id=assignment["replicate_id"],
                    infrastructure_complete=(
                        outcome.status != "infrastructure_failed"
                        and grader.status != "infrastructure_failed"
                    ),
                    offline_validity_pass=validity.overall_status != "blocked",
                    exact_valid_delivery=outcome.delivery.valid,
                    major_defect=review.major_defect if review else False,
                    professional_plausibility_pass=(
                        review.professional_plausibility == "pass"
                        if review
                        else False
                    ),
                    productive_complexity_pass=(
                        utility.productive_complexity_coverage == 1.0
                        and all(
                            item.status != "blocked"
                            for item in utility.productive_complexity
                        )
                    ),
                    skill_causal_pass=utility.skill_causal_coverage == 1.0,
                    effective_rubric_dimensions=(
                        review.effective_rubric_dimensions if review else 0
                    ),
                    weighted_score=review.weighted_score if review else None,
                )
            )
        return rows
