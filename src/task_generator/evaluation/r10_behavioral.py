"""Thin, task-specific behavioral evidence contracts for the R10 four-task pilot.

This module intentionally does not adapt the R9 24-task screening schema.  R10
has mixed DOCX/XLSX deliverables and task-specific decision matrices, so its
only durable responsibilities are identity binding, delivery inspection,
judge-output validation, and deterministic aggregation.
"""

from __future__ import annotations

import hashlib
import json
import subprocess
import tempfile
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from task_generator.core.deliverable_contract import DeliverableContractV1
from task_generator.planning.scenario_task_compiler import (
    TaskCompilationOutputV1,
    TaskSpecificRubricV1,
    tree_sha256,
)


SolverStackId = Literal[
    "gpt-5.6-sol@chatgpt_codex",
    "deepseek-v4-pro@official_opencode",
]
JudgeId = SolverStackId
RunStatus = Literal[
    "not_started", "running", "completed", "task_failed", "infrastructure_failed", "interrupted"
]
DecisionRating = Literal["met", "partial", "not_met"]
TeacherAnchorKind = Literal["percentage_threshold", "sum", "date_order", "quantity_delta"]
TeacherAnchorRelation = Literal["within", "exceeds", "equals", "before", "after"]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sha256_json(value: object) -> str:
    return hashlib.sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


class R10PilotTaskBindingV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    task_id: str = Field(min_length=1)
    domain: Literal["audit_compliance", "procurement_operations"]
    package_root: str = Field(min_length=1)
    package_tree_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    candidate_tree_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    task_compilation_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    deliverable_contract_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    expected_delivery: str = Field(min_length=1)


class R10BehavioralScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    scope_version: Literal["r10.behavioral_scope.1"] = "r10.behavioral_scope.1"
    campaign_id: str = Field(min_length=1)
    source_commit: str = Field(pattern=r"^[0-9a-f]{40}$")
    image: str = Field(min_length=1)
    image_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    solver_stacks: tuple[SolverStackId, SolverStackId] = (
        "gpt-5.6-sol@chatgpt_codex",
        "deepseek-v4-pro@official_opencode",
    )
    judges: tuple[JudgeId, JudgeId] = (
        "gpt-5.6-sol@chatgpt_codex",
        "deepseek-v4-pro@official_opencode",
    )
    task_timeout_seconds: Literal[1800] = 1800
    solver_attempt_limit: Literal[1] = 1
    judge_format_attempt_limit: Literal[2] = 2
    bindings: list[R10PilotTaskBindingV1] = Field(min_length=4, max_length=4)
    professional_validity: Literal["provisional"] = "provisional"
    excluded_actions: list[str] = Field(
        default_factory=lambda: ["task_generation", "candidate_mutation", "release", "training", "promotion"]
    )

    @model_validator(mode="after")
    def require_four_cross_domain_tasks(self) -> "R10BehavioralScopeV1":
        if len({item.task_id for item in self.bindings}) != 4:
            raise ValueError("r10_behavioral_task_ids_not_unique")
        if {item.domain for item in self.bindings} != {"audit_compliance", "procurement_operations"}:
            raise ValueError("r10_behavioral_domains_incomplete")
        if len(set(self.solver_stacks)) != 2 or len(set(self.judges)) != 2:
            raise ValueError("r10_behavioral_requires_two_distinct_stacks")
        return self


class DeliveryInspectionV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    relative_path: str
    format: Literal["xlsx", "docx"]
    exists: bool
    nonempty: bool
    package_valid: bool
    office_openable: bool
    has_content: bool
    differs_from_inputs: bool
    sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    valid: bool
    first_failure: str | None = None


class R10SolverOutcomeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    task_id: str
    solver_id: SolverStackId
    status: RunStatus
    returncode: int | None = None
    timed_out: bool = False
    duration_seconds: float = Field(ge=0)
    stdout_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    stderr_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    delivery: DeliveryInspectionV1
    first_failure: str | None = None


class DecisionAssessmentV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    decision_id: str = Field(min_length=1)
    rating: DecisionRating
    major_error: bool
    evidence_paths: list[str] = Field(min_length=1)
    rationale: str = Field(min_length=1, max_length=1600)
    evidence_insufficient: bool = False

    @classmethod
    def __get_pydantic_json_schema__(cls, core_schema, handler):  # type: ignore[no-untyped-def]
        """Keep the Python default while satisfying Responses strict-schema mode."""
        schema = handler(core_schema)
        schema["required"] = list(schema.get("properties", {}))
        return schema


class R10JudgeDraftV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    task_id: str = Field(min_length=1)
    assessments: list[DecisionAssessmentV1] = Field(min_length=3, max_length=5)

    @model_validator(mode="after")
    def require_unique_decisions(self) -> "R10JudgeDraftV1":
        ids = [item.decision_id for item in self.assessments]
        if len(ids) != len(set(ids)):
            raise ValueError("r10_judge_duplicate_decision")
        return self


class R10JudgeReviewV1(R10JudgeDraftV1):
    judge_id: JudgeId
    weighted_score: float = Field(ge=0, le=1)
    major_defect: bool


class R10ModelTaskResultV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    task_id: str
    solver_id: SolverStackId
    delivery_valid: bool
    reviews: list[R10JudgeReviewV1] = Field(default_factory=list, max_length=2)


class R10TaskDiscriminationV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    task_id: str
    comparable: bool
    composite_score_gap: float | None = Field(default=None, ge=0, le=1)
    delivery_or_major_defect_differs: bool = False
    explainable_difference: bool = False


class R10BehavioralResultV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    result_version: Literal["r10.behavioral_result.1"] = "r10.behavioral_result.1"
    decision: Literal["behaviorally_admitted", "behaviorally_inconclusive", "task_design_revision_candidate"]
    low_model_separation: bool
    common_dual_graded_count: int = Field(ge=0, le=4)
    solver_valid_delivery_counts: dict[SolverStackId, int]
    task_discrimination: list[R10TaskDiscriminationV1] = Field(min_length=4, max_length=4)
    recurring_insufficient_evidence_decisions: list[str] = Field(default_factory=list)
    recurring_major_error_decisions: list[str] = Field(default_factory=list)
    professional_validity: Literal["provisional"] = "provisional"


PilotDiscriminationClassification = Literal[
    "saturated", "near_tie", "cleanly_discriminative", "judge_ambiguous", "incomplete"
]
CompilerDiagnosticSignal = Literal[
    "score_saturation", "decision_points_too_explicit", "evidence_tension_insufficient",
    "judge_boundary_ambiguous",
]


class R10PilotTaskDiscriminationDiagnosisV1(BaseModel):
    """Read-only explanation of one frozen pilot task's observed separation.

    The diagnosis reports evidence quality, rather than claiming that either
    Solver is intrinsically stronger.  A Judge disagreement prevents a
    non-tied score gap from being promoted into a clean capability comparison;
    a true near-tie remains a near-tie rather than fabricated separation.
    """

    model_config = ConfigDict(extra="forbid")

    task_id: str = Field(min_length=1)
    solver_composite_scores: dict[SolverStackId, float | None]
    composite_score_gap: float | None = Field(default=None, ge=0, le=1)
    delivery_or_major_defect_differs: bool
    judge_score_gaps: dict[SolverStackId, float | None]
    judge_major_defect_disagreement: dict[SolverStackId, bool]
    classification: PilotDiscriminationClassification
    compiler_signals: list[CompilerDiagnosticSignal] = Field(default_factory=list)
    first_failure: str | None = None


class R10PilotDiscriminationReportV1(BaseModel):
    """Versioned, read-only diagnosis for the R10 four-task behavioral pilot."""

    model_config = ConfigDict(extra="forbid")

    report_version: Literal["r10.pilot_discrimination_report.1"] = "r10.pilot_discrimination_report.1"
    records_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    behavioral_result_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    bindings_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    input_complete: bool
    task_diagnoses: list[R10PilotTaskDiscriminationDiagnosisV1] = Field(min_length=4, max_length=4)
    cleanly_discriminative_count: int = Field(ge=0, le=4)
    cohort_decision: Literal["scale_discussion_ready", "compiler_revision_candidate"]
    compiler_signals: list[CompilerDiagnosticSignal] = Field(default_factory=list)
    professional_validity: Literal["provisional"] = "provisional"


class TeacherAnchorCheckV1(BaseModel):
    """A small, deterministic check for teacher-side facts that can be calculated.

    It deliberately covers only arithmetic and chronological anchors.  It is not
    an attempt to encode professional judgment; the decision matrix continues to
    own that judgment.
    """

    model_config = ConfigDict(extra="forbid")

    check_id: str = Field(min_length=1)
    decision_id: str = Field(min_length=1)
    kind: TeacherAnchorKind
    relation: TeacherAnchorRelation
    source_paths: list[str] = Field(min_length=1)
    numerator: Decimal | None = None
    denominator: Decimal | None = None
    threshold: Decimal | None = None
    left: Decimal | None = None
    right: Decimal | None = None
    earlier: date | None = None
    later: date | None = None

    @model_validator(mode="after")
    def require_operands_for_kind(self) -> "TeacherAnchorCheckV1":
        if self.kind == "percentage_threshold":
            if self.numerator is None or self.denominator is None or self.threshold is None or self.denominator <= 0:
                raise ValueError("teacher_anchor_percentage_operands_missing")
            if self.relation not in {"within", "exceeds"}:
                raise ValueError("teacher_anchor_percentage_relation_invalid")
        elif self.kind in {"sum", "quantity_delta"}:
            if self.left is None or self.right is None:
                raise ValueError("teacher_anchor_numeric_operands_missing")
            if self.relation != "equals":
                raise ValueError("teacher_anchor_numeric_relation_invalid")
        elif self.kind == "date_order":
            if self.earlier is None or self.later is None or self.relation not in {"before", "after"}:
                raise ValueError("teacher_anchor_date_operands_missing")
        return self


class TeacherAnchorFindingV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    check_id: str
    decision_id: str
    passed: bool
    observed: str
    expected_relation: TeacherAnchorRelation
    source_paths: list[str]


class TeacherAnchorAuditV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    audit_version: Literal["r10.teacher_anchor_audit.1"] = "r10.teacher_anchor_audit.1"
    decision: Literal["pass", "blocked"]
    findings: list[TeacherAnchorFindingV1]


def audit_teacher_anchors(checks: list[TeacherAnchorCheckV1]) -> TeacherAnchorAuditV1:
    """Evaluate declared numerical/date anchors without interpreting their policy meaning."""
    findings: list[TeacherAnchorFindingV1] = []
    for check in checks:
        if check.kind == "percentage_threshold":
            actual = check.numerator / check.denominator * Decimal("100")  # type: ignore[operator]
            passed = actual <= check.threshold if check.relation == "within" else actual > check.threshold  # type: ignore[operator]
            observed = f"{check.numerator}/{check.denominator}={actual}%"
        elif check.kind in {"sum", "quantity_delta"}:
            passed = check.left == check.right
            observed = f"{check.left}={check.right}"
        else:
            actual_before = check.earlier < check.later  # type: ignore[operator]
            passed = actual_before if check.relation == "before" else not actual_before
            observed = f"{check.earlier.isoformat()} {'<' if actual_before else '>='} {check.later.isoformat()}"  # type: ignore[union-attr]
        findings.append(TeacherAnchorFindingV1(
            check_id=check.check_id, decision_id=check.decision_id, passed=passed,
            observed=observed, expected_relation=check.relation, source_paths=check.source_paths,
        ))
    return TeacherAnchorAuditV1(
        decision="pass" if all(item.passed for item in findings) else "blocked", findings=findings,
    )


def binding_from_task(task_root: Path, *, domain: Literal["audit_compliance", "procurement_operations"]) -> R10PilotTaskBindingV1:
    manifest = json.loads((task_root / "task_manifest.json").read_text(encoding="utf-8"))
    contract = DeliverableContractV1.model_validate_json((task_root / "deliverable_contract.json").read_text(encoding="utf-8"))
    if len(contract.deliverables) != 1:
        raise ValueError("r10_behavioral_requires_single_delivery")
    expected = contract.deliverables[0].relative_path
    suffix = Path(expected).suffix.lower()
    if suffix not in {".xlsx", ".docx"}:
        raise ValueError("r10_behavioral_delivery_format_not_supported")
    return R10PilotTaskBindingV1(
        task_id=str(manifest["task_id"]),
        domain=domain,
        package_root=str(task_root),
        package_tree_sha256=tree_sha256(task_root),
        candidate_tree_sha256=tree_sha256(task_root / "reference_files"),
        task_compilation_sha256=str(manifest["task_compilation_sha256"]),
        deliverable_contract_sha256=sha256_file(task_root / "deliverable_contract.json"),
        expected_delivery=expected,
    )


def inspect_delivery(
    workspace: Path,
    *,
    expected: str,
    input_hashes: set[str],
    verify_docx_with_office: bool = True,
    libreoffice_command: str = "libreoffice",
    office_openable_override: bool | None = None,
) -> DeliveryInspectionV1:
    relative = Path(expected)
    if relative.is_absolute() or ".." in relative.parts or relative.parts[:1] != ("deliverable_files",):
        raise ValueError("r10_behavioral_delivery_path_invalid")
    suffix = relative.suffix.casefold()
    if suffix not in {".xlsx", ".docx"}:
        raise ValueError("r10_behavioral_delivery_format_not_supported")
    path = workspace / relative
    exists = path.is_file()
    nonempty = bool(exists and path.stat().st_size > 0)
    digest = sha256_file(path) if nonempty else None
    package_valid = False
    office_openable = False
    has_content = False
    failure: str | None = None
    if not exists:
        failure = "expected_delivery_missing"
    elif not nonempty:
        failure = "expected_delivery_empty"
    elif suffix == ".xlsx":
        try:
            workbook = load_workbook(path, data_only=False, read_only=True)
            package_valid = True
            has_content = any(
                cell.value not in (None, "")
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            )
            workbook.close()
            office_openable = package_valid
            if not has_content:
                failure = "xlsx_has_no_content"
        except Exception:
            failure = "xlsx_not_openable"
    else:
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
                package_valid = "[Content_Types].xml" in names and "word/document.xml" in names
                if package_valid:
                    document = archive.read("word/document.xml")
                    has_content = b"<w:t" in document or b"<w:tbl" in document
            if not package_valid:
                failure = "docx_not_ooxml"
        except (OSError, zipfile.BadZipFile, KeyError):
            failure = "docx_not_ooxml"
        if package_valid and office_openable_override is not None:
            office_openable = office_openable_override
            if not office_openable:
                failure = "docx_not_openable_by_libreoffice"
        elif package_valid and verify_docx_with_office:
            with tempfile.TemporaryDirectory(prefix="r10-docx-open-") as converted:
                completed = subprocess.run(
                    [libreoffice_command, "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", converted, str(path)],
                    capture_output=True,
                    timeout=90,
                    check=False,
                )
                output = Path(converted) / f"{path.stem}.pdf"
                office_openable = completed.returncode == 0 and output.is_file() and output.stat().st_size > 0
            if not office_openable:
                failure = "docx_not_openable_by_libreoffice"
        elif package_valid:
            office_openable = True
        if package_valid and not has_content:
            failure = "docx_has_no_content"
    differs = bool(digest and digest not in input_hashes)
    if package_valid and office_openable and has_content and not differs:
        failure = "delivery_copies_input"
    valid = bool(exists and nonempty and package_valid and office_openable and has_content and differs)
    return DeliveryInspectionV1(
        relative_path=relative.as_posix(), format=suffix.removeprefix("."), exists=exists,
        nonempty=nonempty, package_valid=package_valid, office_openable=office_openable,
        has_content=has_content, differs_from_inputs=differs, sha256=digest, valid=valid,
        first_failure=None if valid else failure,
    )


def finalize_judge_review(
    *,
    judge_id: JudgeId,
    draft: R10JudgeDraftV1,
    rubric: TaskSpecificRubricV1,
) -> R10JudgeReviewV1:
    ratings = {item.decision_id: item for item in draft.assessments}
    rubric_ids = [item.decision_id for item in rubric.criteria]
    if draft.task_id.strip() == "" or set(ratings) != set(rubric_ids):
        raise ValueError("r10_judge_decision_matrix_mismatch")
    value = {"met": 1.0, "partial": 0.5, "not_met": 0.0}
    weighted = sum(item.weight * value[ratings[item.decision_id].rating] for item in rubric.criteria)
    return R10JudgeReviewV1(
        **draft.model_dump(mode="json"), judge_id=judge_id,
        weighted_score=round(weighted, 6),
        major_defect=any(item.major_error for item in draft.assessments),
    )


def aggregate_behavioral_result(records: list[R10ModelTaskResultV1]) -> R10BehavioralResultV1:
    expected_solvers: tuple[SolverStackId, SolverStackId] = (
        "gpt-5.6-sol@chatgpt_codex", "deepseek-v4-pro@official_opencode"
    )
    by_solver = {solver: {item.task_id: item for item in records if item.solver_id == solver} for solver in expected_solvers}
    task_ids = sorted({item.task_id for item in records})
    if len(task_ids) != 4 or any(set(values) != set(task_ids) for values in by_solver.values()):
        raise ValueError("r10_behavioral_records_incomplete")
    valid_counts = {solver: sum(item.delivery_valid for item in values.values()) for solver, values in by_solver.items()}
    discrimination: list[R10TaskDiscriminationV1] = []
    common = 0
    insufficient_by_decision: dict[str, int] = {}
    major_error_by_decision: dict[str, int] = {}
    for task_id in task_ids:
        left, right = (by_solver[solver][task_id] for solver in expected_solvers)
        dual = len(left.reviews) == len(right.reviews) == 2 and left.delivery_valid and right.delivery_valid
        left_score = sum(item.weighted_score for item in left.reviews) / 2 if len(left.reviews) == 2 else None
        right_score = sum(item.weighted_score for item in right.reviews) / 2 if len(right.reviews) == 2 else None
        left_major = any(item.major_defect for item in left.reviews)
        right_major = any(item.major_defect for item in right.reviews)
        if dual:
            common += 1
        for review in [*left.reviews, *right.reviews]:
            for assessment in review.assessments:
                if assessment.evidence_insufficient:
                    insufficient_by_decision[assessment.decision_id] = insufficient_by_decision.get(assessment.decision_id, 0) + 1
                if assessment.major_error:
                    major_error_by_decision[assessment.decision_id] = major_error_by_decision.get(assessment.decision_id, 0) + 1
        gap = abs(left_score - right_score) if dual and left_score is not None and right_score is not None else None
        differs = left.delivery_valid != right.delivery_valid or left_major != right_major
        discrimination.append(R10TaskDiscriminationV1(
            task_id=task_id, comparable=dual, composite_score_gap=gap,
            delivery_or_major_defect_differs=differs,
            explainable_difference=bool((gap is not None and gap >= 0.05) or differs),
        ))
    recurring = sorted(key for key, count in insufficient_by_decision.items() if count >= 4)
    recurring_major = sorted(key for key, count in major_error_by_decision.items() if count >= 4)
    if any(value < 3 for value in valid_counts.values()) or common < 3:
        decision = "behaviorally_inconclusive"
    elif recurring or recurring_major:
        decision = "task_design_revision_candidate"
    else:
        decision = "behaviorally_admitted"
    low_separation = sum(item.explainable_difference for item in discrimination) < 2
    return R10BehavioralResultV1(
        decision=decision, low_model_separation=low_separation, common_dual_graded_count=common,
        solver_valid_delivery_counts=valid_counts, task_discrimination=discrimination,
        recurring_insufficient_evidence_decisions=recurring,
        recurring_major_error_decisions=recurring_major,
    )


def diagnose_pilot_discrimination(
    records: list[R10ModelTaskResultV1],
    *,
    behavioral_result: R10BehavioralResultV1,
    bindings: list[R10PilotTaskBindingV1],
) -> R10PilotDiscriminationReportV1:
    """Diagnose model separation from one frozen R10 pilot result.

    This function is deliberately read-only.  It only checks that the supplied
    aggregate and four bindings still describe the supplied records, then
    classifies the observed evidence.  It does not inspect or mutate task
    packages, candidate files, teacher material, or provider state.
    """
    expected_solvers: tuple[SolverStackId, SolverStackId] = (
        "gpt-5.6-sol@chatgpt_codex", "deepseek-v4-pro@official_opencode"
    )
    if len(bindings) != 4 or len({item.task_id for item in bindings}) != 4:
        raise ValueError("r10_pilot_discrimination_bindings_invalid")
    binding_ids = {item.task_id for item in bindings}
    record_ids = {item.task_id for item in records}
    if record_ids != binding_ids:
        raise ValueError("r10_pilot_discrimination_record_binding_drift")
    try:
        recomputed = aggregate_behavioral_result(records)
    except ValueError as error:
        # A partial record set is still a useful diagnosis, but it can never
        # justify scale discussion.  Preserve the per-task `incomplete` state.
        aggregate_failure = str(error)
    else:
        aggregate_failure = None
        if recomputed.model_dump(mode="json") != behavioral_result.model_dump(mode="json"):
            raise ValueError("r10_pilot_discrimination_behavioral_result_drift")

    by_task: dict[str, dict[SolverStackId, R10ModelTaskResultV1]] = {
        task_id: {} for task_id in sorted(binding_ids)
    }
    for record in records:
        values = by_task[record.task_id]
        if record.solver_id in values:
            raise ValueError("r10_pilot_discrimination_duplicate_solver_record")
        values[record.solver_id] = record

    diagnoses: list[R10PilotTaskDiscriminationDiagnosisV1] = []
    for task_id in sorted(binding_ids):
        per_solver = by_task[task_id]
        composites: dict[SolverStackId, float | None] = {}
        score_gaps: dict[SolverStackId, float | None] = {}
        major_disagreement: dict[SolverStackId, bool] = {}
        incomplete_reason: str | None = aggregate_failure
        majors: dict[SolverStackId, bool | None] = {}
        for solver in expected_solvers:
            record = per_solver.get(solver)
            if record is None:
                composites[solver] = None
                score_gaps[solver] = None
                major_disagreement[solver] = False
                majors[solver] = None
                incomplete_reason = incomplete_reason or "solver_record_missing"
                continue
            if not record.delivery_valid:
                incomplete_reason = incomplete_reason or "delivery_not_valid"
            if len(record.reviews) != 2:
                composites[solver] = None
                score_gaps[solver] = None
                major_disagreement[solver] = False
                majors[solver] = None
                incomplete_reason = incomplete_reason or "judge_reviews_incomplete"
                continue
            judge_ids = {review.judge_id for review in record.reviews}
            if judge_ids != set(expected_solvers):
                composites[solver] = None
                score_gaps[solver] = None
                major_disagreement[solver] = False
                majors[solver] = None
                incomplete_reason = incomplete_reason or "judge_identity_incomplete"
                continue
            scores = [review.weighted_score for review in record.reviews]
            composites[solver] = sum(scores) / len(scores)
            score_gaps[solver] = abs(scores[0] - scores[1])
            major_disagreement[solver] = len({review.major_defect for review in record.reviews}) != 1
            majors[solver] = any(review.major_defect for review in record.reviews)

        left, right = expected_solvers
        left_score, right_score = composites[left], composites[right]
        score_gap = abs(left_score - right_score) if left_score is not None and right_score is not None else None
        delivery_or_major_differs = bool(
            left in per_solver and right in per_solver
            and (
                per_solver[left].delivery_valid != per_solver[right].delivery_valid
                or majors[left] != majors[right]
            )
        )
        judge_ambiguous = any(
            (gap is not None and gap >= 0.125) or major_disagreement[solver]
            for solver, gap in score_gaps.items()
        )
        signals: list[CompilerDiagnosticSignal] = []
        if incomplete_reason is not None:
            classification: PilotDiscriminationClassification = "incomplete"
        elif left_score is not None and right_score is not None and left_score >= 0.95 and right_score >= 0.95 and not any(majors.values()):
            classification = "saturated"
            signals = ["score_saturation", "decision_points_too_explicit"]
        elif score_gap is not None and score_gap < 0.05 and not delivery_or_major_differs:
            classification = "near_tie"
            signals = ["evidence_tension_insufficient"]
        elif judge_ambiguous:
            classification = "judge_ambiguous"
            signals = ["judge_boundary_ambiguous"]
        else:
            classification = "cleanly_discriminative"
        diagnoses.append(R10PilotTaskDiscriminationDiagnosisV1(
            task_id=task_id, solver_composite_scores=composites,
            composite_score_gap=score_gap,
            delivery_or_major_defect_differs=delivery_or_major_differs,
            judge_score_gaps=score_gaps,
            judge_major_defect_disagreement=major_disagreement,
            classification=classification, compiler_signals=signals,
            first_failure=incomplete_reason,
        ))

    clean_count = sum(item.classification == "cleanly_discriminative" for item in diagnoses)
    all_signals = sorted({signal for item in diagnoses for signal in item.compiler_signals})
    return R10PilotDiscriminationReportV1(
        records_sha256=sha256_json([item.model_dump(mode="json") for item in records]),
        behavioral_result_sha256=sha256_json(behavioral_result.model_dump(mode="json")),
        bindings_sha256=sha256_json([item.model_dump(mode="json") for item in bindings]),
        input_complete=all(item.classification != "incomplete" for item in diagnoses),
        task_diagnoses=diagnoses,
        cleanly_discriminative_count=clean_count,
        cohort_decision="scale_discussion_ready" if clean_count >= 2 else "compiler_revision_candidate",
        compiler_signals=all_signals,
    )
