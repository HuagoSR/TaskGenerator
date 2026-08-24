from __future__ import annotations

import hashlib
import json
import os
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Type
from xml.etree import ElementTree

from pydantic import BaseModel

from task_generator.evaluation.semantic_validity import (
    CandidateBlindReview,
    SecondarySemanticReview,
    SemanticFinding,
    TeacherRubricReview,
)
from task_generator.substrate.skill_extractor import ProviderConfig, build_deepseek_config, build_tuzi_config, load_env_file
from task_generator.evaluation.semantic_secondary_cost import SecondaryCostLedgerManager
from task_generator.evaluation.validity_utility import (
    CandidateBlindRealityReviewV2,
    RubricFocusRealityReviewV2,
    RubricFocusRealityReviewV3,
    RubricFocusRealityReviewV4,
)
from task_generator.core.external_model_policy import enforce_external_model_policy


class SemanticReviewExecutionError(RuntimeError):
    def __init__(self, message: str, *, retry_eligible: bool = False, failure_code: str = "provider_failure") -> None:
        super().__init__(message)
        self.retry_eligible = retry_eligible
        self.failure_code = failure_code


class SemanticReviewExecutor:
    """Execute bounded JSON-only semantic reviews without exposing secrets in artifacts."""

    def __init__(
        self,
        config: ProviderConfig,
        max_tokens: int = 8000,
        cost_ledger: Optional[SecondaryCostLedgerManager] = None,
        input_token_hard_limit: int = 20_000,
        max_retries: int = 2,
    ) -> None:
        self.config = config
        enforce_external_model_policy(config.provider_name, config.model)
        self.max_tokens = max_tokens
        self.cost_ledger = cost_ledger
        self.input_token_hard_limit = input_token_hard_limit
        self.max_retries = max_retries
        self.last_diagnostics: Dict[str, Any] = {}
        self.last_raw_response_content: str = ""

    def review_reality_candidate_blind(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> CandidateBlindRealityReviewV2:
        instructions = (
            "Act as an independent professional task reviewer. Review only the "
            "candidate-visible task, deliverable contract, and reference contents. "
            "Return exactly five dimensions: role_realism, information_sufficiency, "
            "natural_difficulty, professional_judgment, and deliverable_realism. "
            "Use concrete evidence locators. Do not infer generation route, teacher "
            "truth, hidden rubric, or repairs."
        )
        return self._call(
            instructions,
            payload,
            CandidateBlindRealityReviewV2,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )

    def review_reality_rubric_focus(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> RubricFocusRealityReviewV2:
        instructions = (
            "Act as an independent rubric auditor. Evaluate only whether the supplied "
            "rubric and bindings focus on real professional capability, avoid duplicate "
            "scoring, and avoid rewarding accidental difficulty. Return only the "
            "rubric_focus dimension with concrete evidence locators. Do not infer or "
            "mention generation route or model."
        )
        return self._call(
            instructions,
            payload,
            RubricFocusRealityReviewV2,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )

    def review_reality_rubric_focus_v3(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> RubricFocusRealityReviewV3:
        instructions = (
            "Act as an independent rubric scoring-authority auditor. Only entries in "
            "rubric_plan.criteria are final weighted scoring criteria. Entries in the "
            "rubric_binding_plan are trace-only annotations because final_weights_assigned "
            "is false and scoring_authority says not_final_weight; never count them as extra "
            "criteria. Assess every unordered pair of the seven final criteria exactly once. "
            "Shared skill IDs, capability IDs, or evidence alone do not establish duplicate "
            "scoring when observable behaviors and independent failure signals are distinct. "
            "Use shared_evidence_distinct_behavior exactly for a pair marked as sharing "
            "evidence by the frozen scoring-authority audit when its behavior and failure "
            "signal remain distinct; otherwise use distinct for a non-risk pair. "
            "Use potential_duplicate or duplicate only when the scored behaviors or failure "
            "conditions materially overlap, and cite both criterion IDs. Keep route and "
            "generator identity out of the review. Return the criterion IDs in lexical order, "
            "and each pair with criterion_a < criterion_b."
        )
        return self._call(
            instructions,
            payload,
            RubricFocusRealityReviewV3,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )

    def estimate_reality_rubric_focus_v3_input_tokens(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> int:
        instructions = (
            "Act as an independent rubric scoring-authority auditor. Only entries in "
            "rubric_plan.criteria are final weighted scoring criteria. Entries in the "
            "rubric_binding_plan are trace-only annotations because final_weights_assigned "
            "is false and scoring_authority says not_final_weight; never count them as extra "
            "criteria. Assess every unordered pair of the seven final criteria exactly once. "
            "Shared skill IDs, capability IDs, or evidence alone do not establish duplicate "
            "scoring when observable behaviors and independent failure signals are distinct. "
            "Use shared_evidence_distinct_behavior exactly for a pair marked as sharing "
            "evidence by the frozen scoring-authority audit when its behavior and failure "
            "signal remain distinct; otherwise use distinct for a non-risk pair. "
            "Use potential_duplicate or duplicate only when the scored behaviors or failure "
            "conditions materially overlap, and cite both criterion IDs. Keep route and "
            "generator identity out of the review. Return the criterion IDs in lexical order, "
            "and each pair with criterion_a < criterion_b."
        )
        messages = self._request_messages(
            instructions,
            payload,
            RubricFocusRealityReviewV3,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )
        return self._estimated_message_tokens(messages)

    def review_reality_rubric_focus_v4(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> RubricFocusRealityReviewV4:
        instructions = self._rubric_focus_v4_instructions()
        return self._call(
            instructions,
            payload,
            RubricFocusRealityReviewV4,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )

    def estimate_reality_rubric_focus_v4_input_tokens(
        self, payload: Dict[str, Any], *, format_feedback: Optional[str] = None
    ) -> int:
        messages = self._request_messages(
            self._rubric_focus_v4_instructions(),
            payload,
            RubricFocusRealityReviewV4,
            {"case_id": payload["case_id"]},
            format_feedback=format_feedback,
        )
        return self._estimated_message_tokens(messages)

    @staticmethod
    def _rubric_focus_v4_instructions() -> str:
        return (
            "Act as an independent rubric scoring-authority auditor. Only entries in "
            "rubric_plan.criteria are final weighted scoring criteria; rubric_binding_plan "
            "entries are trace-only and non-scoring. Classify all 21 unordered pairs once. "
            "For non-risk pairs return only the compact assessment: use "
            "shared_evidence_distinct_behavior exactly when the frozen audit marks shared "
            "evidence, otherwise use distinct. Shared skill, capability, or evidence alone "
            "never establishes duplicate scoring. Use potential_duplicate or duplicate only "
            "for material overlap in scored behavior or failure conditions. Detailed "
            "rationale and evidence locators belong only in risk_findings, one finding for "
            "each risky pair, citing both criterion IDs. Keep the summary under 80 words. "
            "Do not mention route, generator identity, or provider."
        )

    def review_blind(self, package: Dict[str, Any]) -> CandidateBlindReview:
        enriched = dict(package)
        enriched["reference_files"] = [
            {key: value for key, value in item.items() if key != "path"}
            for item in package.get("reference_files") or []
        ]
        enriched["reference_contents"] = [
            self._reference_content(Path(item["path"])) for item in package.get("reference_files") or []
        ]
        instructions = (
            "Act as an adversarial candidate-blind task auditor. Use only the candidate-visible prompt, "
            "deliverable contract, and reference contents. Determine whether every requirement is supported, "
            "uniquely answerable, and free of hidden assumptions. Return JSON matching CandidateBlindReview. "
            "Use only the documented semantic finding codes. Do not infer teacher truth or repair files."
        )
        review = self._call(
            instructions,
            enriched,
            CandidateBlindReview,
            {
                "task_id": package.get("task_id"),
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "prompt_package_sha256": hashlib.sha256(json.dumps(package, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
                "status": "completed",
            },
        )
        expected_ids = {str(item.get("requirement_id")) for item in package.get("semantic_requirements") or []}
        returned_ids = {item.requirement_id for item in review.requirement_reviews}
        if not expected_ids or returned_ids != expected_ids:
            self.last_diagnostics.update({
                "requirement_coverage_valid": False,
                "expected_requirement_count": len(expected_ids),
                "returned_requirement_count": len(returned_ids),
                "missing_requirement_count": len(expected_ids - returned_ids),
                "unexpected_requirement_count": len(returned_ids - expected_ids),
            })
            raise SemanticReviewExecutionError("Provider contract violation: requirement review coverage mismatch.")
        self._clear_provider_corroboration(review)
        return review

    def review_teacher(self, package: Dict[str, Any]) -> TeacherRubricReview:
        expanded = {
            "package_version": package.get("package_version"),
            "task_id": package.get("task_id"),
            "blind_review_sha256": (package.get("blind_review") or {}).get("sha256"),
            "semantic_contract_sha256": (package.get("semantic_contract") or {}).get("sha256"),
        }
        expanded["blind_review_payload"] = self._read_json(Path(package["blind_review"]["path"]))
        expanded["semantic_contract_payload"] = self._read_json(Path(package["semantic_contract"]["path"]))
        expanded["teacher_artifact_payloads"] = {
            name: self._read_or_summarize(Path(item["path"]))
            for name, item in (package.get("teacher_artifacts") or {}).items()
        }
        instructions = (
            "Audit teacher truth, GoldenRun, and rubric against the already-frozen candidate-blind review. "
            "Find unsupported teacher assumptions, uncovered requirements, irrelevant criteria, and weight imbalance. "
            "Return JSON matching TeacherRubricReview. Findings and repair proposals are advisory only."
        )
        review = self._call(
            instructions,
            expanded,
            TeacherRubricReview,
            {
                "task_id": package.get("task_id"),
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "blind_review_sha256": (package.get("blind_review") or {}).get("sha256"),
                "status": "completed",
            },
        )
        self._clear_provider_corroboration(review)
        return review

    def review_secondary(
        self,
        package: Dict[str, Any],
        primary_findings: List[SemanticFinding],
        review_scope: str,
    ) -> SecondarySemanticReview:
        material = [item for item in primary_findings if item.severity == "blocking"][:6]
        if review_scope == "candidate_blind":
            requirement_ids = {item.requirement_id for item in material if item.requirement_id}
            evidence = {
                "requirements": [item for item in package.get("semantic_requirements") or [] if not requirement_ids or item.get("requirement_id") in requirement_ids],
                "deliverable_contract": package.get("deliverable_contract"),
                "relevant_candidate_evidence": [
                    self._compact_reference_content(Path(item["path"]))
                    for item in package.get("reference_files") or []
                ],
            }
        elif review_scope == "teacher_rubric":
            contract = self._read_json(Path(package["semantic_contract"]["path"]))
            claim_ids = {item.claim_id for item in material if item.claim_id}
            requirement_ids = {item.requirement_id for item in material if item.requirement_id}
            claims = [item for item in contract.get("claims") or [] if not claim_ids or item.get("claim_id") in claim_ids]
            if requirement_ids and not claim_ids:
                claims = [item for item in claims if item.get("requirement_id") in requirement_ids]
            teacher_payloads = {
                name: self._read_json(Path(item["path"]))
                for name, item in (package.get("teacher_artifacts") or {}).items()
                if name in {"golden_run", "teacher_truth", "deterministic_answer_key", "rubric"}
                and Path(item["path"]).suffix.lower() == ".json"
            }
            evidence = {
                "resolved_claims": claims,
                "validator_results": contract.get("validator_results") or contract.get("validators") or [],
                "rubric_bindings": contract.get("rubric_bindings") or [],
                "relevant_teacher_artifacts": teacher_payloads,
            }
        else:
            raise SemanticReviewExecutionError(f"Unknown secondary review scope: {review_scope}")
        payload = {
            "task_id": package.get("task_id"),
            "review_scope": review_scope,
            "primary_material_findings": [
                {
                    "finding_code": item.finding_code,
                    "requirement_id": item.requirement_id,
                    "claim_id": item.claim_id,
                    "message": item.message,
                    "evidence_locators": item.evidence_locators,
                }
                for item in material
            ],
            "evidence": evidence,
        }
        instructions = (
            "Act as an independent secondary adjudicator. For each material issue family supported by the evidence, "
            "return one compact decision. Do not recreate the primary review, do not repair artifacts, and do not "
            "treat style suggestions as material. Use only the enumerated finding families."
        )
        self._enforce_secondary_prompt_budget(payload)
        return self._call(
            instructions,
            payload,
            SecondarySemanticReview,
            {
                "task_id": package.get("task_id"),
                "review_scope": review_scope,
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "status": "completed",
            },
            cost_scope=review_scope,
        )

    def _clear_provider_corroboration(self, review: BaseModel) -> None:
        findings = list(getattr(review, "findings", []) or [])
        for requirement in getattr(review, "requirement_reviews", []) or []:
            findings.extend(requirement.findings)
        for finding in findings:
            finding.deterministic_corroboration = False

    def _call(
        self,
        system_prompt: str,
        payload: Dict[str, Any],
        model_type: Type[BaseModel],
        defaults: Dict[str, Any],
        cost_scope: Optional[str] = None,
        format_feedback: Optional[str] = None,
    ):
        started = time.monotonic()
        messages = self._request_messages(
            system_prompt,
            payload,
            model_type,
            defaults,
            format_feedback=format_feedback,
        )
        estimated_tokens = self._estimated_message_tokens(messages)
        self.last_diagnostics = {
            "estimated_input_tokens": estimated_tokens,
            "input_token_hard_limit": self.input_token_hard_limit,
        }
        if estimated_tokens > self.input_token_hard_limit:
            raise SemanticReviewExecutionError(
                "Provider message input token ceiling exceeded.",
                retry_eligible=False,
                failure_code="input_token_ceiling",
            )
        try:
            from openai import OpenAI
        except Exception as exc:  # pragma: no cover
            raise SemanticReviewExecutionError(f"OpenAI SDK unavailable: {type(exc).__name__}") from exc
        client = OpenAI(
            api_key=self.config.api_key,
            base_url=self.config.base_url,
            timeout=self.config.timeout_seconds,
            max_retries=self.max_retries,
        )
        reserved = self.cost_ledger.reserve(self.input_token_hard_limit, self.max_tokens) if self.cost_ledger else 0.0
        request_args: Dict[str, Any] = {
            "model": self.config.model,
            "messages": messages,
            "max_tokens": self.max_tokens,
            "response_format": {"type": "json_object"},
        }
        if self.config.provider_name == "deepseek":
            if self.config.reasoning_mode == "disabled":
                request_args["extra_body"] = {
                    "thinking": {"type": "disabled"}
                }
            elif self.config.reasoning_mode == "high":
                request_args["reasoning_effort"] = "high"
                request_args["extra_body"] = {
                    "thinking": {"type": "enabled"}
                }
            else:
                raise SemanticReviewExecutionError(
                    "Unsupported DeepSeek reasoning mode.",
                    retry_eligible=False,
                    failure_code="provider_configuration_failure",
                )
        else:
            request_args["temperature"] = 0
        try:
            response = client.chat.completions.create(**request_args)
        except Exception as exc:
            status = getattr(exc, "status_code", None)
            eligible = status in {408, 429} or (isinstance(status, int) and 500 <= status <= 599) or isinstance(exc, (TimeoutError, ConnectionError))
            code = (
                f"transport_http_{status}" if status in {408, 429}
                else "transport_http_5xx" if isinstance(status, int) and 500 <= status <= 599
                else "transport_timeout" if isinstance(exc, (TimeoutError, ConnectionError))
                else "transport_failure"
            )
            raise SemanticReviewExecutionError(
                f"Provider request failed: {type(exc).__name__}",
                retry_eligible=eligible,
                failure_code=code,
            ) from exc
        content = response.choices[0].message.content or ""
        self.last_raw_response_content = content
        finish_reason = str(getattr(response.choices[0], "finish_reason", ""))
        usage = getattr(response, "usage", None)
        self.last_diagnostics = {
            "provider": self.config.provider_name,
            "model": self.config.model,
            "finish_reason": finish_reason,
            "response_char_count": len(content),
            "response_sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "prompt_tokens": getattr(usage, "prompt_tokens", None),
            "completion_tokens": getattr(usage, "completion_tokens", None),
            "total_tokens": getattr(usage, "total_tokens", None),
            "duration_seconds": round(time.monotonic() - started, 3),
            "likely_truncated": finish_reason.lower() == "length",
            "key_slot": "backup" if self.cost_ledger else None,
            "estimated_input_tokens": estimated_tokens,
            "input_token_hard_limit": self.input_token_hard_limit,
        }
        if self.cost_ledger and cost_scope:
            self.cost_ledger.record(
                scope=cost_scope,
                attempt=1,
                prompt_tokens=int(getattr(usage, "prompt_tokens", 0) or 0),
                completion_tokens=int(getattr(usage, "completion_tokens", 0) or 0),
                reserved_rmb=reserved,
                status="completed" if content and finish_reason.lower() != "length" else "failed",
            )
        if not content:
            raise SemanticReviewExecutionError("Provider returned empty semantic review output.", retry_eligible=True, failure_code="empty_content")
        if self.last_diagnostics["likely_truncated"]:
            raise SemanticReviewExecutionError("Provider truncated semantic review output.", retry_eligible=True, failure_code="truncated_output")
        try:
            parsed = json.loads(content)
            if not isinstance(parsed, dict):
                raise ValueError("Semantic review response must be a JSON object.")
            parsed.update(defaults)
            return model_type.model_validate(parsed)
        except Exception as exc:
            raise SemanticReviewExecutionError(
                f"Provider contract violation: {type(exc).__name__}",
                retry_eligible=True,
                failure_code="invalid_json" if isinstance(exc, json.JSONDecodeError) else "schema_contract_failure",
            ) from exc

    def _request_messages(
        self,
        system_prompt: str,
        payload: Dict[str, Any],
        model_type: Type[BaseModel],
        defaults: Dict[str, Any],
        *,
        format_feedback: Optional[str] = None,
    ) -> list[Dict[str, str]]:
        schema = model_type.model_json_schema()
        schema_example = self._json_example(model_type, defaults)
        return [
                {
                    "role": "system",
                    "content": system_prompt
                    + " Output exactly one JSON object matching this JSON Schema. Do not add wrapper keys or prose: "
                    + json.dumps(schema, ensure_ascii=False)
                    + " Example JSON shape: "
                    + json.dumps(schema_example, ensure_ascii=False)
                    + ((" Previous attempt failed only this output contract check: " + format_feedback) if format_feedback else ""),
                },
                {
                    "role": "user",
                    "content": json.dumps(
                        payload, ensure_ascii=False, default=str
                    ),
                },
            ]

    @staticmethod
    def _estimated_message_tokens(messages: list[Dict[str, str]]) -> int:
        serialized = json.dumps(messages, ensure_ascii=False, separators=(",", ":"))
        return max(1, (len(serialized) + 2) // 3)

    def _json_example(self, model_type: Type[BaseModel], defaults: Dict[str, Any]) -> Dict[str, Any]:
        if model_type is CandidateBlindRealityReviewV2:
            names = ["role_realism", "information_sufficiency", "natural_difficulty", "professional_judgment", "deliverable_realism"]
            return {
                "case_id": defaults.get("case_id", "case_id"),
                "dimensions": [
                    {"dimension": name, "decision": "revise", "rationale": "Concrete evidence-based rationale.", "evidence_locators": ["file:sheet!A1"]}
                    for name in names
                ],
            }
        if model_type is RubricFocusRealityReviewV2:
            return {
                "case_id": defaults.get("case_id", "case_id"),
                "dimension": {"dimension": "rubric_focus", "decision": "revise", "rationale": "Concrete evidence-based rationale.", "evidence_locators": ["rubric:criterion_id"]},
            }
        if model_type is RubricFocusRealityReviewV3:
            criterion_ids = sorted(
                str(item) for item in defaults.get("criterion_ids", [
                    "criterion_evidence_traceability",
                    "criterion_exception_handling",
                    "criterion_factual_accuracy",
                    "criterion_method_process",
                    "criterion_professional_expression",
                    "criterion_reproducibility",
                    "criterion_structural_usability",
                ])
            )
            pairs = []
            for index, left in enumerate(criterion_ids):
                for right in criterion_ids[index + 1 :]:
                    pairs.append({
                        "criterion_a": left,
                        "criterion_b": right,
                        "assessment": "distinct",
                        "rationale": "The observable behaviors and independent failure signals are distinct.",
                        "evidence_locators": [f"rubric:{left}", f"rubric:{right}"],
                    })
            return {
                "case_id": defaults.get("case_id", "case_id"),
                "final_scoring_criterion_ids": criterion_ids,
                "final_scoring_criteria_count": 7,
                "annotation_bindings_recognized_non_scoring": True,
                "pair_assessments": pairs,
                "decision": "pass",
                "rationale": "All final scoring criteria assess distinct professional behaviors.",
            }
        if model_type is RubricFocusRealityReviewV4:
            criterion_ids = sorted([
                "criterion_evidence_traceability",
                "criterion_exception_handling",
                "criterion_factual_accuracy",
                "criterion_method_process",
                "criterion_professional_expression",
                "criterion_reproducibility",
                "criterion_structural_usability",
            ])
            pairs = [
                {
                    "criterion_a": left,
                    "criterion_b": right,
                    "assessment": "distinct",
                }
                for index, left in enumerate(criterion_ids)
                for right in criterion_ids[index + 1 :]
            ]
            return {
                "case_id": defaults.get("case_id", "case_id"),
                "final_scoring_criterion_ids": criterion_ids,
                "final_scoring_criteria_count": 7,
                "annotation_bindings_recognized_non_scoring": True,
                "pair_assessments": pairs,
                "risk_findings": [],
                "decision": "pass",
                "summary": "All final scoring criteria assess distinct professional behaviors.",
            }
        return defaults

    def _reference_content(self, path: Path) -> Dict[str, Any]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._xlsx_content(path)
        if suffix == ".docx":
            return self._docx_content(path)
        if suffix in {".md", ".txt", ".csv", ".json"}:
            return {"file_name": path.name, "kind": suffix[1:], "text": path.read_text(encoding="utf-8", errors="replace")[:50000]}
        return {"file_name": path.name, "kind": "unsupported_binary", "size_bytes": path.stat().st_size}

    def _compact_reference_content(self, path: Path) -> Dict[str, Any]:
        content = self._reference_content(path)
        if content.get("kind") == "xlsx":
            for sheet in content.get("sheets") or []:
                sheet["rows"] = (sheet.get("rows") or [])[:50]
        elif "text" in content:
            content["text"] = str(content["text"])[:8000]
        return content

    def _enforce_secondary_prompt_budget(self, payload: Dict[str, Any]) -> None:
        serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        estimated_tokens = (len(serialized) + 3) // 4
        self.last_diagnostics = {
            "secondary_payload_chars": len(serialized),
            "estimated_input_tokens": estimated_tokens,
            "input_token_hard_limit": self.input_token_hard_limit,
        }
        if estimated_tokens > self.input_token_hard_limit:
            raise SemanticReviewExecutionError("secondary_prompt_budget_exceeded")

    def _xlsx_content(self, path: Path) -> Dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover
            raise SemanticReviewExecutionError(f"openpyxl unavailable: {type(exc).__name__}") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                rows.append([value for value in row])
                if index >= 500:
                    break
            sheets.append({"sheet_name": sheet.title, "rows": rows})
        workbook.close()
        return {"file_name": path.name, "kind": "xlsx", "sheets": sheets}

    def _docx_content(self, path: Path) -> Dict[str, Any]:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
        return {"file_name": path.name, "kind": "docx", "text": "\n".join(texts)[:50000]}

    def _read_or_summarize(self, path: Path) -> Any:
        if path.suffix.lower() == ".json":
            return self._read_json(path)
        return self._reference_content(path)

    def _read_json(self, path: Path) -> Dict[str, Any]:
        return json.loads(path.read_text(encoding="utf-8"))


def deepseek_semantic_config(key_path: str | Path, timeout_seconds: int = 900) -> ProviderConfig:
    config = build_deepseek_config(key_path, "deepseek-v4-pro", timeout_seconds)
    if config is None:
        raise SemanticReviewExecutionError("DeepSeek semantic-review key is unavailable.")
    return config


def claude_semantic_config(
    env_path: str | Path,
    timeout_seconds: int = 900,
    *,
    allow_expensive_model: bool = False,
) -> ProviderConfig:
    enforce_external_model_policy("tuzi", "claude-sonnet-4-6")
    config = build_tuzi_config(env_path, "claude-sonnet-4-6", timeout_seconds)
    if config is not None:
        return config
    values = load_env_file(env_path)
    api_key = values.get("AGENT_API_KEY") or values.get("GRADER_API_KEY")
    base_url = values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model="claude-sonnet-4-6",
        timeout_seconds=timeout_seconds,
    )


def gpt54_semantic_config(env_path: str | Path, timeout_seconds: int = 900) -> ProviderConfig:
    enforce_external_model_policy("tuzi", "gpt-5.4-pro")
    config = build_tuzi_config(env_path, "gpt-5.4-pro", timeout_seconds)
    if config is not None:
        return config
    values = load_env_file(env_path)
    api_key = values.get("AGENT_API_KEY") or values.get("GRADER_API_KEY")
    base_url = values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model="gpt-5.4-pro",
        timeout_seconds=timeout_seconds,
    )


def tuzi_backup_semantic_config(
    env_path: str | Path,
    model: str = "gpt-5.6-sol",
    timeout_seconds: int = 900,
) -> ProviderConfig:
    enforce_external_model_policy("tuzi", model)
    values = load_env_file(env_path)
    api_key = values.get("OPENAI_API_KEY_BACKUP")
    base_url = (
        values.get("OPENAI_BASE_URL") or values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
        or os.environ.get("OPENAI_BASE_URL")
    )
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi backup semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model=model,
        timeout_seconds=timeout_seconds,
    )


def tuzi_semantic_config(
    env_path: str | Path,
    model: str,
    timeout_seconds: int = 900,
) -> ProviderConfig:
    """Load a Tuzi model without persisting credentials.

    Prefer the standard OPENAI-compatible names, while retaining the project's
    established AGENT/GRADER aliases used by the existing ignored environment.
    """
    enforce_external_model_policy("tuzi", model)
    config = build_tuzi_config(env_path, model, timeout_seconds)
    if config is not None:
        return config
    values = load_env_file(env_path)
    api_key = values.get("AGENT_API_KEY") or values.get("GRADER_API_KEY")
    base_url = values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
    if not api_key or not base_url:
        raise SemanticReviewExecutionError(
            "Tuzi semantic-review configuration is unavailable."
        )
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model=model,
        timeout_seconds=timeout_seconds,
    )
