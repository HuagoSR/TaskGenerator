"""Deterministic, read-only evidence extraction for XLSX grading.

This intentionally records workbook facts instead of deciding business rubric
items.  It is small enough to use as a grading attachment and keeps formula
locations intact, unlike an XML text flattening pass.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from pydantic import Field

from task_generator.core.scenario_first import ScenarioFirstModel


def _sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class SpreadsheetCellEvidenceV1(ScenarioFirstModel):
    address: str
    value: str | int | float | bool | None
    formula: str | None = None
    recalculated_value: str | int | float | bool | None = None
    number_format: str
    style_id: int


class SpreadsheetSheetEvidenceV1(ScenarioFirstModel):
    name: str
    used_range: str
    max_row: int
    max_column: int
    headers: list[str] = Field(default_factory=list)
    cells: list[SpreadsheetCellEvidenceV1] = Field(default_factory=list)
    merged_ranges: list[str] = Field(default_factory=list)
    hidden_rows: list[int] = Field(default_factory=list)
    hidden_columns: list[str] = Field(default_factory=list)
    freeze_panes: str | None = None
    column_widths: dict[str, float | None] = Field(default_factory=dict)
    style_counts: dict[str, int] = Field(default_factory=dict)


class SpreadsheetFormulaDiagnosticV1(ScenarioFirstModel):
    kind: str
    locations: list[str] = Field(min_length=1)
    detail: str


class SpreadsheetEvidenceV1(ScenarioFirstModel):
    evidence_version: str = "r10.spreadsheet_evidence.1"
    relative_path: str
    original_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    recalculated_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    recalculation_status: str
    sheets: list[SpreadsheetSheetEvidenceV1] = Field(min_length=1)
    diagnostics: list[SpreadsheetFormulaDiagnosticV1] = Field(default_factory=list)


def _safe_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _formula_refs(formula: str, sheet: str) -> set[tuple[str, str]]:
    """Return cell dependencies, expanding small same-workbook ranges.

    Unsupported/external references are deliberately omitted: extraction stays
    factual and the original formula remains available to the judge.
    """
    refs: set[tuple[str, str]] = set()
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return refs
    for token in tokens:
        if token.subtype != "RANGE":
            continue
        raw = token.value.replace("$", "")
        if "[" in raw or "!" in raw and raw.count("!") > 1:
            continue
        target_sheet, target = sheet, raw
        if "!" in raw:
            prefix, target = raw.rsplit("!", 1)
            target_sheet = prefix.strip("'")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(target)
        except ValueError:
            continue
        # Evidence workbooks are bounded; never expand a large reference.
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 10_000:
            continue
        from openpyxl.utils import get_column_letter
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                refs.add((target_sheet, f"{get_column_letter(col)}{row}"))
    return refs


def _cycles(edges: dict[tuple[str, str], set[tuple[str, str]]]) -> list[list[tuple[str, str]]]:
    result: list[list[tuple[str, str]]] = []
    active: set[tuple[str, str]] = set()
    visited: set[tuple[str, str]] = set()
    stack: list[tuple[str, str]] = []

    def visit(node: tuple[str, str]) -> None:
        if node in active:
            result.append(stack[stack.index(node):] + [node])
            return
        if node in visited:
            return
        visited.add(node)
        active.add(node)
        stack.append(node)
        for dependency in edges.get(node, set()):
            if dependency in edges:
                visit(dependency)
        stack.pop()
        active.remove(node)

    for node in edges:
        visit(node)
    unique: dict[tuple[tuple[str, str], ...], list[tuple[str, str]]] = {}
    for cycle in result:
        key = tuple(sorted(set(cycle)))
        unique[key] = cycle
    return list(unique.values())


def _meaningful_cells(sheet: Any) -> list[Any]:
    """Return cells that carry a displayed value or formula.

    Excel files commonly retain formatting for whole columns or rows.  Those
    cells are not evidence-bearing data and must not turn a compact workbook
    into a giant LLM packet.  Basic style statistics remain derived from the
    meaningful data cells, while sheet-level formatting metadata is retained.
    """
    return sorted(
        (cell for cell in sheet._cells.values() if cell.value is not None),
        key=lambda cell: (cell.row, cell.column),
    )


def _meaningful_range(cells: list[Any]) -> tuple[str, int, int]:
    if not cells:
        return "A1", 1, 1
    min_row = min(cell.row for cell in cells)
    max_row = max(cell.row for cell in cells)
    min_col = min(cell.column for cell in cells)
    max_col = max(cell.column for cell in cells)
    return f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}", max_row, max_col


def extract_spreadsheet_evidence(
    path: Path,
    *,
    relative_path: str,
    recalculated_path: Path | None = None,
    recalculation_status: str = "not_requested",
    max_cells: int = 20_000,
) -> SpreadsheetEvidenceV1:
    """Read one workbook without modifying it and retain coordinate-level facts."""
    path = path.resolve()
    formula_book = load_workbook(path, read_only=False, data_only=False)
    value_book = load_workbook(path, read_only=False, data_only=True)
    recalc_book = None
    if recalculated_path is not None and recalculated_path.is_file():
        recalc_book = load_workbook(recalculated_path, read_only=False, data_only=True)
    sheets, edges, diagnostics = [], {}, []
    try:
        for sheet in formula_book.worksheets:
            values = value_book[sheet.title]
            recalc_values = recalc_book[sheet.title] if recalc_book and sheet.title in recalc_book.sheetnames else None
            meaningful = _meaningful_cells(sheet)
            used, effective_max_row, effective_max_column = _meaningful_range(meaningful)
            cells: list[SpreadsheetCellEvidenceV1] = []
            header_row: list[str] = []
            for cell in meaningful:
                if len(cells) >= max_cells:
                    raise ValueError("spreadsheet_evidence_cell_budget_exceeded")
                formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                current = values[cell.coordinate].value
                recalculated = recalc_values[cell.coordinate].value if recalc_values else None
                cells.append(SpreadsheetCellEvidenceV1(
                    address=cell.coordinate, value=_safe_value(current), formula=formula,
                    recalculated_value=_safe_value(recalculated), number_format=cell.number_format,
                    style_id=cell.style_id,
                ))
                if cell.row <= 10 and isinstance(current, str) and current.strip():
                    header_row.append(current.strip())
                if formula:
                    node = (sheet.title, cell.coordinate)
                    edges[node] = _formula_refs(formula, sheet.title)
                if isinstance(current, str) and current.startswith("#"):
                    diagnostics.append(SpreadsheetFormulaDiagnosticV1(
                        kind="formula_error", locations=[f"{sheet.title}!{cell.coordinate}"], detail=current,
                    ))
            widths = {key: dimension.width for key, dimension in sheet.column_dimensions.items() if dimension.width is not None}
            styles = defaultdict(int)
            for cell in cells:
                styles[str(cell.style_id)] += 1
            sheets.append(SpreadsheetSheetEvidenceV1(
                name=sheet.title, used_range=used, max_row=effective_max_row, max_column=effective_max_column,
                headers=header_row[:40], cells=cells,
                merged_ranges=sorted(str(value) for value in sheet.merged_cells.ranges),
                hidden_rows=[index for index, dim in sheet.row_dimensions.items() if dim.hidden],
                hidden_columns=[key for key, dim in sheet.column_dimensions.items() if dim.hidden],
                freeze_panes=str(sheet.freeze_panes) if sheet.freeze_panes else None,
                column_widths=widths, style_counts=dict(styles),
            ))
        for cycle in _cycles(edges):
            diagnostics.append(SpreadsheetFormulaDiagnosticV1(
                kind="formula_cycle", locations=[f"{sheet}!{address}" for sheet, address in cycle],
                detail="formula dependency cycle",
            ))
    finally:
        formula_book.close()
        value_book.close()
        if recalc_book:
            recalc_book.close()
    return SpreadsheetEvidenceV1(
        relative_path=relative_path, original_sha256=_sha(path),
        recalculated_sha256=_sha(recalculated_path) if recalculated_path and recalculated_path.is_file() else None,
        recalculation_status=recalculation_status, sheets=sheets, diagnostics=diagnostics,
    )


_LOCATOR = re.compile(r"^(?P<path>[^#]+)#sheet=(?P<sheet>[^&]+)&range=(?P<range>[A-Z]+[1-9][0-9]*(?::[A-Z]+[1-9][0-9]*)?)$")


def validate_spreadsheet_locator(locator: str, evidence: dict[str, SpreadsheetEvidenceV1]) -> None:
    match = _LOCATOR.fullmatch(locator)
    if not match:
        raise ValueError("spreadsheet_evidence_locator_invalid")
    path = match.group("path")
    if path not in evidence:
        raise ValueError("spreadsheet_evidence_locator_path_unknown")
    sheet = match.group("sheet")
    expected = next((row for row in evidence[path].sheets if row.name == sheet), None)
    if expected is None:
        raise ValueError("spreadsheet_evidence_locator_sheet_unknown")
    min_col, min_row, max_col, max_row = range_boundaries(match.group("range"))
    if max_row > expected.max_row or max_col > expected.max_column:
        raise ValueError("spreadsheet_evidence_locator_outside_used_range")
