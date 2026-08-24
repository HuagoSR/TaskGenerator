from __future__ import annotations

import re
from pathlib import Path
from typing import List, Literal, Sequence

from pydantic import BaseModel, Field


ContentDecision = Literal["pass", "blocked"]


class EvidenceContentFileCheckV1(BaseModel):
    check_version: Literal["v3.evidence_content_file_check.1"] = (
        "v3.evidence_content_file_check.1"
    )
    node_id: str
    file_name: str
    decision: ContentDecision
    reason_codes: List[str] = Field(default_factory=list)
    headers: List[str] = Field(default_factory=list)
    row_count: int = 0
    business_field_count: int = 0
    role_header_overlap_count: int = 0
    inadequate_width_columns: List[str] = Field(default_factory=list)


class EvidenceContentQualityReportV1(BaseModel):
    report_version: Literal["v3.evidence_content_quality.1"] = (
        "v3.evidence_content_quality.1"
    )
    proposal_id: str
    decision: ContentDecision
    checks: List[EvidenceContentFileCheckV1] = Field(default_factory=list)
    blocking_reason_codes: List[str] = Field(default_factory=list)
    notes: List[str] = Field(default_factory=list)


class EvidenceContentQualityValidator:
    """Report-first audit for candidate-visible evidence semantics.

    This validator deliberately distinguishes workbook mechanics from business
    content. It does not decide whether a scenario fact is true; it detects
    evidence artifacts that cannot possibly support the proposal's declared
    business judgments because they contain only the legacy placeholder schema,
    model candidate-authored outputs as source evidence, or are visibly clipped.
    """

    PLACEHOLDER_HEADERS = {
        "evidence_id",
        "record_key",
        "observed_value",
        "observation_status",
        "source_ref",
    }
    GENERIC_HEADERS = PLACEHOLDER_HEADERS | {
        "id",
        "record_id",
        "value",
        "status",
        "source",
        "notes",
        "description",
    }
    ROLE_STOPWORDS = {
        "and",
        "assessment",
        "business",
        "candidate",
        "data",
        "documentation",
        "evidence",
        "file",
        "information",
        "record",
        "records",
        "source",
        "supporting",
    }
    OUTPUT_ROLE_PATTERN = re.compile(
        r"\b(?:candidate[- ]authored|candidate[- ]produced|final deliverable|"
        r"assessment synthesis|executive synthesis)\b",
        re.IGNORECASE,
    )
    NEGATED_OUTPUT_ROLE_PATTERN = re.compile(
        r"\b(?:not|never)\s+(?:candidate[- ]authored|candidate[- ]produced|"
        r"final deliverable|assessment synthesis|executive synthesis)\b",
        re.IGNORECASE,
    )
    INPUT_TO_OUTPUT_RELATION_PATTERN = re.compile(
        r"\b(?:used\s+to\s+support|supports?|that\s+requires?|requiring|"
        r"requires?|used\s+for|provided\s+for)\s+(?:the\s+)?"
        r"(?:candidate[- ]authored|candidate[- ]produced)\b",
        re.IGNORECASE,
    )

    def validate(
        self,
        *,
        proposal: object,
        files: Sequence[object],
        reference_root: Path,
    ) -> EvidenceContentQualityReportV1:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        nodes = {
            str(item.node_id): item
            for item in getattr(proposal, "evidence_nodes", [])
        }
        checks: List[EvidenceContentFileCheckV1] = []
        all_reasons: set[str] = set()
        for record in files:
            node_id = str(record.node_id)
            file_name = str(record.file_name)
            node = nodes.get(node_id)
            reasons: set[str] = set()
            headers: List[str] = []
            row_count = 0
            business_fields: List[str] = []
            overlap_count = 0
            inadequate_width_columns: List[str] = []
            try:
                workbook = load_workbook(
                    reference_root / file_name,
                    read_only=False,
                    data_only=False,
                )
                sheet = workbook["Evidence"]
                headers = [
                    str(sheet.cell(row=1, column=index).value or "").strip()
                    for index in range(1, sheet.max_column + 1)
                ]
                normalized_headers = {
                    self._normalize(value) for value in headers if value
                }
                row_count = max(0, sheet.max_row - 1)
                if normalized_headers == self.PLACEHOLDER_HEADERS:
                    reasons.add("generic_placeholder_schema")
                business_fields = [
                    value
                    for value in normalized_headers
                    if value and value not in self.GENERIC_HEADERS
                ]
                if len(business_fields) < 3:
                    reasons.add("insufficient_business_fields")
                artifact_spec = getattr(node, "artifact_spec", None)
                role_tokens = self._semantic_tokens(
                    " ".join(
                        [
                            str(getattr(node, "artifact_role", "")),
                            str(getattr(node, "intended_contents", "")),
                            str(getattr(artifact_spec, "record_type", "")),
                        ]
                    )
                )
                header_tokens = self._semantic_tokens(" ".join(headers))
                overlap_count = len(role_tokens & header_tokens)
                if role_tokens and overlap_count == 0:
                    reasons.add("artifact_role_not_represented_in_headers")
                for column_index in range(1, sheet.max_column + 1):
                    letter = get_column_letter(column_index)
                    width = sheet.column_dimensions[letter].width or 13.0
                    max_length = max(
                        len(str(sheet.cell(row=row, column=column_index).value or ""))
                        for row in range(1, sheet.max_row + 1)
                    )
                    required_width = min(60.0, max(10.0, max_length + 2.0))
                    if width < required_width * 0.8:
                        inadequate_width_columns.append(letter)
                if inadequate_width_columns:
                    reasons.add("visible_column_clipping_risk")
                workbook.close()
            except Exception:
                reasons.add("content_workbook_unreadable")
            role_text = " ".join(
                [
                    str(getattr(node, "artifact_role", "")),
                    str(getattr(node, "intended_contents", "")),
                ]
            )
            positive_role_text = self.NEGATED_OUTPUT_ROLE_PATTERN.sub(
                "",
                role_text,
            )
            positive_role_text = self.INPUT_TO_OUTPUT_RELATION_PATTERN.sub(
                "",
                positive_role_text,
            )
            if self.OUTPUT_ROLE_PATTERN.search(positive_role_text):
                reasons.add("candidate_output_modeled_as_reference_input")
            all_reasons.update(reasons)
            checks.append(
                EvidenceContentFileCheckV1(
                    node_id=node_id,
                    file_name=file_name,
                    decision="blocked" if reasons else "pass",
                    reason_codes=sorted(reasons),
                    headers=headers,
                    row_count=row_count,
                    business_field_count=len(business_fields),
                    role_header_overlap_count=overlap_count,
                    inadequate_width_columns=inadequate_width_columns,
                )
            )
        decision: ContentDecision = (
            "pass" if checks and all(item.decision == "pass" for item in checks) else "blocked"
        )
        return EvidenceContentQualityReportV1(
            proposal_id=str(getattr(proposal, "proposal_id", "unknown")),
            decision=decision,
            checks=checks,
            blocking_reason_codes=sorted(all_reasons),
            notes=[
                "This report is content-semantic evidence, not professional review.",
                "A pass does not establish business truth; provenance and scenario consistency require separate gates.",
            ],
        )

    @classmethod
    def _normalize(cls, value: str) -> str:
        return re.sub(r"_+", "_", re.sub(r"[^a-z0-9]+", "_", value.lower())).strip("_")

    @classmethod
    def _semantic_tokens(cls, text: str) -> set[str]:
        return {
            cls._canonical_semantic_token(token)
            for token in re.findall(r"[a-z][a-z0-9]{2,}", text.lower())
            if token not in cls.ROLE_STOPWORDS
        }

    @staticmethod
    def _canonical_semantic_token(token: str) -> str:
        """Conservatively align common plural role/header word forms.

        This is intentionally narrower than general stemming: content gates
        should recognize ``thresholds`` versus ``Threshold`` without making
        unrelated business terms equivalent.
        """

        narrow_business_equivalents = {
            "corroborating": "support",
            "corroboration": "support",
            "corroborative": "support",
        }
        if token in narrow_business_equivalents:
            return narrow_business_equivalents[token]
        irregular_plurals = {
            "criteria": "criterion",
        }
        if token in irregular_plurals:
            return irregular_plurals[token]
        if len(token) > 4 and token.endswith("ies"):
            return f"{token[:-3]}y"
        if (
            len(token) > 3
            and token.endswith("s")
            and not token.endswith(("ss", "us", "is"))
        ):
            return token[:-1]
        return token
