from __future__ import annotations

import hashlib
import json
import os
import shlex
import shutil
import subprocess
import time
from io import BytesIO
from pathlib import Path, PurePosixPath
from types import SimpleNamespace
from typing import Any, Dict, List, Literal, Optional

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field, model_validator

from task_generator.v3_codex_local_grader import (
    CodexLocalGraderOutcomeV1,
    CodexLocalGraderRunnerV1,
)
from task_generator.v3_codex_local_solver import (
    CodexLocalCLIIdentityV1,
    CodexLocalCampaignManifestV1,
    CodexLocalTooling,
    _atomic_json,
    _copy_projection,
    _expected_deliverable,
    _now,
    _projection_tree,
    _sha_file,
    _terminate_process_tree,
    _tree,
    inspect_delivery,
)
from task_generator.v3_compact_screening_grader import (
    CompactGraderBindingV2,
    CompactGraderReviewV2,
)
from task_generator.v3_external_model_comparison import (
    ExternalComparisonGradeRecordV1,
    ExternalComparisonTaskBindingV1,
    ExternalStackSummaryV1,
    ExternalTaskObservationV1,
    PairwiseStackComparisonV1,
    _pairwise,
    _summary,
)
from task_generator.v3_representative_codex_local import (
    RepresentativeCodexCampaignV1,
)
from task_generator.v3_source_fingerprint import governed_source_fingerprint


MODEL = "deepseek-v4-pro"
BASE_URL = "https://api.deepseek.com"
STACK_ID = "deepseek-v4-pro@official_opencode"
BASELINE_STACK_ID = "gpt-5.6-sol@chatgpt_codex"
MATCHED_TASK_IDS = (
    "rp_003152533de5b652",
    "rp_fb0d66050783b08e",
)
PUBLIC_TIMEOUT_SECONDS = 300
TASK_TIMEOUT_SECONDS = 1800
SYSTEMIC_FAILURE_THRESHOLD = 3
REMOTE_ROOT = PurePosixPath("/home/user/task")
REMOTE_VENV = PurePosixPath("/home/user/.taskgenerator-opencode-venv")


def _sha_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _canonical_bytes(value: object) -> bytes:
    return (
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        + "\n"
    ).encode("utf-8")


def _redact(value: str, secrets: List[str]) -> str:
    result = value
    for secret in secrets:
        if secret:
            result = result.replace(secret, "[REDACTED_SECRET]")
    return result


def _assert_secrets_absent(value: str, secrets: List[str]) -> None:
    if any(secret and secret in value for secret in secrets):
        raise ValueError("official_opencode_secret_persistence_blocked")


class OfficialDeepSeekProviderProbeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    probe_version: Literal[
        "v3.official_deepseek_provider_probe.1"
    ] = "v3.official_deepseek_provider_probe.1"
    status: Literal["pass", "infrastructure_failed"]
    endpoint: Literal["https://api.deepseek.com"] = BASE_URL
    model: Literal["deepseek-v4-pro"] = MODEL
    models_endpoint_pass: bool = False
    stream_tool_call_pass: bool = False
    tool_name: Optional[str] = None
    tool_arguments: Dict[str, Any] = Field(default_factory=dict)
    final_text_present: bool = False
    usage: Dict[str, int] = Field(default_factory=dict)
    duration_seconds: float
    first_failure: Optional[str] = None


class OpenCodeDiagnosticsV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    line_count: int = 0
    invalid_line_count: int = 0
    event_types: List[str] = Field(default_factory=list)
    turn_event_seen: bool = False
    tool_event_seen: bool = False
    shell_or_python_seen: bool = False
    error_event_seen: bool = False
    usage: Dict[str, int] = Field(default_factory=dict)
    native_retry_signals: int = 0


def _walk(value: Any):
    yield value
    if isinstance(value, dict):
        for child in value.values():
            yield from _walk(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk(child)


def parse_opencode_jsonl(text: str) -> OpenCodeDiagnosticsV1:
    kinds: List[str] = []
    usage: Dict[str, int] = {}
    invalid = 0
    turn = False
    tool = False
    shell = False
    error = False
    retries = 0
    lines = [line for line in text.splitlines() if line.strip()]
    for line in lines:
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            invalid += 1
            continue
        if not isinstance(event, dict):
            invalid += 1
            continue
        kind = str(event.get("type", "unknown"))
        kinds.append(kind)
        lowered = kind.lower()
        turn = turn or any(
            marker in lowered
            for marker in ("session", "message", "step", "turn")
        )
        error = error or "error" in lowered or "failed" in lowered
        for node in _walk(event):
            if not isinstance(node, dict):
                continue
            node_kinds = {
                str(node.get(key) or "").lower()
                for key in ("type", "tool", "name")
            }
            if any("tool" in value for value in node_kinds) or bool(
                node_kinds
                & {
                "bash",
                "shell",
                "write",
                "read",
                "edit",
                }
            ):
                tool = True
            if any(
                marker in value
                for marker in ("bash", "shell", "python", "command")
                for value in node_kinds
            ):
                shell = True
            raw_usage = node.get("usage")
            if isinstance(raw_usage, dict):
                for key, raw in raw_usage.items():
                    if isinstance(raw, int) and raw >= 0:
                        usage[key] = max(usage.get(key, 0), raw)
            retry = node.get("retry")
            if retry is True or (
                isinstance(retry, int) and not isinstance(retry, bool)
            ):
                retries += 1
        if "retry" in lowered:
            retries += 1
    return OpenCodeDiagnosticsV1(
        line_count=len(lines),
        invalid_line_count=invalid,
        event_types=sorted(set(kinds)),
        turn_event_seen=turn,
        tool_event_seen=tool,
        shell_or_python_seen=shell,
        error_event_seen=error,
        usage=usage,
        native_retry_signals=retries,
    )


class OpenCodeEnvironmentProbeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    probe_version: Literal[
        "v3.opencode_environment_probe.1"
    ] = "v3.opencode_environment_probe.1"
    environment_kind: Literal[
        "e2b_official_opencode",
        "e2b_derived_opencode",
        "wsl2_ubuntu_opencode",
    ]
    status: Literal["pass", "infrastructure_failed"]
    template_or_distribution: str
    opencode_version: Optional[str] = None
    python_version: Optional[str] = None
    openpyxl_version: Optional[str] = None
    derived_template_build_id: Optional[str] = None
    first_failure: Optional[str] = None


class SolverEnvironmentSelectionV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    selection_version: Literal[
        "v3.solver_environment_selection.1"
    ] = "v3.solver_environment_selection.1"
    selected_environment: Literal[
        "e2b_official_opencode",
        "e2b_derived_opencode",
        "wsl2_ubuntu_opencode",
    ]
    opencode_version: str
    template_or_distribution: str
    environment_probe_sha256: str = Field(min_length=64, max_length=64)
    public_agent_probe_sha256: str = Field(min_length=64, max_length=64)
    homogeneous_cohort_required: Literal[True] = True


class OfficialDeepSeekOpenCodeCampaignV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal[
        "v3.official_deepseek_opencode_campaign.1"
    ] = "v3.official_deepseek_opencode_campaign.1"
    campaign_id: str
    representative_campaign_path: str
    representative_campaign_sha256: str = Field(
        min_length=64, max_length=64
    )
    baseline_solver_manifest_path: str
    baseline_solver_manifest_sha256: str = Field(
        min_length=64, max_length=64
    )
    baseline_grader_outcome_path: str
    baseline_grader_outcome_sha256: str = Field(
        min_length=64, max_length=64
    )
    parity_report_path: str
    parity_report_sha256: str = Field(min_length=64, max_length=64)
    source_fingerprint: str = Field(min_length=64, max_length=64)
    grader_cli: CodexLocalCLIIdentityV1
    endpoint: Literal["https://api.deepseek.com"] = BASE_URL
    model: Literal["deepseek-v4-pro"] = MODEL
    e2b_template: Literal["opencode"] = "opencode"
    wsl_distribution: Literal["Ubuntu"] = "Ubuntu"
    stable_fallback_opencode_version: Literal["1.18.10"] = "1.18.10"
    public_timeout_seconds: Literal[300] = PUBLIC_TIMEOUT_SECONDS
    task_timeout_seconds: Literal[1800] = TASK_TIMEOUT_SECONDS
    orchestration_retry_count: Literal[0] = 0
    task_bindings: List[ExternalComparisonTaskBindingV1] = Field(
        min_length=24, max_length=24
    )
    matched_task_ids: List[str] = Field(min_length=2, max_length=2)
    private_task_upload_authorized: Literal[True] = True
    grader_authorized: Literal[True] = True
    training_authorized: Literal[False] = False
    release_authorized: Literal[False] = False
    promotion_authorized: Literal[False] = False

    @model_validator(mode="after")
    def validate_matrix(self) -> "OfficialDeepSeekOpenCodeCampaignV1":
        if tuple(self.matched_task_ids) != MATCHED_TASK_IDS:
            raise ValueError("official_opencode_matched_pair_mismatch")
        if len({item.blind_task_id for item in self.task_bindings}) != 24:
            raise ValueError("official_opencode_task_identity_duplicate")
        cells = {
            (
                item.domain,
                item.route_id,
                item.motif,
                item.replicate_id,
            )
            for item in self.task_bindings
        }
        if len(cells) != 24:
            raise ValueError("official_opencode_matrix_incomplete")
        by_id = {item.blind_task_id: item for item in self.task_bindings}
        first, second = (by_id[key] for key in MATCHED_TASK_IDS)
        if not (
            first.brief_id == second.brief_id == "brief_f584c00214df"
            and first.domain == second.domain == "audit_compliance"
            and first.motif
            == second.motif
            == "cross_check_validation"
            and {first.route_id, second.route_id}
            == {"skill_guided_llm", "llm_led_hybrid"}
        ):
            raise ValueError("official_opencode_matched_pair_not_matched")
        return self


class OfficialDeepSeekOpenCodeReceiptV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    receipt_version: Literal[
        "v3.official_deepseek_opencode_receipt.1"
    ] = "v3.official_deepseek_opencode_receipt.1"
    scope_sha256: str = Field(min_length=64, max_length=64)
    authorization_basis: Literal[
        "user_official_deepseek_opencode_campaign_authorization"
    ] = "user_official_deepseek_opencode_campaign_authorization"
    private_package_upload_authorized: Literal[True] = True
    issued_at: str
    consumed_at: Optional[str] = None


class OpenCodeProcessOutcomeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    outcome_version: Literal[
        "v3.official_deepseek_opencode_process.1"
    ] = "v3.official_deepseek_opencode_process.1"
    kind: Literal["public_probe", "private_task"]
    blind_task_id: str
    environment_kind: Literal[
        "e2b_official_opencode",
        "e2b_derived_opencode",
        "wsl2_ubuntu_opencode",
    ]
    status: Literal["pass", "task_failed", "infrastructure_failed"]
    started_at: str
    completed_at: str
    duration_seconds: float
    returncode: Optional[int]
    timed_out: bool
    sandbox_id: Optional[str] = None
    command_identity: List[str]
    workspace_path: str
    workspace_input_tree_sha256: str
    jsonl_path: str
    jsonl_sha256: str
    stderr_path: str
    stderr_sha256: str
    diagnostics: OpenCodeDiagnosticsV1
    delivery: Dict[str, Any]
    output_tree_sha256: str
    first_failure: Optional[str] = None


TaskState = Literal["not_started", "running", "completed", "interrupted"]


class OpenCodeTaskStateV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    status: TaskState = "not_started"
    outcome_path: Optional[str] = None
    outcome_sha256: Optional[str] = None


class OfficialDeepSeekOpenCodeExecutionV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    manifest_version: Literal[
        "v3.official_deepseek_opencode_execution.1"
    ] = "v3.official_deepseek_opencode_execution.1"
    scope_path: str
    scope_sha256: str = Field(min_length=64, max_length=64)
    receipt_path: str
    receipt_sha256: str = Field(min_length=64, max_length=64)
    status: Literal[
        "prepared",
        "running",
        "provider_incomplete",
        "stack_incomplete",
        "matched_failed",
        "solver_completed",
        "incomplete",
    ]
    provider_probe_path: Optional[str] = None
    provider_probe_sha256: Optional[str] = None
    environment_probe_paths: List[str] = Field(default_factory=list)
    selection_path: Optional[str] = None
    selection_sha256: Optional[str] = None
    public_probe_path: Optional[str] = None
    public_probe_sha256: Optional[str] = None
    tasks: Dict[str, OpenCodeTaskStateV1]
    created_at: str
    updated_at: str
    systemic_failure_code: Optional[str] = None


class OfficialDeepSeekOpenCodeGraderV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    manifest_version: Literal[
        "v3.official_deepseek_opencode_grader.1"
    ] = "v3.official_deepseek_opencode_grader.1"
    scope_sha256: str = Field(min_length=64, max_length=64)
    execution_sha256: str = Field(min_length=64, max_length=64)
    baseline_grader_outcome_sha256: str = Field(
        min_length=64, max_length=64
    )
    status: Literal["prepared", "running", "completed", "incomplete"]
    records: Dict[str, ExternalComparisonGradeRecordV1]
    created_at: str
    updated_at: str


class OpenCodeExecutionAuditRecordV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blind_task_id: str
    original_status: Literal[
        "pass", "task_failed", "infrastructure_failed"
    ]
    audited_status: Literal[
        "pass",
        "task_failed",
        "infrastructure_failed",
        "out_of_scope_after_freeze",
    ]
    failure_code: Optional[str] = None
    counted_in_cohort: bool


class OfficialDeepSeekOpenCodeExecutionAuditV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    audit_version: Literal[
        "v3.official_deepseek_opencode_execution_audit.1"
    ] = "v3.official_deepseek_opencode_execution_audit.1"
    scope_sha256: str = Field(min_length=64, max_length=64)
    execution_sha256: str = Field(min_length=64, max_length=64)
    records: List[OpenCodeExecutionAuditRecordV1]
    freeze_task_id: Optional[str] = None
    freeze_failure_code: Optional[str] = None
    counted_task_count: int
    post_freeze_attempt_count: int
    governance_breach_detected: bool
    audited_execution_status: Literal[
        "solver_completed", "stack_incomplete"
    ]


class OfficialDeepSeekOpenCodeResultV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    result_version: Literal[
        "v3.official_deepseek_opencode_result.1"
    ] = "v3.official_deepseek_opencode_result.1"
    campaign_id: str
    scope_sha256: str = Field(min_length=64, max_length=64)
    execution_sha256: str = Field(min_length=64, max_length=64)
    grader_sha256: str = Field(min_length=64, max_length=64)
    execution_audit_sha256: Optional[str] = None
    observations: List[ExternalTaskObservationV1]
    stack_summaries: List[ExternalStackSummaryV1]
    pairwise_comparison: PairwiseStackComparisonV1
    decision: Literal[
        "official_deepseek_comparison_ready",
        "official_deepseek_partial",
        "stack_incomplete",
    ]
    practical_winner: Literal[
        "gpt_baseline",
        "official_deepseek",
        "practical_tie",
        "insufficient_common_coverage",
    ]
    professional_validity: Literal["provisional"] = "provisional"
    grader_evidence_kind: Literal[
        "independent_model_proxy"
    ] = "independent_model_proxy"
    expert_evidence_present: Literal[False] = False
    default_solver_change_authorized: Literal[False] = False
    training_authorized: Literal[False] = False
    release_authorized: Literal[False] = False
    promotion_authorized: Literal[False] = False


def compile_official_opencode_scope(
    *,
    representative_campaign_path: str | Path,
    baseline_solver_manifest_path: str | Path,
    baseline_grader_outcome_path: str | Path,
    parity_report_path: str | Path,
    repository_root: str | Path,
    output_root: str | Path,
) -> tuple[OfficialDeepSeekOpenCodeCampaignV1, Path, Path]:
    campaign_path = Path(representative_campaign_path).resolve()
    solver_path = Path(baseline_solver_manifest_path).resolve()
    grader_path = Path(baseline_grader_outcome_path).resolve()
    parity_path = Path(parity_report_path).resolve()
    campaign = RepresentativeCodexCampaignV1.model_validate_json(
        campaign_path.read_text(encoding="utf-8")
    )
    solver = CodexLocalCampaignManifestV1.model_validate_json(
        solver_path.read_text(encoding="utf-8")
    )
    grader = CodexLocalGraderOutcomeV1.model_validate_json(
        grader_path.read_text(encoding="utf-8")
    )
    parity = json.loads(parity_path.read_text(encoding="utf-8"))
    fingerprint = governed_source_fingerprint(repository_root)
    if not (
        campaign.reality_decision == "screening_ready"
        and solver.status == "solver_completed"
        and len(solver.task_outcome_paths) == 24
        and grader.completed_grades == 24
        and parity.get("passed") is True
        and parity.get("network_mode") == "none"
        and parity.get("read_only_root") is True
        and parity.get("provider_credentials_mounted") is False
        and parity.get("cleanup_returncode") == 0
        and parity.get("source_fingerprint") == fingerprint
    ):
        raise ValueError("official_opencode_upstream_not_ready")
    candidate_root = (
        campaign_path.parent
        / "reality"
        / "blind_staging"
        / "candidate_packages"
    )
    bindings: List[ExternalComparisonTaskBindingV1] = []
    for assignment in sorted(
        campaign.assignments, key=lambda item: item.blind_task_id
    ):
        candidate = candidate_root / assignment.blind_task_id
        bindings.append(
            ExternalComparisonTaskBindingV1(
                blind_task_id=assignment.blind_task_id,
                brief_id=assignment.brief_id,
                route_id=assignment.route_id,
                motif=assignment.motif,
                replicate_id=assignment.replicate_id,
                domain=assignment.domain,
                source_package_path=assignment.package_root,
                candidate_package_path=str(candidate),
                package_fingerprint=assignment.package_fingerprint,
                candidate_tree_sha256=_tree(candidate),
                candidate_projection_sha256=_projection_tree(candidate),
                reality_evidence_sha256=assignment.reality_evidence_sha256,
                expected_deliverable=_expected_deliverable(candidate),
            )
        )
    scope = OfficialDeepSeekOpenCodeCampaignV1(
        campaign_id=f"{campaign.campaign_id}_official_deepseek_opencode_v1",
        representative_campaign_path=str(campaign_path),
        representative_campaign_sha256=_sha_file(campaign_path),
        baseline_solver_manifest_path=str(solver_path),
        baseline_solver_manifest_sha256=_sha_file(solver_path),
        baseline_grader_outcome_path=str(grader_path),
        baseline_grader_outcome_sha256=_sha_file(grader_path),
        parity_report_path=str(parity_path),
        parity_report_sha256=_sha_file(parity_path),
        source_fingerprint=fingerprint,
        grader_cli=CodexLocalTooling(
            Path(repository_root).resolve()
            / "artifacts"
            / "pipeline_reconstruction"
            / "codex_local_tooling"
        ).freeze_identity(),
        task_bindings=bindings,
        matched_task_ids=list(MATCHED_TASK_IDS),
    )
    encoded = _canonical_bytes(scope.model_dump(mode="json"))
    digest = _sha_bytes(encoded)
    root = Path(output_root).resolve()
    scope_path = root / "governance" / "scopes" / f"{digest}.json"
    receipt_path = root / "governance" / "receipts" / f"{digest}.json"
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    receipt_path.parent.mkdir(parents=True, exist_ok=True)
    if scope_path.exists() and scope_path.read_bytes() != encoded:
        raise ValueError("official_opencode_scope_immutable_collision")
    scope_path.write_bytes(encoded)
    if receipt_path.exists():
        raise FileExistsError("official_opencode_receipt_exists")
    _atomic_json(
        receipt_path,
        OfficialDeepSeekOpenCodeReceiptV1(
            scope_sha256=digest, issued_at=_now()
        ).model_dump(mode="json"),
    )
    return scope, scope_path, receipt_path


def run_official_provider_probe(
    api_key: str,
    output_path: str | Path,
    *,
    client_factory=None,
) -> OfficialDeepSeekProviderProbeV1:
    started = time.monotonic()
    failure = None
    models_ok = False
    tool_ok = False
    final_text = ""
    usage: Dict[str, int] = {}
    tool_args: Dict[str, Any] = {}
    try:
        if not api_key:
            raise ValueError("deepseek_api_key_missing")
        if client_factory is None:
            from openai import OpenAI

            client = OpenAI(
                api_key=api_key,
                base_url=BASE_URL,
                max_retries=0,
                timeout=120,
            )
        else:
            client = client_factory(api_key=api_key, base_url=BASE_URL)
        models = client.models.list()
        model_ids = {
            str(getattr(item, "id", ""))
            for item in getattr(models, "data", [])
        }
        models_ok = MODEL in model_ids
        if not models_ok:
            raise RuntimeError("deepseek_v4_pro_missing_from_models")
        tools = [
            {
                "type": "function",
                "function": {
                    "name": "multiply",
                    "description": "Multiply two public integers.",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "a": {"type": "integer"},
                            "b": {"type": "integer"},
                        },
                        "required": ["a", "b"],
                        "additionalProperties": False,
                    },
                },
            }
        ]
        first = client.chat.completions.create(
            model=MODEL,
            messages=[
                {
                    "role": "user",
                    "content": "Use the multiply tool to calculate 17 times 23.",
                }
            ],
            tools=tools,
            stream=True,
            extra_body={"thinking": {"type": "disabled"}},
        )
        tool_id = ""
        tool_name = ""
        argument_text = ""
        for chunk in first:
            choice = chunk.choices[0] if getattr(chunk, "choices", []) else None
            delta = getattr(choice, "delta", None) if choice else None
            for call in getattr(delta, "tool_calls", None) or []:
                tool_id += getattr(call, "id", "") or ""
                function = getattr(call, "function", None)
                if function:
                    tool_name += getattr(function, "name", "") or ""
                    argument_text += getattr(function, "arguments", "") or ""
            raw_usage = getattr(chunk, "usage", None)
            if raw_usage:
                for key in ("prompt_tokens", "completion_tokens", "total_tokens"):
                    value = getattr(raw_usage, key, None)
                    if isinstance(value, int):
                        usage[key] = usage.get(key, 0) + value
        tool_args = json.loads(argument_text)
        tool_ok = (
            tool_name == "multiply"
            and tool_args == {"a": 17, "b": 23}
            and bool(tool_id)
        )
        if not tool_ok:
            raise RuntimeError("deepseek_stream_tool_call_invalid")
        second = client.chat.completions.create(
            model=MODEL,
            messages=[
                {
                    "role": "user",
                    "content": "Use the multiply tool to calculate 17 times 23.",
                },
                {
                    "role": "assistant",
                    "content": None,
                    "tool_calls": [
                        {
                            "id": tool_id,
                            "type": "function",
                            "function": {
                                "name": tool_name,
                                "arguments": argument_text,
                            },
                        }
                    ],
                },
                {
                    "role": "tool",
                    "tool_call_id": tool_id,
                    "content": "391",
                },
            ],
            tools=tools,
            stream=True,
            extra_body={"thinking": {"type": "disabled"}},
        )
        for chunk in second:
            choice = chunk.choices[0] if getattr(chunk, "choices", []) else None
            delta = getattr(choice, "delta", None) if choice else None
            final_text += getattr(delta, "content", "") or ""
            raw_usage = getattr(chunk, "usage", None)
            if raw_usage:
                for key in ("prompt_tokens", "completion_tokens", "total_tokens"):
                    value = getattr(raw_usage, key, None)
                    if isinstance(value, int):
                        usage[key] = usage.get(key, 0) + value
        if not final_text.strip():
            raise RuntimeError("deepseek_stream_final_text_missing")
        status = "pass"
    except Exception as exc:
        status = "infrastructure_failed"
        failure = _redact(f"{type(exc).__name__}:{exc}", [api_key])
    result = OfficialDeepSeekProviderProbeV1(
        status=status,
        models_endpoint_pass=models_ok,
        stream_tool_call_pass=tool_ok,
        tool_name="multiply" if tool_ok else None,
        tool_arguments=tool_args if tool_ok else {},
        final_text_present=bool(final_text.strip()),
        usage=usage,
        duration_seconds=round(time.monotonic() - started, 3),
        first_failure=failure,
    )
    _assert_secrets_absent(result.model_dump_json(), [api_key])
    _atomic_json(Path(output_path), result.model_dump(mode="json"))
    return result


def _probe_uploads() -> tuple[Dict[str, bytes], str, str]:
    stream = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet.append(["Item", "Amount"])
    sheet.append(["A", 12])
    sheet.append(["B", 30])
    workbook.save(stream)
    expected = "deliverable_files/opencode_probe.xlsx"
    prompt = (
        "Use shell commands and Python with openpyxl. Read "
        "reference_files/source.xlsx and create exactly "
        f"{expected}. Include the two source rows, a total row, and a Notes "
        "sheet. Verify the workbook with openpyxl. Do not copy the input."
    )
    return {"reference_files/source.xlsx": stream.getvalue()}, expected, prompt


def _private_uploads(
    binding: ExternalComparisonTaskBindingV1,
) -> tuple[Dict[str, bytes], str]:
    candidate = Path(binding.candidate_package_path)
    uploads = {
        "dataset_row.json": (candidate / "dataset_row.json").read_bytes(),
        "deliverable_contract.json": (
            candidate / "deliverable_contract.json"
        ).read_bytes(),
    }
    for path in sorted((candidate / "reference_files").glob("*.xlsx")):
        uploads[f"reference_files/{path.name}"] = path.read_bytes()
    prompt = (
        "Complete this route-blind spreadsheet work task. Read "
        "dataset_row.json, deliverable_contract.json, and every workbook in "
        "reference_files. Use shell commands and Python/openpyxl as needed. "
        f"Produce exactly {binding.expected_deliverable}. The output must be "
        "a genuine, non-empty XLSX that addresses the task, must not copy an "
        "input workbook, and must be verified with openpyxl before finishing."
    )
    return uploads, prompt


def _inspect_bytes(
    data: Optional[bytes], expected: str, input_hashes: set[str]
) -> Dict[str, Any]:
    result = {
        "relative_path": expected,
        "exists": data is not None,
        "nonempty": bool(data),
        "xlsx_openable": False,
        "has_nonempty_sheet": False,
        "differs_from_inputs": False,
        "sha256": _sha_bytes(data) if data else None,
        "size_bytes": len(data) if data else 0,
        "workbook_sheets": [],
        "valid": False,
        "failure": None,
    }
    if data is None:
        result["failure"] = "expected_delivery_missing"
        return result
    if not data:
        result["failure"] = "expected_delivery_empty"
        return result
    result["differs_from_inputs"] = result["sha256"] not in input_hashes
    try:
        workbook = load_workbook(BytesIO(data), data_only=False)
        result["workbook_sheets"] = list(workbook.sheetnames)
        result["has_nonempty_sheet"] = any(
            any(
                cell.value not in (None, "")
                for row in sheet.iter_rows()
                for cell in row
            )
            for sheet in workbook.worksheets
        )
        workbook.close()
        result["xlsx_openable"] = True
    except Exception:
        result["failure"] = "expected_delivery_not_real_xlsx"
    if result["xlsx_openable"] and not result["has_nonempty_sheet"]:
        result["failure"] = "expected_delivery_has_no_content"
    if (
        result["xlsx_openable"]
        and result["has_nonempty_sheet"]
        and not result["differs_from_inputs"]
    ):
        result["failure"] = "expected_delivery_copies_input"
    result["valid"] = bool(
        result["nonempty"]
        and result["xlsx_openable"]
        and result["has_nonempty_sheet"]
        and result["differs_from_inputs"]
    )
    return result


def _failure_code(
    diagnostics: OpenCodeDiagnosticsV1,
    stderr: str,
    *,
    timed_out: bool,
    returncode: Optional[int],
) -> Optional[str]:
    text = stderr.lower()
    if any(marker in text for marker in ("401", "403", "unauthorized", "api key")):
        return "authentication_failure"
    if any(
        marker in text
        for marker in (
            "connection refused",
            "connecterror",
            "getaddrinfo failed",
            "name resolution",
            "remoteprotocolerror",
            "service unavailable",
            "stream disconnected",
            "incomplete chunk",
            "incomplete chunked read",
            "peer closed connection",
            "502",
            "503",
        )
    ):
        return "service_or_stream_failure"
    if timed_out and not diagnostics.tool_event_seen:
        return "timeout_without_tool_activity"
    if returncode not in (0, None) and diagnostics.error_event_seen:
        return "opencode_protocol_failure"
    return None


def audit_official_opencode_execution(
    *,
    scope_path: str | Path,
    execution_path: str | Path,
    output_path: str | Path,
) -> tuple[OfficialDeepSeekOpenCodeExecutionAuditV1, Path]:
    scope_file = Path(scope_path).resolve()
    execution_file = Path(execution_path).resolve()
    scope = OfficialDeepSeekOpenCodeCampaignV1.model_validate_json(
        scope_file.read_text(encoding="utf-8")
    )
    execution = OfficialDeepSeekOpenCodeExecutionV1.model_validate_json(
        execution_file.read_text(encoding="utf-8")
    )
    if execution.scope_sha256 != _sha_file(scope_file):
        raise PermissionError("official_opencode_audit_scope_drift")
    order = task_execution_order(
        [item.blind_task_id for item in scope.task_bindings]
    )
    records: List[OpenCodeExecutionAuditRecordV1] = []
    consecutive_code = None
    consecutive_count = 0
    frozen = False
    freeze_task = None
    freeze_code = None
    post_freeze = 0
    for task_id in order:
        state = execution.tasks[task_id]
        if not state.outcome_path:
            continue
        outcome = OpenCodeProcessOutcomeV1.model_validate_json(
            Path(state.outcome_path).read_text(encoding="utf-8")
        )
        if frozen:
            records.append(
                OpenCodeExecutionAuditRecordV1(
                    blind_task_id=task_id,
                    original_status=outcome.status,
                    audited_status="out_of_scope_after_freeze",
                    failure_code=outcome.first_failure,
                    counted_in_cohort=False,
                )
            )
            post_freeze += 1
            continue
        stderr = Path(outcome.stderr_path).read_text(
            encoding="utf-8", errors="replace"
        )
        code = _failure_code(
            outcome.diagnostics,
            stderr,
            timed_out=outcome.timed_out,
            returncode=outcome.returncode,
        )
        if code:
            audited_status = "infrastructure_failed"
        elif outcome.status == "pass" and outcome.delivery.get("valid", False):
            audited_status = "pass"
        else:
            audited_status = "task_failed"
            code = outcome.first_failure
        records.append(
            OpenCodeExecutionAuditRecordV1(
                blind_task_id=task_id,
                original_status=outcome.status,
                audited_status=audited_status,
                failure_code=code,
                counted_in_cohort=True,
            )
        )
        if audited_status == "infrastructure_failed":
            if code == consecutive_code:
                consecutive_count += 1
            else:
                consecutive_code = code
                consecutive_count = 1
        else:
            consecutive_code = None
            consecutive_count = 0
        if consecutive_count >= SYSTEMIC_FAILURE_THRESHOLD:
            frozen = True
            freeze_task = task_id
            freeze_code = code
    audit = OfficialDeepSeekOpenCodeExecutionAuditV1(
        scope_sha256=_sha_file(scope_file),
        execution_sha256=_sha_file(execution_file),
        records=records,
        freeze_task_id=freeze_task,
        freeze_failure_code=freeze_code,
        counted_task_count=sum(item.counted_in_cohort for item in records),
        post_freeze_attempt_count=post_freeze,
        governance_breach_detected=post_freeze > 0,
        audited_execution_status=(
            "stack_incomplete" if frozen else "solver_completed"
        ),
    )
    destination = Path(output_path).resolve()
    _atomic_json(destination, audit.model_dump(mode="json"))
    return audit, destination


def task_execution_order(task_ids: List[str]) -> List[str]:
    if not set(MATCHED_TASK_IDS).issubset(task_ids):
        raise ValueError("official_opencode_matched_pair_missing")
    return list(MATCHED_TASK_IDS) + sorted(
        set(task_ids) - set(MATCHED_TASK_IDS)
    )


def matched_extension_allowed(
    outcomes: List[OpenCodeProcessOutcomeV1],
) -> bool:
    return (
        len(outcomes) == 2
        and {item.blind_task_id for item in outcomes}
        == set(MATCHED_TASK_IDS)
        and all(
            item.status == "pass" and item.delivery.get("valid", False)
            for item in outcomes
        )
    )


class E2BOpenCodeBackend:
    def __init__(self, *, sandbox_factory=None):
        self.sandbox_factory = sandbox_factory

    def _create(self, template: str, timeout: int, e2b_key: str):
        if self.sandbox_factory is not None:
            return self.sandbox_factory(
                template=template, timeout=timeout, api_key=e2b_key
            )
        from e2b import Sandbox

        return Sandbox.create(
            template=template, timeout=timeout, api_key=e2b_key
        )

    def inspect_environment(
        self,
        *,
        template: str,
        e2b_key: str,
        output_path: Path,
    ) -> OpenCodeEnvironmentProbeV1:
        sandbox = None
        first_failure = None
        version = None
        python_version = None
        openpyxl_version = None
        try:
            sandbox = self._create(template, 600, e2b_key)
            command = (
                "bash -lc 'set +e; "
                "printf \"OPENCODE=\"; opencode --version; "
                "printf \"PYTHON=\"; python3 --version 2>&1; "
                "printf \"PIP=\"; python3 -m pip --version 2>&1; "
                "python3 -m venv --help >/dev/null 2>&1; "
                "printf \"VENV=%s\\n\" \"$?\"; "
                "printf \"OPENPYXL=\"; "
                "python3 -c \"import openpyxl; print(openpyxl.__version__)\" "
                "2>&1; exit 0'"
            )
            result = sandbox.commands.run(command, timeout=120)
            text = (
                (getattr(result, "stdout", "") or "")
                + "\n"
                + (getattr(result, "stderr", "") or "")
            )
            for line in text.splitlines():
                if line.startswith("OPENCODE="):
                    version = line.split("=", 1)[1].strip()
                elif line.startswith("PYTHON="):
                    python_version = line.split("=", 1)[1].strip()
                elif line.startswith("OPENPYXL="):
                    value = line.split("=", 1)[1].strip()
                    if value and "Traceback" not in value:
                        openpyxl_version = value
            if not version or not python_version:
                raise RuntimeError("e2b_opencode_or_python_missing")
            if not openpyxl_version:
                install = sandbox.commands.run(
                    "bash -lc 'set -e; python3 -m venv "
                    f"{REMOTE_VENV}; {REMOTE_VENV}/bin/pip install "
                    "--disable-pip-version-check openpyxl; "
                    f"{REMOTE_VENV}/bin/python -c "
                    "\"import openpyxl; print(openpyxl.__version__)\"'",
                    timeout=300,
                )
                if getattr(install, "exit_code", 1) != 0:
                    raise RuntimeError("e2b_openpyxl_venv_install_failed")
                openpyxl_version = (
                    (getattr(install, "stdout", "") or "")
                    .strip()
                    .splitlines()[-1]
                )
            status = "pass"
        except Exception as exc:
            status = "infrastructure_failed"
            first_failure = _redact(
                f"{type(exc).__name__}:{exc}", [e2b_key]
            )
        finally:
            if sandbox is not None:
                try:
                    sandbox.kill()
                except Exception:
                    pass
        probe = OpenCodeEnvironmentProbeV1(
            environment_kind="e2b_official_opencode",
            status=status,
            template_or_distribution=template,
            opencode_version=version,
            python_version=python_version,
            openpyxl_version=openpyxl_version,
            first_failure=first_failure,
        )
        _assert_secrets_absent(probe.model_dump_json(), [e2b_key])
        _atomic_json(output_path, probe.model_dump(mode="json"))
        return probe

    def run_one(
        self,
        *,
        template: str,
        environment_kind: Literal[
            "e2b_official_opencode", "e2b_derived_opencode"
        ],
        kind: Literal["public_probe", "private_task"],
        task_id: str,
        uploads: Dict[str, bytes],
        prompt: str,
        expected: str,
        timeout: int,
        deepseek_key: str,
        e2b_key: str,
        output_dir: Path,
    ) -> OpenCodeProcessOutcomeV1:
        sandbox = None
        started_at = _now()
        started = time.monotonic()
        stdout = ""
        stderr = ""
        returncode = None
        sandbox_id = None
        timed_out = False
        secrets = [deepseek_key, e2b_key]
        input_hashes = {
            _sha_bytes(value)
            for name, value in uploads.items()
            if name.lower().endswith(".xlsx")
        }
        workspace_input_hash = _sha_bytes(
            _canonical_bytes(
                {
                    name: _sha_bytes(value)
                    for name, value in sorted(uploads.items())
                }
            )
        )
        data = None
        try:
            sandbox = self._create(template, timeout + 180, e2b_key)
            sandbox_id = str(
                getattr(sandbox, "sandbox_id", getattr(sandbox, "id", ""))
            ) or None
            output_dir.mkdir(parents=True, exist_ok=True)
            if sandbox_id:
                (output_dir / "sandbox_id.txt").write_text(
                    sandbox_id, encoding="utf-8"
                )
            sandbox.files.make_dir(str(REMOTE_ROOT))
            sandbox.files.make_dir(str(REMOTE_ROOT / "reference_files"))
            sandbox.files.make_dir(str(REMOTE_ROOT / "deliverable_files"))
            for relative, value in uploads.items():
                target = REMOTE_ROOT / relative
                sandbox.files.make_dir(str(target.parent))
                sandbox.files.write(str(target), value)
            sandbox.files.write(
                str(REMOTE_ROOT / "TASK.md"),
                prompt
                + "\nThe prepared Python executable is "
                + str(REMOTE_VENV / "bin" / "python")
                + ".\n",
            )
            prep = sandbox.commands.run(
                "bash -lc 'set -e; if [ ! -x "
                f"{REMOTE_VENV}/bin/python ]; then python3 -m venv "
                f"{REMOTE_VENV}; fi; if ! {REMOTE_VENV}/bin/python "
                "-c \"import openpyxl\" >/dev/null 2>&1; then "
                f"{REMOTE_VENV}/bin/pip install "
                "--disable-pip-version-check openpyxl; fi'",
                timeout=300,
            )
            if getattr(prep, "exit_code", 1) != 0:
                raise RuntimeError("e2b_openpyxl_preparation_failed")
            command = (
                "bash -lc 'set +e; PROMPT=$(cat /home/user/task/TASK.md); "
                "opencode run --format json --model "
                "deepseek/deepseek-v4-pro --auto --dir /home/user/task "
                "\"$PROMPT\"; code=$?; "
                "printf \"{\\\"type\\\":\\\"runner.exit\\\","
                "\\\"exit_code\\\":%s}\\n\" \"$code\"; exit 0'"
            )
            result = sandbox.commands.run(
                command,
                timeout=timeout,
                envs={
                    "DEEPSEEK_API_KEY": deepseek_key,
                    "VIRTUAL_ENV": str(REMOTE_VENV),
                },
            )
            stdout = getattr(result, "stdout", "") or ""
            stderr = getattr(result, "stderr", "") or ""
            returncode = _runner_exit_code(stdout)
            remote = str(REMOTE_ROOT / expected)
            if sandbox.files.exists(remote):
                data = sandbox.files.read(remote, format="bytes")
        except Exception as exc:
            stderr = stderr or f"{type(exc).__name__}:{exc}"
        finally:
            if sandbox is not None:
                try:
                    sandbox.kill()
                except Exception:
                    pass
        stdout = _redact(stdout, secrets)
        stderr = _redact(stderr, secrets)
        diagnostics = parse_opencode_jsonl(stdout)
        delivery = _inspect_bytes(data, expected, input_hashes)
        failure = _failure_code(
            diagnostics,
            stderr,
            timed_out=timed_out,
            returncode=returncode,
        )
        if failure:
            status = "infrastructure_failed"
        elif (
            returncode == 0
            and diagnostics.turn_event_seen
            and diagnostics.tool_event_seen
            and diagnostics.shell_or_python_seen
            and delivery["valid"]
        ):
            status = "pass"
        else:
            status = "task_failed"
            failure = delivery["failure"] or "opencode_or_protocol_failure"
        output_dir.mkdir(parents=True, exist_ok=True)
        jsonl_path = output_dir / "opencode.jsonl"
        stderr_path = output_dir / "stderr.txt"
        jsonl_path.write_text(stdout, encoding="utf-8")
        stderr_path.write_text(stderr, encoding="utf-8")
        if data:
            delivery_path = output_dir / expected
            delivery_path.parent.mkdir(parents=True, exist_ok=True)
            delivery_path.write_bytes(data)
        output_tree = _tree(output_dir)
        outcome = OpenCodeProcessOutcomeV1(
            kind=kind,
            blind_task_id=task_id,
            environment_kind=environment_kind,
            status=status,
            started_at=started_at,
            completed_at=_now(),
            duration_seconds=round(time.monotonic() - started, 3),
            returncode=returncode,
            timed_out=timed_out,
            sandbox_id=sandbox_id,
            command_identity=[
                "opencode",
                "run",
                "--format",
                "json",
                "--model",
                "deepseek/deepseek-v4-pro",
                "--auto",
            ],
            workspace_path=str(output_dir),
            workspace_input_tree_sha256=workspace_input_hash,
            jsonl_path=str(jsonl_path),
            jsonl_sha256=_sha_file(jsonl_path),
            stderr_path=str(stderr_path),
            stderr_sha256=_sha_file(stderr_path),
            diagnostics=diagnostics,
            delivery=delivery,
            output_tree_sha256=output_tree,
            first_failure=failure,
        )
        _assert_secrets_absent(outcome.model_dump_json(), secrets)
        _atomic_json(
            output_dir / "outcome.json", outcome.model_dump(mode="json")
        )
        return outcome


def _runner_exit_code(text: str) -> Optional[int]:
    for line in reversed(text.splitlines()):
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        if event.get("type") == "runner.exit":
            value = event.get("exit_code")
            return value if isinstance(value, int) else None
    return None


class WSLOpenCodeBackend:
    def __init__(
        self,
        *,
        repository_root: Path,
        popen_factory=subprocess.Popen,
    ):
        self.repository_root = repository_root.resolve()
        self.popen_factory = popen_factory
        self.tool_root = (
            self.repository_root / "artifacts" / "tooling" / "opencode-wsl"
        )

    @staticmethod
    def _wsl_path(path: Path) -> str:
        drive = path.drive.rstrip(":").lower()
        rest = path.resolve().as_posix().split(":", 1)[1]
        return f"/mnt/{drive}{rest}"

    def prepare(
        self,
        *,
        requested_version: str,
        fallback_version: str,
        output_path: Path,
    ) -> OpenCodeEnvironmentProbeV1:
        self.tool_root.mkdir(parents=True, exist_ok=True)
        version = requested_version or fallback_version
        wsl_tool = self._wsl_path(self.tool_root)
        quoted_tool = shlex.quote(wsl_tool)
        command = [
            "wsl.exe",
            "-d",
            "Ubuntu",
            "--",
            "bash",
            "-lc",
            (
                f"set -e; mkdir -p {quoted_tool}; "
                f"npm install --prefix {quoted_tool} "
                f"opencode-ai@{version} >/dev/null; "
                f"{quoted_tool}/node_modules/.bin/opencode --version; "
                "python3 --version; "
                "python3 -c 'import openpyxl; print(openpyxl.__version__)'"
            ),
        ]
        first_failure = None
        detected = None
        python_version = None
        openpyxl_version = None
        try:
            result = subprocess.run(
                command,
                text=True,
                encoding="utf-8",
                errors="replace",
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=600,
                check=False,
            )
            if result.returncode != 0:
                raise RuntimeError(
                    "wsl_opencode_prepare_failed:" + result.stderr[-1000:]
                )
            lines = [line.strip() for line in result.stdout.splitlines() if line.strip()]
            detected = lines[0] if lines else None
            python_version = lines[1] if len(lines) > 1 else None
            openpyxl_version = lines[2] if len(lines) > 2 else None
            if not detected or not python_version or not openpyxl_version:
                raise RuntimeError("wsl_environment_probe_incomplete")
            status = "pass"
        except Exception as exc:
            status = "infrastructure_failed"
            first_failure = f"{type(exc).__name__}:{exc}"
        probe = OpenCodeEnvironmentProbeV1(
            environment_kind="wsl2_ubuntu_opencode",
            status=status,
            template_or_distribution="Ubuntu",
            opencode_version=detected,
            python_version=python_version,
            openpyxl_version=openpyxl_version,
            first_failure=first_failure,
        )
        _atomic_json(output_path, probe.model_dump(mode="json"))
        return probe

    def run_one(
        self,
        *,
        kind: Literal["public_probe", "private_task"],
        task_id: str,
        uploads: Dict[str, bytes],
        prompt: str,
        expected: str,
        timeout: int,
        deepseek_key: str,
        output_dir: Path,
    ) -> OpenCodeProcessOutcomeV1:
        workspace = output_dir / "workspace"
        workspace.mkdir(parents=True, exist_ok=True)
        for relative, data in uploads.items():
            path = workspace / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        (workspace / "deliverable_files").mkdir(exist_ok=True)
        (workspace / "TASK.md").write_text(prompt, encoding="utf-8")
        input_tree = _tree(workspace)
        input_hashes = {
            _sha_bytes(value)
            for name, value in uploads.items()
            if name.lower().endswith(".xlsx")
        }
        wsl_workspace = self._wsl_path(workspace)
        wsl_tool = self._wsl_path(self.tool_root)
        quoted_workspace = shlex.quote(wsl_workspace)
        quoted_tool = shlex.quote(wsl_tool)
        command = [
            "wsl.exe",
            "-d",
            "Ubuntu",
            "--",
            "bash",
            "-lc",
            (
                f"set +e; cd {quoted_workspace}; PROMPT=$(cat TASK.md); "
                f"{quoted_tool}/node_modules/.bin/opencode run --format json "
                "--model deepseek/deepseek-v4-pro --auto "
                f"--dir {quoted_workspace} "
                "\"$PROMPT\"; code=$?; "
                "printf '{\"type\":\"runner.exit\",\"exit_code\":%s}\\n' "
                "\"$code\"; exit 0"
            ),
        ]
        env = os.environ.copy()
        env["DEEPSEEK_API_KEY"] = deepseek_key
        current_wslenv = env.get("WSLENV", "")
        env["WSLENV"] = (
            f"{current_wslenv}:DEEPSEEK_API_KEY"
            if current_wslenv
            else "DEEPSEEK_API_KEY"
        )
        started_at = _now()
        started = time.monotonic()
        process = self.popen_factory(
            command,
            cwd=str(workspace),
            env=env,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            creationflags=(
                subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0
            ),
            start_new_session=os.name != "nt",
        )
        timed_out = False
        try:
            stdout, stderr = process.communicate(timeout=timeout)
        except subprocess.TimeoutExpired as exc:
            timed_out = True
            _terminate_process_tree(process)
            tail_out, tail_err = process.communicate()
            stdout = (exc.stdout or "") + (tail_out or "")
            stderr = (exc.stderr or "") + (tail_err or "")
        stdout = _redact(stdout or "", [deepseek_key])
        stderr = _redact(stderr or "", [deepseek_key])
        returncode = _runner_exit_code(stdout)
        diagnostics = parse_opencode_jsonl(stdout)
        delivery = inspect_delivery(workspace, expected, input_hashes)
        failure = _failure_code(
            diagnostics,
            stderr,
            timed_out=timed_out,
            returncode=returncode,
        )
        if failure:
            status = "infrastructure_failed"
        elif (
            returncode == 0
            and diagnostics.turn_event_seen
            and diagnostics.tool_event_seen
            and diagnostics.shell_or_python_seen
            and delivery.valid
        ):
            status = "pass"
        else:
            status = "task_failed"
            failure = delivery.failure or "opencode_or_protocol_failure"
        output_dir.mkdir(parents=True, exist_ok=True)
        jsonl_path = output_dir / "opencode.jsonl"
        stderr_path = output_dir / "stderr.txt"
        jsonl_path.write_text(stdout, encoding="utf-8")
        stderr_path.write_text(stderr, encoding="utf-8")
        outcome = OpenCodeProcessOutcomeV1(
            kind=kind,
            blind_task_id=task_id,
            environment_kind="wsl2_ubuntu_opencode",
            status=status,
            started_at=started_at,
            completed_at=_now(),
            duration_seconds=round(time.monotonic() - started, 3),
            returncode=returncode,
            timed_out=timed_out,
            command_identity=[
                "opencode",
                "run",
                "--format",
                "json",
                "--model",
                "deepseek/deepseek-v4-pro",
                "--auto",
            ],
            workspace_path=str(workspace),
            workspace_input_tree_sha256=input_tree,
            jsonl_path=str(jsonl_path),
            jsonl_sha256=_sha_file(jsonl_path),
            stderr_path=str(stderr_path),
            stderr_sha256=_sha_file(stderr_path),
            diagnostics=diagnostics,
            delivery=delivery.model_dump(mode="json"),
            output_tree_sha256=_tree(workspace),
            first_failure=failure,
        )
        _assert_secrets_absent(outcome.model_dump_json(), [deepseek_key])
        _atomic_json(
            output_dir / "outcome.json", outcome.model_dump(mode="json")
        )
        return outcome


class OfficialDeepSeekOpenCodeRunner:
    def __init__(
        self,
        *,
        scope_path: str | Path,
        receipt_path: str | Path,
        output_root: str | Path,
        repository_root: str | Path,
        deepseek_key: str,
        e2b_key: str,
        e2b_backend: Optional[E2BOpenCodeBackend] = None,
        wsl_backend: Optional[WSLOpenCodeBackend] = None,
        provider_probe_runner=run_official_provider_probe,
    ):
        self.scope_path = Path(scope_path).resolve()
        self.receipt_path = Path(receipt_path).resolve()
        self.output = Path(output_root).resolve()
        self.repository_root = Path(repository_root).resolve()
        self.deepseek_key = deepseek_key
        self.e2b_key = e2b_key
        self.scope = OfficialDeepSeekOpenCodeCampaignV1.model_validate_json(
            self.scope_path.read_text(encoding="utf-8")
        )
        self.e2b = e2b_backend or E2BOpenCodeBackend()
        self.wsl = wsl_backend or WSLOpenCodeBackend(
            repository_root=self.repository_root
        )
        self.provider_probe_runner = provider_probe_runner
        self._validate()

    def _validate(self) -> None:
        if not self.deepseek_key:
            raise ValueError("deepseek_api_key_missing")
        if not self.e2b_key:
            raise ValueError("e2b_api_key_missing")
        if _sha_file(self.scope_path) != self.scope_path.stem:
            raise PermissionError("official_opencode_scope_hash_mismatch")
        if (
            governed_source_fingerprint(self.repository_root)
            != self.scope.source_fingerprint
        ):
            raise PermissionError("official_opencode_source_drift")
        for path, digest in (
            (
                self.scope.representative_campaign_path,
                self.scope.representative_campaign_sha256,
            ),
            (
                self.scope.baseline_solver_manifest_path,
                self.scope.baseline_solver_manifest_sha256,
            ),
            (
                self.scope.baseline_grader_outcome_path,
                self.scope.baseline_grader_outcome_sha256,
            ),
            (
                self.scope.parity_report_path,
                self.scope.parity_report_sha256,
            ),
        ):
            if _sha_file(Path(path)) != digest:
                raise PermissionError("official_opencode_upstream_drift")
        for binding in self.scope.task_bindings:
            candidate = Path(binding.candidate_package_path)
            if (
                _tree(candidate) != binding.candidate_tree_sha256
                or _projection_tree(candidate)
                != binding.candidate_projection_sha256
            ):
                raise PermissionError("official_opencode_candidate_tree_drift")

    def _new_manifest(self) -> OfficialDeepSeekOpenCodeExecutionV1:
        return OfficialDeepSeekOpenCodeExecutionV1(
            scope_path=str(self.scope_path),
            scope_sha256=_sha_file(self.scope_path),
            receipt_path=str(self.receipt_path),
            receipt_sha256=_sha_file(self.receipt_path),
            status="prepared",
            tasks={
                item.blind_task_id: OpenCodeTaskStateV1()
                for item in self.scope.task_bindings
            },
            created_at=_now(),
            updated_at=_now(),
        )

    def execute(self) -> tuple[OfficialDeepSeekOpenCodeExecutionV1, Path]:
        manifest_path = self.output / "execution_manifest.json"
        if manifest_path.exists():
            manifest = OfficialDeepSeekOpenCodeExecutionV1.model_validate_json(
                manifest_path.read_text(encoding="utf-8")
            )
        else:
            receipt = OfficialDeepSeekOpenCodeReceiptV1.model_validate_json(
                self.receipt_path.read_text(encoding="utf-8")
            )
            if (
                receipt.scope_sha256 != _sha_file(self.scope_path)
                or receipt.consumed_at is not None
            ):
                raise PermissionError("official_opencode_receipt_invalid")
            receipt.consumed_at = _now()
            _atomic_json(
                self.receipt_path, receipt.model_dump(mode="json")
            )
            manifest = self._new_manifest()
            manifest.receipt_sha256 = _sha_file(self.receipt_path)
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        for record in manifest.tasks.values():
            if record.status == "running":
                record.status = "interrupted"
        if any(record.status == "interrupted" for record in manifest.tasks.values()):
            manifest.status = "incomplete"
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            return manifest, manifest_path
        manifest.status = "running"
        manifest.updated_at = _now()
        _atomic_json(manifest_path, manifest.model_dump(mode="json"))

        provider_path = self.output / "provider_probe.json"
        if not provider_path.exists():
            provider = self.provider_probe_runner(
                self.deepseek_key, provider_path
            )
        else:
            provider = OfficialDeepSeekProviderProbeV1.model_validate_json(
                provider_path.read_text(encoding="utf-8")
            )
        manifest.provider_probe_path = str(provider_path)
        manifest.provider_probe_sha256 = _sha_file(provider_path)
        if provider.status != "pass":
            manifest.status = "provider_incomplete"
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            return manifest, manifest_path

        selection = self._select_environment(manifest)
        if selection is None:
            manifest.status = "stack_incomplete"
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            return manifest, manifest_path
        selection_path = self.output / "environment_selection.json"
        _atomic_json(selection_path, selection.model_dump(mode="json"))
        manifest.selection_path = str(selection_path)
        manifest.selection_sha256 = _sha_file(selection_path)
        # The selected probe is stored at the fixed path written below.
        selected_probe_path = self.output / "public_agent_probe" / "outcome.json"
        manifest.public_probe_path = str(selected_probe_path)
        manifest.public_probe_sha256 = _sha_file(selected_probe_path)
        _atomic_json(manifest_path, manifest.model_dump(mode="json"))

        bindings = {
            item.blind_task_id: item for item in self.scope.task_bindings
        }
        ordered = task_execution_order(list(bindings))
        consecutive_code = None
        consecutive_count = 0
        for index, task_id in enumerate(ordered):
            if index == 2:
                matched = [
                    OpenCodeProcessOutcomeV1.model_validate_json(
                        Path(manifest.tasks[key].outcome_path).read_text(
                            encoding="utf-8"
                        )
                    )
                    for key in MATCHED_TASK_IDS
                ]
                if not matched_extension_allowed(matched):
                    manifest.status = "matched_failed"
                    break
            record = manifest.tasks[task_id]
            if record.status != "not_started":
                continue
            binding = bindings[task_id]
            uploads, prompt = _private_uploads(binding)
            record.status = "running"
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
            outcome = self._run_selected(
                selection=selection,
                kind="private_task",
                task_id=task_id,
                uploads=uploads,
                prompt=prompt,
                expected=binding.expected_deliverable,
                timeout=self.scope.task_timeout_seconds,
                output_dir=self.output / "tasks" / task_id,
            )
            outcome_path = self.output / "tasks" / task_id / "outcome.json"
            record.status = "completed"
            record.outcome_path = str(outcome_path)
            record.outcome_sha256 = _sha_file(outcome_path)
            code = (
                outcome.first_failure
                if outcome.status == "infrastructure_failed"
                else None
            )
            if code and code == consecutive_code:
                consecutive_count += 1
            elif code:
                consecutive_code = code
                consecutive_count = 1
            else:
                consecutive_code = None
                consecutive_count = 0
            if consecutive_count >= SYSTEMIC_FAILURE_THRESHOLD:
                manifest.status = "incomplete"
                manifest.systemic_failure_code = code
                break
            manifest.updated_at = _now()
            _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        if all(record.status == "completed" for record in manifest.tasks.values()):
            manifest.status = "solver_completed"
        elif manifest.status == "running":
            manifest.status = "incomplete"
        manifest.updated_at = _now()
        _atomic_json(manifest_path, manifest.model_dump(mode="json"))
        return manifest, manifest_path

    def _select_environment(
        self, manifest: OfficialDeepSeekOpenCodeExecutionV1
    ) -> Optional[SolverEnvironmentSelectionV1]:
        env_dir = self.output / "environment_probes"
        env_dir.mkdir(parents=True, exist_ok=True)
        e2b_env_path = env_dir / "e2b_official.json"
        e2b_env = self.e2b.inspect_environment(
            template=self.scope.e2b_template,
            e2b_key=self.e2b_key,
            output_path=e2b_env_path,
        )
        manifest.environment_probe_paths.append(str(e2b_env_path))
        if e2b_env.status == "pass":
            uploads, expected, prompt = _probe_uploads()
            probe = self.e2b.run_one(
                template=self.scope.e2b_template,
                environment_kind="e2b_official_opencode",
                kind="public_probe",
                task_id="public_official_deepseek_opencode_probe",
                uploads=uploads,
                prompt=prompt,
                expected=expected,
                timeout=self.scope.public_timeout_seconds,
                deepseek_key=self.deepseek_key,
                e2b_key=self.e2b_key,
                output_dir=self.output / "public_agent_probe",
            )
            if probe.status == "pass":
                return SolverEnvironmentSelectionV1(
                    selected_environment="e2b_official_opencode",
                    opencode_version=e2b_env.opencode_version or "unknown",
                    template_or_distribution=self.scope.e2b_template,
                    environment_probe_sha256=_sha_file(e2b_env_path),
                    public_agent_probe_sha256=_sha_file(
                        self.output / "public_agent_probe" / "outcome.json"
                    ),
                )
            if probe.status != "infrastructure_failed":
                return None
        wsl_env_path = env_dir / "wsl2_ubuntu.json"
        wsl_env = self.wsl.prepare(
            requested_version=e2b_env.opencode_version or "",
            fallback_version=self.scope.stable_fallback_opencode_version,
            output_path=wsl_env_path,
        )
        manifest.environment_probe_paths.append(str(wsl_env_path))
        if wsl_env.status != "pass":
            return None
        uploads, expected, prompt = _probe_uploads()
        probe = self.wsl.run_one(
            kind="public_probe",
            task_id="public_official_deepseek_opencode_probe",
            uploads=uploads,
            prompt=prompt,
            expected=expected,
            timeout=self.scope.public_timeout_seconds,
            deepseek_key=self.deepseek_key,
            output_dir=self.output / "public_agent_probe",
        )
        if probe.status != "pass":
            return None
        return SolverEnvironmentSelectionV1(
            selected_environment="wsl2_ubuntu_opencode",
            opencode_version=wsl_env.opencode_version or "unknown",
            template_or_distribution="Ubuntu",
            environment_probe_sha256=_sha_file(wsl_env_path),
            public_agent_probe_sha256=_sha_file(
                self.output / "public_agent_probe" / "outcome.json"
            ),
        )

    def _run_selected(self, *, selection, **kwargs):
        if selection.selected_environment.startswith("e2b_"):
            return self.e2b.run_one(
                template=selection.template_or_distribution,
                environment_kind=selection.selected_environment,
                deepseek_key=self.deepseek_key,
                e2b_key=self.e2b_key,
                **kwargs,
            )
        return self.wsl.run_one(deepseek_key=self.deepseek_key, **kwargs)


def grade_official_opencode_deliveries(
    *,
    scope_path: str | Path,
    execution_path: str | Path,
    output_root: str | Path,
    codex_home: str | Path,
    python_executable: str | Path,
) -> tuple[OfficialDeepSeekOpenCodeGraderV1, Path]:
    scope_file = Path(scope_path).resolve()
    execution_file = Path(execution_path).resolve()
    output = Path(output_root).resolve()
    scope = OfficialDeepSeekOpenCodeCampaignV1.model_validate_json(
        scope_file.read_text(encoding="utf-8")
    )
    execution = OfficialDeepSeekOpenCodeExecutionV1.model_validate_json(
        execution_file.read_text(encoding="utf-8")
    )
    if execution.scope_sha256 != _sha_file(scope_file):
        raise PermissionError("official_opencode_grader_scope_drift")
    path = output / "grader_manifest.json"
    if path.exists():
        manifest = OfficialDeepSeekOpenCodeGraderV1.model_validate_json(
            path.read_text(encoding="utf-8")
        )
    else:
        manifest = OfficialDeepSeekOpenCodeGraderV1(
            scope_sha256=_sha_file(scope_file),
            execution_sha256=_sha_file(execution_file),
            baseline_grader_outcome_sha256=scope.baseline_grader_outcome_sha256,
            status="prepared",
            records={
                item.blind_task_id: ExternalComparisonGradeRecordV1(
                    blind_task_id=item.blind_task_id,
                    attempt_count=0,
                )
                for item in scope.task_bindings
            },
            created_at=_now(),
            updated_at=_now(),
        )
        _atomic_json(path, manifest.model_dump(mode="json"))
    if manifest.execution_sha256 != _sha_file(execution_file):
        raise PermissionError("official_opencode_execution_drift")
    bindings = {item.blind_task_id: item for item in scope.task_bindings}
    manifest.status = "running"
    _atomic_json(path, manifest.model_dump(mode="json"))
    for task_id in sorted(bindings):
        record = manifest.records[task_id]
        if record.status == "running":
            record.status = "interrupted"
            continue
        if record.status != "not_started":
            continue
        task = execution.tasks[task_id]
        if not task.outcome_path:
            record.status = "not_eligible"
            record.first_failure = "solver_outcome_missing_not_eligible"
            continue
        solver = OpenCodeProcessOutcomeV1.model_validate_json(
            Path(task.outcome_path).read_text(encoding="utf-8")
        )
        if solver.status != "pass" or not solver.delivery.get("valid"):
            record.status = "not_eligible"
            record.first_failure = "invalid_delivery_not_eligible"
            continue
        binding = bindings[task_id]
        candidate = Path(binding.candidate_package_path)
        teacher = Path(binding.source_package_path)
        delivery = Path(solver.workspace_path) / solver.delivery["relative_path"]
        compact = CompactGraderBindingV2(
            blind_task_id=task_id,
            route_id=binding.route_id,
            motif=binding.motif,
            replicate_id=binding.replicate_id,
            domain=binding.domain,
            candidate_package_path=str(candidate),
            teacher_package_path=str(teacher),
            solver_outcome_sha256=_sha_file(Path(task.outcome_path)),
            delivery_path=str(delivery),
            delivery_sha256=_sha_file(delivery),
            rubric_sha256=_sha_file(
                teacher / "teacher" / "rubric_plan_v2.json"
            ),
            fact_anchors_sha256=_sha_file(
                teacher / "teacher" / "deterministic_fact_anchors.json"
            ),
        )
        record.status = "running"
        record.attempt_count = 1
        _atomic_json(path, manifest.model_dump(mode="json"))
        runner = object.__new__(CodexLocalGraderRunnerV1)
        runner.output = output / "grader"
        runner.codex_home = Path(codex_home).resolve()
        runner.python = Path(python_executable).resolve()
        runner.popen_factory = subprocess.Popen
        runner.scope = SimpleNamespace(
            cli=None,
            timeout_seconds_per_task=TASK_TIMEOUT_SECONDS,
        )
        runner.scope.cli = scope.grader_cli
        local = runner._run_one(compact)
        manifest.records[task_id] = ExternalComparisonGradeRecordV1(
            **local.model_dump(mode="json")
        )
        manifest.updated_at = _now()
        _atomic_json(path, manifest.model_dump(mode="json"))
    manifest.status = (
        "completed"
        if all(
            item.status
            in {
                "completed",
                "not_eligible",
                "infrastructure_failed",
                "grading_failed",
                "interrupted",
            }
            for item in manifest.records.values()
        )
        else "incomplete"
    )
    manifest.updated_at = _now()
    _atomic_json(path, manifest.model_dump(mode="json"))
    return manifest, path


def aggregate_official_opencode_comparison(
    *,
    scope_path: str | Path,
    execution_path: str | Path,
    grader_path: str | Path,
    output_path: str | Path,
    execution_audit_path: str | Path | None = None,
) -> tuple[OfficialDeepSeekOpenCodeResultV1, Path]:
    scope_file = Path(scope_path).resolve()
    execution_file = Path(execution_path).resolve()
    grader_file = Path(grader_path).resolve()
    scope = OfficialDeepSeekOpenCodeCampaignV1.model_validate_json(
        scope_file.read_text(encoding="utf-8")
    )
    execution = OfficialDeepSeekOpenCodeExecutionV1.model_validate_json(
        execution_file.read_text(encoding="utf-8")
    )
    grader = OfficialDeepSeekOpenCodeGraderV1.model_validate_json(
        grader_file.read_text(encoding="utf-8")
    )
    audit_file = (
        Path(execution_audit_path).resolve()
        if execution_audit_path is not None
        else None
    )
    audit = (
        OfficialDeepSeekOpenCodeExecutionAuditV1.model_validate_json(
            audit_file.read_text(encoding="utf-8")
        )
        if audit_file is not None
        else None
    )
    if not (
        execution.scope_sha256 == _sha_file(scope_file)
        and grader.scope_sha256 == _sha_file(scope_file)
        and grader.execution_sha256 == _sha_file(execution_file)
        and grader.baseline_grader_outcome_sha256
        == scope.baseline_grader_outcome_sha256
        and (
            audit is None
            or (
                audit.scope_sha256 == _sha_file(scope_file)
                and audit.execution_sha256 == _sha_file(execution_file)
            )
        )
    ):
        raise PermissionError("official_opencode_aggregate_hash_drift")
    baseline_outcome = CodexLocalGraderOutcomeV1.model_validate_json(
        Path(scope.baseline_grader_outcome_path).read_text(encoding="utf-8")
    )
    baseline_solver = CodexLocalCampaignManifestV1.model_validate_json(
        Path(scope.baseline_solver_manifest_path).read_text(encoding="utf-8")
    )
    bindings = {item.blind_task_id: item for item in scope.task_bindings}
    baseline_scores = {
        item.blind_task_id: item
        for item in baseline_outcome.screening_result.observations
    }
    observations: List[ExternalTaskObservationV1] = []
    audit_records = (
        {item.blind_task_id: item for item in audit.records}
        if audit is not None
        else {}
    )
    for task_id in sorted(bindings):
        binding = bindings[task_id]
        score = baseline_scores[task_id]
        baseline_process = json.loads(
            Path(baseline_solver.task_outcome_paths[task_id]).read_text(
                encoding="utf-8"
            )
        )
        observations.append(
            ExternalTaskObservationV1(
                stack_id=BASELINE_STACK_ID,
                blind_task_id=task_id,
                domain=binding.domain,
                motif=binding.motif,
                route_id=binding.route_id,
                attempted=True,
                infrastructure_complete=True,
                exact_valid_delivery=True,
                task_failed=False,
                grade_complete=True,
                major_defect=score.major_defect,
                professional_plausibility=(
                    "pass" if score.professional_plausibility_pass else "fail"
                ),
                weighted_score=score.weighted_score,
                duration_seconds=baseline_process.get("duration_seconds", 0.0),
                usage=baseline_process.get("diagnostics", {}).get("usage", {}),
            )
        )
        state = execution.tasks[task_id]
        audited = audit_records.get(task_id)
        counted = audited is None or audited.counted_in_cohort
        solver = (
            OpenCodeProcessOutcomeV1.model_validate_json(
                Path(state.outcome_path).read_text(encoding="utf-8")
            )
            if state.outcome_path
            else None
        )
        grade = grader.records[task_id]
        review = (
            CompactGraderReviewV2.model_validate_json(
                Path(grade.review_path).read_text(encoding="utf-8")
            )
            if grade.status == "completed" and grade.review_path
            else None
        )
        observations.append(
            ExternalTaskObservationV1(
                stack_id=STACK_ID,
                blind_task_id=task_id,
                domain=binding.domain,
                motif=binding.motif,
                route_id=binding.route_id,
                attempted=solver is not None and counted,
                infrastructure_complete=(
                    solver is not None
                    and counted
                    and (
                        audited.audited_status
                        != "infrastructure_failed"
                        if audited
                        else solver.status != "infrastructure_failed"
                    )
                ),
                exact_valid_delivery=(
                    solver is not None
                    and counted
                    and solver.status == "pass"
                    and solver.delivery.get("valid", False)
                ),
                task_failed=(
                    solver is not None
                    and counted
                    and (
                        audited.audited_status == "task_failed"
                        if audited
                        else solver.status == "task_failed"
                    )
                ),
                grade_complete=review is not None and counted,
                major_defect=(
                    review.major_defect if review and counted else None
                ),
                professional_plausibility=(
                    review.professional_plausibility
                    if review and counted
                    else None
                ),
                weighted_score=(
                    review.weighted_score if review and counted else None
                ),
                duration_seconds=(
                    solver.duration_seconds if solver and counted else 0.0
                ),
                usage=(
                    solver.diagnostics.usage if solver and counted else {}
                ),
            )
        )
    provider_pass = False
    if execution.provider_probe_path:
        provider_pass = (
            OfficialDeepSeekProviderProbeV1.model_validate_json(
                Path(execution.provider_probe_path).read_text(encoding="utf-8")
            ).status
            == "pass"
        )
    public_pass = False
    if execution.public_probe_path:
        public_pass = (
            OpenCodeProcessOutcomeV1.model_validate_json(
                Path(execution.public_probe_path).read_text(encoding="utf-8")
            ).status
            == "pass"
        )
    baseline_summary = _summary(
        BASELINE_STACK_ID, observations, public_probe_pass=True
    )
    deepseek_summary = _summary(
        STACK_ID,
        observations,
        public_probe_pass=provider_pass and public_pass,
    )
    pair = _pairwise(
        baseline_summary, deepseek_summary, observations
    )
    ready = (
        execution.status == "solver_completed"
        and (
            audit is None
            or audit.audited_execution_status == "solver_completed"
        )
        and deepseek_summary.capability_comparison_eligible
    )
    if ready:
        decision = "official_deepseek_comparison_ready"
    elif (
        provider_pass
        and public_pass
        and (
            audit is None
            or audit.audited_execution_status == "solver_completed"
        )
    ):
        decision = "official_deepseek_partial"
    else:
        decision = "stack_incomplete"
    winner = {
        "left_substantive_winner": "gpt_baseline",
        "right_substantive_winner": "official_deepseek",
        "practical_tie": "practical_tie",
        "insufficient_common_coverage": "insufficient_common_coverage",
    }[pair.decision]
    result = OfficialDeepSeekOpenCodeResultV1(
        campaign_id=scope.campaign_id,
        scope_sha256=_sha_file(scope_file),
        execution_sha256=_sha_file(execution_file),
        grader_sha256=_sha_file(grader_file),
        execution_audit_sha256=(
            _sha_file(audit_file) if audit_file is not None else None
        ),
        observations=observations,
        stack_summaries=[baseline_summary, deepseek_summary],
        pairwise_comparison=pair,
        decision=decision,
        practical_winner=winner,
    )
    destination = Path(output_path).resolve()
    _atomic_json(destination, result.model_dump(mode="json"))
    return result, destination
