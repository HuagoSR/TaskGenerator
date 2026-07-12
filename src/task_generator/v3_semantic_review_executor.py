from __future__ import annotations

import hashlib
import json
import os
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Type
from xml.etree import ElementTree

from pydantic import BaseModel

from task_generator.v3_semantic_validity import (
    CandidateBlindReview,
    SecondarySemanticReview,
    SemanticFinding,
    TeacherRubricReview,
)
from task_generator.v3_skill_extractor import ProviderConfig, build_deepseek_config, build_tuzi_config, load_env_file
from task_generator.v3_semantic_secondary_cost import SecondaryCostLedgerManager


class SemanticReviewExecutionError(RuntimeError):
    pass


class SemanticReviewExecutor:
    """Execute bounded JSON-only semantic reviews without exposing secrets in artifacts."""

    def __init__(
        self,
        config: ProviderConfig,
        max_tokens: int = 8000,
        cost_ledger: Optional[SecondaryCostLedgerManager] = None,
        input_token_hard_limit: int = 20_000,
    ) -> None:
        self.config = config
        self.max_tokens = max_tokens
        self.cost_ledger = cost_ledger
        self.input_token_hard_limit = input_token_hard_limit
        self.last_diagnostics: Dict[str, Any] = {}

    def review_blind(self, package: Dict[str, Any]) -> CandidateBlindReview:
        enriched = dict(package)
        enriched["reference_files"] = [
            {key: value for key, value in item.items() if key != "path"}
            for item in package.get("reference_files") or []
        ]
        enriched["reference_contents"] = [
            self._reference_content(Path(item["path"])) for item in package.get("reference_files") or []
        ]
        instructions = (
            "Act as an adversarial candidate-blind task auditor. Use only the candidate-visible prompt, "
            "deliverable contract, and reference contents. Determine whether every requirement is supported, "
            "uniquely answerable, and free of hidden assumptions. Return JSON matching CandidateBlindReview. "
            "Use only the documented semantic finding codes. Do not infer teacher truth or repair files."
        )
        review = self._call(
            instructions,
            enriched,
            CandidateBlindReview,
            {
                "task_id": package.get("task_id"),
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "prompt_package_sha256": hashlib.sha256(json.dumps(package, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest(),
                "status": "completed",
            },
        )
        expected_ids = {str(item.get("requirement_id")) for item in package.get("semantic_requirements") or []}
        returned_ids = {item.requirement_id for item in review.requirement_reviews}
        if not expected_ids or returned_ids != expected_ids:
            self.last_diagnostics.update({
                "requirement_coverage_valid": False,
                "expected_requirement_count": len(expected_ids),
                "returned_requirement_count": len(returned_ids),
                "missing_requirement_count": len(expected_ids - returned_ids),
                "unexpected_requirement_count": len(returned_ids - expected_ids),
            })
            raise SemanticReviewExecutionError("Provider contract violation: requirement review coverage mismatch.")
        self._clear_provider_corroboration(review)
        return review

    def review_teacher(self, package: Dict[str, Any]) -> TeacherRubricReview:
        expanded = {
            "package_version": package.get("package_version"),
            "task_id": package.get("task_id"),
            "blind_review_sha256": (package.get("blind_review") or {}).get("sha256"),
            "semantic_contract_sha256": (package.get("semantic_contract") or {}).get("sha256"),
        }
        expanded["blind_review_payload"] = self._read_json(Path(package["blind_review"]["path"]))
        expanded["semantic_contract_payload"] = self._read_json(Path(package["semantic_contract"]["path"]))
        expanded["teacher_artifact_payloads"] = {
            name: self._read_or_summarize(Path(item["path"]))
            for name, item in (package.get("teacher_artifacts") or {}).items()
        }
        instructions = (
            "Audit teacher truth, GoldenRun, and rubric against the already-frozen candidate-blind review. "
            "Find unsupported teacher assumptions, uncovered requirements, irrelevant criteria, and weight imbalance. "
            "Return JSON matching TeacherRubricReview. Findings and repair proposals are advisory only."
        )
        review = self._call(
            instructions,
            expanded,
            TeacherRubricReview,
            {
                "task_id": package.get("task_id"),
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "blind_review_sha256": (package.get("blind_review") or {}).get("sha256"),
                "status": "completed",
            },
        )
        self._clear_provider_corroboration(review)
        return review

    def review_secondary(
        self,
        package: Dict[str, Any],
        primary_findings: List[SemanticFinding],
        review_scope: str,
    ) -> SecondarySemanticReview:
        material = [item for item in primary_findings if item.severity == "blocking"][:6]
        if review_scope == "candidate_blind":
            requirement_ids = {item.requirement_id for item in material if item.requirement_id}
            evidence = {
                "requirements": [item for item in package.get("semantic_requirements") or [] if not requirement_ids or item.get("requirement_id") in requirement_ids],
                "deliverable_contract": package.get("deliverable_contract"),
                "relevant_candidate_evidence": [
                    self._compact_reference_content(Path(item["path"]))
                    for item in package.get("reference_files") or []
                ],
            }
        elif review_scope == "teacher_rubric":
            contract = self._read_json(Path(package["semantic_contract"]["path"]))
            claim_ids = {item.claim_id for item in material if item.claim_id}
            requirement_ids = {item.requirement_id for item in material if item.requirement_id}
            claims = [item for item in contract.get("claims") or [] if not claim_ids or item.get("claim_id") in claim_ids]
            if requirement_ids and not claim_ids:
                claims = [item for item in claims if item.get("requirement_id") in requirement_ids]
            teacher_payloads = {
                name: self._read_json(Path(item["path"]))
                for name, item in (package.get("teacher_artifacts") or {}).items()
                if name in {"golden_run", "teacher_truth", "deterministic_answer_key", "rubric"}
                and Path(item["path"]).suffix.lower() == ".json"
            }
            evidence = {
                "resolved_claims": claims,
                "validator_results": contract.get("validator_results") or contract.get("validators") or [],
                "rubric_bindings": contract.get("rubric_bindings") or [],
                "relevant_teacher_artifacts": teacher_payloads,
            }
        else:
            raise SemanticReviewExecutionError(f"Unknown secondary review scope: {review_scope}")
        payload = {
            "task_id": package.get("task_id"),
            "review_scope": review_scope,
            "primary_material_findings": [
                {
                    "finding_code": item.finding_code,
                    "requirement_id": item.requirement_id,
                    "claim_id": item.claim_id,
                    "message": item.message,
                    "evidence_locators": item.evidence_locators,
                }
                for item in material
            ],
            "evidence": evidence,
        }
        instructions = (
            "Act as an independent secondary adjudicator. For each material issue family supported by the evidence, "
            "return one compact decision. Do not recreate the primary review, do not repair artifacts, and do not "
            "treat style suggestions as material. Use only the enumerated finding families."
        )
        self._enforce_secondary_prompt_budget(payload)
        return self._call(
            instructions,
            payload,
            SecondarySemanticReview,
            {
                "task_id": package.get("task_id"),
                "review_scope": review_scope,
                "model": self.config.model,
                "provider": self.config.provider_name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "status": "completed",
            },
            cost_scope=review_scope,
        )

    def _clear_provider_corroboration(self, review: BaseModel) -> None:
        findings = list(getattr(review, "findings", []) or [])
        for requirement in getattr(review, "requirement_reviews", []) or []:
            findings.extend(requirement.findings)
        for finding in findings:
            finding.deterministic_corroboration = False

    def _call(
        self,
        system_prompt: str,
        payload: Dict[str, Any],
        model_type: Type[BaseModel],
        defaults: Dict[str, Any],
        cost_scope: Optional[str] = None,
    ):
        try:
            from openai import OpenAI
        except Exception as exc:  # pragma: no cover
            raise SemanticReviewExecutionError(f"OpenAI SDK unavailable: {type(exc).__name__}") from exc
        client = OpenAI(api_key=self.config.api_key, base_url=self.config.base_url, timeout=self.config.timeout_seconds)
        started = time.monotonic()
        schema = model_type.model_json_schema()
        reserved = self.cost_ledger.reserve(self.input_token_hard_limit, self.max_tokens) if self.cost_ledger else 0.0
        response = client.chat.completions.create(
            model=self.config.model,
            messages=[
                {
                    "role": "system",
                    "content": system_prompt
                    + " Output exactly one JSON object matching this JSON Schema. Do not add wrapper keys or prose: "
                    + json.dumps(schema, ensure_ascii=False),
                },
                {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
            ],
            temperature=0,
            max_tokens=self.max_tokens,
            response_format={"type": "json_object"},
        )
        content = response.choices[0].message.content or ""
        finish_reason = str(getattr(response.choices[0], "finish_reason", ""))
        usage = getattr(response, "usage", None)
        self.last_diagnostics = {
            "provider": self.config.provider_name,
            "model": self.config.model,
            "finish_reason": finish_reason,
            "response_char_count": len(content),
            "response_sha256": hashlib.sha256(content.encode("utf-8")).hexdigest(),
            "prompt_tokens": getattr(usage, "prompt_tokens", None),
            "completion_tokens": getattr(usage, "completion_tokens", None),
            "total_tokens": getattr(usage, "total_tokens", None),
            "duration_seconds": round(time.monotonic() - started, 3),
            "likely_truncated": finish_reason.lower() == "length",
            "key_slot": "backup" if self.cost_ledger else None,
        }
        if self.cost_ledger and cost_scope:
            self.cost_ledger.record(
                scope=cost_scope,
                attempt=1,
                prompt_tokens=int(getattr(usage, "prompt_tokens", 0) or 0),
                completion_tokens=int(getattr(usage, "completion_tokens", 0) or 0),
                reserved_rmb=reserved,
                status="completed" if content and finish_reason.lower() != "length" else "failed",
            )
        if not content:
            raise SemanticReviewExecutionError("Provider returned empty semantic review output.")
        if self.last_diagnostics["likely_truncated"]:
            raise SemanticReviewExecutionError("Provider truncated semantic review output.")
        try:
            parsed = json.loads(content)
            if not isinstance(parsed, dict):
                raise ValueError("Semantic review response must be a JSON object.")
            parsed.update(defaults)
            return model_type.model_validate(parsed)
        except Exception as exc:
            raise SemanticReviewExecutionError(f"Provider contract violation: {type(exc).__name__}") from exc

    def _reference_content(self, path: Path) -> Dict[str, Any]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._xlsx_content(path)
        if suffix == ".docx":
            return self._docx_content(path)
        if suffix in {".md", ".txt", ".csv", ".json"}:
            return {"file_name": path.name, "kind": suffix[1:], "text": path.read_text(encoding="utf-8", errors="replace")[:50000]}
        return {"file_name": path.name, "kind": "unsupported_binary", "size_bytes": path.stat().st_size}

    def _compact_reference_content(self, path: Path) -> Dict[str, Any]:
        content = self._reference_content(path)
        if content.get("kind") == "xlsx":
            for sheet in content.get("sheets") or []:
                sheet["rows"] = (sheet.get("rows") or [])[:50]
        elif "text" in content:
            content["text"] = str(content["text"])[:8000]
        return content

    def _enforce_secondary_prompt_budget(self, payload: Dict[str, Any]) -> None:
        serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        estimated_tokens = (len(serialized) + 3) // 4
        self.last_diagnostics = {
            "secondary_payload_chars": len(serialized),
            "estimated_input_tokens": estimated_tokens,
            "input_token_hard_limit": self.input_token_hard_limit,
        }
        if estimated_tokens > self.input_token_hard_limit:
            raise SemanticReviewExecutionError("secondary_prompt_budget_exceeded")

    def _xlsx_content(self, path: Path) -> Dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover
            raise SemanticReviewExecutionError(f"openpyxl unavailable: {type(exc).__name__}") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                rows.append([value for value in row])
                if index >= 500:
                    break
            sheets.append({"sheet_name": sheet.title, "rows": rows})
        workbook.close()
        return {"file_name": path.name, "kind": "xlsx", "sheets": sheets}

    def _docx_content(self, path: Path) -> Dict[str, Any]:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        root = ElementTree.fromstring(xml)
        texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t")]
        return {"file_name": path.name, "kind": "docx", "text": "\n".join(texts)[:50000]}

    def _read_or_summarize(self, path: Path) -> Any:
        if path.suffix.lower() == ".json":
            return self._read_json(path)
        return self._reference_content(path)

    def _read_json(self, path: Path) -> Dict[str, Any]:
        return json.loads(path.read_text(encoding="utf-8"))


def deepseek_semantic_config(key_path: str | Path, timeout_seconds: int = 900) -> ProviderConfig:
    config = build_deepseek_config(key_path, "deepseek-v4-pro", timeout_seconds)
    if config is None:
        raise SemanticReviewExecutionError("DeepSeek semantic-review key is unavailable.")
    return config


def claude_semantic_config(env_path: str | Path, timeout_seconds: int = 900) -> ProviderConfig:
    config = build_tuzi_config(env_path, "claude-sonnet-4-6", timeout_seconds)
    if config is not None:
        return config
    values = load_env_file(env_path)
    api_key = values.get("AGENT_API_KEY") or values.get("GRADER_API_KEY")
    base_url = values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model="claude-sonnet-4-6",
        timeout_seconds=timeout_seconds,
    )


def gpt54_semantic_config(env_path: str | Path, timeout_seconds: int = 900) -> ProviderConfig:
    config = build_tuzi_config(env_path, "gpt-5.4-pro", timeout_seconds)
    if config is not None:
        return config
    values = load_env_file(env_path)
    api_key = values.get("AGENT_API_KEY") or values.get("GRADER_API_KEY")
    base_url = values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model="gpt-5.4-pro",
        timeout_seconds=timeout_seconds,
    )


def tuzi_backup_semantic_config(
    env_path: str | Path,
    model: str = "gpt-5.6-sol",
    timeout_seconds: int = 900,
) -> ProviderConfig:
    values = load_env_file(env_path)
    api_key = values.get("OPENAI_API_KEY_BACKUP")
    base_url = (
        values.get("OPENAI_BASE_URL") or values.get("AGENT_BASE_URL") or values.get("GRADER_BASE_URL")
        or os.environ.get("OPENAI_BASE_URL")
    )
    if not api_key or not base_url:
        raise SemanticReviewExecutionError("Tuzi backup semantic-review configuration is unavailable.")
    return ProviderConfig(
        provider_name="tuzi",
        base_url=base_url,
        api_key=api_key,
        model=model,
        timeout_seconds=timeout_seconds,
    )
