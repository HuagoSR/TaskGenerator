from __future__ import annotations

import json
import re
from pathlib import Path, PurePosixPath
from typing import Any, Dict, Iterable, List, Literal, Optional

from pydantic import BaseModel, ConfigDict, Field, model_validator


CreationMode = Literal["create", "copy_then_edit", "edit_provided_copy"]
ContractStatus = Literal["pass", "invalid"]
FindingSeverity = Literal["blocking", "warning", "info"]
DeliveryStatus = Literal["valid", "invalid"]

_FILE_TOKEN = re.compile(
    r"(?<![\w.-])([A-Za-z0-9][A-Za-z0-9_.-]*\.(?:xlsx|docx|json|csv|pdf|txt))(?![\w-])",
    re.IGNORECASE,
)
_OUTPUT_VERBS = re.compile(
    r"\b(save|submit|deliver|create|write|export|produce|populate|complete|replace|return)\b",
    re.IGNORECASE,
)


def _safe_relative_path(value: str, expected_parent: str | None = None) -> PurePosixPath:
    normalized = value.replace("\\", "/")
    path = PurePosixPath(normalized)
    if path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts):
        raise ValueError("deliverable_path_not_allowed")
    if expected_parent and (len(path.parts) != 2 or path.parts[0] != expected_parent):
        raise ValueError(f"deliverable_path_must_be_under:{expected_parent}")
    return path


class DeliverableSpecV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_name: str
    relative_path: str
    format: str
    creation_mode: CreationMode = "create"
    source_template: Optional[str] = None
    must_exist: bool = True
    must_be_nonempty: bool = True
    openability_check: bool = True

    @model_validator(mode="after")
    def validate_contract(self) -> "DeliverableSpecV1":
        relative = _safe_relative_path(self.relative_path, "deliverable_files")
        file_path = PurePosixPath(self.file_name.replace("\\", "/"))
        if len(file_path.parts) != 1 or file_path.name in {"", ".", ".."}:
            raise ValueError("deliverable_file_name_not_allowed")
        if relative.name != file_path.name:
            raise ValueError("deliverable_file_name_path_mismatch")
        suffix = Path(file_path.name).suffix.lower().lstrip(".")
        if not suffix or suffix != self.format.lower().lstrip("."):
            raise ValueError("deliverable_format_extension_mismatch")
        if self.creation_mode in {"copy_then_edit", "edit_provided_copy"}:
            if not self.source_template:
                raise ValueError("source_template_required")
            template = _safe_relative_path(self.source_template, "reference_files")
            if self.creation_mode == "edit_provided_copy" and template.name != file_path.name:
                raise ValueError("edit_provided_copy_requires_matching_basename")
        elif self.source_template:
            raise ValueError("source_template_forbidden_for_create")
        return self


class DeliverableContractV1(BaseModel):
    model_config = ConfigDict(extra="forbid")

    contract_version: str = "v3.deliverable_contract.1"
    case_id: str
    deliverables: List[DeliverableSpecV1] = Field(min_length=1)
    prompt_submission_clause: str = "compiled_by_program"
    allow_unlisted_deliverables: bool = False

    @model_validator(mode="after")
    def validate_unique_outputs(self) -> "DeliverableContractV1":
        names = [item.file_name.casefold() for item in self.deliverables]
        paths = [item.relative_path.casefold() for item in self.deliverables]
        if len(names) != len(set(names)):
            raise ValueError("duplicate_deliverable_file_name")
        if len(paths) != len(set(paths)):
            raise ValueError("duplicate_deliverable_relative_path")
        return self


class DeliverableContractFinding(BaseModel):
    finding_id: str
    check_name: str
    severity: FindingSeverity
    passed: bool
    message: str
    details: Dict[str, Any] = Field(default_factory=dict)


class DeliverableContractValidationReport(BaseModel):
    validation_version: str = "v3.deliverable_contract_validation.1"
    case_id: str
    validation_status: ContractStatus
    blocking_count: int = 0
    warning_count: int = 0
    findings: List[DeliverableContractFinding] = Field(default_factory=list)


class DeliveredFileRecord(BaseModel):
    expected_relative_path: str
    resolved_path: Optional[str] = None
    exists: bool = False
    nonempty: bool = False
    openable: bool = False
    valid: bool = False
    reason_codes: List[str] = Field(default_factory=list)


class DeliveryInspectionReport(BaseModel):
    inspection_version: str = "v3.delivery_inspection.1"
    case_id: str
    output_root: str
    delivery_status: DeliveryStatus
    expected_count: int
    valid_count: int
    wrong_name_or_path_files: List[str] = Field(default_factory=list)
    records: List[DeliveredFileRecord] = Field(default_factory=list)

    @property
    def valid_paths(self) -> List[Path]:
        return [
            Path(record.resolved_path)
            for record in self.records
            if record.valid and record.resolved_path
        ]


class DeliverableContractCompiler:
    """Build the single source of truth for prompt, export, and delivery inspection."""

    def build(
        self,
        case_id: str,
        deliverable_specs: Iterable[Dict[str, Any] | str],
        reference_files: Iterable[str] = (),
    ) -> DeliverableContractV1:
        deliverables: List[DeliverableSpecV1] = []
        for raw in deliverable_specs:
            item = {"file_name": raw} if isinstance(raw, str) else dict(raw)
            file_name = PurePosixPath(
                str(item.get("file_name") or item.get("relative_path") or "").replace("\\", "/")
            ).name
            if not file_name:
                raise ValueError("deliverable_file_name_missing")
            relative_path = str(item.get("relative_path") or f"deliverable_files/{file_name}")
            source_template = item.get("source_template") or item.get("template_file")
            creation_mode = str(item.get("creation_mode") or "").strip()
            if source_template:
                normalized_template = str(source_template).replace("\\", "/")
                if not normalized_template.startswith("reference_files/"):
                    normalized_template = f"reference_files/{PurePosixPath(normalized_template).name}"
                source_template = normalized_template
                if not creation_mode:
                    creation_mode = (
                        "edit_provided_copy"
                        if PurePosixPath(normalized_template).name.casefold() == file_name.casefold()
                        else "copy_then_edit"
                    )
            else:
                creation_mode = creation_mode or "create"
            deliverables.append(
                DeliverableSpecV1(
                    file_name=file_name,
                    relative_path=relative_path,
                    format=str(item.get("format") or Path(file_name).suffix.lstrip(".")),
                    creation_mode=creation_mode,
                    source_template=source_template,
                    must_exist=bool(item.get("must_exist", True)),
                    must_be_nonempty=bool(item.get("must_be_nonempty", True)),
                    openability_check=bool(item.get("openability_check", True)),
                )
            )
        return DeliverableContractV1(case_id=case_id, deliverables=deliverables)

    def compile_prompt(self, base_prompt: str, contract: DeliverableContractV1) -> str:
        base = base_prompt.strip()
        lines = [
            "Submission requirements (authoritative):",
            "- Save final deliverables under the `deliverable_files/` directory.",
        ]
        for item in contract.deliverables:
            if item.creation_mode == "create":
                action = "Create"
            elif item.creation_mode == "copy_then_edit":
                action = f"Copy `{item.source_template}` and edit the copy"
            else:
                action = f"Use the provided template `{item.source_template}` as the starting file"
            lines.append(f"- {action}; submit exactly `{item.relative_path}`.")
        lines.append("- Do not submit the completed file only in `reference_files/` or under another name.")
        clause = "\n".join(lines)
        return f"{base}\n\n{clause}\n" if base else f"{clause}\n"


class DeliverableContractValidator:
    def validate(
        self,
        contract: DeliverableContractV1,
        prompt: str,
        reference_files: Iterable[str],
    ) -> DeliverableContractValidationReport:
        findings: List[DeliverableContractFinding] = []
        references = {
            str(value).replace("\\", "/").casefold(): str(value).replace("\\", "/")
            for value in reference_files
        }
        reference_names = {PurePosixPath(value).name.casefold() for value in references.values()}
        prompt_targets = self._prompt_output_targets(prompt)
        expected_names = {item.file_name.casefold() for item in contract.deliverables}

        for item in contract.deliverables:
            prompt_has_path = item.relative_path.casefold() in prompt.casefold()
            findings.append(
                self._finding(
                    f"prompt_exact_path:{item.file_name}",
                    "blocking",
                    prompt_has_path,
                    (
                        f"Prompt contains the authoritative output path `{item.relative_path}`."
                        if prompt_has_path
                        else f"Prompt does not contain the authoritative output path `{item.relative_path}`."
                    ),
                    {"relative_path": item.relative_path},
                )
            )
            collision = item.file_name.casefold() in reference_names
            collision_declared = (
                item.creation_mode in {"copy_then_edit", "edit_provided_copy"}
                and bool(item.source_template)
            )
            findings.append(
                self._finding(
                    f"template_collision_declared:{item.file_name}",
                    "blocking",
                    not collision or collision_declared,
                    (
                        "Candidate template collision is explicitly governed."
                        if collision and collision_declared
                        else "No candidate template basename collision exists."
                        if not collision
                        else "A candidate template has the deliverable basename but no explicit copy/edit contract."
                    ),
                    {
                        "collision": collision,
                        "creation_mode": item.creation_mode,
                        "source_template": item.source_template,
                    },
                )
            )

        conflicting = sorted(prompt_targets - expected_names)
        findings.append(
            self._finding(
                "prompt_output_name_conflicts",
                "blocking",
                not conflicting,
                (
                    "Prompt output file names agree with the deliverable contract."
                    if not conflicting
                    else "Prompt contains output file names outside the deliverable contract."
                ),
                {"conflicting_file_names": conflicting, "prompt_output_targets": sorted(prompt_targets)},
            )
        )
        blocking_count = sum(
            item.severity == "blocking" and not item.passed for item in findings
        )
        warning_count = sum(
            item.severity == "warning" and not item.passed for item in findings
        )
        return DeliverableContractValidationReport(
            case_id=contract.case_id,
            validation_status="invalid" if blocking_count else "pass",
            blocking_count=blocking_count,
            warning_count=warning_count,
            findings=findings,
        )

    def _prompt_output_targets(self, prompt: str) -> set[str]:
        targets: set[str] = set()
        for line in prompt.splitlines():
            for match in _FILE_TOKEN.finditer(line):
                prefix = line[max(0, match.start() - 48):match.start()]
                if _OUTPUT_VERBS.search(prefix):
                    targets.add(match.group(1).strip().casefold())
        return targets

    def _finding(
        self,
        check_name: str,
        severity: FindingSeverity,
        passed: bool,
        message: str,
        details: Optional[Dict[str, Any]] = None,
    ) -> DeliverableContractFinding:
        import hashlib

        raw = f"{check_name}|{message}"
        return DeliverableContractFinding(
            finding_id=f"delcon_{hashlib.sha1(raw.encode('utf-8')).hexdigest()[:10]}",
            check_name=check_name,
            severity=severity,
            passed=passed,
            message=message,
            details=details or {},
        )


def contract_from_dataset_row(
    dataset_row: Dict[str, Any],
    contract_path: str | Path | None = None,
) -> DeliverableContractV1:
    if contract_path and Path(contract_path).exists():
        return DeliverableContractV1.model_validate(
            json.loads(Path(contract_path).read_text(encoding="utf-8"))
        )
    extra_contract = (dataset_row.get("extra") or {}).get("deliverable_contract")
    if isinstance(extra_contract, dict):
        return DeliverableContractV1.model_validate(extra_contract)
    return DeliverableContractCompiler().build(
        case_id=str(dataset_row.get("task_id") or "unknown"),
        deliverable_specs=list(dataset_row.get("deliverable_files") or []),
        reference_files=list(dataset_row.get("reference_files") or []),
    )


def inspect_delivery(
    output_root: str | Path,
    contract: DeliverableContractV1,
) -> DeliveryInspectionReport:
    root = Path(output_root)
    delivery_dirs = sorted(
        path for path in root.rglob("deliverable_files") if path.is_dir()
    )
    if root.name == "deliverable_files" and root.is_dir():
        delivery_dirs.insert(0, root)
    all_files = sorted(
        path
        for directory in delivery_dirs
        for path in directory.rglob("*")
        if path.is_file() and path.name != "expected_deliverables.json"
    )
    records: List[DeliveredFileRecord] = []
    matched: set[Path] = set()
    for spec in contract.deliverables:
        candidates = [directory / spec.file_name for directory in delivery_dirs]
        resolved = next((path for path in candidates if path.is_file()), None)
        reasons: List[str] = []
        exists = resolved is not None
        if spec.must_exist and not exists:
            reasons.append("expected_deliverable_missing")
        nonempty = bool(resolved and resolved.stat().st_size > 0)
        if exists and spec.must_be_nonempty and not nonempty:
            reasons.append("expected_deliverable_empty")
        openable = bool(resolved and (not spec.openability_check or _is_openable(resolved)))
        if exists and spec.openability_check and not openable:
            reasons.append("expected_deliverable_unopenable")
        valid = (
            (exists or not spec.must_exist)
            and (nonempty or not spec.must_be_nonempty)
            and (openable or not spec.openability_check)
        )
        if resolved:
            matched.add(resolved)
        records.append(
            DeliveredFileRecord(
                expected_relative_path=spec.relative_path,
                resolved_path=str(resolved) if resolved else None,
                exists=exists,
                nonempty=nonempty,
                openable=openable,
                valid=valid,
                reason_codes=reasons,
            )
        )
    unlisted = [str(path) for path in all_files if path not in matched]
    valid_count = sum(record.valid for record in records)
    valid = valid_count == len(records) and (contract.allow_unlisted_deliverables or not unlisted)
    return DeliveryInspectionReport(
        case_id=contract.case_id,
        output_root=str(root),
        delivery_status="valid" if valid else "invalid",
        expected_count=len(records),
        valid_count=valid_count,
        wrong_name_or_path_files=unlisted,
        records=records,
    )


def _is_openable(path: Path) -> bool:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=False)
            workbook.close()
        elif suffix == ".docx":
            from docx import Document

            Document(path)
        elif suffix == ".json":
            json.loads(path.read_text(encoding="utf-8"))
        else:
            with path.open("rb") as handle:
                handle.read(1)
        return True
    except Exception:
        return False
