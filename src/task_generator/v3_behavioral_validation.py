from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from pydantic import BaseModel, ConfigDict, Field

from task_generator.v3_deliverable_contract import (
    DeliverableContractCompiler,
    DeliveryInspectionReport,
    inspect_delivery,
)


PreflightStatus = Literal["not_evaluated", "pass", "fail"]
ProcessStatus = Literal["not_run", "succeeded", "failed", "timeout"]
EvidenceStatus = Literal["not_evaluated", "pass", "fail", "provisional"]
FailureCategory = Literal[
    "none",
    "not_evaluated",
    "provider_failure",
    "tool_failure",
    "non_delivery",
    "wrong_path",
    "invalid_file",
    "business_error",
    "professional_quality_failure",
]


class SolverPreflightOperationV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    operation: Literal["create", "copy", "edit", "save", "submit"]
    passed: bool
    evidence_paths: List[str] = Field(default_factory=list)
    details: str = ""


class SolverPreflightAttemptV1(BaseModel):
    attempt: int = Field(ge=1)
    process_status: ProcessStatus
    retry_reason: Optional[str] = None
    output_fingerprint: Optional[str] = None
    failure_reason: Optional[str] = None


class SolverToolPreflightReportV1(BaseModel):
    report_version: Literal["v3.solver_tool_preflight.1"] = (
        "v3.solver_tool_preflight.1"
    )
    solver_model: str
    environment_id: str
    status: PreflightStatus
    eligible_for_business_eval: bool = False
    operations: List[SolverPreflightOperationV1] = Field(default_factory=list)
    delivery_inspection: Optional[DeliveryInspectionReport] = None
    attempts: List[SolverPreflightAttemptV1] = Field(default_factory=list)
    first_failure: Optional[str] = None
    output_fingerprints: List[str] = Field(default_factory=list)
    notes: List[str] = Field(default_factory=list)


class BehavioralAttemptEvidenceV1(BaseModel):
    attempt: int = Field(ge=1)
    process_status: ProcessStatus
    failure_stage: Optional[str] = None
    retry_reason: Optional[str] = None
    output_fingerprint: Optional[str] = None
    stdout_sha256: Optional[str] = None
    stderr_sha256: Optional[str] = None


class BehavioralExecutionReportV1(BaseModel):
    report_version: Literal["v3.behavioral_execution.1"] = (
        "v3.behavioral_execution.1"
    )
    case_id: str
    solver_model: str
    environment_id: str
    preflight_required: bool
    preflight_status: PreflightStatus
    process_status: ProcessStatus
    delivery_status: Literal["not_evaluated", "valid", "invalid"]
    file_validity_status: EvidenceStatus
    business_validity_status: EvidenceStatus = "not_evaluated"
    professional_quality_status: EvidenceStatus = "not_evaluated"
    grader_eligible: bool = False
    grader_executed: bool = False
    failure_category: FailureCategory = "not_evaluated"
    process_completed_count: int = 0
    exact_delivery_count: int = 0
    valid_file_count: int = 0
    business_score: Optional[float] = None
    first_failure: Optional[str] = None
    attempts: List[BehavioralAttemptEvidenceV1] = Field(default_factory=list)
    input_package_root: Optional[str] = None
    input_package_fingerprint: Optional[str] = None
    output_root: Optional[str] = None
    output_tree_fingerprint: Optional[str] = None
    output_fingerprints: List[str] = Field(default_factory=list)
    solver_preflight_report_path: Optional[str] = None
    solver_preflight_report_sha256: Optional[str] = None
    delivery_inspection_path: Optional[str] = None
    delivery_inspection_sha256: Optional[str] = None
    evidence_paths: List[str] = Field(default_factory=list)
    notes: List[str] = Field(default_factory=list)


class SolverToolPreflight:
    EXPECTED_CREATED = "created_workbook.xlsx"
    EXPECTED_EDITED = "edited_template.xlsx"

    def create_fixture(self, fixture_root: str | Path) -> Dict[str, str]:
        from openpyxl import Workbook

        root = Path(fixture_root)
        reference_root = root / "reference_files"
        deliverable_root = root / "deliverable_files"
        reference_root.mkdir(parents=True, exist_ok=True)
        deliverable_root.mkdir(parents=True, exist_ok=True)
        template_path = reference_root / "copy_template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Template"
        sheet["A1"] = "template_marker"
        workbook.save(template_path)
        contract = DeliverableContractCompiler().build(
            case_id="solver_tool_preflight",
            deliverable_specs=[
                {"file_name": self.EXPECTED_CREATED},
                {
                    "file_name": self.EXPECTED_EDITED,
                    "creation_mode": "copy_then_edit",
                    "source_template": "reference_files/copy_template.xlsx",
                },
            ],
            reference_files=["reference_files/copy_template.xlsx"],
        )
        contract_path = root / "deliverable_contract.json"
        contract_path.write_text(
            contract.model_dump_json(indent=2),
            encoding="utf-8",
        )
        prompt = DeliverableContractCompiler().compile_prompt(
            (
                "Create one new workbook containing `created_marker`. Copy the provided "
                "template, retain `template_marker`, add `edited_marker`, and save both "
                "files as instructed."
            ),
            contract,
        )
        prompt_path = root / "prompt.md"
        prompt_path.write_text(prompt, encoding="utf-8")
        return {
            "fixture_root": str(root),
            "prompt_path": str(prompt_path),
            "contract_path": str(contract_path),
            "template_path": str(template_path),
            "deliverable_dir": str(deliverable_root),
        }

    def inspect(
        self,
        fixture_root: str | Path,
        solver_model: str,
        environment_id: str,
        attempts: Optional[List[SolverPreflightAttemptV1]] = None,
    ) -> SolverToolPreflightReportV1:
        from openpyxl import load_workbook

        root = Path(fixture_root)
        contract_payload = json.loads(
            (root / "deliverable_contract.json").read_text(encoding="utf-8")
        )
        from task_generator.v3_deliverable_contract import DeliverableContractV1

        contract = DeliverableContractV1.model_validate(contract_payload)
        delivery = inspect_delivery(root, contract)
        created_path = root / "deliverable_files" / self.EXPECTED_CREATED
        edited_path = root / "deliverable_files" / self.EXPECTED_EDITED

        created_marker = False
        edited_retains_template = False
        edited_marker = False
        if created_path.is_file():
            try:
                workbook = load_workbook(created_path, read_only=True, data_only=False)
                created_marker = any(
                    cell.value == "created_marker"
                    for row in workbook.active.iter_rows()
                    for cell in row
                )
                workbook.close()
            except Exception:
                created_marker = False
        if edited_path.is_file():
            try:
                workbook = load_workbook(edited_path, read_only=True, data_only=False)
                values = {
                    cell.value
                    for row in workbook["Template"].iter_rows()
                    for cell in row
                }
                workbook.close()
                edited_retains_template = "template_marker" in values
                edited_marker = "edited_marker" in values
            except Exception:
                pass

        operations = [
            SolverPreflightOperationV1(
                operation="create",
                passed=created_marker,
                evidence_paths=[str(created_path)] if created_path.exists() else [],
                details="New workbook contains the required created marker.",
            ),
            SolverPreflightOperationV1(
                operation="copy",
                passed=edited_retains_template,
                evidence_paths=[str(edited_path)] if edited_path.exists() else [],
                details="Edited workbook retains the original template marker.",
            ),
            SolverPreflightOperationV1(
                operation="edit",
                passed=edited_marker,
                evidence_paths=[str(edited_path)] if edited_path.exists() else [],
                details="Copied workbook contains the required edit marker.",
            ),
            SolverPreflightOperationV1(
                operation="save",
                passed=all(record.openable for record in delivery.records),
                evidence_paths=[
                    record.resolved_path
                    for record in delivery.records
                    if record.resolved_path
                ],
                details="All expected workbooks are nonempty and openable.",
            ),
            SolverPreflightOperationV1(
                operation="submit",
                passed=delivery.delivery_status == "valid",
                evidence_paths=[
                    record.resolved_path
                    for record in delivery.records
                    if record.valid and record.resolved_path
                ],
                details="All expected files are under the exact governed deliverable paths.",
            ),
        ]
        status: PreflightStatus = (
            "pass" if all(item.passed for item in operations) else "fail"
        )
        fingerprints = self.output_fingerprints(root / "deliverable_files")
        attempt_records = list(attempts or [])
        if not attempt_records:
            attempt_records = [
                SolverPreflightAttemptV1(
                    attempt=1,
                    process_status="succeeded" if status == "pass" else "failed",
                    output_fingerprint=self.aggregate_fingerprint(fingerprints),
                    failure_reason=(
                        None
                        if status == "pass"
                        else next(
                            (
                                item.operation
                                for item in operations
                                if not item.passed
                            ),
                            "unknown",
                        )
                    ),
                )
            ]
        first_failure = next(
            (
                f"preflight_operation:{item.operation}"
                for item in operations
                if not item.passed
            ),
            None,
        )
        return SolverToolPreflightReportV1(
            solver_model=solver_model,
            environment_id=environment_id,
            status=status,
            eligible_for_business_eval=status == "pass",
            operations=operations,
            delivery_inspection=delivery,
            attempts=attempt_records,
            first_failure=first_failure,
            output_fingerprints=fingerprints,
            notes=[
                "This preflight measures file-tool capability only, not business reasoning.",
                "A failed model remains tool-ineligible for governed business comparison.",
            ],
        )

    @staticmethod
    def output_fingerprints(root: str | Path) -> List[str]:
        path = Path(root)
        if not path.exists():
            return []
        records = []
        for file in sorted(item for item in path.rglob("*") if item.is_file()):
            relative = file.relative_to(path).as_posix()
            records.append(
                f"{relative}:{hashlib.sha256(file.read_bytes()).hexdigest()}"
            )
        return records

    @staticmethod
    def aggregate_fingerprint(records: List[str]) -> Optional[str]:
        if not records:
            return None
        return hashlib.sha256("\n".join(records).encode("utf-8")).hexdigest()

    @staticmethod
    def tree_fingerprint(records: List[str]) -> str:
        return hashlib.sha256("\n".join(records).encode("utf-8")).hexdigest()


class BehavioralExecutionBuilder:
    PROVIDER_ERROR_PATTERN = re.compile(
        r"(?:rate.?limit|quota|provider|http\s*[45]\d\d|authentication|"
        r"connection\s*(?:error|reset)|service unavailable)",
        re.IGNORECASE,
    )

    def build(
        self,
        *,
        case_id: str,
        solver_model: str,
        environment_id: str,
        command_records: List[Any],
        delivery_inspection: Optional[DeliveryInspectionReport],
        output_root: str | Path,
        input_package_root: str | Path | None = None,
        preflight_report: Optional[SolverToolPreflightReportV1] = None,
        preflight_report_path: str | Path | None = None,
        preflight_required: bool = True,
        delivery_inspection_path: Optional[str] = None,
        business_validity_status: EvidenceStatus = "not_evaluated",
        professional_quality_status: EvidenceStatus = "not_evaluated",
        business_score: Optional[float] = None,
    ) -> BehavioralExecutionReportV1:
        solver_record = next(
            (
                record
                for record in command_records
                if self._value(record, "command_name")
                != "bench_standalone.grade_deliverables"
            ),
            None,
        )
        grader_record = next(
            (
                record
                for record in command_records
                if self._value(record, "command_name")
                == "bench_standalone.grade_deliverables"
            ),
            None,
        )
        process_status: ProcessStatus = (
            self._value(solver_record, "status") if solver_record else "not_run"
        )
        if process_status not in {"not_run", "succeeded", "failed", "timeout"}:
            process_status = "failed"
        preflight_status: PreflightStatus = (
            preflight_report.status if preflight_report else "not_evaluated"
        )
        delivery_status = (
            delivery_inspection.delivery_status
            if delivery_inspection
            else "not_evaluated"
        )
        file_validity: EvidenceStatus = (
            "pass"
            if delivery_inspection and delivery_inspection.delivery_status == "valid"
            else "fail"
            if delivery_inspection
            else "not_evaluated"
        )
        preflight_satisfied = (
            preflight_status == "pass"
            if preflight_required
            else preflight_status in {"not_evaluated", "pass"}
        )
        grader_eligible = (
            preflight_satisfied
            and process_status == "succeeded"
            and delivery_status == "valid"
        )
        grader_executed = bool(
            grader_record and self._value(grader_record, "status") == "succeeded"
        )
        failure_category, first_failure = self._failure_category(
            preflight_report=preflight_report,
            preflight_required=preflight_required,
            solver_record=solver_record,
            delivery=delivery_inspection,
            business_status=business_validity_status,
            professional_status=professional_quality_status,
        )
        fingerprints = SolverToolPreflight.output_fingerprints(output_root)
        input_fingerprints = (
            SolverToolPreflight.output_fingerprints(input_package_root)
            if input_package_root is not None
            else []
        )
        attempts = []
        if solver_record:
            stdout = str(self._value(solver_record, "stdout_excerpt") or "")
            stderr = str(self._value(solver_record, "stderr_excerpt") or "")
            attempts.append(
                BehavioralAttemptEvidenceV1(
                    attempt=1,
                    process_status=process_status,
                    failure_stage=self._value(solver_record, "failure_stage"),
                    output_fingerprint=(
                        SolverToolPreflight.aggregate_fingerprint(fingerprints)
                    ),
                    stdout_sha256=hashlib.sha256(stdout.encode("utf-8")).hexdigest(),
                    stderr_sha256=hashlib.sha256(stderr.encode("utf-8")).hexdigest(),
                )
            )
        evidence_paths = [
            path
            for path in [
                delivery_inspection_path,
            ]
            if path
        ]
        return BehavioralExecutionReportV1(
            case_id=case_id,
            solver_model=solver_model,
            environment_id=environment_id,
            preflight_required=preflight_required,
            preflight_status=preflight_status,
            process_status=process_status,
            delivery_status=delivery_status,
            file_validity_status=file_validity,
            business_validity_status=business_validity_status,
            professional_quality_status=professional_quality_status,
            grader_eligible=grader_eligible,
            grader_executed=grader_executed,
            failure_category=failure_category,
            process_completed_count=int(process_status == "succeeded"),
            exact_delivery_count=(
                delivery_inspection.valid_count if delivery_inspection else 0
            ),
            valid_file_count=(
                delivery_inspection.valid_count if delivery_inspection else 0
            ),
            business_score=business_score,
            first_failure=first_failure,
            attempts=attempts,
            input_package_root=(
                str(Path(input_package_root).resolve())
                if input_package_root is not None
                else None
            ),
            input_package_fingerprint=(
                SolverToolPreflight.tree_fingerprint(input_fingerprints)
                if input_package_root is not None
                else None
            ),
            output_root=str(Path(output_root).resolve()),
            output_tree_fingerprint=SolverToolPreflight.tree_fingerprint(
                fingerprints
            ),
            output_fingerprints=fingerprints,
            solver_preflight_report_path=(
                str(Path(preflight_report_path).resolve())
                if preflight_report_path is not None
                else None
            ),
            solver_preflight_report_sha256=(
                hashlib.sha256(Path(preflight_report_path).read_bytes()).hexdigest()
                if preflight_report_path is not None
                and Path(preflight_report_path).is_file()
                else None
            ),
            delivery_inspection_path=delivery_inspection_path,
            delivery_inspection_sha256=(
                hashlib.sha256(Path(delivery_inspection_path).read_bytes()).hexdigest()
                if delivery_inspection_path is not None
                and Path(delivery_inspection_path).is_file()
                else None
            ),
            evidence_paths=evidence_paths,
            notes=[
                "Process completion, exact delivery, file validity, business validity, "
                "and professional quality are separate evidence axes.",
                "Only a tool-eligible solver with exact valid delivery is grader-eligible.",
            ],
        )

    def _failure_category(
        self,
        *,
        preflight_report: Optional[SolverToolPreflightReportV1],
        preflight_required: bool,
        solver_record: Any,
        delivery: Optional[DeliveryInspectionReport],
        business_status: EvidenceStatus,
        professional_status: EvidenceStatus,
    ) -> tuple[FailureCategory, Optional[str]]:
        if preflight_required and (
            preflight_report is None or preflight_report.status != "pass"
        ):
            first = (
                preflight_report.first_failure
                if preflight_report
                else "solver_preflight_not_evaluated"
            )
            return "tool_failure", first
        status = self._value(solver_record, "status") if solver_record else "not_run"
        if status in {"failed", "timeout"}:
            stderr = str(self._value(solver_record, "stderr_excerpt") or "")
            if self.PROVIDER_ERROR_PATTERN.search(stderr):
                return "provider_failure", (
                    self._value(solver_record, "failure_stage")
                    or f"solver_process:{status}"
                )
            return "tool_failure", (
                self._value(solver_record, "failure_stage")
                or f"solver_process:{status}"
            )
        if status != "succeeded":
            return "not_evaluated", "solver_process_not_run"
        if delivery is None:
            return "non_delivery", "delivery_not_inspected"
        if delivery.delivery_status != "valid":
            if delivery.wrong_name_or_path_files:
                return "wrong_path", "unexpected_or_wrong_path_deliverable"
            reason_codes = {
                reason
                for record in delivery.records
                for reason in record.reason_codes
            }
            if reason_codes & {
                "expected_deliverable_empty",
                "expected_deliverable_unopenable",
            }:
                return "invalid_file", sorted(reason_codes)[0]
            return "non_delivery", (
                sorted(reason_codes)[0]
                if reason_codes
                else "expected_deliverable_missing"
            )
        if business_status == "fail":
            return "business_error", "business_validity_failed"
        if professional_status == "fail":
            return "professional_quality_failure", "professional_quality_failed"
        return "none", None

    @staticmethod
    def _value(record: Any, name: str) -> Any:
        if record is None:
            return None
        if isinstance(record, dict):
            return record.get(name)
        return getattr(record, name, None)


def write_behavioral_report(
    report: BehavioralExecutionReportV1,
    output_path: str | Path,
) -> None:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
