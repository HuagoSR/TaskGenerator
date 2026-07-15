from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional
from xml.etree import ElementTree

from pydantic import BaseModel, Field, model_validator


ContractLifecycle = Literal["design", "resolved", "verified"]
ContractOrigin = Literal["generator_owned_v2", "legacy_inferred_v1"]
LocatorKind = Literal["prompt", "xlsx_field", "docx_clause", "candidate_rule"]
ResultKind = Literal["exact", "tolerance", "acceptable_set", "judgmental"]
ValidatorStatus = Literal["pending", "passed", "failed", "not_applicable"]


class TypedEvidenceLocator(BaseModel):
    locator_id: str
    kind: LocatorKind
    file_name: Optional[str] = None
    sheet_name: Optional[str] = None
    field_names: List[str] = Field(default_factory=list)
    rule_id: Optional[str] = None
    prompt_anchor: Optional[str] = None
    candidate_visible: bool = True


class ExpectedResultSpec(BaseModel):
    result_kind: ResultKind
    resolved: bool = False
    value: Any = None
    acceptable_values: List[Any] = Field(default_factory=list)
    tolerance: Optional[float] = None


class DeterministicValidatorSpec(BaseModel):
    validator_id: str
    input_locator_ids: List[str] = Field(default_factory=list)
    status: ValidatorStatus = "pending"
    recomputed_value: Any = None
    input_sha256: Dict[str, str] = Field(default_factory=dict)
    failure_reason: Optional[str] = None


class SemanticDependencyV2(BaseModel):
    dependency_id: str
    description: str
    locator_ids: List[str] = Field(default_factory=list)
    dependency_kind: Literal["data", "rule", "prompt"]


class SemanticRequirementV2(BaseModel):
    requirement_id: str
    prompt_text: str
    deliverable_file: str
    deliverable_location: str
    claim_ids: List[str] = Field(min_length=1)


class SemanticClaimV2(BaseModel):
    claim_id: str
    requirement_id: str
    claim_type: Literal["amount", "quantity", "matching", "classification", "status", "judgment"]
    description: str
    dependency_ids: List[str] = Field(min_length=1)
    expected_result: ExpectedResultSpec
    validator: DeterministicValidatorSpec
    explicit_assumptions: List[str] = Field(default_factory=list)
    prohibited_hidden_assumptions: List[str] = Field(default_factory=list)


class RubricBindingV2(BaseModel):
    criterion_id: str
    claim_ids: List[str] = Field(min_length=1)
    criterion_type: Literal["fact", "deliverable", "reasoning", "advisory"]
    weight: float = Field(gt=0.0, le=1.0)
    check_mode: Literal["exact", "tolerance", "set_membership", "presence", "judgment"]


class TaskSemanticContractV2(BaseModel):
    contract_version: str = "v3.task_semantic_contract.2"
    contract_origin: ContractOrigin = "generator_owned_v2"
    lifecycle: ContractLifecycle = "design"
    task_id: str
    motif: str
    template_family: str
    requirements: List[SemanticRequirementV2] = Field(min_length=1)
    claims: List[SemanticClaimV2] = Field(min_length=1)
    dependencies: List[SemanticDependencyV2] = Field(min_length=1)
    locators: List[TypedEvidenceLocator] = Field(min_length=1)
    rubric_bindings: List[RubricBindingV2] = Field(default_factory=list)
    notes: List[str] = Field(default_factory=list)

    @model_validator(mode="after")
    def validate_links(self) -> "TaskSemanticContractV2":
        requirement_ids = {item.requirement_id for item in self.requirements}
        claim_ids = {item.claim_id for item in self.claims}
        dependency_ids = {item.dependency_id for item in self.dependencies}
        locator_ids = {item.locator_id for item in self.locators}
        for requirement in self.requirements:
            if not set(requirement.claim_ids).issubset(claim_ids):
                raise ValueError(f"Requirement {requirement.requirement_id} references unknown claims.")
        for claim in self.claims:
            if claim.requirement_id not in requirement_ids:
                raise ValueError(f"Claim {claim.claim_id} references an unknown requirement.")
            if not set(claim.dependency_ids).issubset(dependency_ids):
                raise ValueError(f"Claim {claim.claim_id} references unknown dependencies.")
            if not set(claim.validator.input_locator_ids).issubset(locator_ids):
                raise ValueError(f"Claim {claim.claim_id} validator references unknown locators.")
        for dependency in self.dependencies:
            if not set(dependency.locator_ids).issubset(locator_ids):
                raise ValueError(f"Dependency {dependency.dependency_id} references unknown locators.")
        for binding in self.rubric_bindings:
            if not set(binding.claim_ids).issubset(claim_ids):
                raise ValueError(f"Rubric binding {binding.criterion_id} references unknown claims.")
        return self


class ContractConsistencyReport(BaseModel):
    report_version: str = "v3.semantic_contract_consistency.1"
    task_id: str
    lifecycle: ContractLifecycle
    decision: Literal["pass", "revise", "blocked"]
    reason_codes: List[str] = Field(default_factory=list)
    fact_weight_ratio: float = 0.0
    validator_counts: Dict[str, int] = Field(default_factory=dict)


class LegacyFinanceAuditReport(BaseModel):
    report_version: str = "v3.legacy_finance_semantic_audit.1"
    task_id: str
    motif: str
    decision: Literal["pass", "blocked"]
    reason_codes: List[str] = Field(default_factory=list)


class LegacyFinanceSemanticAuditor:
    """Read-only negative-control audit for tasks created before contract V2."""

    def audit(
        self,
        dataset: Dict[str, Any],
        reference_dir: str | Path,
        rubric: Dict[str, Any] | List[Dict[str, Any]],
    ) -> LegacyFinanceAuditReport:
        root = Path(reference_dir)
        prompt = str(dataset.get("prompt") or "").lower()
        motif = str(dataset.get("motif") or "")
        reasons: List[str] = []
        if "cash_reconciliation" in motif or "adjusted bank" in prompt:
            bank_headers = _xlsx_headers(root / "bank_statement.xlsx", "Bank_Activity")
            ledger_headers = _xlsx_headers(root / "cash_ledger.xlsx", "Cash_Ledger")
            if "adjusted" in prompt and not ({"Opening_Balance", "Closing_Balance"} & (bank_headers | ledger_headers)):
                reasons.append("missing_candidate_input")
        elif "three_way_match" in motif or "clear, hold, or investigate" in prompt:
            rule_text = " ".join(
                _docx_text(path).lower() for path in root.glob("*.docx")
            )
            required = {"clear", "hold", "investigate", "precedence"}
            if not required.issubset(set(rule_text.replace("/", " ").replace(",", " ").split())):
                reasons.append("underdefined_decision_rule")
        elif "expense_policy" in motif or "per person" in prompt:
            headers = _xlsx_headers(root / "expense_transactions.xlsx", "Transactions")
            rule_text = " ".join(_docx_text(path).lower() for path in root.glob("*.docx"))
            if "per person" in rule_text and "Attendee_Count" not in headers:
                reasons.append("missing_candidate_input")
        criteria = rubric if isinstance(rubric, list) else [
            criterion
            for section in rubric.get("sections") or []
            for criterion in section.get("criteria") or []
        ]
        if not any(
            "fact_checks" in (item.get("tags") or [])
            or (item.get("criterion_type") == "fact" and item.get("semantic_claim_ids"))
            for item in criteria
        ):
            reasons.append("rubric_missing_fact_coverage")
        return LegacyFinanceAuditReport(
            task_id=str(dataset.get("task_id") or "legacy_task"),
            motif=motif,
            decision="blocked" if reasons else "pass",
            reason_codes=sorted(set(reasons)),
        )


class FinanceSemanticContractAdapter:
    """Build generator-owned contracts for the governed finance motifs."""

    SUPPORTED_MOTIFS = {
        "fan_in_reconciliation", "cross_check_validation", "policy_application",
        "evidence_to_deliverable",
    }

    def prepare_blueprint(self, blueprint: Dict[str, Any], motif: str) -> Dict[str, Any]:
        payload = json.loads(json.dumps(blueprint))
        if motif not in self.SUPPORTED_MOTIFS:
            raise ValueError(f"Unsupported finance semantic motif: {motif}")
        files = (payload.setdefault("data_spec", {})).setdefault("reference_files", [])
        deliverables = payload.setdefault("deliverable_spec", [])
        if motif == "fan_in_reconciliation":
            self._ensure_docx(files, "reconciliation_rules.docx")
            requirements = [
                "match bank and ledger items by reference and amount",
                "list bank-only and ledger-only items with source identifiers",
                "report bank and ledger period-activity totals and explain the remaining difference",
            ]
        elif motif == "cross_check_validation":
            self._ensure_docx(files, "three_way_match_rules.docx")
            requirements = [
                "join every invoice line to purchase orders and goods receipts",
                "calculate quantity and unit-price variances using the supplied rules",
                "classify every invoice line as clear, hold, or investigate and cite source row identifiers",
            ]
        elif motif == "policy_application":
            transaction = next((item for item in files if item.get("file_name") == "expense_transactions.xlsx"), None)
            if transaction:
                columns = transaction["sheet_specs"][0].setdefault("columns", [])
                if not any(item.get("name") == "Attendee_Count" for item in columns):
                    columns.append({"name": "Attendee_Count", "semantic_type": "count"})
            requirements = [
                "classify every transaction against the supplied policy clauses",
                "state the exception reason and required follow-up with transaction and clause citations",
                "summarize exception counts and amounts without inventing missing facts",
            ]
        else:
            self._ensure_docx(files, "control_reporting_rules.docx")
            requirements = [
                "report tested, passed, failed, and unresolved-evidence counts",
                "list every confirmed exception and evidence gap with Test_ID, Control_ID, and Evidence_ID",
                "state the required follow-up and cite the applicable reporting rule for every non-pass item",
            ]
        if deliverables:
            deliverables[0]["requirements"] = requirements
        payload.setdefault("prompt_spec", {})["visible_requirements"] = requirements
        payload.setdefault("golden_plan", {})["required_final_checks"] = requirements
        return payload

    def design(self, task_id: str, motif: str, blueprint: Dict[str, Any]) -> TaskSemanticContractV2:
        prepared = self.prepare_blueprint(blueprint, motif)
        deliverable = prepared["deliverable_spec"][0]
        specs = self._motif_specs(motif)
        locators = [TypedEvidenceLocator(**item) for item in specs["locators"]]
        dependencies = [SemanticDependencyV2(**item) for item in specs["dependencies"]]
        requirements: List[SemanticRequirementV2] = []
        claims: List[SemanticClaimV2] = []
        bindings: List[RubricBindingV2] = []
        for index, spec in enumerate(specs["claims"], start=1):
            req_id = f"req_{index:03d}"
            claim_id = f"claim_{index:03d}"
            criterion_id = f"FACT_{index:03d}"
            requirements.append(SemanticRequirementV2(
                requirement_id=req_id,
                prompt_text=spec["requirement"],
                deliverable_file=deliverable["file_name"],
                deliverable_location=spec["deliverable_location"],
                claim_ids=[claim_id],
            ))
            claims.append(SemanticClaimV2(
                claim_id=claim_id,
                requirement_id=req_id,
                claim_type=spec["claim_type"],
                description=spec["description"],
                dependency_ids=spec["dependency_ids"],
                expected_result=ExpectedResultSpec(result_kind=spec["result_kind"]),
                validator=DeterministicValidatorSpec(
                    validator_id=spec["validator_id"],
                    input_locator_ids=spec["locator_ids"],
                ),
                prohibited_hidden_assumptions=["Use only candidate-visible fields and rules declared by this contract."],
            ))
            bindings.append(RubricBindingV2(
                criterion_id=criterion_id,
                claim_ids=[claim_id],
                criterion_type="fact",
                weight=round(0.7 / len(specs["claims"]), 6),
                check_mode=spec["check_mode"],
            ))
        bindings.append(RubricBindingV2(
            criterion_id="DELIVERABLE_001",
            claim_ids=[item.claim_id for item in claims],
            criterion_type="deliverable",
            weight=0.3,
            check_mode="presence",
        ))
        return TaskSemanticContractV2(
            task_id=task_id,
            motif=motif,
            template_family=str(prepared.get("template_family") or ""),
            requirements=requirements,
            claims=claims,
            dependencies=dependencies,
            locators=locators,
            rubric_bindings=bindings,
            notes=["Generator-owned before reference materialization.", "LLM findings cannot modify this contract directly."],
        )

    def _ensure_docx(self, files: List[Dict[str, Any]], name: str) -> None:
        if not any(item.get("file_name") == name for item in files):
            files.append({"file_name": name, "file_role": "reference_table", "sheet_specs": []})

    def _motif_specs(self, motif: str) -> Dict[str, Any]:
        if motif == "fan_in_reconciliation":
            return _cash_specs()
        if motif == "cross_check_validation":
            return _three_way_specs()
        if motif == "policy_application":
            return _expense_specs()
        return _evidence_to_deliverable_specs()


class FinanceSemanticContractResolver:
    def resolve(self, contract: TaskSemanticContractV2, reference_dir: str | Path) -> TaskSemanticContractV2:
        root = Path(reference_dir)
        payload = contract.model_copy(deep=True)
        if payload.contract_origin != "generator_owned_v2":
            return payload
        calculators = {
            "finance.cash.activity_reconciliation.v1": self._cash,
            "finance.ap.three_way_match.v1": self._three_way,
            "finance.expense.policy_application.v1": self._expense,
            "finance.audit.control_testing_summary.v1": self._evidence_to_deliverable,
        }
        for claim in payload.claims:
            calculator = calculators.get(claim.validator.validator_id)
            try:
                value = calculator(root) if calculator else None
                claim.expected_result.resolved = value is not None
                claim.expected_result.value = value
                claim.validator.recomputed_value = value
                claim.validator.status = "passed" if value is not None else "not_applicable"
                claim.validator.input_sha256 = self._input_hashes(root, claim.validator.input_locator_ids, payload)
            except Exception as exc:
                claim.validator.status = "failed"
                claim.validator.failure_reason = type(exc).__name__
        payload.lifecycle = "resolved"
        return payload

    def verify(self, contract: TaskSemanticContractV2) -> tuple[TaskSemanticContractV2, ContractConsistencyReport]:
        payload = contract.model_copy(deep=True)
        reasons: List[str] = []
        deterministic = [item for item in payload.claims if item.expected_result.result_kind != "judgmental"]
        if payload.contract_origin != "generator_owned_v2":
            reasons.append("legacy_contract_not_production_eligible")
        if any(not locator.candidate_visible for locator in payload.locators):
            reasons.append("hidden_candidate_dependency")
        if any(item.validator.status != "passed" or not item.expected_result.resolved for item in deterministic):
            reasons.append("deterministic_validator_not_passed")
        fact_weight = sum(item.weight for item in payload.rubric_bindings if item.criterion_type in {"fact", "deliverable"})
        covered = {claim for item in payload.rubric_bindings if item.criterion_type == "fact" for claim in item.claim_ids}
        if not {item.claim_id for item in deterministic}.issubset(covered):
            reasons.append("rubric_missing_fact_coverage")
        if fact_weight < 0.60:
            reasons.append("rubric_weight_imbalance")
        decision = "pass" if not reasons else "revise"
        if decision == "pass":
            payload.lifecycle = "verified"
        counts: Dict[str, int] = {}
        for claim in payload.claims:
            counts[claim.validator.status] = counts.get(claim.validator.status, 0) + 1
        return payload, ContractConsistencyReport(
            task_id=payload.task_id,
            lifecycle=payload.lifecycle,
            decision=decision,
            reason_codes=sorted(set(reasons)),
            fact_weight_ratio=round(fact_weight, 6),
            validator_counts=dict(sorted(counts.items())),
        )

    def _cash(self, root: Path) -> Dict[str, Any]:
        bank = _xlsx_rows(root / "bank_statement.xlsx", "Bank_Activity")
        ledger = _xlsx_rows(root / "cash_ledger.xlsx", "Cash_Ledger")
        bank_keys = {(str(row["Reference"]), float(row["Amount"])) for row in bank}
        ledger_keys = {(str(row["Reference"]), float(row["Amount"])) for row in ledger}
        return {
            "matched_count": len(bank_keys & ledger_keys),
            "bank_only_count": len(bank_keys - ledger_keys),
            "ledger_only_count": len(ledger_keys - bank_keys),
            "bank_activity_total": round(sum(float(row["Amount"]) for row in bank), 2),
            "ledger_activity_total": round(sum(float(row["Amount"]) for row in ledger), 2),
        }

    def _three_way(self, root: Path) -> Dict[str, Any]:
        pos = {(str(row["PO_ID"]), str(row["Item_ID"])): row for row in _xlsx_rows(root / "purchase_orders.xlsx", "PO_Lines")}
        receipts = {(str(row["PO_ID"]), str(row["Item_ID"])): row for row in _xlsx_rows(root / "goods_receipts.xlsx", "Receipt_Lines")}
        invoices = _xlsx_rows(root / "supplier_invoices.xlsx", "Invoice_Lines")
        business_key_counts: Dict[tuple, int] = {}
        for row in invoices:
            key = (str(row["PO_ID"]), str(row["Item_ID"]), float(row["Invoiced_Qty"]), float(row["Unit_Price"]))
            business_key_counts[key] = business_key_counts.get(key, 0) + 1
        statuses = []
        for row in invoices:
            key = (str(row["PO_ID"]), str(row["Item_ID"]))
            business_key = (
                str(row["PO_ID"]), str(row["Item_ID"]),
                float(row["Invoiced_Qty"]), float(row["Unit_Price"]),
            )
            duplicate = business_key_counts[business_key] > 1
            po = pos.get(key)
            receipt = receipts.get(key)
            if not po or not receipt:
                status = "investigate"
            else:
                receipt_quantity_variance = float(receipt["Received_Qty"]) - float(po["Ordered_Qty"])
                invoice_quantity_variance = float(row["Invoiced_Qty"]) - float(receipt["Received_Qty"])
                unit_price_variance = float(row["Unit_Price"]) - float(po["Unit_Price"])
                status = "hold" if (
                    duplicate or receipt_quantity_variance != 0 or invoice_quantity_variance != 0
                    or abs(unit_price_variance) > 0.01
                ) else "clear"
            statuses.append(status)
        return {status: statuses.count(status) for status in ("clear", "hold", "investigate")}

    def _expense(self, root: Path) -> Dict[str, Any]:
        rows = _xlsx_rows(root / "expense_transactions.xlsx", "Transactions")
        exception_count = 0
        exception_amount = 0.0
        for row in rows:
            amount = float(row["Amount"])
            attendees = max(int(float(row.get("Attendee_Count") or 1)), 1)
            category = str(row["Category"]).lower()
            missing_receipt = amount >= 75 and str(row["Receipt_Available"]).lower() != "yes"
            meal = category == "meals" and amount / attendees > 100 and str(row["Approval_Level"]).lower() != "director"
            entertainment = category == "entertainment" and (
                str(row["Approval_Level"]).lower() != "director" or not str(row["Business_Purpose"]).strip()
            )
            if missing_receipt or meal or entertainment:
                exception_count += 1
                exception_amount += amount
        return {"exception_count": exception_count, "exception_amount": round(exception_amount, 2)}

    def _evidence_to_deliverable(self, root: Path) -> Dict[str, Any]:
        tests = _xlsx_rows(root / "control_test_results.xlsx", "Control_Tests")
        evidence_rows = _xlsx_rows(root / "audit_evidence_register.xlsx", "Evidence_Register")
        evidence = {str(row["Evidence_ID"]): row for row in evidence_rows}
        passed: List[str] = []
        failed: List[str] = []
        unresolved: List[str] = []
        follow_up: List[Dict[str, Any]] = []
        for row in tests:
            test_id = str(row["Test_ID"])
            evidence_id = str(row["Evidence_ID"])
            linked = evidence.get(evidence_id)
            evidence_supported = linked is not None and str(linked.get("Supports_Test_ID")) == test_id
            result = str(row.get("Result") or "").strip().lower()
            exception_count = int(float(row.get("Exceptions_Found") or 0))
            if result == "evidence gap" or not evidence_supported:
                unresolved.append(test_id)
                action = "collect evidence and retest before conclusion"
                rule = "E2D-003; E2D-004"
            elif result == "fail" or exception_count > 0:
                failed.append(test_id)
                action = f"remediate by {row.get('Due_Date')} and retest"
                rule = "E2D-002; E2D-004"
            else:
                passed.append(test_id)
                continue
            follow_up.append({
                "Test_ID": test_id,
                "Control_ID": str(row["Control_ID"]),
                "Evidence_ID": evidence_id,
                "Owner": str(row["Owner"]),
                "Due_Date": str(row["Due_Date"]),
                "Required_Action": action,
                "Rule": rule,
            })
        return {
            "tested_count": len(tests),
            "passed_count": len(passed),
            "failed_count": len(failed),
            "unresolved_count": len(unresolved),
            "passed_test_ids": passed,
            "failed_test_ids": failed,
            "unresolved_test_ids": unresolved,
            "follow_up": follow_up,
        }

    def _input_hashes(self, root: Path, locator_ids: List[str], contract: TaskSemanticContractV2) -> Dict[str, str]:
        locator_by_id = {item.locator_id: item for item in contract.locators}
        result = {}
        for locator_id in locator_ids:
            locator = locator_by_id[locator_id]
            if locator.file_name:
                path = root / locator.file_name
                if path.exists():
                    result[locator.file_name] = hashlib.sha256(path.read_bytes()).hexdigest()
        return result


def _xlsx_rows(path: Path, sheet_name: str) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    headers = [str(value) for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def _xlsx_headers(path: Path, sheet_name: str) -> set[str]:
    return set(_xlsx_rows(path, sheet_name)[0].keys()) if _xlsx_rows(path, sheet_name) else set()


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    return " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))


def _loc(locator_id: str, file_name: str, sheet: Optional[str], fields: List[str], kind: LocatorKind = "xlsx_field", rule_id: Optional[str] = None) -> Dict[str, Any]:
    return {"locator_id": locator_id, "kind": kind, "file_name": file_name, "sheet_name": sheet, "field_names": fields, "rule_id": rule_id}


def _cash_specs() -> Dict[str, Any]:
    return {
        "locators": [
            _loc("loc_bank", "bank_statement.xlsx", "Bank_Activity", ["Bank_ID", "Reference", "Amount"]),
            _loc("loc_ledger", "cash_ledger.xlsx", "Cash_Ledger", ["Ledger_ID", "Reference", "Amount"]),
            _loc("loc_rules", "reconciliation_rules.docx", None, [], "docx_clause", "REC-001"),
        ],
        "dependencies": [
            {"dependency_id": "dep_bank", "description": "Bank activity", "locator_ids": ["loc_bank"], "dependency_kind": "data"},
            {"dependency_id": "dep_ledger", "description": "Cash ledger activity", "locator_ids": ["loc_ledger"], "dependency_kind": "data"},
            {"dependency_id": "dep_rules", "description": "Matching and total rules", "locator_ids": ["loc_rules"], "dependency_kind": "rule"},
        ],
        "claims": [{"requirement": "Reconcile bank and ledger activity using the supplied matching rules.", "deliverable_location": "workbook:reconciliation", "claim_type": "matching", "description": "Matched and unmatched activity plus period totals", "dependency_ids": ["dep_bank", "dep_ledger", "dep_rules"], "result_kind": "exact", "validator_id": "finance.cash.activity_reconciliation.v1", "locator_ids": ["loc_bank", "loc_ledger", "loc_rules"], "check_mode": "exact"}],
    }


def _three_way_specs() -> Dict[str, Any]:
    return {
        "locators": [
            _loc("loc_po", "purchase_orders.xlsx", "PO_Lines", ["PO_ID", "Item_ID", "Ordered_Qty", "Unit_Price"]),
            _loc("loc_receipt", "goods_receipts.xlsx", "Receipt_Lines", ["Receipt_ID", "PO_ID", "Item_ID", "Received_Qty"]),
            _loc("loc_invoice", "supplier_invoices.xlsx", "Invoice_Lines", ["Invoice_ID", "PO_ID", "Item_ID", "Invoiced_Qty", "Unit_Price"]),
            _loc("loc_rules", "three_way_match_rules.docx", None, [], "docx_clause", "AP-001"),
        ],
        "dependencies": [
            {"dependency_id": "dep_po", "description": "Purchase orders", "locator_ids": ["loc_po"], "dependency_kind": "data"},
            {"dependency_id": "dep_receipt", "description": "Goods receipts", "locator_ids": ["loc_receipt"], "dependency_kind": "data"},
            {"dependency_id": "dep_invoice", "description": "Supplier invoices", "locator_ids": ["loc_invoice"], "dependency_kind": "data"},
            {"dependency_id": "dep_rules", "description": "Variance and disposition rules", "locator_ids": ["loc_rules"], "dependency_kind": "rule"},
        ],
        "claims": [{"requirement": "Classify each invoice line using the supplied three-way-match rules.", "deliverable_location": "workbook:line_review", "claim_type": "classification", "description": "Clear, hold, and investigate counts", "dependency_ids": ["dep_po", "dep_receipt", "dep_invoice", "dep_rules"], "result_kind": "exact", "validator_id": "finance.ap.three_way_match.v1", "locator_ids": ["loc_po", "loc_receipt", "loc_invoice", "loc_rules"], "check_mode": "exact"}],
    }


def _expense_specs() -> Dict[str, Any]:
    return {
        "locators": [
            _loc("loc_transactions", "expense_transactions.xlsx", "Transactions", ["Transaction_ID", "Category", "Amount", "Receipt_Available", "Approval_Level", "Business_Purpose", "Attendee_Count"]),
            _loc("loc_policy", "expense_policy.docx", None, [], "docx_clause", "POL-001:POL-004"),
        ],
        "dependencies": [
            {"dependency_id": "dep_transactions", "description": "Expense transactions including attendee count", "locator_ids": ["loc_transactions"], "dependency_kind": "data"},
            {"dependency_id": "dep_policy", "description": "Candidate-visible expense policy", "locator_ids": ["loc_policy"], "dependency_kind": "rule"},
        ],
        "claims": [{"requirement": "Classify every expense using the supplied policy and report exception totals.", "deliverable_location": "memo:exception_register", "claim_type": "classification", "description": "Exception count and amount", "dependency_ids": ["dep_transactions", "dep_policy"], "result_kind": "exact", "validator_id": "finance.expense.policy_application.v1", "locator_ids": ["loc_transactions", "loc_policy"], "check_mode": "exact"}],
    }


def _evidence_to_deliverable_specs() -> Dict[str, Any]:
    return {
        "locators": [
            _loc("loc_control_tests", "control_test_results.xlsx", "Control_Tests", [
                "Test_ID", "Control_ID", "Evidence_ID", "Procedure", "Sample_Size",
                "Exceptions_Found", "Result", "Owner", "Due_Date",
            ]),
            _loc("loc_evidence_register", "audit_evidence_register.xlsx", "Evidence_Register", [
                "Evidence_ID", "Source", "Period", "Reliability", "Supports_Test_ID",
            ]),
            _loc("loc_reporting_rules", "control_reporting_rules.docx", None, [], "docx_clause", "E2D-001:E2D-004"),
        ],
        "dependencies": [
            {"dependency_id": "dep_control_tests", "description": "Control test results", "locator_ids": ["loc_control_tests"], "dependency_kind": "data"},
            {"dependency_id": "dep_evidence_register", "description": "Evidence-to-test mapping", "locator_ids": ["loc_evidence_register"], "dependency_kind": "data"},
            {"dependency_id": "dep_reporting_rules", "description": "Candidate-visible conclusion and follow-up rules", "locator_ids": ["loc_reporting_rules"], "dependency_kind": "rule"},
        ],
        "claims": [{
            "requirement": "Prepare a control-testing summary grounded in the supplied tests, evidence register, and reporting rules.",
            "deliverable_location": "document:control_testing_summary",
            "claim_type": "classification",
            "description": "Passed, failed, unresolved, exception, evidence-gap, and follow-up results",
            "dependency_ids": ["dep_control_tests", "dep_evidence_register", "dep_reporting_rules"],
            "result_kind": "exact",
            "validator_id": "finance.audit.control_testing_summary.v1",
            "locator_ids": ["loc_control_tests", "loc_evidence_register", "loc_reporting_rules"],
            "check_mode": "exact",
        }],
    }


def write_contract(path: str | Path, model: BaseModel) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".tmp")
    temporary.write_text(model.model_dump_json(indent=2), encoding="utf-8")
    temporary.replace(target)


def to_review_contract(contract: TaskSemanticContractV2):
    """Bridge V2 into the stable LLM review package without downgrading its evidence."""
    from task_generator.v3_semantic_validity import (
        SemanticClaim,
        SemanticDependency,
        SemanticRequirement,
        TaskSemanticContract,
    )

    locator_by_id = {item.locator_id: item for item in contract.locators}
    dependencies = []
    for item in contract.dependencies:
        locators = [locator_by_id[value] for value in item.locator_ids]
        first = locators[0]
        dependencies.append(SemanticDependency(
            dependency_id=item.dependency_id,
            description=item.description,
            source_kind="candidate_rule" if item.dependency_kind == "rule" else ("prompt" if item.dependency_kind == "prompt" else "candidate_file"),
            file_name=first.file_name,
            locator=first.rule_id or first.sheet_name or first.prompt_anchor or "candidate-visible",
            field_names=sorted({field for locator in locators for field in locator.field_names}),
            candidate_visible=all(locator.candidate_visible for locator in locators),
        ))
    requirements = [SemanticRequirement(
        requirement_id=item.requirement_id,
        prompt_text=item.prompt_text,
        deliverable_file=item.deliverable_file,
        deliverable_location=item.deliverable_location,
        claim_ids=item.claim_ids,
    ) for item in contract.requirements]
    claims = [SemanticClaim(
        claim_id=item.claim_id,
        requirement_id=item.requirement_id,
        claim_type=item.claim_type,
        determinism=item.expected_result.result_kind,
        description=item.description,
        dependency_ids=item.dependency_ids,
        expected_value=item.expected_result.value,
        acceptable_values=item.expected_result.acceptable_values,
        tolerance=item.expected_result.tolerance,
        explicit_assumptions=item.explicit_assumptions,
        prohibited_hidden_assumptions=item.prohibited_hidden_assumptions,
        validator_id=item.validator.validator_id,
        rubric_criterion_ids=[binding.criterion_id for binding in contract.rubric_bindings if item.claim_id in binding.claim_ids],
    ) for item in contract.claims]
    return TaskSemanticContract(
        contract_version="v3.task_semantic_contract.2.review_bridge",
        task_id=contract.task_id,
        created_at="generator_owned_v2",
        requirements=requirements,
        claims=claims,
        dependencies=dependencies,
        notes=["Review bridge for a generator-owned V2 contract."],
    )
