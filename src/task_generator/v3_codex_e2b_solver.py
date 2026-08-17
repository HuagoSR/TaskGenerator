from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, Callable, Dict, List, Literal, Optional

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ConfigDict, Field

from task_generator.v3_skill_extractor import ProviderConfig


TASK_IDS = ("ms_2c770ac494eb53b8", "ms_7e4e35a9be2befda")
EXPECTED_FINGERPRINTS = {
    "ms_2c770ac494eb53b8": "18a814145c6308e0977ba08b85080fae07e7af00085d3fddd74b80d5cde26da6",
    "ms_7e4e35a9be2befda": "0bdf180661bc51dd16e545efe7a8a5dbbece27108a35b8cbb501446f766e08d7",
}
REMOTE_ROOT = PurePosixPath("/home/user/task")
CODEX_HOME = PurePosixPath("/home/user/.codex-tuzi")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _sha_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _sha(path: Path) -> str:
    return _sha_bytes(path.read_bytes())


def _tree(root: Path) -> str:
    rows = [
        f"{path.relative_to(root).as_posix()}:{_sha(path)}"
        for path in sorted(item for item in root.rglob("*") if item.is_file())
    ]
    return _sha_bytes("\n".join(rows).encode())


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


class CodexE2BTaskBindingV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    blind_task_id: str
    package_fingerprint: str = Field(min_length=64, max_length=64)
    candidate_tree_sha256: str = Field(min_length=64, max_length=64)
    expected_deliverable: str = Field(pattern=r"^deliverable_files/[A-Za-z0-9_.-]+\.xlsx$")


class CodexE2BSliceScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal["v3.codex_e2b_slice_scope.2"] = "v3.codex_e2b_slice_scope.2"
    e2b_auth_contract: Literal["explicit_sdk_api_key_process_only"] = (
        "explicit_sdk_api_key_process_only"
    )
    runtime_prepare_contract: Literal["single_attempt_with_exit_marker"] = (
        "single_attempt_with_exit_marker"
    )
    campaign_id: str
    provider: Literal["tuzi"] = "tuzi"
    model: Literal["gpt-5.6-sol"] = "gpt-5.6-sol"
    wire_api: Literal["responses"] = "responses"
    e2b_template: Literal["codex"] = "codex"
    public_timeout_seconds: Literal[300] = 300
    task_timeout_seconds: Literal[1800] = 1800
    retries: Literal[0] = 0
    task_bindings: List[CodexE2BTaskBindingV1]
    allowed_uploads: List[str] = ["candidate_prompt", "deliverable_contract", "reference_xlsx"]
    grader_authorized: Literal[False] = False
    remaining_ten_tasks_authorized: Literal[False] = False


class CodexE2BSliceReceiptV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    receipt_version: Literal["v3.codex_e2b_slice_receipt.1"] = "v3.codex_e2b_slice_receipt.1"
    scope_sha256: str = Field(min_length=64, max_length=64)
    authorization_basis: Literal["user_project_level_external_model_authorization"] = (
        "user_project_level_external_model_authorization"
    )
    private_package_upload_authorized: Literal[True] = True
    consumed_at: Optional[str] = None


class DeliveryInspectionV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    exists: bool = False
    nonempty: bool = False
    xlsx_openable: bool = False
    has_nonempty_sheet: bool = False
    differs_from_inputs: bool = False
    sha256: Optional[str] = None
    size_bytes: int = 0


class CodexE2BOutcomeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    kind: Literal["public_probe", "private_task"]
    task_id: str
    status: Literal["pass", "delivery_failed", "infrastructure_failed"]
    sandbox_id: Optional[str] = None
    exit_code: Optional[int] = None
    event_types: List[str] = []
    command_event_seen: bool = False
    turn_event_seen: bool = False
    usage: Dict[str, int] = {}
    delivery: DeliveryInspectionV1 = DeliveryInspectionV1()
    first_failure: Optional[str] = None
    jsonl_sha256: Optional[str] = None
    stderr_sha256: Optional[str] = None


class CodexE2BSliceResultV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    result_version: Literal["v3.codex_e2b_slice_result.1"] = "v3.codex_e2b_slice_result.1"
    scope_sha256: str
    completed_at: str
    public_probe: CodexE2BOutcomeV1
    private_tasks: List[CodexE2BOutcomeV1] = []
    decision: Literal["compatibility_pass", "compatibility_failed", "incomplete"]
    grader_called: Literal[False] = False
    route_comparison_made: Literal[False] = False


class CodexE2BRecoveryScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal["v3.codex_e2b_recovery_scope.1"] = "v3.codex_e2b_recovery_scope.1"
    campaign_id: str
    provider: Literal["tuzi"] = "tuzi"
    model: Literal["gpt-5.6-sol"] = "gpt-5.6-sol"
    wire_api: Literal["responses"] = "responses"
    e2b_template: Literal["codex"] = "codex"
    task_timeout_seconds: Literal[1800] = 1800
    retries: Literal[0] = 0
    recovered_task: CodexE2BTaskBindingV1
    retained_result_sha256: str = Field(min_length=64, max_length=64)
    retained_public_outcome_sha256: str = Field(min_length=64, max_length=64)
    retained_passed_task_outcome_sha256: str = Field(min_length=64, max_length=64)
    parity_report_sha256: str = Field(min_length=64, max_length=64)
    authorization_basis: Literal["user_project_level_external_model_authorization"] = (
        "user_project_level_external_model_authorization"
    )
    grader_authorized: Literal[False] = False
    remaining_ten_tasks_authorized: Literal[False] = False


class CodexE2BScreeningSolverScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal["v3.codex_e2b_screening_solver_scope.1"] = "v3.codex_e2b_screening_solver_scope.1"
    campaign_id: str
    provider: Literal["tuzi"] = "tuzi"
    model: Literal["gpt-5.6-sol"] = "gpt-5.6-sol"
    wire_api: Literal["responses"] = "responses"
    e2b_template: Literal["codex"] = "codex"
    task_timeout_seconds: Literal[1800] = 1800
    retries: Literal[0] = 0
    task_bindings: List[CodexE2BTaskBindingV1]
    retained_task_ids: List[str]
    authorized_new_task_ids: List[str]
    retained_compatibility_result_sha256: str = Field(min_length=64, max_length=64)
    parity_report_sha256: str = Field(min_length=64, max_length=64)
    authorization_basis: Literal["user_project_level_external_model_authorization"] = (
        "user_project_level_external_model_authorization"
    )
    grader_authorized: Literal[False] = False


class CodexE2BScreeningSolverResultV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    result_version: Literal["v3.codex_e2b_screening_solver_result.1"] = "v3.codex_e2b_screening_solver_result.1"
    scope_sha256: str
    completed_at: str
    retained_task_ids: List[str]
    newly_executed_task_ids: List[str]
    task_outcomes: List[CodexE2BOutcomeV1]
    valid_delivery_count: int
    decision: Literal["solver_complete", "incomplete"]
    grader_called: Literal[False] = False


class CodexE2BScreeningRecoveryScopeV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scope_version: Literal["v3.codex_e2b_screening_recovery_scope.1"] = "v3.codex_e2b_screening_recovery_scope.1"
    campaign_id: str
    provider: Literal["tuzi"] = "tuzi"
    model: Literal["gpt-5.6-sol"] = "gpt-5.6-sol"
    e2b_template: Literal["codex"] = "codex"
    task_timeout_seconds: Literal[1800] = 1800
    retries: Literal[0] = 0
    retained_task_ids: List[str]
    recovered_task_bindings: List[CodexE2BTaskBindingV1]
    frozen_solver_result_sha256: str = Field(min_length=64, max_length=64)
    parity_report_sha256: str = Field(min_length=64, max_length=64)
    authorization_basis: Literal["user_project_level_external_model_authorization"] = (
        "user_project_level_external_model_authorization"
    )
    grader_authorized: Literal[False] = False


def compile_scope(campaign_root: Path, output_root: Path) -> tuple[CodexE2BSliceScopeV1, Path, Path]:
    manifest_path = campaign_root / "matched_screening_campaign.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assignments = {item["blind_task_id"]: item for item in manifest["assignments"]}
    bindings: List[CodexE2BTaskBindingV1] = []
    for task_id in TASK_IDS:
        item = assignments.get(task_id)
        if item is None or item.get("package_fingerprint") != EXPECTED_FINGERPRINTS[task_id]:
            raise ValueError(f"codex_slice_binding_mismatch:{task_id}")
        candidate = campaign_root / "blind_staging" / "candidate_packages" / task_id
        if not candidate.is_dir():
            raise ValueError(f"codex_slice_candidate_missing:{task_id}")
        bindings.append(CodexE2BTaskBindingV1(
            blind_task_id=task_id,
            package_fingerprint=EXPECTED_FINGERPRINTS[task_id],
            candidate_tree_sha256=_tree(candidate),
            expected_deliverable=_expected_deliverable(candidate),
        ))
    scope = CodexE2BSliceScopeV1(campaign_id=manifest["campaign_id"], task_bindings=bindings)
    encoded = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(encoded)
    scope_path = output_root / "governance" / "scopes" / f"{scope_sha}.json"
    receipt_path = output_root / "governance" / "receipts" / f"{scope_sha}.json"
    if scope_path.exists() and scope_path.read_bytes() != encoded + b"\n":
        raise ValueError("codex_slice_scope_immutable_collision")
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    scope_path.write_bytes(encoded + b"\n")
    receipt = CodexE2BSliceReceiptV1(scope_sha256=scope_sha)
    _write_json(receipt_path, receipt.model_dump(mode="json"))
    return scope, scope_path, receipt_path


def _config_toml(config: ProviderConfig) -> str:
    base = config.base_url.rstrip("/")
    return (
        'model = "gpt-5.6-sol"\n'
        'model_provider = "tuzi"\n'
        'model_reasoning_effort = "medium"\n'
        '[model_providers.tuzi]\n'
        'name = "Tuzi"\n'
        f'base_url = {json.dumps(base)}\n'
        'env_key = "TUZI_API_KEY"\n'
        'wire_api = "responses"\n'
        'request_max_retries = 0\n'
        'stream_max_retries = 0\n'
    )


def _redact(value: str, secrets: List[str]) -> str:
    redacted = value
    for secret in secrets:
        if secret:
            redacted = redacted.replace(secret, "[REDACTED_SECRET]")
    return redacted


def _assert_secrets_absent(value: str, secrets: List[str]) -> None:
    if any(secret and secret in value for secret in secrets):
        raise ValueError("codex_slice_secret_persistence_blocked")


def _probe_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet.append(["item", "amount"])
    sheet.append(["alpha", 7])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _parse_jsonl(text: str) -> tuple[List[str], bool, bool, Dict[str, int]]:
    event_types: List[str] = []
    command = False
    turn = False
    usage: Dict[str, int] = {}
    for line in text.splitlines():
        if not line.strip():
            continue
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        event_type = str(event.get("type", ""))
        if event_type:
            event_types.append(event_type)
        item = event.get("item") or {}
        command = command or item.get("type") in {"command_execution", "command"}
        turn = turn or event_type.startswith("turn.")
        candidate_usage = event.get("usage") or (event.get("turn") or {}).get("usage") or {}
        for key in ("input_tokens", "output_tokens", "total_tokens"):
            value = candidate_usage.get(key)
            if isinstance(value, int):
                usage[key] = value
    return sorted(set(event_types)), command, turn, usage


def _codex_exit_code(text: str) -> Optional[int]:
    for line in reversed(text.splitlines()):
        try:
            event = json.loads(line)
        except json.JSONDecodeError:
            continue
        if event.get("type") == "runner.exit" and isinstance(event.get("exit_code"), int):
            return event["exit_code"]
    return None


def _inspect(data: Optional[bytes], input_hashes: set[str]) -> DeliveryInspectionV1:
    if not data:
        return DeliveryInspectionV1(exists=data is not None)
    digest = _sha_bytes(data)
    result = DeliveryInspectionV1(exists=True, nonempty=True, sha256=digest, size_bytes=len(data), differs_from_inputs=digest not in input_hashes)
    try:
        workbook = load_workbook(BytesIO(data), data_only=False)
        result.xlsx_openable = True
        result.has_nonempty_sheet = any(
            any(cell.value not in (None, "") for row in sheet.iter_rows() for cell in row)
            for sheet in workbook.worksheets
        )
    except Exception:
        pass
    return result


def _expected_deliverable(package: Path) -> str:
    contract = json.loads((package / "deliverable_contract.json").read_text(encoding="utf-8"))
    deliverables = contract.get("deliverables") or []
    if len(deliverables) != 1:
        raise ValueError("codex_requires_single_deliverable")
    relative = deliverables[0].get("relative_path")
    if not isinstance(relative, str):
        raise ValueError("codex_deliverable_path_missing")
    return relative


class CodexE2BRunner:
    def __init__(self, sandbox_factory: Optional[Callable[..., Any]] = None):
        self._factory = sandbox_factory

    def _factory_call(self, template: str, timeout: int, e2b_api_key: str) -> Any:
        if self._factory is None:
            from e2b import Sandbox
            return Sandbox.create(template=template, timeout=timeout, api_key=e2b_api_key)
        return self._factory(template=template, timeout=timeout, api_key=e2b_api_key)

    def _run_one(self, *, kind: Literal["public_probe", "private_task"], task_id: str,
                 prompt: str, uploads: Dict[str, bytes], expected: str, timeout: int,
                 config: ProviderConfig, e2b_api_key: str, output_dir: Path) -> CodexE2BOutcomeV1:
        sandbox = None
        stdout = ""
        stderr = ""
        secrets = [config.api_key, e2b_api_key]
        try:
            sandbox = self._factory_call("codex", timeout + 120, e2b_api_key)
            sandbox.files.make_dir(str(REMOTE_ROOT / "reference_files"))
            sandbox.files.make_dir(str(REMOTE_ROOT / "deliverable_files"))
            sandbox.files.make_dir(str(CODEX_HOME))
            install = sandbox.commands.run(
                "bash -lc 'python3 -m pip install --disable-pip-version-check openpyxl; code=$?; echo __INSTALL_EXIT__:$code; exit 0'",
                timeout=300,
            )
            install_stdout = _redact(getattr(install, "stdout", "") or "", secrets)
            install_stderr = _redact(getattr(install, "stderr", "") or "", secrets)
            if "__INSTALL_EXIT__:0" not in install_stdout:
                raise RuntimeError(
                    "openpyxl_install_failed:"
                    + install_stdout[-2000:]
                    + install_stderr[-2000:]
                )
            sandbox.files.write(str(CODEX_HOME / "config.toml"), _config_toml(config))
            sandbox.files.write(str(REMOTE_ROOT / "TASK.md"), prompt)
            for relative, data in uploads.items():
                target = REMOTE_ROOT / relative
                sandbox.files.make_dir(str(target.parent))
                sandbox.files.write(str(target), data)
            command = (
                "bash -lc 'set +e; codex exec --json --sandbox workspace-write "
                "--skip-git-repo-check -C /home/user/task - < /home/user/task/TASK.md; "
                "code=$?; printf \"{\\\"type\\\":\\\"runner.exit\\\",\\\"exit_code\\\":%s}\\n\" \"$code\"; exit 0'"
            )
            process = sandbox.commands.run(
                command,
                timeout=timeout,
                envs={"TUZI_API_KEY": config.api_key, "CODEX_HOME": str(CODEX_HOME)},
            )
            stdout = _redact(getattr(process, "stdout", "") or "", secrets)
            stderr = _redact(getattr(process, "stderr", "") or "", secrets)
            exit_code = _codex_exit_code(stdout)
            events, command_seen, turn_seen, usage = _parse_jsonl(stdout)
            remote_expected = str(REMOTE_ROOT / expected)
            data = sandbox.files.read(remote_expected, format="bytes") if sandbox.files.exists(remote_expected) else None
            input_hashes = {_sha_bytes(value) for name, value in uploads.items() if name.lower().endswith(".xlsx")}
            delivery = _inspect(data, input_hashes)
            valid = all((delivery.nonempty, delivery.xlsx_openable, delivery.has_nonempty_sheet, delivery.differs_from_inputs))
            protocol_valid = turn_seen and command_seen
            infrastructure_markers = (
                "/responses", "authentication", "unauthorized", "401", "403", "protocol",
                "stream error", "stream disconnected", "response.completed",
                "incomplete chunked read", "peer closed connection",
            )
            diagnostic_text = (stdout + "\n" + stderr).lower()
            infrastructure = exit_code not in (0, None) and any(marker in diagnostic_text for marker in infrastructure_markers)
            status = "pass" if exit_code == 0 and protocol_valid and valid else ("infrastructure_failed" if infrastructure or not protocol_valid else "delivery_failed")
            failure = None if status == "pass" else ("codex_protocol_or_provider_failure" if status == "infrastructure_failed" else "invalid_or_missing_delivery")
            outcome = CodexE2BOutcomeV1(kind=kind, task_id=task_id, status=status, sandbox_id=str(getattr(sandbox, "sandbox_id", getattr(sandbox, "id", ""))) or None, exit_code=exit_code, event_types=events, command_event_seen=command_seen, turn_event_seen=turn_seen, usage=usage, delivery=delivery, first_failure=failure)
            if data:
                delivery_path = output_dir / expected
                delivery_path.parent.mkdir(parents=True, exist_ok=True)
                delivery_path.write_bytes(data)
        except Exception as exc:
            failure = _redact(f"{type(exc).__name__}:{exc}", secrets)
            stderr = stderr or failure
            outcome = CodexE2BOutcomeV1(kind=kind, task_id=task_id, status="infrastructure_failed", sandbox_id=str(getattr(sandbox, "sandbox_id", getattr(sandbox, "id", ""))) or None if sandbox else None, first_failure=failure)
        finally:
            if sandbox is not None:
                try:
                    sandbox.kill()
                except Exception:
                    pass
        output_dir.mkdir(parents=True, exist_ok=True)
        _assert_secrets_absent(stdout + stderr + outcome.model_dump_json(), secrets)
        (output_dir / "codex.jsonl").write_text(stdout, encoding="utf-8")
        (output_dir / "stderr.txt").write_text(stderr, encoding="utf-8")
        outcome.jsonl_sha256 = _sha(output_dir / "codex.jsonl")
        outcome.stderr_sha256 = _sha(output_dir / "stderr.txt")
        _write_json(output_dir / "outcome.json", outcome.model_dump(mode="json"))
        return outcome

    def execute(self, *, scope_path: Path, receipt_path: Path, campaign_root: Path,
                output_root: Path, config: ProviderConfig, e2b_api_key: str) -> tuple[CodexE2BSliceResultV1, Path]:
        if not e2b_api_key:
            raise ValueError("e2b_api_key_missing")
        scope = CodexE2BSliceScopeV1.model_validate_json(scope_path.read_text(encoding="utf-8"))
        canonical = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
        scope_sha = _sha_bytes(canonical)
        receipt = CodexE2BSliceReceiptV1.model_validate_json(receipt_path.read_text(encoding="utf-8"))
        if receipt.scope_sha256 != scope_sha or receipt.consumed_at is not None:
            raise ValueError("codex_slice_receipt_invalid_or_consumed")
        receipt.consumed_at = _now()
        _write_json(receipt_path, receipt.model_dump(mode="json"))
        probe_prompt = "Use Python and openpyxl to read reference_files/source.xlsx, then create exactly deliverable_files/codex_probe.xlsx with a non-empty sheet named Result and include the source row plus a new status column. Do not merely copy the input."
        public = self._run_one(kind="public_probe", task_id="public_codex_xlsx_probe", prompt=probe_prompt, uploads={"reference_files/source.xlsx": _probe_xlsx()}, expected="deliverable_files/codex_probe.xlsx", timeout=scope.public_timeout_seconds, config=config, e2b_api_key=e2b_api_key, output_dir=output_root / "public_probe")
        private: List[CodexE2BOutcomeV1] = []
        if public.status == "pass":
            for binding in scope.task_bindings:
                package = campaign_root / "blind_staging" / "candidate_packages" / binding.blind_task_id
                if _tree(package) != binding.candidate_tree_sha256:
                    raise ValueError(f"codex_slice_candidate_tree_drift:{binding.blind_task_id}")
                row = json.loads((package / "dataset_row.json").read_text(encoding="utf-8"))
                contract = (package / "deliverable_contract.json").read_bytes()
                uploads = {"deliverable_contract.json": contract}
                for reference in sorted((package / "reference_files").glob("*.xlsx")):
                    uploads[f"reference_files/{reference.name}"] = reference.read_bytes()
                prompt = row["prompt"] + "\nThe authoritative deliverable contract is available at deliverable_contract.json. Use Python/openpyxl as needed."
                outcome = self._run_one(kind="private_task", task_id=binding.blind_task_id, prompt=prompt, uploads=uploads, expected=binding.expected_deliverable, timeout=scope.task_timeout_seconds, config=config, e2b_api_key=e2b_api_key, output_dir=output_root / binding.blind_task_id)
                private.append(outcome)
                if outcome.status == "infrastructure_failed":
                    break
        decision = "incomplete" if public.status == "infrastructure_failed" or any(item.status == "infrastructure_failed" for item in private) else ("compatibility_pass" if len(private) == 2 and all(item.status == "pass" for item in private) else "compatibility_failed")
        result = CodexE2BSliceResultV1(scope_sha256=scope_sha, completed_at=_now(), public_probe=public, private_tasks=private, decision=decision)
        result_path = output_root / "slice_result.json"
        _assert_secrets_absent(result.model_dump_json(), [config.api_key, e2b_api_key])
        _write_json(result_path, result.model_dump(mode="json"))
        return result, result_path


def compile_recovery_scope(*, campaign_root: Path, retained_result_path: Path,
                           parity_report_path: Path, output_root: Path) -> tuple[CodexE2BRecoveryScopeV1, Path, Path]:
    retained = CodexE2BSliceResultV1.model_validate_json(retained_result_path.read_text(encoding="utf-8"))
    if retained.public_probe.status != "pass":
        raise ValueError("codex_recovery_public_probe_not_passed")
    passed = [item for item in retained.private_tasks if item.status == "pass"]
    failed = [item for item in retained.private_tasks if item.status == "infrastructure_failed"]
    if [item.task_id for item in passed] != [TASK_IDS[0]] or [item.task_id for item in failed] != [TASK_IDS[1]]:
        raise ValueError("codex_recovery_retained_shape_invalid")
    manifest = json.loads((campaign_root / "matched_screening_campaign.json").read_text(encoding="utf-8"))
    assignment = next(item for item in manifest["assignments"] if item["blind_task_id"] == TASK_IDS[1])
    if assignment["package_fingerprint"] != EXPECTED_FINGERPRINTS[TASK_IDS[1]]:
        raise ValueError("codex_recovery_package_fingerprint_drift")
    candidate = campaign_root / "blind_staging" / "candidate_packages" / TASK_IDS[1]
    binding = CodexE2BTaskBindingV1(
        blind_task_id=TASK_IDS[1], package_fingerprint=EXPECTED_FINGERPRINTS[TASK_IDS[1]],
        candidate_tree_sha256=_tree(candidate),
        expected_deliverable=_expected_deliverable(candidate),
    )
    scope = CodexE2BRecoveryScopeV1(
        campaign_id=manifest["campaign_id"], recovered_task=binding,
        retained_result_sha256=_sha(retained_result_path),
        retained_public_outcome_sha256=_sha_bytes(retained.public_probe.model_dump_json().encode()),
        retained_passed_task_outcome_sha256=_sha_bytes(passed[0].model_dump_json().encode()),
        parity_report_sha256=_sha(parity_report_path),
    )
    encoded = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(encoded)
    scope_path = output_root / "governance" / "scopes" / f"{scope_sha}.json"
    receipt_path = output_root / "governance" / "receipts" / f"{scope_sha}.json"
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    scope_path.write_bytes(encoded + b"\n")
    _write_json(receipt_path, CodexE2BSliceReceiptV1(scope_sha256=scope_sha).model_dump(mode="json"))
    return scope, scope_path, receipt_path


def execute_recovery(*, scope_path: Path, receipt_path: Path, retained_result_path: Path,
                     campaign_root: Path, output_root: Path, config: ProviderConfig,
                     e2b_api_key: str, sandbox_factory: Optional[Callable[..., Any]] = None
                     ) -> tuple[CodexE2BSliceResultV1, Path]:
    scope = CodexE2BRecoveryScopeV1.model_validate_json(scope_path.read_text(encoding="utf-8"))
    canonical = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(canonical)
    receipt = CodexE2BSliceReceiptV1.model_validate_json(receipt_path.read_text(encoding="utf-8"))
    if receipt.scope_sha256 != scope_sha or receipt.consumed_at is not None:
        raise ValueError("codex_recovery_receipt_invalid_or_consumed")
    if _sha(retained_result_path) != scope.retained_result_sha256:
        raise ValueError("codex_recovery_retained_result_drift")
    retained = CodexE2BSliceResultV1.model_validate_json(retained_result_path.read_text(encoding="utf-8"))
    passed = next(item for item in retained.private_tasks if item.status == "pass")
    if _sha_bytes(retained.public_probe.model_dump_json().encode()) != scope.retained_public_outcome_sha256:
        raise ValueError("codex_recovery_public_evidence_drift")
    if _sha_bytes(passed.model_dump_json().encode()) != scope.retained_passed_task_outcome_sha256:
        raise ValueError("codex_recovery_passed_evidence_drift")
    receipt.consumed_at = _now()
    _write_json(receipt_path, receipt.model_dump(mode="json"))
    binding = scope.recovered_task
    package = campaign_root / "blind_staging" / "candidate_packages" / binding.blind_task_id
    if _tree(package) != binding.candidate_tree_sha256:
        raise ValueError("codex_recovery_candidate_tree_drift")
    row = json.loads((package / "dataset_row.json").read_text(encoding="utf-8"))
    uploads = {"deliverable_contract.json": (package / "deliverable_contract.json").read_bytes()}
    for reference in sorted((package / "reference_files").glob("*.xlsx")):
        uploads[f"reference_files/{reference.name}"] = reference.read_bytes()
    prompt = row["prompt"] + "\nThe authoritative deliverable contract is available at deliverable_contract.json. Use Python/openpyxl as needed."
    outcome = CodexE2BRunner(sandbox_factory)._run_one(
        kind="private_task", task_id=binding.blind_task_id, prompt=prompt, uploads=uploads,
        expected=binding.expected_deliverable, timeout=scope.task_timeout_seconds,
        config=config, e2b_api_key=e2b_api_key, output_dir=output_root / binding.blind_task_id,
    )
    tasks = [passed, outcome]
    decision = "compatibility_pass" if outcome.status == "pass" else (
        "incomplete" if outcome.status == "infrastructure_failed" else "compatibility_failed"
    )
    result = CodexE2BSliceResultV1(
        scope_sha256=scope_sha, completed_at=_now(), public_probe=retained.public_probe,
        private_tasks=tasks, decision=decision,
    )
    result_path = output_root / "recovery_result.json"
    _assert_secrets_absent(result.model_dump_json(), [config.api_key, e2b_api_key])
    _write_json(result_path, result.model_dump(mode="json"))
    return result, result_path


def compile_screening_solver_scope(*, campaign_root: Path, retained_result_path: Path,
                                   parity_report_path: Path, output_root: Path
                                   ) -> tuple[CodexE2BScreeningSolverScopeV1, Path, Path]:
    retained = CodexE2BSliceResultV1.model_validate_json(retained_result_path.read_text(encoding="utf-8"))
    if retained.decision != "compatibility_pass" or len(retained.private_tasks) != 2:
        raise ValueError("codex_screening_retained_compatibility_not_passed")
    manifest = json.loads((campaign_root / "matched_screening_campaign.json").read_text(encoding="utf-8"))
    bindings: List[CodexE2BTaskBindingV1] = []
    for assignment in manifest["assignments"]:
        task_id = assignment["blind_task_id"]
        candidate = campaign_root / "blind_staging" / "candidate_packages" / task_id
        if not candidate.is_dir():
            raise ValueError(f"codex_screening_candidate_missing:{task_id}")
        bindings.append(CodexE2BTaskBindingV1(
            blind_task_id=task_id,
            package_fingerprint=assignment["package_fingerprint"],
            candidate_tree_sha256=_tree(candidate),
            expected_deliverable=_expected_deliverable(candidate),
        ))
    if len(bindings) != 12 or len({item.blind_task_id for item in bindings}) != 12:
        raise ValueError("codex_screening_requires_twelve_unique_tasks")
    retained_ids = [item.task_id for item in retained.private_tasks]
    new_ids = [item.blind_task_id for item in bindings if item.blind_task_id not in retained_ids]
    if len(retained_ids) != 2 or len(new_ids) != 10:
        raise ValueError("codex_screening_retained_new_partition_invalid")
    scope = CodexE2BScreeningSolverScopeV1(
        campaign_id=manifest["campaign_id"], task_bindings=bindings,
        retained_task_ids=retained_ids, authorized_new_task_ids=new_ids,
        retained_compatibility_result_sha256=_sha(retained_result_path),
        parity_report_sha256=_sha(parity_report_path),
    )
    encoded = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(encoded)
    scope_path = output_root / "governance" / "scopes" / f"{scope_sha}.json"
    receipt_path = output_root / "governance" / "receipts" / f"{scope_sha}.json"
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    scope_path.write_bytes(encoded + b"\n")
    _write_json(receipt_path, CodexE2BSliceReceiptV1(scope_sha256=scope_sha).model_dump(mode="json"))
    return scope, scope_path, receipt_path


def execute_screening_solver(*, scope_path: Path, receipt_path: Path, retained_result_path: Path,
                             campaign_root: Path, output_root: Path, config: ProviderConfig,
                             e2b_api_key: str, sandbox_factory: Optional[Callable[..., Any]] = None
                             ) -> tuple[CodexE2BScreeningSolverResultV1, Path]:
    scope = CodexE2BScreeningSolverScopeV1.model_validate_json(scope_path.read_text(encoding="utf-8"))
    canonical = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(canonical)
    receipt = CodexE2BSliceReceiptV1.model_validate_json(receipt_path.read_text(encoding="utf-8"))
    if receipt.scope_sha256 != scope_sha or receipt.consumed_at is not None:
        raise ValueError("codex_screening_receipt_invalid_or_consumed")
    if _sha(retained_result_path) != scope.retained_compatibility_result_sha256:
        raise ValueError("codex_screening_retained_result_drift")
    retained = CodexE2BSliceResultV1.model_validate_json(retained_result_path.read_text(encoding="utf-8"))
    retained_outcomes = retained.private_tasks
    if [item.task_id for item in retained_outcomes] != scope.retained_task_ids or any(item.status != "pass" for item in retained_outcomes):
        raise ValueError("codex_screening_retained_outcomes_invalid")
    receipt.consumed_at = _now()
    _write_json(receipt_path, receipt.model_dump(mode="json"))
    runner = CodexE2BRunner(sandbox_factory)
    outcomes = list(retained_outcomes)
    binding_map = {item.blind_task_id: item for item in scope.task_bindings}
    for task_id in scope.authorized_new_task_ids:
        binding = binding_map[task_id]
        package = campaign_root / "blind_staging" / "candidate_packages" / task_id
        if _tree(package) != binding.candidate_tree_sha256:
            raise ValueError(f"codex_screening_candidate_tree_drift:{task_id}")
        row = json.loads((package / "dataset_row.json").read_text(encoding="utf-8"))
        uploads = {"deliverable_contract.json": (package / "deliverable_contract.json").read_bytes()}
        for reference in sorted((package / "reference_files").glob("*.xlsx")):
            uploads[f"reference_files/{reference.name}"] = reference.read_bytes()
        prompt = row["prompt"] + "\nThe authoritative deliverable contract is available at deliverable_contract.json. Use Python/openpyxl as needed."
        outcomes.append(runner._run_one(
            kind="private_task", task_id=task_id, prompt=prompt, uploads=uploads,
            expected=binding.expected_deliverable, timeout=scope.task_timeout_seconds,
            config=config, e2b_api_key=e2b_api_key, output_dir=output_root / task_id,
        ))
    valid_count = sum(item.status == "pass" for item in outcomes)
    decision = "incomplete" if any(item.status == "infrastructure_failed" for item in outcomes) else "solver_complete"
    result = CodexE2BScreeningSolverResultV1(
        scope_sha256=scope_sha, completed_at=_now(), retained_task_ids=scope.retained_task_ids,
        newly_executed_task_ids=scope.authorized_new_task_ids, task_outcomes=outcomes,
        valid_delivery_count=valid_count, decision=decision,
    )
    result_path = output_root / "screening_solver_result.json"
    _assert_secrets_absent(result.model_dump_json(), [config.api_key, e2b_api_key])
    _write_json(result_path, result.model_dump(mode="json"))
    return result, result_path


def compile_screening_recovery_scope(*, campaign_root: Path, frozen_result_path: Path,
                                     parity_report_path: Path, output_root: Path
                                     ) -> tuple[CodexE2BScreeningRecoveryScopeV1, Path, Path]:
    frozen = CodexE2BScreeningSolverResultV1.model_validate_json(frozen_result_path.read_text(encoding="utf-8"))
    retained_ids = [item.task_id for item in frozen.task_outcomes if item.status == "pass"]
    failed_ids = [item.task_id for item in frozen.task_outcomes if item.status != "pass"]
    if len(retained_ids) != 3 or len(failed_ids) != 9:
        raise ValueError("codex_screening_recovery_expected_three_pass_nine_failed")
    manifest = json.loads((campaign_root / "matched_screening_campaign.json").read_text(encoding="utf-8"))
    assignments = {item["blind_task_id"]: item for item in manifest["assignments"]}
    bindings: List[CodexE2BTaskBindingV1] = []
    for task_id in failed_ids:
        package = campaign_root / "blind_staging" / "candidate_packages" / task_id
        bindings.append(CodexE2BTaskBindingV1(
            blind_task_id=task_id, package_fingerprint=assignments[task_id]["package_fingerprint"],
            candidate_tree_sha256=_tree(package), expected_deliverable=_expected_deliverable(package),
        ))
    scope = CodexE2BScreeningRecoveryScopeV1(
        campaign_id=manifest["campaign_id"], retained_task_ids=retained_ids,
        recovered_task_bindings=bindings, frozen_solver_result_sha256=_sha(frozen_result_path),
        parity_report_sha256=_sha(parity_report_path),
    )
    encoded = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(encoded)
    scope_path = output_root / "governance" / "scopes" / f"{scope_sha}.json"
    receipt_path = output_root / "governance" / "receipts" / f"{scope_sha}.json"
    scope_path.parent.mkdir(parents=True, exist_ok=True)
    scope_path.write_bytes(encoded + b"\n")
    _write_json(receipt_path, CodexE2BSliceReceiptV1(scope_sha256=scope_sha).model_dump(mode="json"))
    return scope, scope_path, receipt_path


def execute_screening_recovery(*, scope_path: Path, receipt_path: Path, frozen_result_path: Path,
                               campaign_root: Path, output_root: Path, config: ProviderConfig,
                               e2b_api_key: str, sandbox_factory: Optional[Callable[..., Any]] = None
                               ) -> tuple[CodexE2BScreeningSolverResultV1, Path]:
    scope = CodexE2BScreeningRecoveryScopeV1.model_validate_json(scope_path.read_text(encoding="utf-8"))
    canonical = json.dumps(scope.model_dump(mode="json"), sort_keys=True, separators=(",", ":")).encode()
    scope_sha = _sha_bytes(canonical)
    receipt = CodexE2BSliceReceiptV1.model_validate_json(receipt_path.read_text(encoding="utf-8"))
    if receipt.scope_sha256 != scope_sha or receipt.consumed_at is not None:
        raise ValueError("codex_screening_recovery_receipt_invalid_or_consumed")
    if _sha(frozen_result_path) != scope.frozen_solver_result_sha256:
        raise ValueError("codex_screening_recovery_frozen_result_drift")
    frozen = CodexE2BScreeningSolverResultV1.model_validate_json(frozen_result_path.read_text(encoding="utf-8"))
    retained = [item for item in frozen.task_outcomes if item.task_id in scope.retained_task_ids]
    if len(retained) != 3 or any(item.status != "pass" for item in retained):
        raise ValueError("codex_screening_recovery_retained_invalid")
    receipt.consumed_at = _now()
    _write_json(receipt_path, receipt.model_dump(mode="json"))
    runner = CodexE2BRunner(sandbox_factory)
    outcomes = list(retained)
    for binding in scope.recovered_task_bindings:
        package = campaign_root / "blind_staging" / "candidate_packages" / binding.blind_task_id
        if _tree(package) != binding.candidate_tree_sha256:
            raise ValueError(f"codex_screening_recovery_tree_drift:{binding.blind_task_id}")
        row = json.loads((package / "dataset_row.json").read_text(encoding="utf-8"))
        uploads = {"deliverable_contract.json": (package / "deliverable_contract.json").read_bytes()}
        for reference in sorted((package / "reference_files").glob("*.xlsx")):
            uploads[f"reference_files/{reference.name}"] = reference.read_bytes()
        prompt = row["prompt"] + "\nThe authoritative deliverable contract is available at deliverable_contract.json. Use Python/openpyxl as needed."
        outcomes.append(runner._run_one(
            kind="private_task", task_id=binding.blind_task_id, prompt=prompt, uploads=uploads,
            expected=binding.expected_deliverable, timeout=scope.task_timeout_seconds,
            config=config, e2b_api_key=e2b_api_key, output_dir=output_root / binding.blind_task_id,
        ))
    valid_count = sum(item.status == "pass" for item in outcomes)
    decision = "incomplete" if any(item.status == "infrastructure_failed" for item in outcomes) else "solver_complete"
    result = CodexE2BScreeningSolverResultV1(
        scope_sha256=scope_sha, completed_at=_now(), retained_task_ids=scope.retained_task_ids,
        newly_executed_task_ids=[item.blind_task_id for item in scope.recovered_task_bindings],
        task_outcomes=outcomes, valid_delivery_count=valid_count, decision=decision,
    )
    result_path = output_root / "screening_recovery_result.json"
    _assert_secrets_absent(result.model_dump_json(), [config.api_key, e2b_api_key])
    _write_json(result_path, result.model_dump(mode="json"))
    return result, result_path
