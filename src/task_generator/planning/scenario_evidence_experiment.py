"""Minimal R10 professional-Skill A/B evidence-bundle experiment support."""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from typing import Any, Callable, Literal

from openpyxl import load_workbook
from pydantic import Field, model_validator

from task_generator.core.scenario_first import ProfessionalRuleSetV1, ScenarioBibleV1, ScenarioFirstModel, WorkSeedV1
from task_generator.substrate.professional_skills import LoadedProfessionalSkillV1


Condition = Literal["without_skill", "with_skill"]
AdmissionDecision = Literal["pass", "blocked", "incomplete"]
ExperimentDecision = Literal["skill_effect_supported", "skill_revision_required", "incomplete"]


class ScenarioEvidenceSessionV1(ScenarioFirstModel):
    session_id: str = Field(min_length=1)
    scenario_id: str = Field(min_length=1)
    domain: Literal["audit_compliance", "procurement_operations"]
    condition: Condition
    bible_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    rule_set_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    skill_id: str | None = None
    skill_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    skill_source_ids: list[str] = Field(default_factory=list)
    model: Literal["gpt-5.6-terra"] = "gpt-5.6-terra"
    reasoning_effort: Literal["medium"] = "medium"
    image: str = Field(min_length=1)
    image_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    timeout_seconds: int = Field(default=1800, ge=1, le=1800)

    @model_validator(mode="after")
    def _skill_condition_matches(self) -> "ScenarioEvidenceSessionV1":
        if self.condition == "with_skill" and not (self.skill_id and self.skill_sha256):
            raise ValueError("with_skill_session_requires_skill_identity")
        if self.condition == "with_skill" and not self.skill_source_ids:
            raise ValueError("with_skill_session_requires_curated_source_ids")
        if self.condition == "without_skill" and (self.skill_id or self.skill_sha256):
            raise ValueError("without_skill_session_must_not_include_skill")
        if self.condition == "without_skill" and self.skill_source_ids:
            raise ValueError("without_skill_session_must_not_include_skill_sources")
        if len(set(self.skill_source_ids)) != len(self.skill_source_ids):
            raise ValueError("duplicate_skill_source_id")
        return self


class ScenarioEvidenceExperimentPlanV1(ScenarioFirstModel):
    plan_version: Literal["r10.scenario_evidence_experiment_plan.1"] = "r10.scenario_evidence_experiment_plan.1"
    sessions: list[ScenarioEvidenceSessionV1] = Field(min_length=4, max_length=4)

    @model_validator(mode="after")
    def _four_fixed_sessions(self) -> "ScenarioEvidenceExperimentPlanV1":
        identities = {(item.scenario_id, item.condition) for item in self.sessions}
        if len(identities) != 4 or {item.condition for item in self.sessions} != {"without_skill", "with_skill"}:
            raise ValueError("experiment_requires_two_complete_ab_pairs")
        if len({item.session_id for item in self.sessions}) != 4:
            raise ValueError("duplicate_evidence_session_id")
        if len({item.scenario_id for item in self.sessions}) != 2:
            raise ValueError("experiment_requires_two_scenarios")
        return self


class EvidenceAdmissionFindingV1(ScenarioFirstModel):
    code: str = Field(min_length=1)
    passed: bool
    details: dict[str, Any] = Field(default_factory=dict)


class ScenarioEvidenceAdmissionReportV1(ScenarioFirstModel):
    report_version: Literal["r10.scenario_evidence_admission_report.1"] = "r10.scenario_evidence_admission_report.1"
    session_id: str = Field(min_length=1)
    output_tree_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    decision: AdmissionDecision
    findings: list[EvidenceAdmissionFindingV1]
    first_failure: str | None = None


class ConditionBlindFindingV1(ScenarioFirstModel):
    favored_bundle: Literal["bundle_1", "bundle_2", "neither"]
    dimension: Literal["professional_realism", "evidence_naturalness", "judgment_depth", "leakage_risk"]
    file_paths: list[str] = Field(min_length=1, max_length=4)
    reason: str = Field(min_length=20, max_length=1000)


class ConditionBlindPairReviewV1(ScenarioFirstModel):
    review_version: Literal["r10.condition_blind_pair_review.1"] = "r10.condition_blind_pair_review.1"
    scenario_id: str = Field(min_length=1)
    bundle_order: tuple[Literal["without_skill", "with_skill"], Literal["without_skill", "with_skill"]]
    provider_call_count: int = Field(ge=0, le=2)
    provider_retry_count: int = Field(default=0, ge=0, le=1)
    response_sha256: str | None = Field(default=None, pattern=r"^[0-9a-f]{64}$")
    decision: AdmissionDecision
    findings: list[ConditionBlindFindingV1] = Field(default_factory=list)
    summary: str | None = None
    first_failure: str | None = None
    usage: dict[str, int | None] = Field(default_factory=dict)


class ScenarioEvidenceExperimentResultV1(ScenarioFirstModel):
    result_version: Literal["r10.scenario_evidence_experiment_result.1"] = "r10.scenario_evidence_experiment_result.1"
    plan_sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    admission_reports: list[ScenarioEvidenceAdmissionReportV1]
    blind_reviews: list[ConditionBlindPairReviewV1]
    decision: ExperimentDecision
    reasons: list[str] = Field(default_factory=list)


class ScenarioEvidenceExperiment:
    """Stages a frozen A/B pair and performs only lightweight admission."""

    answer_labels = ("questionable", "exception", "requires follow-up")
    allowed_suffixes = {".xlsx", ".csv", ".txt", ".pdf", ".docx"}

    @staticmethod
    def sha256_path(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def sha256_json(value: Any) -> str:
        return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()

    def stage_session(
        self,
        *,
        workspace: Path,
        session: ScenarioEvidenceSessionV1,
        bible: ScenarioBibleV1,
        rules: ProfessionalRuleSetV1,
        skill: LoadedProfessionalSkillV1 | None,
    ) -> None:
        if workspace.exists():
            raise FileExistsError("scenario_evidence_workspace_exists")
        if bible.canonical_sha256() != session.bible_sha256 or rules.canonical_sha256() != session.rule_set_sha256:
            raise ValueError("scenario_evidence_frozen_input_drift")
        if session.condition == "with_skill":
            if skill is None or skill.entry.skill_id != session.skill_id:
                raise ValueError("scenario_evidence_skill_selection_mismatch")
            if self.sha256_json(skill.model_dump(mode="json")) != session.skill_sha256:
                raise ValueError("scenario_evidence_skill_content_drift")
        elif skill is not None:
            raise ValueError("scenario_evidence_without_skill_must_not_stage_skill")

        (workspace / "teacher").mkdir(parents=True)
        (workspace / "candidate").mkdir()
        self._write(workspace / "teacher" / "scenario_bible.json", bible.model_dump(mode="json"))
        self._write(workspace / "teacher" / "professional_rules.json", rules.model_dump(mode="json"))
        self._write(workspace / "teacher" / "condition.json", session.model_dump(mode="json"))
        if skill is not None:
            package = workspace / "teacher" / "professional_skill"
            package.mkdir()
            (package / "SKILL.md").write_text(skill.skill_markdown, encoding="utf-8")
            for relative, content in skill.reference_markdown.items():
                target = package / relative
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_text(content, encoding="utf-8")
        (workspace / "TASK.md").write_text(self._task_prompt(session.condition), encoding="utf-8")

    def admit(
        self,
        *,
        session: ScenarioEvidenceSessionV1,
        workspace: Path,
        bible: ScenarioBibleV1,
        require_libreoffice: bool = False,
    ) -> ScenarioEvidenceAdmissionReportV1:
        findings: list[EvidenceAdmissionFindingV1] = []
        candidate = workspace / "candidate"
        teacher = workspace / "teacher"
        files = sorted(path for path in candidate.rglob("*") if path.is_file()) if candidate.is_dir() else []
        findings.append(self._finding("candidate_teacher_isolation", candidate.is_dir() and teacher.is_dir() and not any("teacher" in path.parts for path in files)))
        tabular = [path for path in files if path.suffix.lower() in {".xlsx", ".csv"}]
        narrative = [path for path in files if path.suffix.lower() in {".txt", ".docx", ".pdf"}]
        findings.append(self._finding("natural_file_mix", bool(tabular) and bool(narrative), {"file_count": len(files)}))
        findings.append(self._finding("candidate_files_nonempty", bool(files) and all(path.stat().st_size > 0 for path in files)))
        format_errors: list[str] = []
        text_parts: list[str] = []
        for path in files:
            if path.suffix.lower() not in self.allowed_suffixes:
                format_errors.append(f"unsupported_suffix:{path.name}")
                continue
            try:
                text_parts.append(self._extract_candidate_text(path, require_libreoffice=require_libreoffice))
            except Exception as exc:
                format_errors.append(f"invalid_file:{path.name}:{type(exc).__name__}")
        findings.append(self._finding("candidate_files_openable", not format_errors, {"errors": format_errors}))

        extension = teacher / "scenario_extension.md"
        evidence_map_path = teacher / "evidence_map.json"
        map_payload: dict[str, Any] | None = None
        try:
            map_payload = json.loads(evidence_map_path.read_text(encoding="utf-8"))
        except Exception:
            pass
        findings.append(self._finding("teacher_outputs_present", extension.is_file() and extension.stat().st_size > 0 and isinstance(map_payload, dict)))
        map_errors = self._validate_evidence_map(map_payload, candidate, files, bible, session)
        findings.append(self._finding("evidence_map_closed", not map_errors, {"errors": map_errors}))

        candidate_text = "\n".join(text_parts).casefold()
        forbidden = list(self.answer_labels)
        forbidden.extend(text.casefold() for text in bible.correct_treatments)
        forbidden.extend(item.statement.casefold() for item in bible.facts if item.kind == "treatment")
        leaks = sorted({term for term in forbidden if term and term in candidate_text})
        findings.append(self._finding("answer_leakage_absent", not leaks, {"matches": leaks[:10]}))
        tree = [{"path": str(path.relative_to(workspace)).replace("\\", "/"), "sha256": self.sha256_path(path)} for path in sorted(workspace.rglob("*")) if path.is_file() and ".codex" not in path.parts]
        decision: AdmissionDecision = "pass" if all(item.passed for item in findings) else "blocked"
        return ScenarioEvidenceAdmissionReportV1(
            session_id=session.session_id, output_tree_sha256=self.sha256_json(tree), decision=decision, findings=findings,
        )

    def blind_payload(
        self,
        *,
        seed: WorkSeedV1,
        rules: ProfessionalRuleSetV1,
        workspaces: dict[Condition, Path],
        scenario_id: str,
    ) -> tuple[tuple[Condition, Condition], dict[str, Any]]:
        if set(workspaces) != {"without_skill", "with_skill"}:
            raise ValueError("blind_review_requires_complete_pair")
        order: tuple[Condition, Condition] = ("without_skill", "with_skill") if int(hashlib.sha256(scenario_id.encode()).hexdigest()[-1], 16) % 2 == 0 else ("with_skill", "without_skill")
        bundles: list[dict[str, Any]] = []
        for index, condition in enumerate(order, start=1):
            candidate = workspaces[condition] / "candidate"
            contents = []
            for path in sorted(item for item in candidate.rglob("*") if item.is_file()):
                contents.append({"path": str(path.relative_to(candidate)).replace("\\", "/"), "text": self._extract_candidate_text(path, require_libreoffice=False)[:12000]})
            bundles.append({"bundle_slot": f"bundle_{index}", "files": contents})
        public_rules = [{"title": rule.title, "applicability_conditions": rule.applicability_conditions, "evidence_requirements": rule.evidence_requirements, "exceptions": rule.exceptions} for rule in rules.rules]
        return order, {"work_seed": {"domain": seed.domain, "role": seed.role, "trigger_event": seed.trigger_event, "business_goal": seed.business_goal, "expected_deliverable": seed.expected_deliverable, "audience": seed.audience}, "professional_rules": public_rules, "bundles": bundles}

    @staticmethod
    def _finding(code: str, passed: bool, details: dict[str, Any] | None = None) -> EvidenceAdmissionFindingV1:
        return EvidenceAdmissionFindingV1(code=code, passed=passed, details=details or {})

    @staticmethod
    def _write(path: Path, payload: dict[str, Any]) -> None:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    @staticmethod
    def _task_prompt(condition: Condition) -> str:
        skill_instruction = "Read teacher/professional_skill/SKILL.md and its source map before planning." if condition == "with_skill" else "Do not load or infer any professional Skill package."
        return f"""You are a factory-side evidence author. Read teacher/scenario_bible.json, teacher/professional_rules.json, and teacher/condition.json. {skill_instruction}

Create a plausible candidate-visible evidence bundle only in candidate/. Choose the natural business file types yourself, but include at least one table-like file (XLSX or CSV) and one narrative file (TXT, DOCX, or PDF). Candidate files must express underlying facts, not answer labels. Never put any of these in candidate/: the Scenario Bible, correct treatments, professional rules, Skill text, source citations, Questionable, Exception, Requires Follow-Up, or a direct final disposition.

Create teacher/scenario_extension.md with only new facts that are compatible with the parent Bible. Create teacher/evidence_map.json exactly as JSON object {{"artifacts":[{{"path":"candidate/relative-file","fact_ids":["parent-fact-id"],"professional_judgments":["short judgment point"],"skill_source_ids":["source-id only when a loaded Skill supports this artifact"]}}]}}. Every candidate file needs one artifacts entry. Do not create a task prompt, deliverable contract, rubric, teacher truth, solver answer, or package.
"""

    @classmethod
    def _validate_evidence_map(cls, payload: dict[str, Any] | None, candidate_root: Path, files: list[Path], bible: ScenarioBibleV1, session: ScenarioEvidenceSessionV1) -> list[str]:
        if not isinstance(payload, dict) or set(payload) != {"artifacts"} or not isinstance(payload.get("artifacts"), list):
            return ["evidence_map_shape_invalid"]
        seen: set[str] = set()
        errors: list[str] = []
        fact_ids = {item.fact_id for item in bible.facts}
        for item in payload["artifacts"]:
            if not isinstance(item, dict) or set(item) != {"path", "fact_ids", "professional_judgments", "skill_source_ids"}:
                errors.append("evidence_map_entry_shape_invalid")
                continue
            path = item.get("path")
            if not isinstance(path, str) or not path.startswith("candidate/") or ".." in Path(path).parts:
                errors.append("evidence_map_path_invalid")
            elif path in seen:
                errors.append("evidence_map_duplicate_path")
            else:
                seen.add(path)
            if not isinstance(item.get("fact_ids"), list) or not item["fact_ids"] or not set(item["fact_ids"]) <= fact_ids:
                errors.append("evidence_map_fact_reference_invalid")
            if not isinstance(item.get("professional_judgments"), list) or not item["professional_judgments"]:
                errors.append("evidence_map_judgment_missing")
            if not isinstance(item.get("skill_source_ids"), list):
                errors.append("evidence_map_skill_source_invalid")
            elif session.condition == "without_skill" and item["skill_source_ids"]:
                errors.append("without_skill_evidence_map_has_skill_source")
            elif not all(isinstance(value, str) and value for value in item["skill_source_ids"]):
                errors.append("evidence_map_skill_source_invalid")
            elif not set(item["skill_source_ids"]) <= set(session.skill_source_ids):
                errors.append("evidence_map_skill_source_not_curated")
        actual = {f"candidate/{path.relative_to(candidate_root).as_posix()}" for path in files}
        if seen != actual:
            errors.append("evidence_map_candidate_coverage_incomplete")
        if session.condition == "with_skill":
            source_bound = sum(bool(item.get("skill_source_ids")) for item in payload["artifacts"] if isinstance(item, dict))
            if source_bound < 2:
                errors.append("with_skill_requires_two_source_bound_artifacts")
        return sorted(set(errors))

    @staticmethod
    def _extract_candidate_text(path: Path, *, require_libreoffice: bool) -> str:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                if not book.sheetnames or not any(sheet.max_row and sheet.max_column for sheet in book.worksheets):
                    raise ValueError("empty_workbook")
                return "\n".join(str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
            finally:
                book.close()
        if suffix in {".csv", ".txt"}:
            return path.read_text(encoding="utf-8")
        if suffix == ".docx":
            with zipfile.ZipFile(path) as archive:
                document = archive.read("word/document.xml").decode("utf-8")
            if require_libreoffice:
                if shutil.which("libreoffice") is None:
                    raise RuntimeError("libreoffice_unavailable")
                with subprocess.Popen(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(path.parent), str(path)], stdout=subprocess.PIPE, stderr=subprocess.PIPE) as proc:
                    proc.communicate(timeout=90)
                    if proc.returncode != 0:
                        raise RuntimeError("libreoffice_docx_open_failed")
            return re.sub(r"<[^>]+>", " ", document)
        if suffix == ".pdf":
            if not path.read_bytes().startswith(b"%PDF"):
                raise ValueError("invalid_pdf_header")
            if shutil.which("pdftotext") is None:
                return ""
            result = subprocess.run(["pdftotext", str(path), "-"], text=True, encoding="utf-8", errors="replace", capture_output=True, check=False)
            if result.returncode != 0:
                raise RuntimeError("pdf_text_extraction_failed")
            return result.stdout
        raise ValueError("unsupported_candidate_file")


class OfficialDeepSeekConditionBlindReviewer:
    """One substantive pair review; retry strictly for transport or JSON format."""

    def __init__(self, request_executor: Callable[[dict[str, Any], str, int], tuple[int, dict[str, Any]]] | None = None) -> None:
        self.request_executor = request_executor or self._official_request

    def review(self, *, scenario_id: str, bundle_order: tuple[Condition, Condition], payload: dict[str, Any], api_key: str, output_root: Path, timeout_seconds: int = 300) -> ConditionBlindPairReviewV1:
        output_root.mkdir(parents=True, exist_ok=False)
        input_sha = ScenarioEvidenceExperiment.sha256_json(payload)
        self._write(output_root / "input_manifest.json", {"input_sha256": input_sha, "provider": "deepseek", "model": "deepseek-v4-pro", "provider_retry_count": 0})
        first_failure: str | None = None
        usage: dict[str, int | None] = {}
        call_count = 0
        for attempt in range(2):
            try:
                call_count += 1
                status, response = self.request_executor(self._request(payload), api_key, timeout_seconds)
                content, usage = self._content(response)
                failure = None
                if not 200 <= status < 300:
                    failure = f"provider_http_{status}"
                elif not content or (response.get("choices") or [{}])[0].get("finish_reason") == "length":
                    failure = "empty_or_truncated_provider_content"
                else:
                    response_sha = hashlib.sha256(content.encode("utf-8")).hexdigest()
                    self._write(output_root / f"provider_response_attempt_{attempt + 1}.json", {"response_sha256": response_sha, "content": content})
                    try:
                        value = json.loads(content)
                        findings = [ConditionBlindFindingV1.model_validate(item) for item in value["findings"]]
                        summary = str(value["summary"])
                    except Exception:
                        failure = "provider_json_or_schema_invalid"
                    else:
                        report = ConditionBlindPairReviewV1(scenario_id=scenario_id, bundle_order=bundle_order, provider_call_count=attempt + 1, provider_retry_count=attempt, response_sha256=response_sha, decision="pass", findings=findings, summary=summary, usage=usage)
                        self._write(output_root / "condition_blind_pair_review.json", report.model_dump(mode="json"))
                        return report
            except Exception as exc:
                failure = f"provider_transport_failure:{type(exc).__name__}"
            if first_failure is None:
                first_failure = failure
            retryable = failure in {"empty_or_truncated_provider_content", "provider_json_or_schema_invalid", "provider_http_408", "provider_http_429"} or str(failure).startswith("provider_http_5") or str(failure).startswith("provider_transport_failure")
            if attempt == 0 and retryable:
                continue
            break
        report = ConditionBlindPairReviewV1(scenario_id=scenario_id, bundle_order=bundle_order, provider_call_count=call_count, provider_retry_count=max(0, call_count - 1), decision="incomplete", first_failure=first_failure, usage=usage)
        self._write(output_root / "condition_blind_pair_review.json", report.model_dump(mode="json"))
        return report

    @staticmethod
    def _request(payload: dict[str, Any]) -> dict[str, Any]:
        system = """You are a condition-blind reviewer of two candidate-visible professional evidence bundles. Return JSON only: {\"findings\":[{\"favored_bundle\":\"bundle_1|bundle_2|neither\",\"dimension\":\"professional_realism|evidence_naturalness|judgment_depth|leakage_risk\",\"file_paths\":[\"relative file path\"],\"reason\":\"specific comparison\"}],\"summary\":\"brief comparison\"}. Do not infer, mention, or score a hidden generation method. Do not invent missing facts. Assess whether the files feel like natural work products, create connected professional judgment, and reveal answers through labels. Give evidence-based findings only."""
        return {"model": "deepseek-v4-pro", "messages": [{"role": "system", "content": system}, {"role": "user", "content": json.dumps(payload, ensure_ascii=False)}], "response_format": {"type": "json_object"}, "thinking": {"type": "disabled"}, "max_tokens": 4000, "stream": False}

    @staticmethod
    def _official_request(body: dict[str, Any], api_key: str, timeout_seconds: int) -> tuple[int, dict[str, Any]]:
        request = urllib.request.Request("https://api.deepseek.com/chat/completions", data=json.dumps(body, ensure_ascii=False).encode("utf-8"), headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"}, method="POST")
        try:
            with urllib.request.urlopen(request, timeout=timeout_seconds) as response:
                return int(response.status), json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            return int(exc.code), {}

    @staticmethod
    def _content(response: dict[str, Any]) -> tuple[str | None, dict[str, int | None]]:
        choice = (response.get("choices") or [{}])[0] if isinstance(response, dict) else {}
        message = choice.get("message") if isinstance(choice, dict) else {}
        usage = response.get("usage") if isinstance(response.get("usage"), dict) else {}
        return message.get("content") if isinstance(message, dict) else None, {"prompt_tokens": usage.get("prompt_tokens"), "completion_tokens": usage.get("completion_tokens"), "total_tokens": usage.get("total_tokens")}

    @staticmethod
    def _write(path: Path, payload: dict[str, Any]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def aggregate_experiment(*, plan: ScenarioEvidenceExperimentPlanV1, admissions: list[ScenarioEvidenceAdmissionReportV1], reviews: list[ConditionBlindPairReviewV1]) -> ScenarioEvidenceExperimentResultV1:
    reasons: list[str] = []
    if len(admissions) != 4 or len(reviews) != 2 or any(item.decision == "incomplete" for item in admissions + reviews):
        return ScenarioEvidenceExperimentResultV1(plan_sha256=plan.canonical_sha256(), admission_reports=admissions, blind_reviews=reviews, decision="incomplete", reasons=["execution_or_review_incomplete"])
    if any(item.decision != "pass" for item in admissions):
        return ScenarioEvidenceExperimentResultV1(plan_sha256=plan.canonical_sha256(), admission_reports=admissions, blind_reviews=reviews, decision="skill_revision_required", reasons=["candidate_admission_failed"])
    sessions = {(item.scenario_id, item.condition): item for item in plan.sessions}
    for review in reviews:
        with_skill = sessions[(review.scenario_id, "with_skill")]
        slot = "bundle_1" if review.bundle_order[0] == "with_skill" else "bundle_2"
        supported = [item for item in review.findings if item.favored_bundle == slot and item.dimension != "leakage_risk"]
        if len(supported) < 2:
            reasons.append(f"{review.scenario_id}:insufficient_skill_supported_improvements")
    return ScenarioEvidenceExperimentResultV1(plan_sha256=plan.canonical_sha256(), admission_reports=admissions, blind_reviews=reviews, decision="skill_effect_supported" if not reasons else "skill_revision_required", reasons=reasons)
