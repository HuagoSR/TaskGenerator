"""Local file tools and a scope-bound Unix socket broker; no provider credentials."""
from __future__ import annotations

import json
import contextlib
import io
import itertools
import os
import re
from pathlib import Path, PurePosixPath
import shutil
import socket
import socketserver
import subprocess
import sys
import tempfile
import time
import uuid

from task_generator.production import agent_factory as f


def issue_for(error):
    message = str(error)
    head, separator, tail = message.partition(':')
    path = head if separator and ('.' in head or '[' in head) else None
    expected = None
    actual = None
    if 'expected ' in message:
        expected = message.split('expected ', 1)[1].split(', got ', 1)[0]
    if ', got ' in message:
        actual = message.split(', got ', 1)[1][:500]
    code = (head if not path else head.rsplit('.', 1)[-1]).replace('[', '_').replace(']', '')
    artifact = next((name for prefix, name in (
        ('calculation_evidence.json', 'calculation_evidence'), ('new_rubric.json', 'rubric'),
        ('basis_draft.json', 'basis'), ('comparison.json', 'comparison'),
        ('supervision.json', 'supervision'), ('consultation.json', 'consultation'),
        ('review.json', 'review'), ('diagnostic.json', 'development_diagnostic'))
        if message.startswith(prefix)), None)
    repair = (f'run factory-tools schema with artifact {artifact}, revise the located field, then rerun check'
              if artifact else 'query factory-tools schema for the affected artifact, revise the current draft, then rerun check')
    if any(token in message for token in ('submission_prerequisites_missing', 'next_production_launch_budget_unavailable')):
        repair = 'Current role cannot supply these prerequisites. Read status and report the blocker; do not repeat draft edits or fabricate evidence.'
    subject = re.search(r'\.(criteria|executions|calculations)\[([^]]+)\]', message)
    context = {}
    if subject:
        name = {'criteria': 'criterion', 'executions': 'execution', 'calculations': 'calculation'}[subject[1]]
        context[name + ('_index' if subject[2].isdigit() else '_id')] = subject[2]
    for known in ('invalid_supervision', 'rubric_basis_not_candidate_visible', 'reference_not_in_stage_allowlist'):
        if known in message:
            code = known
    return {'code': code, 'path': path, 'artifact': artifact, 'message': message, 'expected': expected,
            'actual': actual, 'repair': repair, **context}


def issues_for(error):
    messages = getattr(error, 'messages', None)
    return [issue_for(ValueError(message)) for message in messages] if messages else [issue_for(error)]


class Broker:
    def __init__(self, config):
        self.config = config
        if config.get('protocol') == 'task_factory_harness_v1':
            from task_generator.production import task_factory_harness
            self.factory = task_factory_harness
        else:
            self.factory = f
        self.root = Path(config['root'])
        self.draft = self.root / 'draft'
        self.store = self.root / 'snapshots'
        self.store.mkdir(exist_ok=True)
        self.events = []
        self.process = config.get('process')
        self.pending = None
        self.dispositions = list(config.get('dispositions', []))

    def current_snapshot(self, reason):
        hashes = self.factory.files(self.draft)
        for event in reversed(self.events):
            if event['action'] == 'snapshot' and event['result']['hashes'] == hashes:
                return event['result'], 'reused'
        result = self.factory.snapshot(self.draft, self.store, self.config['role'],
                                       self.config['parents'], reason, self.process)
        self.events.append({'ordinal': len(self.events) + 1, 'action': 'snapshot',
                            'request': {'reason': reason, 'automatic': True}, 'result': result})
        return result, 'created'

    def checked(self, hashes, mode):
        return any(event['action'] == 'log-check'
                   and event['result']['recorded'].get('tool') == 'check'
                   and event['result']['recorded'].get('hashes') == hashes
                   and event['result']['recorded'].get('mode', 'ready') in (
                       ('draft', 'ready') if mode == 'draft' else ('ready',))
                   for event in self.events)

    def validate_disposition(self, request, consultation, identity):
        if request.get('finding_id') not in consultation['finding_ids']:
            raise ValueError('unknown_consultation_finding')
        if request.get('decision') not in ('accept', 'partial', 'reject') or not request.get('reason', '').strip():
            raise ValueError('disposition_decision_and_reason_required')
        evidence = request.get('evidence', [])
        if not evidence:
            raise ValueError('disposition_evidence_required')
        for ref in evidence:
            root = self.draft if ref.get('area') == 'draft' else self.root / 'workspace/inputs'
            if ref.get('area') not in ('draft', 'inputs') or not ref.get('locator'):
                raise ValueError('disposition_area_and_locator_required')
            if not self.factory.safe_path(root, ref.get('path', '')).is_file():
                raise ValueError('disposition_evidence_missing:' + str(ref.get('path')))
        return {**{key: request[key] for key in ('finding_id', 'decision', 'reason', 'evidence')},
                'request_id': consultation['request_id'], 'snapshot': identity}

    def handle(self, request):
        from task_generator.production.quality_diagnostics import micro_enabled
        if micro_enabled(self.config) and request.get('action') in ('consult', 'replay', 'submit', 'request-development-trial', 'handoff'):
            raise ValueError('diagnostic_mode_requires_finish')
        action = request['action']
        role = self.config['role']
        if self.config.get('purpose') == 'tool_microtest' and (
                action in ('submit', 'request-development-trial')
                or action == 'handoff' and request.get('target') != 'stop'
                or action == 'replay' and request.get('next_action') != 'stop'):
            raise ValueError('action_outside_microtest_scope')
        if self.config.get('allowed_actions') is not None and action not in self.config['allowed_actions']:
            raise ValueError('action_outside_scope')
        if action in ('consult', 'replay', 'submit', 'request-development-trial', 'handoff') and not (action == 'handoff' and request.get('target') == 'stop'):
            deadline = self.config.get('launch_deadline_epoch')
            if deadline is not None and time.time() >= deadline:
                raise ValueError('launch_time_budget_exhausted')
        if action == 'status':
            workflow = json.loads(json.dumps(self.config.get('workflow_status', {})))
            hashes = self.factory.files(self.draft)
            snapshots = [event['result'] for event in self.events if event['action'] == 'snapshot'
                         and event['result']['hashes'] == hashes]
            starting = self.config.get('starting_snapshot')
            if not snapshots and starting and starting.get('hashes') == hashes and starting.get('parents') == self.config.get('parents'):
                snapshots = [starting]
            checks = [event['result']['recorded'] for event in self.events if event['action'] == 'log-check'
                      and event['result']['recorded'].get('tool') == 'check'
                      and event['result']['recorded'].get('hashes') == hashes]
            unresolved = []
            current_snapshot = snapshots[-1]['id'] if snapshots else None
            for consultation in self.config.get('consultations', []):
                for finding in consultation['finding_ids']:
                    if not any(row['request_id'] == consultation['request_id'] and row['finding_id'] == finding
                               and row['snapshot'] == current_snapshot for row in self.dispositions):
                        unresolved.append({'request_id': consultation['request_id'], 'finding_id': finding})
            workflow.update(pending=bool(self.pending), current_draft_hashes=hashes,
                            latest_matching_snapshot=current_snapshot,
                            current_check_modes=sorted({row.get('mode', 'ready') for row in checks}),
                            unresolved_findings=unresolved)
            if self.config.get('quality_diagnostics_version') is not None:
                required = {
                    'world': 'hidden/material_relations.json', 'compile': 'rubric_diagnostic.json',
                    'review': 'rubric_diagnostic.json and record_relations.json',
                }.get(role)
                present = [] if required is None else [name for name in required.split(' and ')
                                                       if (self.draft / name).is_file()]
                diagnostics = {
                    'version': self.config['quality_diagnostics_version'],
                    'required': required, 'present': present,
                    'current_draft_hashes': hashes,
                    'note': 'Presence and file identity are structural only; semantic status comes from the versioned diagnostic artifact.',
                }
                try:
                    from task_generator.production.quality_diagnostics import validate_record_relations, validate_rubric_diagnostic
                    inputs = self.root / 'workspace/inputs'
                    if role == 'world' and (self.draft / 'hidden/material_relations.json').is_file():
                        assessment = validate_record_relations(
                            f.read(self.draft / 'hidden/material_relations.json'), self.draft / 'candidate', candidate_only=False)
                        diagnostics['record_relations'] = assessment
                    elif role == 'compile' and (self.draft / 'rubric_diagnostic.json').is_file():
                        assessment = validate_rubric_diagnostic(
                            f.read(self.draft / 'rubric_diagnostic.json'), f.read(self.draft / 'new_rubric.json'),
                            f.read(self.draft / 'basis_draft.json'), f.read(self.draft / 'calculation_evidence.json'), inputs)
                        diagnostics['rubric_decision'] = assessment['decision']
                    elif role == 'review' and all((self.draft / name).is_file() for name in ('rubric_diagnostic.json', 'record_relations.json')):
                        assessment = validate_rubric_diagnostic(
                            f.read(self.draft / 'rubric_diagnostic.json'), f.read(inputs / 'new_rubric.json'),
                            f.read(inputs / 'basis_draft.json'), f.read(inputs / 'calculation_evidence.json'), inputs)
                        diagnostics['rubric_decision'] = assessment['decision']
                        diagnostics['record_relations'] = validate_record_relations(
                            f.read(self.draft / 'record_relations.json'), inputs)
                except (ValueError, FileNotFoundError, KeyError, TypeError, json.JSONDecodeError) as error:
                    diagnostics['validation_error'] = str(error)
                workflow['quality_diagnostics'] = diagnostics
            workflow['pending_action'] = self.pending.get('action') if self.pending else None
            workflow['passing_replay_current'] = current_snapshot in self.config.get('passed_replay_snapshots', [])
            if not workflow['passing_replay_current']:
                workflow['legal_next_actions'] = [x for x in workflow.get('legal_next_actions', []) if x != 'submit']
            if (role == 'compile' and self.config.get('quality_diagnostics_version') is not None
                    and workflow.get('quality_diagnostics', {}).get('rubric_decision') != 'pass'):
                workflow['legal_next_actions'] = [x for x in workflow.get('legal_next_actions', []) if x != 'submit']
            if self.config.get('purpose') == 'tool_microtest':
                workflow['legal_next_actions'] = [x for x in workflow.get('legal_next_actions', [])
                    if x in ('status', 'inspect', 'render', 'check', 'snapshot', 'diff', 'handoff:stop', 'consult', 'finish')]
                if role == 'compile' and self.config.get('consultations'):
                    workflow['legal_next_actions'].append('replay:stop')
            if self.pending:
                workflow['legal_next_actions'] = ['status']
            deadline = self.config.get('deadline_epoch')
            if deadline is not None:
                workflow.setdefault('budget', {})['case_seconds_remaining'] = max(
                    0, int(deadline - time.time()))
            if self.config.get('launch_deadline_epoch') is not None:
                workflow.setdefault('budget', {})['launch_seconds_remaining'] = max(
                    0, int(self.config['launch_deadline_epoch'] - time.time()))
            result = workflow
            if self.config.get('diagnostic_micro_version') is not None and role == 'compile':
                from task_generator.production.quality_diagnostics import micro_enabled
                micro_enabled(self.config)
                result['legal_next_actions'] = ['status'] if self.pending else ['schema', 'inspect', 'diagnose', 'check', 'finish']
                result['missing_prerequisites'] = []
                result['submission_base_missing'] = ['production_submit_not_available_in_diagnostic_mode']
                try:
                    from task_generator.production.quality_diagnostics import micro_result
                    result['diagnostic_result'] = micro_result(self.draft, self.root / 'workspace/inputs', self.config)
                except (ValueError, OSError, KeyError, TypeError) as error:
                    result['missing_prerequisites'] = [str(error)]
        elif action == 'log-call':
            result = {'recorded': request['record']}
        elif action == 'log-check':
            result = {'recorded': request['result']}
        elif action == 'finish' and (role in ('devsolve', 'consult', 'review', 'solve') or
                (role == 'compile' and self.config.get('diagnostic_micro_version') == 1)):
            if self.pending:
                raise ValueError('request_pending_end_turn')
            hashes = self.factory.files(self.draft)
            if not any(e['action'] == 'log-check' and e['result']['recorded'].get('tool') == 'check'
                       and e['result']['recorded'].get('hashes') == hashes for e in self.events):
                raise ValueError('current_version_check_required: run factory-tools check')
            if role == 'compile':
                from task_generator.production.quality_diagnostics import micro_result
                outcome = micro_result(self.draft, self.root / 'workspace/inputs', self.config)
                entry, _ = self.current_snapshot('Frozen synthetic diagnostic result')
                self.pending = {'action': 'finish', 'hashes': hashes, 'snapshot_entry': entry, 'diagnostic_outcome': outcome}
            else:
                self.pending = {'action': 'finish', 'hashes': hashes}
            result = {'status': 'checked_end_turn_now'}
        elif role not in self.factory.AUTHOR_ROLES:
            raise ValueError('read_only_role')
        elif self.pending:
            raise ValueError('request_pending_end_turn')
        elif action in ('record-disposition', 'record-dispositions'):
            if action == 'record-dispositions':
                consultation = next((c for c in self.config.get('consultations', [])
                                     if c['request_id'] == request.get('request_id')), None)
                if not consultation:
                    raise ValueError('unknown_consultation_finding')
                items = request.get('items')
                if not isinstance(items, list) or not items:
                    raise ValueError('disposition_items_required')
                ids = [item.get('finding_id') for item in items]
                if len(ids) != len(set(ids)) or set(ids) != set(consultation['finding_ids']):
                    raise ValueError('disposition_items_must_cover_every_finding_once')
                prepared = [self.validate_disposition(item, consultation, '__pending__') for item in items]
                snapshot, reuse = self.current_snapshot(request.get('snapshot_reason', 'Consultation dispositions'))
                prepared = [{**item, 'snapshot': snapshot['id']} for item in prepared]
                self.dispositions.extend(prepared)
                result = {'request_id': consultation['request_id'], 'snapshot': snapshot['id'],
                          'snapshot_status': reuse, 'recorded': len(prepared), 'unresolved': []}
            else:
                consultation = next((c for c in self.config.get('consultations', [])
                                     if c['request_id'] == request.get('request_id')), None)
                if not consultation:
                    raise ValueError('unknown_consultation_finding')
                identity = request.get('snapshot')
                matching = [e['result'] for e in self.events if e['action'] == 'snapshot' and e['result']['id'] == identity]
                if not matching or self.factory.files(self.draft) != matching[-1]['hashes']:
                    raise ValueError('snapshot_current_draft_required')
                result = self.validate_disposition(request, consultation, identity)
                self.dispositions.append(result)
        elif action == 'handoff' and request.get('target') == 'stop':
            if not request.get('reason', '').strip():
                raise ValueError('reason_required')
            self.pending = {**request, 'request_id': f.fingerprint(request)}
            result = {'request_id': self.pending['request_id'], 'status': 'queued_end_turn_now'}
        elif action == 'save-process':
            if role != 'world':
                raise ValueError('world_only')
            path = self.draft / 'hidden/process.md'
            value = self.factory.digest(path)
            if not self.process and any((self.draft / 'candidate').rglob('*')):
                raise ValueError('initial_candidate_precedes_process')
            target = self.root / 'processes' / (value + '.md')
            target.parent.mkdir(exist_ok=True)
            if target.exists():
                raise ValueError('process_already_saved')
            shutil.copyfile(path, target)
            self.process = {'hash': value, 'event': len(self.events) + 1}
            result = self.process
        elif action == 'snapshot':
            result = self.factory.snapshot(self.draft, self.store, role, self.config['parents'],
                                request.get('reason', ''), self.process)
        elif action == 'diff':
            previous = self.factory.safe_path(self.store, request['snapshot'])
            old, new = self.factory.files(previous), self.factory.files(self.draft)
            if not previous.is_dir():
                raise ValueError('unknown_snapshot')
            result = {'added': sorted(new.keys() - old.keys()), 'removed': sorted(old.keys() - new.keys()),
                      'changed': sorted(k for k in old.keys() & new.keys() if old[k] != new[k])}
        elif action in ('consult', 'handoff', 'submit', 'replay', 'request-development-trial'):
            if request.get('snapshot') is None:
                automatic, snapshot_status = self.current_snapshot(request.get('reason', 'Stage transition'))
                request = {**request, 'snapshot': automatic['id'], 'snapshot_automatic': snapshot_status}
            identity = request['snapshot']
            matching = [e['result'] for e in self.events if e['action'] == 'snapshot' and e['result']['id'] == identity]
            if not matching or self.factory.files(self.draft) != matching[-1]['hashes']:
                raise ValueError('snapshot_current_draft_required')
            check_mode = 'draft' if action == 'consult' else 'ready'
            if not self.checked(matching[-1]['hashes'], check_mode):
                raise ValueError('current_version_check_required: run factory-tools check')
            if action != 'consult':
                self.require_dispositions(identity)
            if not request.get('reason', '').strip():
                raise ValueError('reason_required')
            if action == 'submit' and role != 'compile':
                raise ValueError('compiler_submits_only')
            if (self.config.get('protocol') == 'task_factory_harness_v1'
                    and (action == 'submit' or action == 'replay' and request.get('next_action') == 'submit')):
                missing = self.config.get('workflow_status', {}).get('submission_base_missing', [])
                if missing:
                    raise ValueError('submission_prerequisites_missing:' + ','.join(missing))
            if action in ('replay', 'request-development-trial'):
                if self.config.get('protocol') != 'task_factory_harness_v1' or role != 'compile':
                    raise ValueError('harness_compiler_only')
                passed = set(self.config.get('passed_replay_snapshots', []))
                if action == 'request-development-trial' and self.config.get('development_trial_current'):
                    raise ValueError('development_trial_already_completed: proceed with comparison, final replay, and submit unless an upstream revision invalidates the trial')
                if action == 'request-development-trial' and identity not in passed:
                    raise ValueError('current_calculation_replay_required')
                if action == 'replay':
                    next_action = request.get('next_action', 'return')
                    allowed_next = ('return', 'development_trial', 'submit') + (('stop',) if self.config.get('purpose') == 'tool_microtest' else ())
                    if next_action not in allowed_next:
                        raise ValueError('invalid_replay_next_action')
                    if next_action == 'development_trial' and self.config.get('development_trial_current'):
                        raise ValueError('development_trial_already_completed')
                    request = {**request, 'next_action': next_action}
            if (self.config.get('protocol') == 'task_factory_harness_v1' and action == 'submit'
                    and identity not in set(self.config.get('passed_replay_snapshots', []))):
                raise ValueError('current_calculation_replay_required')
            if self.config.get('protocol') == 'task_factory_harness_v1':
                self.factory.validate_next_action(
                    self.config['workflow_status'], role, action, request)
            if action == 'handoff':
                quality = bool(self.config.get('candidate_edit_version'))
                allowed = {'world': ('mine', 'stop'),
                           'mine': ('world', 'edit', 'stop') if quality else ('world', 'compile', 'stop'),
                           'edit': ('world', 'mine', 'compile', 'stop'),
                           'compile': ('world', 'mine', 'edit', 'stop') if quality else ('world', 'mine', 'stop')}
                if request.get('target') not in allowed[role]:
                    raise ValueError('invalid_handoff')
            self.pending = {**request, 'snapshot_entry': matching[-1], 'request_id': f.fingerprint(request)}
            result = {'request_id': self.pending['request_id'], 'status': 'queued_end_turn_now'}
        else:
            raise ValueError('unknown_broker_action')
        self.events.append({'ordinal': len(self.events) + 1, 'action': action, 'request': request, 'result': result})
        self.factory.write(self.root / 'broker.json', {'events': self.events, 'pending': self.pending, 'process': self.process,
                                          'dispositions': self.dispositions})
        return result

    def require_dispositions(self, identity):
        for c in self.config.get('consultations', []):
            for finding in c['finding_ids']:
                if not any(d['request_id'] == c['request_id'] and d['finding_id'] == finding
                           and d['snapshot'] == identity for d in self.dispositions):
                    raise ValueError('disposition_required:' + c['request_id'] + ':' + finding)


def serve(config_path):
    broker = Broker(f.read(config_path))

    class Handler(socketserver.StreamRequestHandler):
        def handle(self):
            try:
                data = self.rfile.readline(1024 * 1024)
                result = {'ok': True, 'result': broker.handle(json.loads(data))}
            except Exception as error:
                result = {'ok': False, 'error': type(error).__name__ + ':' + str(error),
                          'issues': issues_for(error)}
            self.wfile.write((json.dumps(result) + '\n').encode())

    address = broker.root / 'socket/tool.sock'
    address.parent.mkdir(exist_ok=True)
    # Bind relative to the scope directory: Linux sockaddr_un limits path bytes.
    # Clients access the same inode through their short, read-only socket mount.
    os.chdir(broker.root)
    with socketserver.UnixStreamServer('socket/tool.sock', Handler) as server:
        os.chmod(address, 0o600)
        server.serve_forever()


def broker_call(action, args):
    with socket.socket(socket.AF_UNIX, socket.SOCK_STREAM) as client:
        client.settimeout(60)
        client.connect('/run/factory/tool.sock')
        client.sendall((json.dumps({'action': action, **args}) + '\n').encode())
        result = json.loads(client.makefile('rb').readline(1024 * 1024))
    if not result['ok']:
        raise ValueError(result['error'] + '; details=' + json.dumps(result.get('issues', []), ensure_ascii=False))
    return result['result']


def teacher_citation_paths(inputs, protocol):
    if protocol != 'task_factory_harness_v1':
        return ()
    return tuple(path for path in f.files(inputs) if path in (
        'basis_draft.json', 'calculation_evidence.json', 'replay_results.json')
        or path.startswith('calculation_scripts/'))


def resolve_material_check_path(raw, candidate_paths, index):
    """Resolve one material_checks path to exactly one candidate-relative file.

    Accepts a candidate-relative path ('exports/sales.csv') or the same file
    referenced with an explicit 'candidate/' prefix ('candidate/exports/sales.csv')
    when exactly one interpretation exists in candidate_paths; ambiguity, escapes
    and missing files are rejected with the entry index and the received path.
    """
    if not isinstance(raw, str) or not raw.strip():
        raise ValueError(f'hidden/material_checks.json.checks[{index}]: candidate path required')
    pure = PurePosixPath(raw)
    if pure.is_absolute() or '..' in pure.parts:
        raise ValueError(f'hidden/material_checks.json.checks[{index}]: path {raw!r} escapes candidate/')
    candidate_form = raw[len('candidate/'):] if raw.startswith('candidate/') else raw
    matches = sorted({form for form in (candidate_form, 'candidate/' + candidate_form)
                      if form in candidate_paths})
    if len(matches) > 1:
        raise ValueError(f'hidden/material_checks.json.checks[{index}]: path {raw!r} is ambiguous '
                         'between candidate-relative and candidate/-prefixed references')
    if not matches:
        raise ValueError(f'hidden/material_checks.json.checks[{index}]: path {raw!r} does not resolve '
                         'to a candidate file; use the path relative to candidate/ or prefixed with candidate/')
    return matches[0]


def check(role, draft, inputs, protocol=None, mode='ready', calculation_version=None,
          obligation_trace_version=None, initial_basis_snapshot=None, atomic_rubric_version=None,
          quality_diagnostics_version=None):
    from task_generator.production import task_method_pilot as m
    from r10_process_first_tools import validate
    if mode not in ('draft', 'ready'):
        raise ValueError('check_mode_must_be_draft_or_ready')
    if quality_diagnostics_version is not None and str(quality_diagnostics_version) != '1':
        raise ValueError('unsupported_quality_diagnostics_version')
    if protocol == 'task_factory_harness_v1':
        f.files(draft)
        from docx import Document
        from openpyxl import load_workbook
        from pypdf import PdfReader
        allowed = {'.docx', '.xlsx', '.pdf', '.csv', '.txt', '.md', '.json', '.sha256', '.py', '.sql'}
        for path in Path(draft).rglob('*'):
            if not path.is_file():
                continue
            relative = path.relative_to(Path(draft)).as_posix()
            if path.suffix.lower() not in allowed:
                raise ValueError(f'unsupported_output_type:{relative}')
            if not path.stat().st_size:
                raise ValueError(f'empty_output:{relative}')
            if path.suffix.lower() == '.docx':
                Document(path)
            elif path.suffix.lower() == '.xlsx':
                load_workbook(path).close()
            elif path.suffix.lower() == '.pdf':
                if not PdfReader(path).pages:
                    raise ValueError('empty_pdf')
            elif path.suffix.lower() == '.sql':
                try:
                    path.read_bytes().decode('utf-8-sig')  # text record: encoding-valid, never executed
                except UnicodeDecodeError as error:
                    raise ValueError(f'invalid_text_encoding:{relative}') from error
            else:
                path.read_text(encoding='utf-8')
    else:
        with contextlib.redirect_stdout(io.StringIO()):
            validate(draft)
    if mode == 'draft':
        return {'status': 'completed', 'check_mode': 'draft', 'files': f.files(draft)}
    if role == 'world':
        records = f.read(draft / 'hidden/manifest.json')['records']
        actual = f.files(draft / 'candidate')
        if len(records) != len(actual) or {r['path'] for r in records} != set(actual):
            raise ValueError('manifest_coverage')
        for r in records:
            if not r.get('producer') or not r.get('business_purpose') or not isinstance(r.get('source_dependencies'), list):
                raise ValueError('manifest_provenance')
        if not (draft / 'hidden/world.md').is_file():
            raise ValueError('world_record_missing')
        if atomic_rubric_version:
            checks = f.read(draft / 'hidden/material_checks.json').get('checks')
            if not isinstance(checks, list) or not checks:
                raise ValueError('hidden/material_checks.json.checks: targeted material checks required')
            candidate = f.files(draft / 'candidate')
            for index, row in enumerate(checks):
                if not isinstance(row, dict):
                    raise ValueError(f'hidden/material_checks.json.checks[{index}]: dict entry required')
                resolve_material_check_path(row.get('path'), candidate, index)
                if any(not str(row.get(key, '')).strip() for key in (
                        'locator', 'assertion', 'method', 'result', 'limitation')):
                    raise ValueError(f'hidden/material_checks.json.checks[{index}]: complete check evidence required')
        if quality_diagnostics_version is not None:
            from task_generator.production.quality_diagnostics import validate_record_relations
            relations = f.read(draft / 'hidden/material_relations.json')
            relation_result = validate_record_relations(relations, draft / 'candidate', candidate_only=False)
            if any(row['status'] == 'issue' for row in relation_result['relations']):
                raise ValueError('hidden/material_relations.json: declared record relation contradiction requires revision or an upstream issue')
        return {'status': 'completed', 'files': actual}
    if role == 'mine':
        outcome = m.task_result(draft, inputs)
        if protocol == 'task_factory_harness_v1':
            intent = f.read(draft / 'design_intent.json')
            if not all(intent.get(key) for key in ('occupational_use', 'analysis_points', 'likely_difficulties')):
                raise ValueError('design_intent_fields_required')
            for ref in intent.get('evidence', []):
                m.reference(inputs, ref.get('path'))
                if not ref.get('locator'):
                    raise ValueError('design_intent_locator_required')
        return outcome
    if role == 'edit':
        if not atomic_rubric_version:
            raise ValueError('candidate_editor_requires_versioned_quality_scope')
        return m.editor_result(draft, inputs, quality_diagnostics_version=quality_diagnostics_version)
    if role == 'compile':
        if protocol != 'task_factory_harness_v1':
            return m.compilation_result(draft, inputs)
        from task_generator.production import task_factory_harness as harness
        errors = []
        outcome = None
        try:
            outcome = (m.atomic_compilation_result(draft, inputs) if atomic_rubric_version
                       else m.compilation_result(draft, inputs))
        except (ValueError, KeyError, TypeError, FileNotFoundError, json.JSONDecodeError) as error:
            errors.extend(getattr(error, 'messages', [str(error)]))
        if outcome and outcome.get('status') == 'upstream_issue':
            return outcome
        try:
            basis = f.read(draft / 'basis_draft.json')
            if not isinstance(basis.get('requirements'), list) or not basis['requirements']:
                raise ValueError('basis_draft.json.requirements: expected a nonempty list')
            if obligation_trace_version is not None:
                harness.validate_obligation_trace(
                    draft, inputs,
                    require_comparison=(inputs / 'development_trial').is_dir(),
                    initial_basis_snapshot=initial_basis_snapshot, version=obligation_trace_version)
            if quality_diagnostics_version is not None:
                from task_generator.production.quality_diagnostics import validate_rubric_diagnostic
                diagnostic_outcome = validate_rubric_diagnostic(
                    f.read(draft / 'rubric_diagnostic.json'), f.read(draft / 'new_rubric.json'),
                    f.read(draft / 'basis_draft.json'), f.read(draft / 'calculation_evidence.json'), inputs)
                if diagnostic_outcome['decision'] != 'pass':
                    raise ValueError('rubric_diagnostic.json: unresolved diagnostics block ready and submit')
        except (ValueError, KeyError, TypeError, FileNotFoundError, json.JSONDecodeError) as error:
            errors.extend(getattr(error, 'messages', [str(error)]))
        try:
            manifest = harness.validate_calculation_bundle(draft, inputs)
            if calculation_version is not None and str(manifest.get('version')) != str(calculation_version):
                raise ValueError(f'calculation_evidence.json.version: expected {calculation_version} for this scope, got {manifest.get("version")!r}')
        except (ValueError, KeyError, TypeError, FileNotFoundError, json.JSONDecodeError) as error:
            errors.extend(getattr(error, 'messages', [str(error)]))
        if (inputs / 'development_trial').is_dir():
            try:
                comparison = f.read(draft / 'comparison.json')
                if not all(key in comparison for key in ('trial_errors', 'compiler_omissions', 'alternatives', 'upstream_defects')):
                    raise ValueError('comparison.json: fields trial_errors, compiler_omissions, alternatives and upstream_defects are required after development trial')
            except (ValueError, KeyError, TypeError, FileNotFoundError, json.JSONDecodeError) as error:
                errors.extend(getattr(error, 'messages', [str(error)]))
        if errors:
            raise harness.ContractIssues(dict.fromkeys(errors))
        return outcome
    if role == 'review':
        review = f.read(draft / 'review.json')
        extra_allowed = teacher_citation_paths(inputs, protocol)
        for i, dimension in enumerate(review.get('checks', [])):
            for j, finding in enumerate(dimension.get('findings', [])):
                try:
                    m.reference(inputs, finding.get('path'), extra_allowed=extra_allowed)
                except ValueError as error:
                    raise ValueError(f'review.json checks[{i}].findings[{j}].path: use one exact file from schema citation_inputs; separate multiple sources into separate findings') from error
        outcome = m.review_result(draft, inputs, extra_allowed=extra_allowed)
        if quality_diagnostics_version is not None:
            from task_generator.production.quality_diagnostics import validate_record_relations, validate_rubric_diagnostic
            relations = f.read(draft / 'record_relations.json')
            relation_result = validate_record_relations(relations, inputs)
            diagnostic = f.read(draft / 'rubric_diagnostic.json')
            diagnostic_outcome = validate_rubric_diagnostic(
                diagnostic, f.read(inputs / 'new_rubric.json'), f.read(inputs / 'basis_draft.json'),
                f.read(inputs / 'calculation_evidence.json'), inputs)
            if outcome['quality'] == 'pass' and any(row['status'] != 'pass' for row in relation_result['relations']):
                raise ValueError('record_relations.json: independent review must not pass with a declared record-relation issue or uncertainty')
            if outcome['quality'] == 'pass' and diagnostic_outcome['decision'] != 'pass':
                raise ValueError('rubric_diagnostic.json: independent review must not pass with unresolved rubric diagnostics')
        return outcome
    if role == 'consult':
        result = f.read(draft / 'consultation.json')
        if not isinstance(result.get('summary'), str) or not isinstance(result.get('findings'), list) or not isinstance(result.get('questions'), list):
            raise ValueError('invalid_consultation')
        extra_allowed = teacher_citation_paths(inputs, protocol)
        for index, r in enumerate(result['findings']):
            try:
                m.reference(inputs, r['path'], extra_allowed=extra_allowed)
            except ValueError as error:
                raise ValueError(f'consultation.json.findings[{index}].path: {error}; expected one exact allowed file from schema consultation citation_inputs') from error
            if not all(r.get(k) for k in ('locator', 'observation', 'limitation')):
                raise ValueError('consultation_evidence_required')
        return {'status': 'completed'}
    contract = f.read(inputs / 'deliverable_contract.json')
    expected = {r['relative_path'] for r in contract['deliverables']}
    actual = set(f.files(draft))
    if role == 'devsolve':
        if 'diagnostic.json' not in actual:
            raise ValueError('development_diagnostic_required')
        diagnostic = f.read(draft / 'diagnostic.json')
        if not all(key in diagnostic for key in ('evidence_used', 'calculations', 'ambiguities', 'barriers')):
            raise ValueError('development_diagnostic_fields_required')
        actual.remove('diagnostic.json')
    return {'status': 'completed', 'delivery_status': 'valid' if expected == actual else 'invalid',
            'missing': sorted(expected - actual), 'extra': sorted(actual - expected), 'professional_correctness': 'not_scored'}


def _inventory(root, path, args=None):
    if not path.is_dir():
        raise ValueError('inventory_path_must_be_directory')
    rows = []
    args = args or {}
    offset = args.get('offset', 0)
    limit = args.get('max_chars', 30000)
    paths = (item for item in sorted(path.rglob('*')) if item.is_file())
    truncated = False
    used = 0
    for index, item in enumerate(itertools.islice(paths, offset, None), offset):
        if not item.is_file():
            continue
        row = {'path': item.relative_to(root).as_posix(), 'sha256': f.digest(item),
               'bytes': item.stat().st_size, 'extension': item.suffix.lower()}
        if item.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(item, read_only=True, data_only=False)
            row['sheets'] = [{'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column}
                             for sheet in book.worksheets]
            book.close()
        elif item.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(item)
            row.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables))
        elif item.suffix.lower() == '.pdf':
            from pypdf import PdfReader
            row['pages'] = len(PdfReader(item).pages)
        size = len(json.dumps(row, ensure_ascii=False).encode('utf-8'))
        if rows and (len(rows) >= 100 or used + size > limit):
            truncated = True
            break
        rows.append(row)
        used += size
    return {'root': path.relative_to(root).as_posix() or '.', 'files': rows,
            'truncated': truncated, 'offset': offset, 'next_offset': offset + len(rows) if truncated else None,
            'read_limit_bytes': limit}


def _page_records(records, args, limit):
    """Bound serialized UTF-8 content; oversized records have resumable fragments."""
    offset = args.get('offset', 0)
    fragment = args.get('fragment_offset', 0)
    if type(fragment) is not int or fragment < 0:
        raise ValueError('fragment_offset_must_be_nonnegative_integer')
    iterator = iter(itertools.islice(records, offset, None))
    selected, used = [], 0
    for index, record in enumerate(iterator, offset):
        encoded = json.dumps(record, ensure_ascii=False).encode('utf-8')
        if fragment or len(encoded) > limit:
            if selected:
                return {'records': selected, 'truncated': True, 'next_offset': index, 'next_fragment_offset': 0}
            # Fragment offsets are characters in the serialized record, never bytes.
            text = encoded.decode('utf-8')
            piece = text[fragment:fragment + max(1, limit // 4)]
            complete = fragment + len(piece) >= len(text)
            return {'records': [], 'record_json_fragment': piece, 'record_complete': complete,
                    'truncated': True, 'next_offset': index + int(complete),
                    'next_fragment_offset': 0 if complete else fragment + len(piece)}
        if used + len(encoded) > limit or len(selected) >= 200:
            return {'records': selected, 'truncated': True, 'next_offset': index, 'next_fragment_offset': 0}
        selected.append(record)
        used += len(encoded)
    return {'records': selected, 'truncated': False, 'next_offset': None, 'next_fragment_offset': None}


def _inspect_one(area, relative, args):
    root = Path('/draft' if area == 'draft' else '/workspace/inputs')
    path = f.safe_path(root, relative)
    f.files(root)
    limit = args.get('max_chars', 30000)
    offset = args.get('offset', 0)
    if isinstance(limit, bool) or not isinstance(limit, int) or limit < 1 or limit > 100000:
        raise ValueError('max_chars_must_be_integer_1_to_100000')
    if isinstance(offset, bool) or not isinstance(offset, int) or offset < 0:
        raise ValueError('offset_must_be_nonnegative_integer')
    if args.get('mode') == 'inventory':
        return _inventory(root, path, args)
    if path.suffix == '.xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        book = load_workbook(path, read_only=True, data_only=args.get('cached', False))
        sheet = book[args['sheet']] if args.get('sheet') else book.active
        bounds = range_boundaries(args.get('range', sheet.calculate_dimension()))
        cells = sheet.iter_rows(min_col=bounds[0], min_row=bounds[1], max_col=bounds[2], max_row=bounds[3])
        records = ({'cell': c.coordinate, 'value': str(c.value), 'type': c.data_type}
                   for row in cells for c in row if c.value is not None)
        try:
            page = _page_records(records, args, limit)
            return {'sheets': book.sheetnames, 'sheet': sheet.title, 'cells': page.pop('records'),
                    'read_limit_bytes': limit, 'offset': offset, **page}
        finally:
            book.close()
    if path.suffix == '.docx':
        from docx import Document
        doc = Document(path)
        section = args.get('section', 'paragraphs')
        if section not in ('paragraphs', 'tables'):
            raise ValueError('docx_section_must_be_paragraphs_or_tables')
        records = ({'paragraph': i, 'text': p.text} for i, p in enumerate(doc.paragraphs, 1)) if section == 'paragraphs' else (
            {'table': t, 'row': r, 'cells': [c.text for c in row.cells]}
            for t, table in enumerate(doc.tables, 1) for r, row in enumerate(table.rows, 1))
        page = _page_records(records, args, limit)
        return {section: page.pop('records'), 'section': section, 'paragraph_count': len(doc.paragraphs),
                'table_count': len(doc.tables), 'read_limit_bytes': limit, 'offset': offset, **page}
    if path.suffix == '.pdf':
        from pypdf import PdfReader
        pdf = PdfReader(path)
        page = args.get('page', 1)
        text = pdf.pages[page-1].extract_text() or ''
        selected = text[offset:offset + limit]
        return {'pages': len(pdf.pages), 'page': page, 'text': selected,
                'read_limit': limit, 'offset': offset,
                'truncated': offset + len(selected) < len(text),
                'next_offset': offset + len(selected) if offset + len(selected) < len(text) else None}
    if not path.is_file():
        raise ValueError('inspect_path_must_be_file_or_inventory_directory')
    text = path.read_text(encoding='utf-8')
    selected = text[offset:offset + limit]
    return {'text': selected, 'read_limit': limit, 'offset': offset,
            'truncated': offset + len(selected) < len(text),
            'next_offset': offset + len(selected) if offset + len(selected) < len(text) else None}


def inspect(args):
    args = dict(args)
    if args.get('mode') == 'inventory' and 'path' not in args and 'paths' not in args:
        args['path'] = '.'
    if isinstance(args.get('path'), list) and 'paths' not in args:
        args['paths'] = args.pop('path')
    area = args.get('area', 'draft')
    if area not in ('draft', 'inputs'):
        raise ValueError('unknown_area')
    paths = args.get('paths')
    if paths is not None:
        if not isinstance(paths, list) or not paths or len(paths) > 20 or not all(isinstance(p, str) for p in paths):
            raise ValueError('paths_must_be_nonempty_string_list_max_20')
        per_file = dict(args, max_chars=min(args.get('max_chars', 30000), 100000 // len(paths)))
        items = {path: _inspect_one(area, path, per_file) for path in paths}
        return {'items': items,
                'requested_paths': len(paths),
                'truncated_paths': [path for path, result in items.items()
                                    if result.get('truncated')]}
    if not isinstance(args.get('path'), str):
        raise ValueError('inspect.path: path_or_paths_required; expected path string or paths array; for root inventory use mode inventory')
    return _inspect_one(area, args['path'], args)


def render(args):
    root = Path('/draft' if args.get('area', 'draft') == 'draft' else '/workspace/inputs')
    path = f.safe_path(root, args['path'])
    target = Path(tempfile.mkdtemp(prefix='factory-render-'))
    if args.get('recalculate'):
        if path.suffix != '.xlsx':
            raise ValueError('recalculate_xlsx_only')
        fmt = 'xlsx'
    else:
        fmt = 'pdf'
    if path.suffix != '.pdf':
        subprocess.run(['libreoffice', '-env:UserInstallation=file://' + str(target / 'profile'), '--headless',
                        '--convert-to', fmt, '--outdir', str(target), str(path)], check=True, timeout=120, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        path = target / (path.stem + '.' + fmt)
    if not path.is_file():
        raise ValueError('render_output_missing')
    if fmt == 'pdf':
        subprocess.run(['pdftoppm', '-scale-to', '1400', '-png', str(path), str(target / 'page')], check=True, timeout=120, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    return {'output': str(path), 'pages': [str(p) for p in sorted(target.glob('page-*.png'))], 'original_unchanged': True}


def diagnose(args, cfg):
    """Expose structural evidence; Agents retain semantic professional judgment."""
    subject = args.get('subject')
    if subject == 'edit':
        from task_generator.production import task_method_pilot as method
        from task_generator.production.quality_diagnostics import edit_difference
        if cfg.get('diagnostic_micro_version') == 1:
            if cfg['role'] != 'compile':
                raise ValueError('edit_history_not_visible_to_review')
            versions = f.read('/workspace/inputs/diagnostic_context/edit_versions.json')
            return edit_difference(versions['before'], versions['after'], versions['before_result'], versions['after_result'])
        source = f.read('/workspace/inputs/task.json')
        edited = f.read('/draft/task.json')
        return edit_difference(source, edited,
                               method.task_result(Path('/workspace/inputs'), Path('/workspace/inputs')),
                               method.task_result(Path('/draft'), Path('/workspace/inputs')))
    if subject == 'rubric':
        from task_generator.production.quality_diagnostics import rubric_view
        root = Path('/draft') if (Path('/draft') / 'new_rubric.json').is_file() else Path('/workspace/inputs')
        required = ('new_rubric.json', 'basis_draft.json', 'calculation_evidence.json')
        if not all((root / name).is_file() for name in required):
            raise ValueError('rubric_diagnose_requires_rubric_basis_and_calculation_evidence')
        candidate_root = Path('/workspace/inputs')
        return rubric_view(*(f.read(root / name) for name in required), candidate_root=candidate_root)
    if subject == 'record-relations':
        from task_generator.production.quality_diagnostics import (
            CONSTRAINTS, KINDS, PRECISIONS, VERSION, validate_record_relations,
        )
        root = Path('/draft/candidate') if cfg['role'] == 'world' else Path('/workspace/inputs')
        result = {'version': VERSION,
                  'contract': {'observations': [{'observation_id': 'unique ID', 'path': 'candidate file path',
                      'locator': 'exact row/cell/paragraph', 'quote': 'verbatim visible text or value',
                      'sha256': 'current file hash', 'timestamp': 'ISO-8601 date/time',
                      'precision': sorted(PRECISIONS), 'kind': sorted(KINDS)}],
                      'relations': [{'relation_id': 'unique ID', 'subject_id': 'observation ID',
                          'related_id': 'observation ID', 'constraint': sorted(CONSTRAINTS),
                          'statement': 'what this order tests'}],
                      'not_applicable_reason': 'required only when no relation is submitted'},
                  'candidate_files': f.files(root)}
        if 'record' in args:
            result['assessment'] = validate_record_relations(
                args['record'], root, candidate_only=cfg['role'] != 'world')
        return result
    raise ValueError('unknown_diagnose_subject:' + str(subject))


_CALL_STACK = []


def client(action, args):
    cfg = f.read('/workspace/role.json')
    if cfg.get('event_version') != 2:
        return _client(action, args)
    record = {'call_id': uuid.uuid4().hex, 'parent_call_id': _CALL_STACK[-1] if _CALL_STACK else None,
              'origin': 'internal' if _CALL_STACK else 'explicit', 'action': action, 'ok': False}
    _CALL_STACK.append(record['call_id'])
    started = time.monotonic()
    try:
        result = _client(action, args)
        record['ok'] = True
        return result
    except Exception as error:
        record['error_type'] = type(error).__name__
        raise
    finally:
        _CALL_STACK.pop()
        record['elapsed_seconds'] = time.monotonic() - started
        broker_call('log-call', {'record': record})


def _client(action, args):
    cfg = f.read('/workspace/role.json')
    protocol = cfg.get('protocol')
    if action == 'help':
        result = {'invocation': ('/workspace/bin/factory-tools <command> \'{"json":"object"}\', pipe JSON on stdin, '
                               'or use --input-file <JSON file> inside /tmp/factory-requests (scratch), /workspace/inputs or /draft; explicit input does not consume stdin'),
                'examples': [
                    "/workspace/bin/factory-tools schema",
                    "/workspace/bin/factory-tools inspect '{\"area\":\"inputs\",\"path\":\"SKILL.md\"}'"
                ],
                'commands': {'status': 'no arguments; current versions, completed stages, legal next actions, missing prerequisites and exact remaining budget',
                'inspect': 'area draft|inputs; path or paths (max 20); mode inventory for a directory; optional sheet/range/page/cached/max_chars',
                'render': 'area, path, optional recalculate:true; writes scratch only',
                'check': 'optional mode draft|ready (default ready)',
                'schema': 'optional artifact calculation_evidence|task|supervision|rubric|basis|comparison|consultation|review|development_diagnostic|quality_diagnostic|record_relations',
                'diagnose': 'JSON object with subject rubric|edit|record-relations; returns current structural evidence or validates a supplied record-relations object',
                'save-process': 'world only; first save before any candidate files',
                'snapshot': 'reason', 'diff': 'snapshot', 'consult': 'optional snapshot, reason (visible evidence question)',
                'record-disposition': 'request_id, finding_id, snapshot, decision accept|partial|reject, reason, evidence:[{area:draft|inputs,path,locator}]',
                'record-dispositions': 'request_id, snapshot_reason, items:[{finding_id,decision,reason,evidence:[...]}]; covers every finding once',
                'replay': 'harness compiler only; optional snapshot, reason, optional next_action return|development_trial|submit',
                'request-development-trial': 'harness compiler only after a passing replay; snapshot, reason',
                'finish': 'devsolve/consult/review/solve only; checks current output, then end normally',
                'handoff': 'optional snapshot, target world|mine|compile|stop, reason', 'submit': 'optional snapshot, reason'},
                'role': cfg['role'], 'inputs': sorted(f.files('/workspace/inputs')),
                'workflow_status': cfg.get('workflow_status'),
                'consultations': cfg.get('consultations', [])}
        if cfg.get('purpose') == 'tool_microtest':
            result['commands']['replay'] = 'microtest compiler only: optional snapshot, reason, next_action stop; requires ready draft and dispositions'
            for command in ('submit', 'request-development-trial'):
                result['commands'].pop(command, None)
            result['commands']['handoff'] = 'target stop only'
        return result
    if action == 'schema':
        if cfg.get('diagnostic_micro_version') is not None and args.get('artifact') == 'diagnostic_result':
            from task_generator.production.quality_diagnostics import micro_enabled, candidate_input_hashes
            micro_enabled(cfg)
            return {'artifact': 'diagnostic_result', 'version': 1, 'status': ['completed', 'upstream_issue'],
                    'reason': 'Evidence-based explanation; unresolved findings require upstream_issue',
                    'candidate_input_hashes': candidate_input_hashes('/workspace/inputs'),
                    'required_subjects': cfg['diagnostic_subjects'], 'end_action': 'finish',
                    'edit_context': 'diagnostic_context/edit_versions.json (compile only)',
                    'required_outputs': ['diagnostic_result.json', 'rubric_diagnostic.json', 'edit_record.json']}
        from task_generator.production import task_method_pilot as m
        from task_generator.planning.rubric_compiler_v2 import TaskSpecificAtomicRubricV1, TaskSpecificRubricV2
        rubric_model = TaskSpecificAtomicRubricV1 if cfg.get('atomic_rubric_version') else TaskSpecificRubricV2
        result = {'mine': m.MINING.replace('/output', '/draft'), 'compile': m.COMPILATION.replace('/output', '/draft'), 'rubric': TaskSpecificRubricV2.model_json_schema(),
                'review': m.REVIEW.replace('/output', '/draft'),
                'evidence_path_contract': 'Each finding.path is one exact existing file from citation_inputs. No joined paths, prefixes, or directories. Use separate findings for multiple sources.',
                'reserved_decision_id': 'deliverable_structure',
                'citation_inputs': sorted(f.files('/workspace/inputs'))}
        if protocol == 'task_factory_harness_v1':
            from task_generator.production import task_factory_harness as harness
            result['harness'] = {'mine_extra': 'design_intent.json: occupational_use, analysis_points, likely_difficulties, evidence[{path,locator}]',
                'compile_extra': 'basis_draft.json plus calculation_evidence.json and calculation_scripts/*.py',
                'calculation_contract': harness.calculation_contract_schema(),
                'development_diagnostic': 'diagnostic.json: evidence_used, calculations, ambiguities, barriers',
                'calculation_script_io': 'version 2 scripts receive --inputs <read-only input root> and emit one JSON object; calculations select scalar leaves by JSON Pointer; no writes or network'}
        artifact = args.get('artifact')
        if artifact:
            extra_citations = teacher_citation_paths(Path('/workspace/inputs'), protocol) if artifact == 'consultation' else ()
            allowed = {
                'task': {'contract': result['mine']},
                'supervision': {'contract': {
                    'file': 'supervision.json', 'status': 'string: compiled or upstream_issue',
                    'upstream_issues': 'empty array for compiled; nonempty located issues for upstream_issue',
                    'decisions': [{'decision_id': 'unique string', 'requirement_ids': ['IDs from task.json, complete coverage'],
                        'supported_judgment': 'nonempty supported judgment', 'conditional_completion': 'nonempty conditional path',
                        'gaps': 'array', 'follow_up': 'array', 'evidence': [{'path': 'candidate-visible file',
                            'locator': 'exact location', 'explanation': 'nonempty support explanation'}]}]},
                    'reserved_decision_id': 'deliverable_structure needs no fabricated professional decision'},
                'consultation': {'contract': {'file': 'consultation.json', 'summary': 'string',
                    'questions': 'array', 'findings': [{'path': 'one exact allowed input file', 'locator': 'nonempty string',
                        'observation': 'nonempty string', 'limitation': 'nonempty string'}]},
                    'citation_inputs': [path for path in result['citation_inputs'] if path.startswith('reference_files/')
                        or path in ('candidate_task.md', 'deliverable_contract.json', 'task.json', 'public_context.json',
                            'professional_rules.json', 'sources.json', 'supervision.json', 'new_rubric.json')
                        or path in extra_citations]},
                'rubric': {'contract': rubric_model.model_json_schema(),
                    'scoring_authority': 'new_rubric.json is the sole scoring authority'},
                'task_edit': {'contract': {
                    'files': ['task.json', 'edit_record.json'],
                    'edit_record': {'version': 'r10.candidate_edit.1',
                        'source_task_sha256': 'canonical JSON SHA-256 from schema response',
                        'edited_task_sha256': 'canonical JSON SHA-256 of draft task.json',
                        'changes': [{'change_id': 'unique', 'area': 'prompt|requirements|deliverables',
                            'original_locator': 'located source', 'edit_summary': 'what changed',
                            'actual_change_paths': 'exact paths from diagnose {subject:"edit"}',
                            'business_reason': 'why this returns judgment to candidate',
                            'preserved_evidence': [{'path': 'candidate-visible path', 'locator': 'location',
                                'explanation': 'what remains supported'}],
                            'returned_judgment': 'professional judgment no longer disclosed'}],
                        'no_change_reason': 'required only when changes is empty'},
                    'source_task_sha256': __import__('hashlib').sha256(json.dumps(
                        f.read('/workspace/inputs/task.json'), ensure_ascii=False, sort_keys=True,
                        separators=(',', ':')).encode()).hexdigest() if Path('/workspace/inputs/task.json').is_file() else None}},
                'basis': {'contract': ('basis_draft.json: requirements [{requirement_id, obligation, '
                    'candidate_obligation_refs:[{path:candidate_task.md|deliverable_contract.json,locator,explanation}], '
                    'rubric_ids:[criterion_id]}], plus supported and conditional judgments, calculations, '
                    'observable results and scoring boundaries. Candidate evidence supports analysis but does not create obligations.')},
                'comparison': {'contract': ('comparison.json: initial_basis_snapshot, trial_errors, compiler_omissions, '
                    'alternatives, upstream_defects, requirement_changes [{change_type:clarify|add|remove, description, '
                    'rubric_ids, candidate_obligation_refs}]. New scoring work requires candidate-visible support or upstream revision.')},
                'review': {'contract': result['review'], 'citation_inputs': result['citation_inputs']},
                'development_diagnostic': {'contract': 'diagnostic.json: evidence_used, calculations, ambiguities, barriers'},
                'quality_diagnostic': {'contract': {
                    'file': 'rubric_diagnostic.json', 'version': 'r10.quality_diagnostics.1',
                    'rubric_sha256': 'current new_rubric.json canonical SHA-256',
                    'basis_sha256': 'current basis_draft.json canonical SHA-256',
                    'calculation_evidence_sha256': 'current calculation_evidence.json canonical SHA-256',
                    'candidate_input_hashes': 'exact candidate-visible input hash map from diagnose',
                    'criteria': [{'criterion_id': 'every current rubric ID',
                                  'criterion_sha256': 'current criterion identity from diagnose',
                                  'candidate_requirement_ids': 'current basis mapping',
                                  'atomicity': 'pass|issue|uncertain', 'atomicity_rationale': 'nonempty',
                                  'alternative_consistency': 'pass|issue|uncertain', 'alternative_consistency_rationale': 'nonempty',
                                  'overlap': 'pass|issue|uncertain', 'overlap_rationale': 'nonempty',
                                  'clause_refs': {'atomicity': ['current rubric JSON pointer'],
                                      'alternative_consistency': ['current rubric JSON pointer'],
                                      'overlap': ['current rubric JSON pointer']},
                                  'overlap_criterion_ids': 'current peer criterion IDs, empty when none'}],
                    'decision': 'pass only when every criterion check is pass; otherwise revision_required or upstream_issue',
                    'authority': 'diagnostic explains or blocks a draft; only new_rubric.json defines scoring'}},
                'record_relations': {'contract': 'Use diagnose {subject:"record-relations"}. World writes hidden/material_relations.json; final reviewer writes record_relations.json from candidate-visible files only.'},
            }
            if cfg.get('atomic_rubric_version'):
                allowed['supervision'] = {'contract': {
                    'file': 'supervision.json', 'status': 'compiled or upstream_issue',
                    'upstream_issues': 'empty for compiled; located issues for upstream_issue',
                    'decisions': [{'decision_id': 'unique string',
                        'requirement_ids': ['task requirement IDs; complete coverage'],
                        'reference_analysis': 'teacher reference derivation, not a scoring rule',
                        'known_facts': ['candidate-visible supported facts'],
                        'uncertainties': ['unresolved limits'], 'follow_up': ['useful actions'],
                        'evidence': [{'path': 'candidate-visible file', 'locator': 'exact location',
                                     'explanation': 'support and limitation'}]}],
                    'forbidden_scoring_fields': ['conditional_completion', 'score_boundaries'],
                    'authority': 'new_rubric.json alone defines every scoring condition'}}
            if protocol == 'task_factory_harness_v1':
                rubric_ids = []
                if 'new_rubric.json' in f.files('/draft'):
                    try:
                        rubric_ids = [row['criterion_id'] for row in f.read('/draft/new_rubric.json').get('criteria', [])]
                    except (KeyError, TypeError, json.JSONDecodeError):
                        rubric_ids = []
                references = [{'path': 'reference_files/' + path, 'sha256': sha}
                              for path, sha in sorted(f.files('/workspace/inputs/reference_files').items())]
                allowed['calculation_evidence'] = {
                    'contract': result['harness']['calculation_contract'],
                    'current_rubric_ids': rubric_ids,
                    'candidate_reference_files': references,
                }
                if str(cfg.get('obligation_trace_version')) == '2':
                    allowed['basis']['contract'] = (
                        'requirements:[{requirement_id,obligation,basis_kind:explicit|material_instruction|necessary_derivation|optional,'
                        'candidate_obligation_refs:[{path,locator,explanation,sha256}],rubric_ids}]. '
                        'Sources: candidate_task.md, deliverable_contract.json, public_context.json, reference_files/*. '
                        'Necessary derivation also needs serves_requirement_id (explicit/material_instruction), necessity, alternative_paths. '
                        'Optional analysis has empty rubric_ids; it cannot be a full-credit condition. '
                        'Facts alone do not establish obligations. deliverable_structure is the contract-grounded exception.')
                    allowed['basis']['candidate_sources'] = {
                        path: sha for path, sha in f.files('/workspace/inputs').items()
                        if path in ('candidate_task.md', 'deliverable_contract.json', 'public_context.json') or path.startswith('reference_files/')}
                    allowed['comparison']['contract'] = (
                        'initial_basis_snapshot plus trial_errors, compiler_omissions, alternatives, upstream_defects, '
                        'requirement_changes:[{change_id,change_type,description,rubric_ids,candidate_obligation_refs}]. '
                        'Cover every actual_changes ID exactly once; explain professional support and alternative paths.')
                    allowed['comparison']['initial_basis_snapshot'] = cfg.get('initial_basis_snapshot')
                    if Path('/workspace/inputs/initial_basis').is_dir():
                        from task_generator.production.obligation_trace import actual_changes
                        try:
                            allowed['comparison']['actual_changes'] = actual_changes('/workspace/inputs/initial_basis', '/draft')
                        except (ValueError, KeyError, OSError) as error:
                            allowed['comparison']['diff_error'] = str(error)
            if artifact not in allowed:
                raise ValueError('unknown_schema_artifact:' + str(artifact))
            return {'artifact': artifact, **allowed[artifact]}
        return result
    if action == 'diagnose':
        result = diagnose(args, cfg)
        broker_call('log-check', {'result': {'tool': action, 'arguments': args, 'ok': True,
                                             'hashes': f.files('/draft'), 'role': cfg['role']}})
        return result
    if action == 'status':
        return broker_call('status', {})
    if cfg.get('diagnostic_micro_version') is not None and action in ('consult', 'handoff', 'submit', 'replay', 'request-development-trial'):
        raise ValueError('diagnostic_mode_requires_finish: save diagnosis and use finish; production actions are unavailable')
    if action in ('check', 'inspect', 'render'):
        mode = args.get('mode', 'ready') if action == 'check' else None
        if action == 'check' and cfg['role'] == 'compile' and cfg.get('diagnostic_micro_version') is not None:
            from task_generator.production.quality_diagnostics import micro_result
            result = micro_result('/draft', '/workspace/inputs', cfg)
            broker_call('log-check', {'result': {'tool': 'check', 'hashes': f.files('/draft'),
                                               'role': cfg['role'], 'ok': True, 'mode': mode}})
            return result
        result = check(cfg['role'], Path('/draft'), Path('/workspace/inputs'), protocol, mode,
                       cfg.get('calculation_contract_version'), cfg.get('obligation_trace_version'),
                       cfg.get('initial_basis_snapshot'), cfg.get('atomic_rubric_version'),
                       cfg.get('quality_diagnostics_version')) if action == 'check' else (inspect(args) if action == 'inspect' else render(args))
        broker_call('log-check', {'result': {'tool': action, 'arguments': args, 'ok': True,
                                           'hashes': f.files('/draft'), 'role': cfg['role'],
                                           **({'mode': mode} if action == 'check' else {})}})
        return result
    if action in ('consult', 'handoff', 'submit', 'replay', 'request-development-trial') and not (action == 'handoff' and args.get('target') == 'stop'):
        outcome = client('check', {'mode': 'draft' if action == 'consult' else 'ready'})
        if outcome.get('status') != 'completed' and not (action == 'handoff' and args.get('target') in ('world', 'mine')):
            raise ValueError('output_not_ready: request upstream revision or handoff to stop; ' + outcome['status'])
    if action == 'finish':
        if cfg['role'] not in ('devsolve', 'consult', 'review', 'solve') and not (
                cfg['role'] == 'compile' and cfg.get('diagnostic_micro_version') == 1):
            raise ValueError('readonly_roles_finish_only')
        outcome = client('check', {})
        if outcome.get('delivery_status') == 'invalid':
            raise ValueError('delivery_contract_invalid: run check for missing and extra paths')
    return broker_call(action, args)


def cli_payload(argv, stdin, allowed_roots=None):
    trailing = argv[2:]
    input_file = None
    positional = []
    if '--input-file' in trailing:
        if trailing.count('--input-file') != 1:
            raise ValueError('input_file_option_must_appear_once')
        index = trailing.index('--input-file')
        if index + 1 >= len(trailing):
            raise ValueError('input_file_path_required')
        input_file = trailing[index + 1]
        positional = trailing[:index] + trailing[index + 2:]
    else:
        positional = trailing
    # Explicit input must never wait for an inherited, open pipe to close.
    stdin_text = stdin.read(1024 * 1024 + 1) if not positional and input_file is None and not stdin.isatty() else ''
    provided = int(bool(positional)) + int(input_file is not None) + int(bool(stdin_text.strip()))
    if provided > 1 or len(positional) > 1:
        raise ValueError('use_exactly_one_of_positional_json_stdin_or_input_file')
    if input_file is not None:
        path = Path(input_file).resolve()
        roots = tuple(Path(root).resolve() for root in (allowed_roots or ('/workspace/inputs', '/draft', '/tmp/factory-requests')))
        if not any(path.is_relative_to(root) for root in roots) or not path.is_file():
            raise ValueError('input_file_must_be_readable_within_role_directories')
        if path.stat().st_size > 1024 * 1024:
            raise ValueError('request_exceeds_1_mib')
        text = path.read_text(encoding='utf-8')
    else:
        text = positional[0] if positional else stdin_text
    if len(text.encode('utf-8')) > 1024 * 1024:
        raise ValueError('request_exceeds_1_mib')
    value = json.loads(text) if text.strip() else {}
    if not isinstance(value, dict):
        raise ValueError('request_must_be_json_object')
    return value


if __name__ == '__main__':
    if sys.argv[1] == 'serve':
        serve(sys.argv[2])
    elif sys.argv[1] == 'validate':
        print(json.dumps(check(sys.argv[2], Path(sys.argv[3]), Path(sys.argv[4])), ensure_ascii=False))
    else:
        try:
            args = cli_payload(sys.argv, sys.stdin)
            print(json.dumps({'ok': True, 'result': client(sys.argv[1], args)}, ensure_ascii=False))
        except Exception as error:
            print(json.dumps({'ok': False, 'error': type(error).__name__ + ':' + str(error),
                              'issues': issues_for(error)}, ensure_ascii=False))
            sys.exit(1)
