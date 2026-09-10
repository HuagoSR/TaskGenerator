"""Versioned structural evidence for task-quality diagnostics.

The functions here deliberately do not decide whether a professional judgment is
good.  They make the evidence for that judgment inspectable: which candidate
input really changed, which rubric item was reviewed, and which dated records
were compared.
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from task_generator.production import agent_factory as factory


VERSION = "r10.quality_diagnostics.1"
STATES = {"pass", "issue", "uncertain"}
PRECISIONS = {"day", "minute", "second"}
KINDS = {"record_created", "record_updated", "record_distributed", "event", "query", "excerpt"}
CONSTRAINTS = {"subject_not_before_related", "subject_not_after_related", "derived_from"}


def micro_enabled(config):
    version = config.get("diagnostic_micro_version")
    if version is None:
        return False
    if version != 1 or config.get("purpose") != "quality_diagnostic_microtest":
        raise ValueError("invalid_diagnostic_micro_scope")
    return True


def micro_result(draft, inputs, config):
    """Accept diagnostic evidence, never waive production submission gates."""
    from task_generator.production import task_method_pilot as method
    from task_generator.production import task_factory_harness as harness
    if not micro_enabled(config):
        raise ValueError("diagnostic_micro_scope_required")
    draft, inputs = Path(draft), Path(inputs)
    result = factory.read(draft / "diagnostic_result.json")
    if result.get("version") != 1 or result.get("status") not in {"completed", "upstream_issue"}:
        raise ValueError("diagnostic_result.json: expected version 1 and completed|upstream_issue")
    if not isinstance(result.get("reason"), str) or not result["reason"].strip():
        raise ValueError("diagnostic_result.json.reason: explanation required")
    if result.get("candidate_input_hashes") != candidate_input_hashes(inputs):
        raise ValueError("diagnostic_result.json: candidate identity mismatch")
    required = config.get("diagnostic_subjects")
    if not isinstance(required, list) or not required or set(required) - {"rubric", "edit", "record-relations"}:
        raise ValueError("diagnostic_subjects_invalid")
    compilation = method.atomic_compilation_result(draft, inputs)
    compilation_upstream = compilation.get("status") == "upstream_issue"
    if compilation_upstream and result["status"] != "upstream_issue":
        raise ValueError("diagnostic_result.json: compilation upstream_issue requires upstream_issue status")
    calculations = {}
    calculation_error = None
    try:
        calculations = factory.read(draft / "calculation_evidence.json")
        harness.validate_calculation_bundle(draft, inputs)
    except (ValueError, KeyError, TypeError, OSError, json.JSONDecodeError) as error:
        calculation_error = str(error)
        if not compilation_upstream and result["status"] != "upstream_issue":
            raise
    outcomes = {}
    if compilation_upstream:
        outcomes["compilation"] = compilation
    if calculation_error is not None:
        outcomes["calculation"] = {"status": "not_verified", "reason": calculation_error}
    if "rubric" in required:
        outcomes["rubric"] = validate_rubric_diagnostic(
            factory.read(draft / "rubric_diagnostic.json"), factory.read(draft / "new_rubric.json"),
            factory.read(draft / "basis_draft.json"), calculations, inputs)
    if "edit" in required:
        context = factory.read(inputs / "diagnostic_context/edit_versions.json")
        difference = edit_difference(context["before"], context["after"],
                                     context["before_result"], context["after_result"])
        validate_edit_claims(factory.read(draft / "edit_record.json"), difference)
        outcomes["edit"] = difference
    if "record-relations" in required:
        relations = factory.read(draft / "record_relations.json")
        outcomes["record-relations"] = validate_record_relations(relations, inputs)
        if not relations.get("observations") or not relations.get("relations"):
            raise ValueError("record_relations.json: explicit observations and relations required")
    unresolved = (outcomes.get("rubric", {}).get("decision", "pass") != "pass" or
                  any(row["status"] != "pass" for row in outcomes.get("record-relations", {}).get("relations", [])))
    if unresolved and result["status"] != "upstream_issue":
        raise ValueError("diagnostic_result.json: unresolved findings require upstream_issue")
    return {"status": "upstream_issue" if compilation_upstream or result["status"] == "upstream_issue" else "completed",
            "reason": result["reason"], "diagnostics": outcomes,
            "candidate_input_hashes": candidate_input_hashes(inputs)}


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def json_hash(value):
    import hashlib
    return hashlib.sha256(canonical(value).encode("utf-8")).hexdigest()


def candidate_input_hashes(root):
    """Hash only files that are part of the candidate-visible assignment."""
    root = Path(root)
    names = []
    for name in factory.files(root):
        if (name in {"candidate_task.md", "deliverable_contract.json", "public_context.json"}
                or name.startswith("reference_files/")):
            names.append(name)
    return {name: factory.digest(root / name) for name in sorted(names)}


def _pointer_part(value):
    return str(value).replace("~", "~0").replace("/", "~1")


def json_differences(before, after, pointer=""):
    """Return deterministic RFC-6901-like field differences."""
    if type(before) is not type(after):
        return [{"path": pointer or "/", "before": before, "after": after}]
    if isinstance(before, dict):
        rows = []
        for key in sorted(set(before) | set(after)):
            child = pointer + "/" + _pointer_part(key)
            if key not in before:
                rows.append({"path": child, "before": None, "after": after[key]})
            elif key not in after:
                rows.append({"path": child, "before": before[key], "after": None})
            else:
                rows.extend(json_differences(before[key], after[key], child))
        return rows
    if isinstance(before, list):
        rows = []
        for index in range(max(len(before), len(after))):
            child = pointer + "/" + str(index)
            if index >= len(before):
                rows.append({"path": child, "before": None, "after": after[index]})
            elif index >= len(after):
                rows.append({"path": child, "before": before[index], "after": None})
            else:
                rows.extend(json_differences(before[index], after[index], child))
        return rows
    return [] if before == after else [{"path": pointer or "/", "before": before, "after": after}]


def edit_difference(source_task, edited_task, source_result, edited_result):
    """Separate candidate-visible effects from task-factory metadata changes."""
    candidate_task = json_differences(source_result["candidate_task"], edited_result["candidate_task"])
    contract = json_differences(source_result["contract"], edited_result["contract"])
    metadata = json_differences(source_task, edited_task)
    # The compiled candidate task includes the delivery appendix.  Contract-only
    # changes must not create a second, fictitious prompt edit obligation.
    prompt_source = source_task.get("prompt")
    prompt_edited = edited_task.get("prompt")
    prompt = json_differences(prompt_source, prompt_edited, "/prompt")
    candidate_changes = ([{"area": "candidate_task", **row} for row in prompt] +
                         [{"area": "deliverable_contract", **row} for row in contract])
    if candidate_changes:
        status = "candidate_input_changed"
    elif metadata:
        status = "internal_only_change"
    else:
        status = "no_change"
    return {
        "version": VERSION,
        "source_task_sha256": json_hash(source_task),
        "edited_task_sha256": json_hash(edited_task),
        "source_candidate_task_sha256": json_hash(source_result["candidate_task"]),
        "edited_candidate_task_sha256": json_hash(edited_result["candidate_task"]),
        "source_contract_sha256": json_hash(source_result["contract"]),
        "edited_contract_sha256": json_hash(edited_result["contract"]),
        "effect": status,
        "candidate_changes": candidate_changes,
        "metadata_changes": [{"area": "task_metadata", **row} for row in metadata],
        "rendered_candidate_task_changes": candidate_task,
    }


def validate_edit_claims(record, diagnostic):
    """Bind editor declarations to actual, candidate-visible differences."""
    candidate = {(row["area"], row["path"]) for row in diagnostic["candidate_changes"]}
    metadata = {row["path"] for row in diagnostic["metadata_changes"]}
    actual = candidate | {("task_metadata", path) for path in metadata}
    claimed = set()
    for index, row in enumerate(record.get("changes", [])):
        paths = row.get("actual_change_paths")
        if not isinstance(paths, list) or not paths or not all(isinstance(x, str) for x in paths):
            raise ValueError(f"edit_record.json.changes[{index}].actual_change_paths: expected nonempty actual path list")
        area = row.get("area")
        for path in paths:
            if area == "prompt":
                if ("candidate_task", path) not in candidate:
                    raise ValueError(f"edit_record.json.changes[{index}].actual_change_paths: prompt claim must name an actual candidate_task change")
                claimed.add(("candidate_task", path))
                if path in metadata:
                    claimed.add(("task_metadata", path))
            elif area == "deliverables":
                if (("deliverable_contract", path) not in candidate
                        and not (path.startswith("/deliverables") and path in metadata)):
                    raise ValueError(f"edit_record.json.changes[{index}].actual_change_paths: deliverable claim must name an actual contract or deliverables change")
                if ("deliverable_contract", path) in candidate:
                    claimed.add(("deliverable_contract", path))
                if path in metadata:
                    claimed.add(("task_metadata", path))
            elif area == "requirements":
                if not path.startswith("/requirements") or path not in metadata:
                    raise ValueError(f"edit_record.json.changes[{index}].actual_change_paths: requirements claim must name an actual requirements change")
                claimed.add(("task_metadata", path))
            elif area == "metadata":
                if path not in metadata or path.startswith(("/prompt", "/requirements", "/deliverables")):
                    raise ValueError(f"edit_record.json.changes[{index}].actual_change_paths: metadata claim must name an actual internal-only change")
                claimed.add(("task_metadata", path))
            else:
                raise ValueError(f"edit_record.json.changes[{index}].area: unsupported edit area")
    missing = actual - claimed
    if missing:
        raise ValueError("edit_record.json.changes: candidate-visible changes lack an edit claim: " +
                         ",".join(f"{area}:{path}" for area, path in sorted(missing)))


def rubric_view(rubric, basis=None, calculations=None, candidate_root=None):
    requirement_map = {}
    if isinstance(basis, dict):
        for row in basis.get("requirements", []):
            for criterion_id in row.get("rubric_ids", []):
                requirement_map.setdefault(criterion_id, []).append(row.get("requirement_id"))
    anchors = {}
    if isinstance(calculations, dict):
        for calculation in calculations.get("calculations", []):
            for criterion_id in calculation.get("rubric_ids", []):
                anchors.setdefault(criterion_id, []).append(calculation.get("calculation_id"))
    return {
        "version": VERSION,
        "rubric_sha256": json_hash(rubric),
        "basis_sha256": json_hash(basis),
        "calculation_evidence_sha256": json_hash(calculations),
        "candidate_input_hashes": candidate_input_hashes(candidate_root) if candidate_root is not None else None,
        "criteria": [{
            "criterion_id": row["criterion_id"],
            "criterion_sha256": json_hash(row),
            "criterion_pointer": "/criteria/" + str(index),
            "requirement": row.get("requirement"),
            "weight": row.get("max_points"),
            "candidate_requirement_ids": sorted(filter(None, requirement_map.get(row["criterion_id"], []))),
            "full_credit_condition": row.get("full_credit_condition"),
            "applicability": row.get("applicability"),
            "acceptable_alternatives": row.get("acceptable_alternatives"),
            "tolerance": row.get("tolerance"),
            "calculation_ids": sorted(filter(None, anchors.get(row["criterion_id"], []))),
        } for index, row in enumerate(rubric.get("criteria", []))],
    }


def validate_rubric_diagnostic(record, rubric, basis, calculations, candidate_root=None):
    if record.get("version") != VERSION:
        raise ValueError("rubric_diagnostic.json.version: expected " + VERSION)
    view = rubric_view(rubric, basis, calculations, candidate_root)
    for field in ("rubric_sha256", "basis_sha256", "calculation_evidence_sha256"):
        if record.get(field) != view[field]:
            raise ValueError(f"rubric_diagnostic.json.{field}: current evidence identity mismatch")
    if candidate_root is not None and record.get("candidate_input_hashes") != view["candidate_input_hashes"]:
        raise ValueError("rubric_diagnostic.json.candidate_input_hashes: current candidate identity mismatch")
    rows = record.get("criteria")
    expected = {row["criterion_id"] for row in view["criteria"]}
    if not isinstance(rows, list) or {row.get("criterion_id") for row in rows if isinstance(row, dict)} != expected or len(rows) != len(expected):
        raise ValueError("rubric_diagnostic.json.criteria: expected exactly one row per criterion")
    for row in rows:
        index = row["criterion_id"]
        expected_row = next(item for item in view["criteria"] if item["criterion_id"] == index)
        if row.get("criterion_sha256") != expected_row["criterion_sha256"]:
            raise ValueError(f"rubric_diagnostic.json.criteria[{index}].criterion_sha256: criterion identity mismatch")
        for field in ("atomicity", "alternative_consistency", "overlap"):
            if row.get(field) not in STATES or not isinstance(row.get(field + "_rationale"), str) or not row[field + "_rationale"].strip():
                raise ValueError(f"rubric_diagnostic.json.criteria[{index}].{field}: expected state and nonempty rationale")
        ids = row.get("candidate_requirement_ids")
        expected_ids = expected_row["candidate_requirement_ids"]
        if ids != expected_ids:
            raise ValueError(f"rubric_diagnostic.json.criteria[{index}].candidate_requirement_ids: expected current basis mapping")
        refs = row.get("clause_refs")
        allowed = {expected_row["criterion_pointer"] + "/" + name for name in (
            "requirement", "full_credit_condition", "applicability", "acceptable_alternatives", "tolerance")}
        if not isinstance(refs, dict):
            raise ValueError(f"rubric_diagnostic.json.criteria[{index}].clause_refs: expected located clause references")
        for field in ("atomicity", "alternative_consistency", "overlap"):
            values = refs.get(field)
            if not isinstance(values, list) or not values or any(value not in allowed for value in values):
                raise ValueError(f"rubric_diagnostic.json.criteria[{index}].clause_refs.{field}: expected current rubric pointers")
        overlaps = row.get("overlap_criterion_ids")
        if not isinstance(overlaps, list) or len(overlaps) != len(set(overlaps)) or any(
                value not in expected or value == index for value in overlaps):
            raise ValueError(f"rubric_diagnostic.json.criteria[{index}].overlap_criterion_ids: expected distinct current peer IDs")
    decision = record.get("decision")
    states = [row[field] for row in rows for field in ("atomicity", "alternative_consistency", "overlap")]
    if decision not in {"pass", "revision_required", "upstream_issue"}:
        raise ValueError("rubric_diagnostic.json.decision: invalid decision")
    if (decision == "pass") != all(state == "pass" for state in states):
        raise ValueError("rubric_diagnostic.json.decision: pass requires every criterion check to pass")
    return {"decision": decision, "view": view}


def _candidate_text(path):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        from docx import Document
        doc = Document(path)
        return "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return "\n".join(str(cell.value) for sheet in book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
        finally:
            book.close()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    raise ValueError("record_relations: unsupported candidate file type")


def _time_interval(value, precision):
    if not isinstance(value, str) or precision not in PRECISIONS:
        raise ValueError("record_relations: timestamp and precision are required")
    if precision == "day":
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("record_relations: day precision must use YYYY-MM-DD")
        start = datetime.fromisoformat(value).replace(tzinfo=timezone.utc)
        return start, start + timedelta(days=1), False
    pattern = (r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?:Z|[+-]\d{2}:\d{2})" if precision == "minute"
               else r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:Z|[+-]\d{2}:\d{2})")
    if not re.fullmatch(pattern, value):
        raise ValueError(f"record_relations: {precision} precision must match its stated ISO-8601 precision")
    normalized = value[:-1] + "+00:00" if value.endswith("Z") else value
    try:
        start = datetime.fromisoformat(normalized)
    except ValueError as error:
        raise ValueError("record_relations: timestamp must be ISO-8601") from error
    if start.tzinfo is None:
        raise ValueError("record_relations: minute and second precision require an explicit timezone")
    start = start.astimezone(timezone.utc)
    return start, start + (timedelta(minutes=1) if precision == "minute" else timedelta(seconds=1)), True


def validate_record_relations(record, root, *, candidate_only=True):
    """Validate submitted temporal/source comparisons without inventing relations."""
    if record.get("version") != VERSION:
        raise ValueError("record_relations.version: expected " + VERSION)
    observations = record.get("observations", [])
    relations = record.get("relations", [])
    if not observations and not str(record.get("not_applicable_reason", "")).strip():
        raise ValueError("record_relations: observations or not_applicable_reason required")
    by_id = {}
    for index, row in enumerate(observations):
        required = ("observation_id", "path", "locator", "quote", "sha256", "timestamp", "precision", "kind")
        if not isinstance(row, dict) or any(not isinstance(row.get(key), str) or not row[key].strip() for key in required):
            raise ValueError(f"record_relations.observations[{index}]: complete located observation required")
        if row["observation_id"] in by_id or row["kind"] not in KINDS:
            raise ValueError(f"record_relations.observations[{index}]: duplicate ID or invalid kind")
        path = factory.safe_path(root, row["path"])
        if not path.is_file() or factory.digest(path) != row["sha256"]:
            raise ValueError(f"record_relations.observations[{index}].sha256: candidate file identity mismatch")
        if candidate_only and not row["path"].startswith("reference_files/"):
            raise ValueError(f"record_relations.observations[{index}].path: candidate records only")
        if row["quote"] not in _candidate_text(path):
            raise ValueError(f"record_relations.observations[{index}].quote: quote not present in cited file")
        if path.suffix.lower() in {".txt", ".md", ".csv"} and re.fullmatch(r"line \d+", row["locator"]):
            line_number = int(row["locator"].split()[1])
            lines = path.read_text(encoding="utf-8").splitlines()
            if line_number < 1 or line_number > len(lines) or row["quote"] not in lines[line_number - 1]:
                raise ValueError(f"record_relations.observations[{index}].locator: quote not present at cited line")
        by_id[row["observation_id"]] = {**row, "interval": _time_interval(row["timestamp"], row["precision"])}
    seen = set()
    assessed = []
    for index, row in enumerate(relations):
        required = ("relation_id", "subject_id", "related_id", "constraint", "statement")
        if not isinstance(row, dict) or any(not isinstance(row.get(key), str) or not row[key].strip() for key in required):
            raise ValueError(f"record_relations.relations[{index}]: complete relation required")
        if row["relation_id"] in seen or row["constraint"] not in CONSTRAINTS:
            raise ValueError(f"record_relations.relations[{index}]: duplicate ID or invalid constraint")
        seen.add(row["relation_id"])
        if row["subject_id"] not in by_id or row["related_id"] not in by_id:
            raise ValueError(f"record_relations.relations[{index}]: unknown observation ID")
        if row["subject_id"] == row["related_id"]:
            raise ValueError(f"record_relations.relations[{index}]: relation cannot cite the same observation twice")
        subject, related = by_id[row["subject_id"]]["interval"], by_id[row["related_id"]]["interval"]
        if row["constraint"] == "derived_from":
            status = "agent_review_required"
        elif subject[2] != related[2]:
            status = "uncertain"
        elif row["constraint"] == "subject_not_before_related":
            status = "pass" if subject[0] >= related[1] else "issue" if subject[1] < related[0] else "uncertain"
        else:
            status = "pass" if subject[1] <= related[0] else "issue" if subject[0] > related[1] else "uncertain"
        assessed.append({"relation_id": row["relation_id"], "status": status,
                         "limitation": "Reported from declared timestamp precision; it does not infer undocumented updates or intent."})
    return {"version": VERSION, "observations": len(observations), "relations": assessed,
            "all_relations_pass": bool(assessed) and all(row["status"] == "pass" for row in assessed)}
