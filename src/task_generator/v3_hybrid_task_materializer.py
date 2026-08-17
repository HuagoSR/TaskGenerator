from __future__ import annotations

import hashlib
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Literal

from pydantic import BaseModel, ConfigDict, Field

from task_generator.v3_deliverable_contract import (
    DeliverableContractCompiler,
    DeliverableContractValidator,
)
from task_generator.v3_evidence_content_quality import (
    EvidenceContentQualityValidator,
)
from task_generator.v3_task_design_frontend import (
    CapabilityBriefV1,
    TaskDesignFrontend,
    TaskDesignProposalV1,
)
from task_generator.v3_validity_utility import (
    R5GovernanceBundleV1,
    RubricPlanV2,
    ValidityUtilityCompiler,
)
from task_generator.v3_rw_task_export_validator import RwTaskExportValidator


class HybridEvidenceFileRecordV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    node_id: str
    file_name: str
    relative_path: str
    source_ref_ids: List[str]
    linked_skill_ids: List[str] = Field(default_factory=list)
    linked_capability_ids: List[str] = Field(default_factory=list)
    row_count: int
    sha256: str
    candidate_visible: Literal[True] = True


class HybridEvidenceRelationRecordV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    relation_id: str
    from_node_id: str
    to_node_id: str
    relation_type: str
    from_file_name: str
    to_file_name: str
    solver_must_infer: bool


class HybridEvidenceDossierV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    dossier_version: Literal["v3.hybrid_evidence_dossier.1"] = (
        "v3.hybrid_evidence_dossier.1"
    )
    proposal_id: str
    files: List[HybridEvidenceFileRecordV1]
    relations: List[HybridEvidenceRelationRecordV1]
    source_ref_ids: List[str]


class DeterministicFactAnchorV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    anchor_id: str
    anchor_type: Literal["file_row_count", "relation_crosscheck"]
    candidate_visible_inputs: List[str]
    recomputation_method: str
    expected_value: Any
    linked_judgment_ids: List[str] = Field(default_factory=list)
    truth_authority: Literal["deterministic_candidate_visible_recomputation"] = (
        "deterministic_candidate_visible_recomputation"
    )


class DeterministicFactAnchorReportV1(BaseModel):
    report_version: Literal["v3.deterministic_fact_anchor.1"] = (
        "v3.deterministic_fact_anchor.1"
    )
    proposal_id: str
    anchors: List[DeterministicFactAnchorV1]
    all_inputs_candidate_visible: bool = True


class HybridRubricBindingV1(BaseModel):
    model_config = ConfigDict(extra="forbid")
    criterion_id: str
    criterion_type: Literal["fact", "reasoning", "deliverable"]
    observable_behavior: str
    skill_ids: List[str]
    capability_ids: List[str]
    evidence_or_judgment_ids: List[str]
    scoring_authority: Literal["program_compiled_binding_not_final_weight"] = (
        "program_compiled_binding_not_final_weight"
    )


class HybridRubricBindingPlanV1(BaseModel):
    plan_version: Literal["v3.hybrid_rubric_binding_plan.1"] = (
        "v3.hybrid_rubric_binding_plan.1"
    )
    proposal_id: str
    bindings: List[HybridRubricBindingV1]
    final_weights_assigned: Literal[False] = False


class ComplexityPreservationCheckV1(BaseModel):
    check_name: str
    passed: bool
    details: Dict[str, Any] = Field(default_factory=dict)


class ComplexityPreservationReportV1(BaseModel):
    report_version: Literal["v3.complexity_preservation.1"] = (
        "v3.complexity_preservation.1"
    )
    proposal_id: str
    decision: Literal["pass", "blocked"]
    checks: List[ComplexityPreservationCheckV1]
    preserved_relation_count: int = 0
    preserved_judgment_count: int = 0
    preserved_skill_binding_count: int = 0


class HybridFileVisualCheckV1(BaseModel):
    file_name: str
    openable: bool
    freeze_panes_present: bool
    auto_filter_present: bool
    fit_to_width: bool


class HybridVisualValidationReportV1(BaseModel):
    report_version: Literal["v3.hybrid_visual_validation.1"] = (
        "v3.hybrid_visual_validation.1"
    )
    proposal_id: str
    decision: Literal["pass", "blocked"]
    checks: List[HybridFileVisualCheckV1]


class HybridCandidatePackageManifestV1(BaseModel):
    manifest_version: Literal["v3.hybrid_candidate_package.1"] = (
        "v3.hybrid_candidate_package.1"
    )
    case_id: str
    proposal_id: str
    prompt_path: str
    reference_files: List[str]
    deliverable_contract_path: str
    teacher_artifacts_included: Literal[False] = False


class HybridMaterializationReportV1(BaseModel):
    report_version: Literal["v3.hybrid_task_materialization.1"] = (
        "v3.hybrid_task_materialization.1"
    )
    case_id: str
    brief_id: str
    proposal_id: str
    decision: Literal["pass", "blocked"]
    reason_codes: List[str] = Field(default_factory=list)
    proposal_validation_pass: bool = False
    evidence_file_count: int = 0
    deterministic_anchor_count: int = 0
    deliverable_contract_valid: bool = False
    complexity_preservation_pass: bool = False
    r5_offline_governance_pass: bool = False
    validity_overall_status: str = "not_evaluated"
    utility_profile_status: str = "not_evaluated"
    rubric_plan_decision: str = "not_evaluated"
    evidence_content_quality_decision: str = "not_evaluated"
    visual_validation_pass: bool = False
    rw_task_export_validation_status: Literal[
        "not_run", "invalid", "draft_compatible", "candidate_ready_compatible"
    ] = "not_run"
    candidate_teacher_isolation_pass: bool = False
    legacy_template_route_used: Literal[False] = False
    materialization_backend: Literal[
        "hybrid_deterministic_fixture_v1",
        "hybrid_semantic_artifact_v2",
    ] = "hybrid_deterministic_fixture_v1"


class WholeTaskRepairContextV1(BaseModel):
    context_version: Literal["v3.whole_task_repair_context.1"] = (
        "v3.whole_task_repair_context.1"
    )
    case_id: str
    proposal_id: str
    candidate_prompt: str
    evidence_dossier: Dict[str, Any]
    deliverable_contract: Dict[str, Any]
    deterministic_fact_anchors: Dict[str, Any]
    rubric_binding_plan: Dict[str, Any]
    complexity_findings: List[Dict[str, Any]] = Field(default_factory=list)
    additional_findings: List[str] = Field(default_factory=list)
    immutable_authorities: List[str] = Field(
        default_factory=lambda: [
            "deterministic fact anchors are read-only",
            "the final submission clause is compiled only from DeliverableContract",
            "registry and promotion state are outside editor authority",
            "any revision must be rematerialized from scratch",
        ]
    )


class HybridRepairAttemptV1(BaseModel):
    attempt: int
    proposal_id: str
    materialization_dir: str
    decision: Literal["pass", "blocked"]
    reason_codes: List[str] = Field(default_factory=list)
    repair_context_path: str


class HybridRepairLoopReportV1(BaseModel):
    report_version: Literal["v3.hybrid_repair_loop.1"] = "v3.hybrid_repair_loop.1"
    case_id: str
    decision: Literal["pass", "blocked"]
    attempt_count: int
    max_attempts: Literal[2] = 2
    selected_proposal_id: str | None = None
    stop_reason: str
    attempts: List[HybridRepairAttemptV1] = Field(default_factory=list)
    external_editor_calls: Literal[False] = False


class HybridTaskMaterializer:
    SHORTCUT_PATTERN = re.compile(
        r"\b(?:step\s*1|first,\s+open|use the following join key|"
        r"copy the value|the correct answer is|classify .* as (?:pass|fail))\b",
        re.IGNORECASE,
    )

    def materialize(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
        output_dir: str | Path,
    ) -> HybridMaterializationReportV1:
        output_root = Path(output_dir)
        if output_root.exists():
            raise FileExistsError("hybrid_materialization_output_exists")
        candidate_root = output_root / "candidate"
        reference_root = candidate_root / "reference_files"
        teacher_root = output_root / "teacher"
        governance_root = output_root / "governance"
        reference_root.mkdir(parents=True)
        teacher_root.mkdir(parents=True)
        governance_root.mkdir(parents=True)

        frontend = TaskDesignFrontend()
        proposal_validation = frontend.validate_proposal(brief, proposal)
        self._write_json(
            governance_root / "capability_brief.json",
            brief.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "task_design_proposal.json",
            proposal.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "task_design_validation_report.json",
            proposal_validation.model_dump(mode="json"),
        )
        if proposal_validation.decision != "pass":
            report = HybridMaterializationReportV1(
                case_id=brief.case_id,
                brief_id=brief.brief_id,
                proposal_id=proposal.proposal_id,
                decision="blocked",
                reason_codes=["task_design_proposal_invalid"],
                proposal_validation_pass=False,
            )
            self._write_json(
                output_root / "hybrid_materialization_report.json",
                report.model_dump(mode="json"),
            )
            return report

        file_records = self._materialize_evidence_files(
            brief,
            proposal,
            reference_root,
        )
        dossier = self._build_dossier(proposal, file_records)
        self._write_json(
            governance_root / "evidence_dossier_plan.json",
            dossier.model_dump(mode="json"),
        )

        anchors = self._recompute_fact_anchors(proposal, dossier, reference_root)
        self._write_json(
            teacher_root / "deterministic_fact_anchors.json",
            anchors.model_dump(mode="json"),
        )
        rubric_bindings = self._build_rubric_bindings(brief, proposal)
        self._write_json(
            teacher_root / "rubric_binding_plan.json",
            rubric_bindings.model_dump(mode="json"),
        )

        deliverable_spec = self._program_owned_deliverable_spec(brief, proposal)
        contract_compiler = DeliverableContractCompiler()
        deliverable_contract = contract_compiler.build(
            case_id=brief.case_id,
            deliverable_specs=[deliverable_spec],
            reference_files=[
                f"reference_files/{item.file_name}" for item in file_records
            ],
        )
        base_prompt = self._compile_candidate_prompt(brief, proposal)
        compiled_prompt = contract_compiler.compile_prompt(
            base_prompt,
            deliverable_contract,
        )
        contract_validation = DeliverableContractValidator().validate(
            deliverable_contract,
            compiled_prompt,
            [f"reference_files/{item.file_name}" for item in file_records],
        )
        (candidate_root / "prompt.md").write_text(
            compiled_prompt,
            encoding="utf-8",
        )
        self._write_json(
            candidate_root / "deliverable_contract.json",
            deliverable_contract.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "deliverable_contract_validation_report.json",
            contract_validation.model_dump(mode="json"),
        )

        complexity = self._complexity_report(
            brief=brief,
            proposal=proposal,
            dossier=dossier,
            rubric_bindings=rubric_bindings,
            compiled_prompt=compiled_prompt,
            contract_valid=contract_validation.validation_status == "pass",
        )
        self._write_json(
            governance_root / "complexity_preservation_report.json",
            complexity.model_dump(mode="json"),
        )
        content_quality = EvidenceContentQualityValidator().validate(
            proposal=proposal,
            files=file_records,
            reference_root=reference_root,
        )
        self._write_json(
            governance_root / "evidence_content_quality_report.json",
            content_quality.model_dump(mode="json"),
        )
        r5_bundle = ValidityUtilityCompiler().compile(
            brief=brief,
            proposal=proposal,
            proposal_validation=proposal_validation,
            deterministic_fact_anchors=anchors,
            deliverable_contract_valid=(
                contract_validation.validation_status == "pass"
            ),
            evidence_content_quality=(
                content_quality
                if proposal.proposal_version == "v3.task_design_proposal.2"
                else None
            ),
            evidence_paths={
                "proposal": str(
                    governance_root / "task_design_proposal.json"
                ),
                "proposal_validation": str(
                    governance_root / "task_design_validation_report.json"
                ),
                "evidence_dossier": str(
                    governance_root / "evidence_dossier_plan.json"
                ),
                "fact_anchors": str(
                    teacher_root / "deterministic_fact_anchors.json"
                ),
                "evidence_content_quality": str(
                    governance_root / "evidence_content_quality_report.json"
                ),
                "deliverable_contract_validation": str(
                    governance_root
                    / "deliverable_contract_validation_report.json"
                ),
            },
        )
        self._write_r5_bundle(
            r5_bundle=r5_bundle,
            governance_root=governance_root,
            teacher_root=teacher_root,
        )
        visual_validation = self._visual_validation(
            proposal.proposal_id,
            file_records,
            reference_root,
        )
        self._write_json(
            governance_root / "visual_validation_report.json",
            visual_validation.model_dump(mode="json"),
        )
        package_manifest = HybridCandidatePackageManifestV1(
            case_id=brief.case_id,
            proposal_id=proposal.proposal_id,
            prompt_path="prompt.md",
            reference_files=[
                f"reference_files/{item.file_name}" for item in file_records
            ],
            deliverable_contract_path="deliverable_contract.json",
        )
        self._write_json(
            candidate_root / "candidate_package_manifest.json",
            package_manifest.model_dump(mode="json"),
        )
        rw_export_status = self._build_rw_task_export(
            brief=brief,
            proposal=proposal,
            candidate_root=candidate_root,
            output_root=output_root,
            file_records=file_records,
            deliverable_contract=deliverable_contract.model_dump(mode="json"),
            compiled_prompt=compiled_prompt,
            rubric_bindings=rubric_bindings,
            rubric_plan=r5_bundle.rubric_plan,
            validity_overall_status=(
                r5_bundle.validity_vector.overall_status
            ),
            utility_profile_status=(
                r5_bundle.utility_profile.profile_status
            ),
        )
        isolation_pass = not any(
            token in compiled_prompt.lower()
            for token in (
                "deterministic_fact_anchor",
                "expected_value",
                "rubric_binding_plan",
                "teacher/",
            )
        ) and not any(
            "teacher" in path.lower()
            for path in package_manifest.reference_files
        )

        reasons: List[str] = []
        if contract_validation.validation_status != "pass":
            reasons.append("deliverable_contract_invalid")
        if complexity.decision != "pass":
            reasons.append("productive_complexity_not_preserved")
        if r5_bundle.offline_decision != "pass":
            reasons.append("r5_offline_governance_blocked")
        if r5_bundle.rubric_plan.decision != "pass":
            reasons.append("rubric_plan_v2_blocked")
        if visual_validation.decision != "pass":
            reasons.append("candidate_file_visual_validation_failed")
        if (
            proposal.proposal_version == "v3.task_design_proposal.2"
            and content_quality.decision != "pass"
        ):
            reasons.append("candidate_evidence_content_quality_failed")
        if rw_export_status != "draft_compatible":
            reasons.append("rw_task_draft_export_incompatible")
        if not isolation_pass:
            reasons.append("candidate_teacher_isolation_failed")
        if not anchors.all_inputs_candidate_visible:
            reasons.append("fact_anchor_uses_non_candidate_input")
        decision = "pass" if not reasons else "blocked"
        report = HybridMaterializationReportV1(
            case_id=brief.case_id,
            brief_id=brief.brief_id,
            proposal_id=proposal.proposal_id,
            decision=decision,
            reason_codes=reasons,
            proposal_validation_pass=True,
            evidence_file_count=len(file_records),
            deterministic_anchor_count=len(anchors.anchors),
            deliverable_contract_valid=(
                contract_validation.validation_status == "pass"
            ),
            complexity_preservation_pass=complexity.decision == "pass",
            r5_offline_governance_pass=r5_bundle.offline_decision == "pass",
            validity_overall_status=r5_bundle.validity_vector.overall_status,
            utility_profile_status=r5_bundle.utility_profile.profile_status,
            rubric_plan_decision=r5_bundle.rubric_plan.decision,
            evidence_content_quality_decision=content_quality.decision,
            visual_validation_pass=visual_validation.decision == "pass",
            rw_task_export_validation_status=rw_export_status,
            candidate_teacher_isolation_pass=isolation_pass,
            materialization_backend=(
                "hybrid_semantic_artifact_v2"
                if proposal.proposal_version == "v3.task_design_proposal.2"
                else "hybrid_deterministic_fixture_v1"
            ),
        )
        self._write_json(
            output_root / "hybrid_materialization_report.json",
            report.model_dump(mode="json"),
        )
        return report

    def _build_rw_task_export(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
        candidate_root: Path,
        output_root: Path,
        file_records: List[HybridEvidenceFileRecordV1],
        deliverable_contract: Dict[str, Any],
        compiled_prompt: str,
        rubric_bindings: HybridRubricBindingPlanV1,
        rubric_plan: RubricPlanV2,
        validity_overall_status: str,
        utility_profile_status: str,
    ) -> str:
        export_root = output_root / "rw_task_export"
        reference_root = export_root / "reference_files"
        deliverable_root = export_root / "deliverable_files"
        artifacts_root = export_root / "artifacts"
        reference_root.mkdir(parents=True)
        deliverable_root.mkdir(parents=True)
        artifacts_root.mkdir(parents=True)
        for file in file_records:
            shutil.copy2(
                candidate_root / "reference_files" / file.file_name,
                reference_root / file.file_name,
            )
        contract_paths = [
            item["relative_path"] for item in deliverable_contract["deliverables"]
        ]
        rubric_items = [
            {
                "id": item.criterion_id,
                "description": item.observable_behavior,
                "dimension": item.dimension,
                "axis": item.axis,
                "skill_ids": item.skill_ids,
                "capability_ids": item.capability_ids,
                "evidence_or_judgment_ids": item.evidence_or_judgment_ids,
                "independent_failure_signal": item.independent_failure_signal,
                "evidence_requirements": item.evidence_requirements,
                "partial_score_bands": [
                    band.model_dump(mode="json")
                    for band in item.partial_score_bands
                ],
                "professional_evidence_ceiling": (
                    item.professional_evidence_ceiling
                ),
                "weight": item.weight,
            }
            for item in rubric_plan.criteria
        ]
        dataset_row = {
            "task_id": brief.case_id,
            "sector": brief.domain_profile_id,
            "occupation": brief.business_role,
            "motif": brief.motif,
            "prompt": compiled_prompt,
            "reference_files": [
                f"reference_files/{item.file_name}" for item in file_records
            ],
            "deliverable_files": contract_paths,
            "rubric": (
                "Frozen R5 multidimensional rubric. Validity and Utility remain "
                "separate; professional proxy evidence is provisional."
            ),
            "rubric_json": json.dumps(rubric_items, ensure_ascii=False),
            "extra": {
                "export_status": "draft_revise_only",
                "not_final_training_data": True,
                "rw_task_export_ready": False,
                "deliverable_contract_mode": "blocking",
                "deliverable_contract": deliverable_contract,
                "task_design_proposal_id": proposal.proposal_id,
                "capability_brief_id": brief.brief_id,
                "materialization_route": "hybrid_deterministic_fixture_v1",
                "validity_overall_status": validity_overall_status,
                "utility_profile_status": utility_profile_status,
                "rubric_plan_version": rubric_plan.rubric_version,
                "rubric_effective_dimension_count": (
                    rubric_plan.effective_dimension_count
                ),
                "rubric_factual_weight_ratio": (
                    rubric_plan.factual_weight_ratio
                ),
                "behavioral_preflight_required": True,
                "rw_task_exporter": {
                    "export_version": "v3.hybrid_rw_task_export.2",
                    "export_decision": "draft_exported",
                },
            },
        }
        self._write_json(export_root / "dataset_row.json", dataset_row)
        self._write_json(
            export_root / "deliverable_contract.json",
            deliverable_contract,
        )
        self._write_json(
            deliverable_root / "expected_deliverables.json",
            {
                "case_id": brief.case_id,
                "deliverables": contract_paths,
                "deliverable_contract_version": deliverable_contract[
                    "contract_version"
                ],
                "contract_path": "deliverable_contract.json",
            },
        )
        self._write_json(
            export_root / "rw_task_export_report.json",
            {
                "export_version": "v3.hybrid_rw_task_export.2",
                "case_id": brief.case_id,
                "export_decision": "draft_exported",
                "not_final_training_data": True,
                "notes": [
                    "This is an R5 governed draft export.",
                    "Rubric weights are frozen, but behavioral and professional evidence remain pending.",
                    "This export is not training-admission eligible.",
                ],
            },
        )
        validation = RwTaskExportValidator().validate(export_root)
        return validation.validation_status

    def _write_r5_bundle(
        self,
        *,
        r5_bundle: R5GovernanceBundleV1,
        governance_root: Path,
        teacher_root: Path,
    ) -> None:
        self._write_json(
            governance_root / "validity_vector.json",
            r5_bundle.validity_vector.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "utility_profile.json",
            r5_bundle.utility_profile.model_dump(mode="json"),
        )
        self._write_json(
            teacher_root / "rubric_plan_v2.json",
            r5_bundle.rubric_plan.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "evaluation_calibration_contract.json",
            r5_bundle.calibration_contract.model_dump(mode="json"),
        )
        self._write_json(
            governance_root / "r5_governance_bundle.json",
            r5_bundle.model_dump(mode="json"),
        )

    def build_repair_context(
        self,
        materialization_dir: str | Path,
        additional_findings: List[str] | None = None,
    ) -> WholeTaskRepairContextV1:
        root = Path(materialization_dir)
        report = HybridMaterializationReportV1.model_validate_json(
            (root / "hybrid_materialization_report.json").read_text(
                encoding="utf-8"
            )
        )
        candidate_prompt_path = root / "candidate" / "prompt.md"
        candidate_prompt = (
            candidate_prompt_path.read_text(encoding="utf-8")
            if candidate_prompt_path.exists()
            else ""
        )
        dossier = self._read_optional_json(
            root / "governance" / "evidence_dossier_plan.json"
        )
        deliverable_contract = self._read_optional_json(
            root / "candidate" / "deliverable_contract.json"
        )
        fact_anchors = self._read_optional_json(
            root / "teacher" / "deterministic_fact_anchors.json"
        )
        rubric_bindings = self._read_optional_json(
            root / "teacher" / "rubric_binding_plan.json"
        )
        complexity = self._read_optional_json(
            root / "governance" / "complexity_preservation_report.json"
        )
        return WholeTaskRepairContextV1(
            case_id=report.case_id,
            proposal_id=report.proposal_id,
            candidate_prompt=candidate_prompt,
            evidence_dossier=dossier,
            deliverable_contract=deliverable_contract,
            deterministic_fact_anchors=fact_anchors,
            rubric_binding_plan=rubric_bindings,
            complexity_findings=[
                item
                for item in complexity.get("checks", [])
                if not item.get("passed")
            ],
            additional_findings=list(additional_findings or []),
        )

    def run_bounded_repair_loop(
        self,
        brief: CapabilityBriefV1,
        proposals: List[TaskDesignProposalV1],
        output_dir: str | Path,
    ) -> HybridRepairLoopReportV1:
        if not proposals:
            raise ValueError("repair_loop_requires_proposal")
        if len(proposals) > 2:
            raise ValueError("repair_loop_attempt_limit_exceeded")
        root = Path(output_dir)
        if root.exists():
            raise FileExistsError("repair_loop_output_exists")
        root.mkdir(parents=True)
        attempts: List[HybridRepairAttemptV1] = []
        selected_proposal_id = None
        for index, proposal in enumerate(proposals, start=1):
            attempt_root = root / f"attempt_{index:02d}"
            materialization = self.materialize(
                brief,
                proposal,
                attempt_root,
            )
            context = self.build_repair_context(
                attempt_root,
                additional_findings=materialization.reason_codes,
            )
            context_path = attempt_root / "governance" / "whole_task_repair_context.json"
            self._write_json(context_path, context.model_dump(mode="json"))
            attempts.append(
                HybridRepairAttemptV1(
                    attempt=index,
                    proposal_id=proposal.proposal_id,
                    materialization_dir=str(attempt_root),
                    decision=materialization.decision,
                    reason_codes=list(materialization.reason_codes),
                    repair_context_path=str(context_path),
                )
            )
            if materialization.decision == "pass":
                selected_proposal_id = proposal.proposal_id
                break
        decision = "pass" if selected_proposal_id else "blocked"
        stop_reason = (
            "materialization_passed"
            if selected_proposal_id
            else "bounded_repair_attempts_exhausted"
        )
        report = HybridRepairLoopReportV1(
            case_id=brief.case_id,
            decision=decision,
            attempt_count=len(attempts),
            selected_proposal_id=selected_proposal_id,
            stop_reason=stop_reason,
            attempts=attempts,
        )
        self._write_json(
            root / "hybrid_repair_loop_report.json",
            report.model_dump(mode="json"),
        )
        return report

    def _materialize_evidence_files(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
        reference_root: Path,
    ) -> List[HybridEvidenceFileRecordV1]:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        records: List[HybridEvidenceFileRecordV1] = []
        used_names: set[str] = set()
        capabilities_by_skill = {
            item.skill_id: set(item.required_capability_ids)
            for item in brief.selected_skills
        }
        relations_by_id = {
            item.relation_id: item for item in proposal.evidence_relations
        }
        judgments_by_id = {
            item.judgment_id: item for item in proposal.required_judgments
        }
        for node_index, node in enumerate(proposal.evidence_nodes):
            stem = self._slug(node.artifact_role) or self._slug(node.node_id)
            file_name = f"{stem}.xlsx"
            suffix = 2
            while file_name in used_names:
                file_name = f"{stem}_{suffix}.xlsx"
                suffix += 1
            used_names.add(file_name)
            path = reference_root / file_name
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Evidence"
            if node.artifact_spec is not None:
                fields = node.artifact_spec.fields
                sheet.append([field.display_name for field in fields])
                for scenario_record in node.artifact_spec.records:
                    sheet.append(
                        [
                            self._coerce_evidence_value(
                                scenario_record.values.get(field.field_name),
                                field.data_type,
                            )
                            for field in fields
                        ]
                    )
                header_fill = PatternFill("solid", fgColor="1F4E78")
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                for column_index, field in enumerate(fields, start=1):
                    letter = get_column_letter(column_index)
                    max_length = max(
                        len(str(sheet.cell(row=row, column=column_index).value or ""))
                        for row in range(1, sheet.max_row + 1)
                    )
                    format_width_floor = 0.0
                    if field.data_type == "date":
                        # openpyxl reloads formatted dates as midnight datetimes;
                        # size for that deterministic representation as well as
                        # the visible yyyy-mm-dd format.
                        format_width_floor = 21.0
                    elif field.data_type == "datetime":
                        format_width_floor = 24.0
                    sheet.column_dimensions[letter].width = min(
                        60.0,
                        max(12.0, max_length + 2.0, format_width_floor),
                    )
                    if field.data_type == "currency":
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(row=row, column=column_index).number_format = (
                                '$#,##0.00;[Red]($#,##0.00);-'
                            )
                    elif field.data_type == "percentage":
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(row=row, column=column_index).number_format = (
                                "0.0%;[Red](0.0%);-"
                            )
                    elif field.data_type in {"date", "datetime"}:
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(row=row, column=column_index).number_format = (
                                "yyyy-mm-dd"
                                if field.data_type == "date"
                                else "yyyy-mm-dd hh:mm"
                            )
                    if field.data_type == "text" and max_length > 58:
                        for row in range(2, sheet.max_row + 1):
                            sheet.cell(
                                row=row,
                                column=column_index,
                            ).alignment = Alignment(
                                wrap_text=True,
                                vertical="top",
                            )
                for row in range(2, sheet.max_row + 1):
                    required_lines = 1
                    for column_index, field in enumerate(fields, start=1):
                        if field.data_type != "text":
                            continue
                        value = sheet.cell(row=row, column=column_index).value
                        if value is None:
                            continue
                        required_lines = max(
                            required_lines,
                            (len(str(value)) + 57) // 58,
                        )
                    if required_lines > 1:
                        sheet.row_dimensions[row].height = min(
                            90.0,
                            18.0 * required_lines,
                        )
                provenance = workbook.create_sheet("Provenance")
                provenance.append(["Property", "Value"])
                provenance.append(
                    [
                        "Fact Origin",
                        {
                            "governed_scenario_fact": "Governed fictional scenario facts",
                            "public_source_fact": "Public-source facts",
                            "methodological_context": "Methodological context",
                        }.get(
                            node.artifact_spec.fact_origin,
                            node.artifact_spec.fact_origin,
                        ),
                    ]
                )
                provenance.append(
                    [
                        "Methodological Source Refs",
                        ", ".join(
                            node.artifact_spec.methodological_source_ref_ids
                        ),
                    ]
                )
                provenance.append(
                    [
                        "Source Boundary",
                        (
                            "Values are governed fictional scenario facts; source refs "
                            "inform methodology only."
                            if node.artifact_spec.fact_origin
                            == "governed_scenario_fact"
                            else "Values follow the declared artifact fact-origin mode."
                        ),
                    ]
                )
                for cell in provenance[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = header_fill
                provenance.freeze_panes = "A2"
                provenance.auto_filter.ref = f"A1:B{provenance.max_row}"
                provenance["B4"].alignment = Alignment(wrap_text=True, vertical="top")
                provenance.column_dimensions["A"].width = 30
                provenance.column_dimensions["B"].width = 80
                provenance.freeze_panes = "A2"
                provenance.auto_filter.ref = provenance.dimensions
            else:
                sheet.append(
                    [
                        "Evidence_ID",
                        "Record_Key",
                        "Observed_Value",
                        "Observation_Status",
                        "Source_Ref",
                    ]
                )
                source_ref = node.source_ref_ids[0]
                for row_index in range(1, 5):
                    value = row_index * 100
                    if node_index > 0 and row_index == 2:
                        value += node_index * 10
                    if node_index > 1 and row_index == 4:
                        continue
                    sheet.append(
                        [
                            f"{self._slug(node.node_id).upper()}-{row_index:03d}",
                            f"REC-{row_index:03d}",
                            value,
                            "observed",
                            source_ref,
                        ]
                    )
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.page_setup.fitToWidth = 1
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            workbook.save(path)
            linked_skills: set[str] = set()
            linked_capabilities: set[str] = set()
            for binding in proposal.skill_bindings:
                touches_node = False
                for element_id in binding.bound_element_ids:
                    relation = relations_by_id.get(element_id)
                    judgment = judgments_by_id.get(element_id)
                    if relation and node.node_id in {
                        relation.from_node_id,
                        relation.to_node_id,
                    }:
                        touches_node = True
                    if judgment and node.node_id in judgment.input_node_ids:
                        touches_node = True
                if touches_node:
                    linked_skills.add(binding.skill_id)
                    linked_capabilities.update(
                        capabilities_by_skill.get(binding.skill_id, set())
                    )
            records.append(
                HybridEvidenceFileRecordV1(
                    node_id=node.node_id,
                    file_name=file_name,
                    relative_path=f"candidate/reference_files/{file_name}",
                    source_ref_ids=list(node.source_ref_ids),
                    linked_skill_ids=sorted(linked_skills),
                    linked_capability_ids=sorted(linked_capabilities),
                    row_count=sheet.max_row - 1,
                    sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                )
            )
        return records

    @staticmethod
    def _coerce_evidence_value(value: Any, data_type: str) -> Any:
        if value is None:
            return None
        if data_type == "date" and isinstance(value, str):
            try:
                return date.fromisoformat(value)
            except ValueError:
                return value
        if data_type == "datetime" and isinstance(value, str):
            try:
                return datetime.fromisoformat(value)
            except ValueError:
                return value
        return value

    def _visual_validation(
        self,
        proposal_id: str,
        files: List[HybridEvidenceFileRecordV1],
        reference_root: Path,
    ) -> HybridVisualValidationReportV1:
        from openpyxl import load_workbook

        checks: List[HybridFileVisualCheckV1] = []
        for file in files:
            openable = True
            freeze_panes_present = False
            auto_filter_present = False
            fit_to_width = False
            try:
                workbook = load_workbook(
                    reference_root / file.file_name,
                    read_only=False,
                    data_only=False,
                )
                sheet = workbook["Evidence"]
                freeze_panes_present = bool(sheet.freeze_panes)
                auto_filter_present = bool(sheet.auto_filter.ref)
                fit_to_width = sheet.page_setup.fitToWidth == 1
                workbook.close()
            except Exception:
                openable = False
            checks.append(
                HybridFileVisualCheckV1(
                    file_name=file.file_name,
                    openable=openable,
                    freeze_panes_present=freeze_panes_present,
                    auto_filter_present=auto_filter_present,
                    fit_to_width=fit_to_width,
                )
            )
        decision = (
            "pass"
            if checks
            and all(
                item.openable
                and item.freeze_panes_present
                and item.auto_filter_present
                and item.fit_to_width
                for item in checks
            )
            else "blocked"
        )
        return HybridVisualValidationReportV1(
            proposal_id=proposal_id,
            decision=decision,
            checks=checks,
        )

    def _build_dossier(
        self,
        proposal: TaskDesignProposalV1,
        files: List[HybridEvidenceFileRecordV1],
    ) -> HybridEvidenceDossierV1:
        file_by_node = {item.node_id: item for item in files}
        return HybridEvidenceDossierV1(
            proposal_id=proposal.proposal_id,
            files=files,
            relations=[
                HybridEvidenceRelationRecordV1(
                    relation_id=relation.relation_id,
                    from_node_id=relation.from_node_id,
                    to_node_id=relation.to_node_id,
                    relation_type=relation.relation_type,
                    from_file_name=file_by_node[relation.from_node_id].file_name,
                    to_file_name=file_by_node[relation.to_node_id].file_name,
                    solver_must_infer=relation.solver_must_infer,
                )
                for relation in proposal.evidence_relations
            ],
            source_ref_ids=sorted(set(proposal.source_ref_ids)),
        )

    def _recompute_fact_anchors(
        self,
        proposal: TaskDesignProposalV1,
        dossier: HybridEvidenceDossierV1,
        reference_root: Path,
    ) -> DeterministicFactAnchorReportV1:
        from openpyxl import load_workbook

        node_by_id = {item.node_id: item for item in proposal.evidence_nodes}
        rows_by_node: Dict[str, List[Dict[str, Any]]] = {}
        anchors: List[DeterministicFactAnchorV1] = []
        for file in dossier.files:
            workbook = load_workbook(
                reference_root / file.file_name,
                read_only=True,
                data_only=True,
            )
            sheet = workbook["Evidence"]
            node = node_by_id[file.node_id]
            rows: List[Dict[str, Any]] = []
            if node.artifact_spec is not None:
                headers = [
                    str(value or "")
                    for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                ]
                field_by_display = {
                    field.display_name: field.field_name
                    for field in node.artifact_spec.fields
                }
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    rows.append(
                        {
                            field_by_display[headers[index]]: value
                            for index, value in enumerate(row)
                            if index < len(headers)
                            and headers[index] in field_by_display
                        }
                    )
            else:
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    rows.append(
                        {
                            "record_key": str(row[1]),
                            "observed_value": float(row[2]),
                        }
                    )
            workbook.close()
            rows_by_node[file.node_id] = rows
            anchors.append(
                DeterministicFactAnchorV1(
                    anchor_id=f"anchor_rows_{self._slug(file.node_id)}",
                    anchor_type="file_row_count",
                    candidate_visible_inputs=[
                        f"candidate/reference_files/{file.file_name}:Evidence"
                    ],
                    recomputation_method="count non-header rows in the Evidence sheet",
                    expected_value=len(rows),
                )
            )
        for relation in dossier.relations:
            relation_proposal = next(
                item
                for item in proposal.evidence_relations
                if item.relation_id == relation.relation_id
            )
            left_rows = rows_by_node[relation.from_node_id]
            right_rows = rows_by_node[relation.to_node_id]
            if (
                proposal.proposal_version == "v3.task_design_proposal.2"
                and relation_proposal.from_join_field
                and relation_proposal.to_join_field
            ):
                left = {
                    str(row.get(relation_proposal.from_join_field, "")): row
                    for row in left_rows
                    if row.get(relation_proposal.from_join_field) not in {None, ""}
                }
                right = {
                    str(row.get(relation_proposal.to_join_field, "")): row
                    for row in right_rows
                    if row.get(relation_proposal.to_join_field) not in {None, ""}
                }
                shared = sorted(set(left) & set(right))
                mismatch_details: List[Dict[str, Any]] = []
                for key in shared:
                    for comparison in relation_proposal.comparison_fields:
                        left_value = left[key].get(comparison.from_field)
                        right_value = right[key].get(comparison.to_field)
                        if not self._comparison_passes(
                            left_value,
                            right_value,
                            comparison.operator,
                        ):
                            mismatch_details.append(
                                {
                                    "join_key": key,
                                    "from_field": comparison.from_field,
                                    "to_field": comparison.to_field,
                                    "operator": comparison.operator,
                                    "from_value": left_value,
                                    "to_value": right_value,
                                }
                            )
                mismatches: List[Any] = mismatch_details
                recomputation_method = (
                    f"join {relation.from_file_name}.{relation_proposal.from_join_field} "
                    f"to {relation.to_file_name}.{relation_proposal.to_join_field}; "
                    "evaluate the declared field comparisons and separately list missing keys"
                )
            else:
                left = {
                    str(row["record_key"]): row["observed_value"]
                    for row in left_rows
                }
                right = {
                    str(row["record_key"]): row["observed_value"]
                    for row in right_rows
                }
                shared = sorted(set(left) & set(right))
                mismatches = [
                    key for key in shared if abs(left[key] - right[key]) > 1e-9
                ]
                recomputation_method = (
                    "join on Record_Key and compare Observed_Value; separately list "
                    "keys missing from either candidate-visible file"
                )
            mismatch_key_count = (
                len(
                    {
                        str(item.get("join_key"))
                        for item in mismatches
                        if isinstance(item, dict)
                    }
                )
                if mismatches and isinstance(mismatches[0], dict)
                else len(mismatches)
            )
            missing_left = sorted(set(right) - set(left))
            missing_right = sorted(set(left) - set(right))
            linked_judgments = [
                judgment.judgment_id
                for judgment in proposal.required_judgments
                if relation.from_node_id in judgment.input_node_ids
                and relation.to_node_id in judgment.input_node_ids
            ]
            anchors.append(
                DeterministicFactAnchorV1(
                    anchor_id=f"anchor_relation_{self._slug(relation.relation_id)}",
                    anchor_type="relation_crosscheck",
                    candidate_visible_inputs=[
                        f"candidate/reference_files/{relation.from_file_name}:Evidence",
                        f"candidate/reference_files/{relation.to_file_name}:Evidence",
                    ],
                    recomputation_method=recomputation_method,
                    expected_value={
                        "shared_key_count": len(shared),
                        "exact_match_count": len(shared) - mismatch_key_count,
                        "mismatch_keys": mismatches,
                        "missing_from_left": missing_left,
                        "missing_from_right": missing_right,
                    },
                    linked_judgment_ids=linked_judgments,
                )
            )
        return DeterministicFactAnchorReportV1(
            proposal_id=proposal.proposal_id,
            anchors=anchors,
        )

    @staticmethod
    def _comparison_passes(left: Any, right: Any, operator: str) -> bool:
        if operator == "equal":
            return left == right
        if operator == "not_equal":
            return left != right
        try:
            if operator == "less_than":
                return left < right
            if operator == "greater_than":
                return left > right
        except TypeError:
            return False
        return False

    def _build_rubric_bindings(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
    ) -> HybridRubricBindingPlanV1:
        capability_by_skill = {
            item.skill_id: list(item.required_capability_ids)
            for item in brief.selected_skills
        }
        bindings = [
            HybridRubricBindingV1(
                criterion_id=f"criterion_{self._slug(binding.skill_id)}",
                criterion_type=(
                    "reasoning"
                    if "required_judgment" in binding.binding_types
                    else "deliverable"
                ),
                observable_behavior=binding.observable_behavior,
                skill_ids=[binding.skill_id],
                capability_ids=capability_by_skill.get(binding.skill_id, []),
                evidence_or_judgment_ids=list(binding.bound_element_ids),
            )
            for binding in proposal.skill_bindings
        ]
        for judgment in proposal.required_judgments:
            bindings.append(
                HybridRubricBindingV1(
                    criterion_id=f"criterion_{self._slug(judgment.judgment_id)}",
                    criterion_type="fact",
                    observable_behavior=judgment.observable_output,
                    skill_ids=[
                        binding.skill_id
                        for binding in proposal.skill_bindings
                        if judgment.judgment_id in binding.bound_element_ids
                    ],
                    capability_ids=list(judgment.capability_ids),
                    evidence_or_judgment_ids=[judgment.judgment_id],
                )
            )
        return HybridRubricBindingPlanV1(
            proposal_id=proposal.proposal_id,
            bindings=bindings,
        )

    def _program_owned_deliverable_spec(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
    ) -> Dict[str, Any]:
        allowed = [
            item.lower().lstrip(".")
            for item in proposal.deliverable_intent.allowed_formats
            if item.lower().lstrip(".") in brief.allowed_output_file_types
        ]
        if not allowed:
            raise ValueError("no_governed_deliverable_format")
        extension = "xlsx" if "xlsx" in allowed else allowed[0]
        artifact_kind = proposal.deliverable_intent.artifact_kind.lower()
        if "workpaper" in artifact_kind:
            stem = "validation_workpaper"
        elif "memo" in artifact_kind:
            stem = "review_memo"
        else:
            stem = "task_deliverable"
        return {
            "file_name": f"{stem}.{extension}",
            "format": extension,
            "creation_mode": "create",
        }

    def _compile_candidate_prompt(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
    ) -> str:
        lines = [
            proposal.scenario,
            "",
            f"Role: {proposal.actor_role}",
            f"Trigger: {proposal.trigger_event}",
            "",
            "Use the supplied candidate-visible evidence to make and document these judgments:",
        ]
        lines.extend(
            f"- {judgment.description} The observable result should be: "
            f"{judgment.observable_output}"
            for judgment in proposal.required_judgments
        )
        lines.extend(
            [
                "",
                f"Prepare a {proposal.deliverable_intent.artifact_kind} for "
                f"{proposal.deliverable_intent.intended_audience}.",
                f"Business use: {proposal.deliverable_intent.business_use}",
                "Include these sections or views:",
            ]
        )
        lines.extend(
            f"- {section}"
            for section in proposal.deliverable_intent.required_sections_or_views
        )
        lines.extend(
            [
                "",
                "Do not invent missing facts, policies, thresholds, or support. "
                "Keep confirmed exceptions separate from unresolved evidence gaps.",
                "The evidence relationships and professional judgments are for you to determine; "
                "the task does not prescribe a step-by-step workflow.",
            ]
        )
        return "\n".join(lines)

    def _complexity_report(
        self,
        brief: CapabilityBriefV1,
        proposal: TaskDesignProposalV1,
        dossier: HybridEvidenceDossierV1,
        rubric_bindings: HybridRubricBindingPlanV1,
        compiled_prompt: str,
        contract_valid: bool,
    ) -> ComplexityPreservationReportV1:
        dossier_relation_ids = {item.relation_id for item in dossier.relations}
        proposal_relation_ids = {
            item.relation_id for item in proposal.evidence_relations
        }
        prompt_judgment_count = sum(
            judgment.description in compiled_prompt
            for judgment in proposal.required_judgments
        )
        bound_skills = {
            skill_id
            for binding in rubric_bindings.bindings
            for skill_id in binding.skill_ids
        }
        expected_skills = {item.skill_id for item in brief.selected_skills}
        checks = [
            ComplexityPreservationCheckV1(
                check_name="evidence_relations_preserved",
                passed=proposal_relation_ids <= dossier_relation_ids,
                details={
                    "proposal_relation_ids": sorted(proposal_relation_ids),
                    "dossier_relation_ids": sorted(dossier_relation_ids),
                },
            ),
            ComplexityPreservationCheckV1(
                check_name="required_judgments_preserved",
                passed=prompt_judgment_count == len(proposal.required_judgments),
                details={
                    "expected_count": len(proposal.required_judgments),
                    "prompt_count": prompt_judgment_count,
                },
            ),
            ComplexityPreservationCheckV1(
                check_name="skill_bindings_traceable",
                passed=expected_skills <= bound_skills,
                details={
                    "missing_skill_ids": sorted(expected_skills - bound_skills),
                },
            ),
            ComplexityPreservationCheckV1(
                check_name="step_by_step_shortcut_absent",
                passed=not bool(self.SHORTCUT_PATTERN.search(compiled_prompt)),
            ),
            ComplexityPreservationCheckV1(
                check_name="deliverable_contract_valid",
                passed=contract_valid,
            ),
        ]
        decision = "pass" if all(item.passed for item in checks) else "blocked"
        return ComplexityPreservationReportV1(
            proposal_id=proposal.proposal_id,
            decision=decision,
            checks=checks,
            preserved_relation_count=len(
                proposal_relation_ids & dossier_relation_ids
            ),
            preserved_judgment_count=prompt_judgment_count,
            preserved_skill_binding_count=len(expected_skills & bound_skills),
        )

    @staticmethod
    def _slug(value: str) -> str:
        slug = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
        return slug[:48]

    @staticmethod
    def _write_json(path: Path, payload: Dict[str, Any]) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    @staticmethod
    def _read_optional_json(path: Path) -> Dict[str, Any]:
        if not path.exists():
            return {}
        payload = json.loads(path.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else {}
