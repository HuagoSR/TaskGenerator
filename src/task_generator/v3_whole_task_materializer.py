from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path, PurePosixPath
from typing import Any, Dict, Literal

from pydantic import BaseModel, ConfigDict, Field, model_validator

from task_generator.v3_rw_task_export_validator import RwTaskExportValidator
from task_generator.v3_semantic_contract_v2 import (
    FinanceSemanticContractAdapter,
    FinanceSemanticContractResolver,
)


FileAction = Literal["retain", "add", "replace"]
ALLOWED_EXTENSIONS = {".xlsx", ".docx", ".json"}


class CandidateFileRevisionV2(BaseModel):
    model_config = ConfigDict(extra="forbid")
    file_name: str
    change_type: FileAction
    content: Any = None
    content_note: str | None = None

    @model_validator(mode="after")
    def validate_file_contract(self) -> "CandidateFileRevisionV2":
        path = PurePosixPath(self.file_name.replace("\\", "/"))
        if path.is_absolute() or len(path.parts) != 1 or path.name in {"", ".", ".."}:
            raise ValueError("candidate_file_path_not_allowed")
        if Path(path.name).suffix.lower() not in ALLOWED_EXTENSIONS:
            raise ValueError("candidate_file_extension_not_allowed")
        if self.change_type in {"add", "replace"} and self.content is None:
            raise ValueError("candidate_file_content_required")
        if self.change_type == "retain" and self.content is not None:
            raise ValueError("retain_must_not_include_content")
        return self


class WholeTaskRevisionBundleV2(BaseModel):
    model_config = ConfigDict(extra="forbid")
    bundle_version: str = "v3.whole_task_revision_bundle.2"
    task_id: str
    motif: str
    overall_decision: Literal["revise", "already_valid", "cannot_repair"]
    issue_summary: list[str] = Field(default_factory=list, max_length=10)
    revised_prompt: str = Field(min_length=20, max_length=30000)
    candidate_files: list[CandidateFileRevisionV2] = Field(default_factory=list, max_length=12)
    expected_result_advisory: Dict[str, Any] = Field(default_factory=dict)
    calculation_explanation: list[str] = Field(default_factory=list, max_length=16)
    rubric_advisory: list[Dict[str, Any]] = Field(default_factory=list, max_length=12)
    unresolved_questions: list[str] = Field(default_factory=list, max_length=6)

    @model_validator(mode="after")
    def validate_unique_files(self) -> "WholeTaskRevisionBundleV2":
        names = [item.file_name.lower() for item in self.candidate_files]
        if len(names) != len(set(names)):
            raise ValueError("duplicate_candidate_file_revision")
        return self

    @classmethod
    def from_legacy(cls, payload: Dict[str, Any], motif: str) -> "WholeTaskRevisionBundleV2":
        revisions = []
        for item in payload.get("candidate_file_changes") or []:
            revisions.append({
                "file_name": item.get("file_name"),
                "change_type": item.get("change_type"),
                **({"content": item.get("content")} if item.get("change_type") != "retain" else {}),
                "content_note": item.get("content_note"),
            })
        return cls(
            task_id=str(payload["task_id"]), motif=motif,
            overall_decision=payload["overall_decision"],
            issue_summary=list(payload.get("issue_summary") or []),
            revised_prompt=str(payload.get("revised_prompt") or ""),
            candidate_files=revisions,
            expected_result_advisory=dict(payload.get("expected_result") or {}),
            calculation_explanation=list(payload.get("calculation_explanation") or []),
            rubric_advisory=list(payload.get("rubric") or []),
            unresolved_questions=list(payload.get("unresolved_questions") or []),
        )


class MaterializationReport(BaseModel):
    report_version: str = "v3.whole_task_materialization.1"
    task_id: str
    motif: str
    decision: Literal["pass", "revise_system", "blocked"]
    reason_codes: list[str] = Field(default_factory=list)
    materialized_file_count: int = 0
    retained_file_count: int = 0
    deterministic_recomputation_pass: bool = False
    semantic_contract_verified: bool = False
    fact_weight_ratio: float = 0.0
    rw_task_export_compatible: bool = False
    candidate_teacher_isolation_pass: bool = False
    file_sha256: Dict[str, str] = Field(default_factory=dict)


class WholeTaskMaterializer:
    MAX_FILES = 12
    MAX_SHEETS = 8
    MAX_ROWS = 500
    MAX_COLUMNS = 40
    MAX_TEXT_CHARS = 100_000
    MAX_JSON_BYTES = 1_000_000

    def materialize(
        self,
        bundle: WholeTaskRevisionBundleV2,
        source_reference_dir: str | Path,
        output_dir: str | Path,
        blueprint: Dict[str, Any],
        sector: str = "Finance",
        occupation: str = "Financial auditor",
    ) -> MaterializationReport:
        source_root = Path(source_reference_dir)
        output_root = Path(output_dir)
        if not source_root.is_dir():
            raise ValueError("source_reference_dir_missing")
        if output_root.exists():
            raise FileExistsError("revision_output_already_exists")
        reference_root = output_root / "reference_files"
        teacher_root = output_root / "teacher"
        reference_root.mkdir(parents=True)
        teacher_root.mkdir(parents=True)

        source_files = [item for item in source_root.iterdir() if item.is_file()]
        if len(source_files) > self.MAX_FILES:
            raise ValueError("source_reference_file_limit_exceeded")
        for source in source_files:
            if source.suffix.lower() not in ALLOWED_EXTENSIONS:
                continue
            shutil.copy2(source, reference_root / source.name)

        retained = 0
        materialized = 0
        for revision in bundle.candidate_files:
            target = reference_root / revision.file_name
            if revision.change_type == "retain":
                if not target.exists():
                    raise ValueError(f"retained_candidate_file_missing:{revision.file_name}")
                retained += 1
                continue
            if revision.file_name.lower().endswith(".xlsx"):
                self._write_xlsx(target, revision.content)
            elif revision.file_name.lower().endswith(".docx"):
                self._write_docx(target, revision.content)
            else:
                self._write_json(target, revision.content)
            materialized += 1

        if len([item for item in reference_root.iterdir() if item.is_file()]) > self.MAX_FILES:
            raise ValueError("materialized_reference_file_limit_exceeded")
        for workbook_path in reference_root.glob("*.xlsx"):
            self._normalize_xlsx_print_layout(workbook_path)

        adapter = FinanceSemanticContractAdapter()
        contract = adapter.design(bundle.task_id, bundle.motif, blueprint)
        resolver = FinanceSemanticContractResolver()
        resolved = resolver.resolve(contract, reference_root)
        verified, consistency = resolver.verify(resolved)
        expected = {
            claim.claim_id: claim.expected_result.value
            for claim in verified.claims
            if claim.expected_result.resolved
        }
        rubric = self._build_rubric(verified)
        fact_weight = sum(
            float(item.get("weight") or 0)
            for section in rubric["sections"]
            for item in section["criteria"]
            if item.get("criterion_type") in {"fact", "deliverable"}
        )

        (output_root / "prompt.md").write_text(bundle.revised_prompt.strip() + "\n", encoding="utf-8")
        self._atomic_json(teacher_root / "task_semantic_contract.json", verified.model_dump(mode="json"))
        self._atomic_json(teacher_root / "semantic_contract_consistency_report.json", consistency.model_dump(mode="json"))
        self._atomic_json(teacher_root / "deterministic_answer_key.json", expected)
        self._atomic_json(teacher_root / "golden_run.json", {
            "golden_run_version": "v3.whole_task_materialized.1",
            "task_id": bundle.task_id,
            "motif": bundle.motif,
            "deterministic_results": expected,
            "calculation_explanation": bundle.calculation_explanation,
            "truth_source": "deterministic_resolver_from_candidate_visible_files",
        })
        self._atomic_json(teacher_root / "training_annotation.json", {
            "annotation_version": "v3.whole_task_materialized.1",
            "task_id": bundle.task_id,
            "semantic_claims": [item.model_dump(mode="json") for item in verified.claims],
            "teacher_only": True,
        })
        self._atomic_json(teacher_root / "rubric.json", rubric)

        export_dir = output_root / "rw_task_export"
        self._build_rw_task_export(
            export_dir=export_dir,
            task_id=bundle.task_id,
            motif=bundle.motif,
            prompt=bundle.revised_prompt,
            reference_root=reference_root,
            blueprint=blueprint,
            rubric=rubric,
            sector=sector,
            occupation=occupation,
        )
        export_report = RwTaskExportValidator().validate(export_dir)
        isolation = not any(
            token in bundle.revised_prompt.lower()
            for token in ("goldenrun", "golden_run", "answer key", "teacher truth", "training annotation")
        )
        reasons = list(consistency.reason_codes)
        if bundle.unresolved_questions:
            reasons.append("unresolved_revision_questions")
        if fact_weight < 0.60:
            reasons.append("rubric_weight_imbalance")
        if export_report.validation_status != "candidate_ready_compatible":
            reasons.append("rw_task_export_incompatible")
        if not isolation:
            reasons.append("candidate_teacher_isolation_failed")
        decision = "pass" if not reasons else "revise_system"
        hashes = {
            str(path.relative_to(output_root)).replace("\\", "/"): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in output_root.rglob("*") if path.is_file()
        }
        report = MaterializationReport(
            task_id=bundle.task_id,
            motif=bundle.motif,
            decision=decision,
            reason_codes=sorted(set(reasons)),
            materialized_file_count=materialized,
            retained_file_count=retained,
            deterministic_recomputation_pass=all(item.validator.status == "passed" for item in verified.claims),
            semantic_contract_verified=verified.lifecycle == "verified",
            fact_weight_ratio=round(fact_weight, 6),
            rw_task_export_compatible=export_report.validation_status == "candidate_ready_compatible",
            candidate_teacher_isolation_pass=isolation,
            file_sha256=hashes,
        )
        self._atomic_json(output_root / "materialization_report.json", report.model_dump(mode="json"))
        return report

    def _write_xlsx(self, target: Path, content: Any) -> None:
        if not isinstance(content, dict) or not content or len(content) > self.MAX_SHEETS:
            raise ValueError("xlsx_content_must_be_sheet_mapping")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, rows in content.items():
            if not isinstance(sheet_name, str) or not sheet_name.strip() or len(sheet_name) > 31:
                raise ValueError("xlsx_sheet_name_invalid")
            if not isinstance(rows, list) or not rows or len(rows) > self.MAX_ROWS:
                raise ValueError("xlsx_row_limit_or_shape_invalid")
            width = max((len(row) for row in rows if isinstance(row, list)), default=0)
            if width == 0 or width > self.MAX_COLUMNS or any(not isinstance(row, list) or len(row) != width for row in rows):
                raise ValueError("xlsx_column_limit_or_shape_invalid")
            sheet = workbook.create_sheet(sheet_name)
            for row in rows:
                sheet.append(row)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for index in range(1, width + 1):
                values = [sheet.cell(row, index).value for row in range(1, min(sheet.max_row, 80) + 1)]
                sheet.column_dimensions[get_column_letter(index)].width = min(42, max(12, *(len(str(value or "")) + 2 for value in values)))
        workbook.save(target)

    def _write_docx(self, target: Path, content: Any) -> None:
        from docx import Document
        from docx.shared import Inches, Pt
        if isinstance(content, str):
            paragraphs = [line.strip() for line in content.splitlines() if line.strip()]
        elif isinstance(content, dict):
            paragraphs = []
            if content.get("title"):
                paragraphs.append(str(content["title"]))
            for section in content.get("sections") or []:
                paragraphs.append(str(section.get("heading") or "Section"))
                paragraphs.extend(str(item) for item in (section.get("paragraphs") or []))
        else:
            raise ValueError("docx_content_invalid")
        if not paragraphs or sum(len(item) for item in paragraphs) > self.MAX_TEXT_CHARS:
            raise ValueError("docx_content_limit_or_shape_invalid")
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.85)
        for index, text in enumerate(paragraphs):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(14 if index == 0 else 10.5)
            run.bold = index == 0 or bool(_looks_like_heading(text))
        document.save(target)

    def _write_json(self, target: Path, content: Any) -> None:
        encoded = json.dumps(content, ensure_ascii=False, indent=2).encode("utf-8")
        if len(encoded) > self.MAX_JSON_BYTES:
            raise ValueError("json_content_limit_exceeded")
        target.write_bytes(encoded)

    @staticmethod
    def _normalize_xlsx_print_layout(path: Path) -> None:
        """Make every candidate workbook readable as a complete rendered sheet."""
        from openpyxl import load_workbook
        from openpyxl.worksheet.page import PageMargins

        workbook = load_workbook(path)
        for sheet in workbook.worksheets:
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_area = sheet.dimensions
            sheet.page_margins = PageMargins(
                left=0.2, right=0.2, top=0.35, bottom=0.35, header=0.1, footer=0.1
            )
        workbook.save(path)

    def _build_rubric(self, contract) -> Dict[str, Any]:
        claims = {item.claim_id: item for item in contract.claims}
        criteria = []
        for binding in contract.rubric_bindings:
            linked = [claims[item] for item in binding.claim_ids]
            criteria.append({
                "criterion_id": binding.criterion_id,
                "section": "fact_checks",
                "criterion_type": "fact" if binding.criterion_type in {"fact", "deliverable"} else binding.criterion_type,
                "audience": "candidate",
                "export_to_rw_task": True,
                "description": "; ".join(item.description for item in linked),
                "pass_condition": "Deliverable agrees with deterministic candidate-visible result: " + json.dumps(
                    [item.expected_result.value for item in linked], ensure_ascii=False, sort_keys=True
                ),
                "failure_signals": ["semantic_claim_mismatch"],
                "severity": "high",
                "status_hint": "pass",
                "source_ids": list(binding.claim_ids),
                "notes": ["Generated from verified whole-task semantic claims."],
                "weight": binding.weight,
                "semantic_claim_ids": list(binding.claim_ids),
                "evidence_requirements": [{"semantic_claim_id": item} for item in binding.claim_ids],
            })
        return {
            "rubric_version": "v3.whole_task_materialized.1",
            "blueprint_id": contract.task_id,
            "golden_run_id": f"golden_{contract.task_id}",
            "annotation_id": f"annotation_{contract.task_id}",
            "readiness": "rubric_ready",
            "sections": [{"section_name": "fact_checks", "criteria": criteria, "summary": ["Fact-centered rubric from verified semantic claims."]}],
            "unresolved_gaps": [],
            "warning_reason_codes": [],
            "notes": ["No cross-template supervision criteria are retained."],
        }

    def _build_rw_task_export(
        self, export_dir: Path, task_id: str, motif: str, prompt: str, reference_root: Path,
        blueprint: Dict[str, Any], rubric: Dict[str, Any], sector: str, occupation: str,
    ) -> None:
        refs = export_dir / "reference_files"
        deliverables = export_dir / "deliverable_files"
        artifacts = export_dir / "artifacts"
        refs.mkdir(parents=True)
        deliverables.mkdir()
        artifacts.mkdir()
        for path in reference_root.iterdir():
            if path.is_file():
                shutil.copy2(path, refs / path.name)
        deliverable_specs = blueprint.get("deliverable_spec") or [{"file_name": "deliverable.xlsx"}]
        deliverable_files = [f"deliverable_files/{item['file_name']}" for item in deliverable_specs]
        rubric_items = [
            {
                "score": max(1, round(float(item["weight"]) * 10)),
                "criterion": item["description"],
                "required": True,
                "rubric_item_id": f"R_{index:03d}",
                "author_type": "model",
                "tags": ["fact_checks", "fact", "outcome"],
                "read_only": None,
                "form_content": None,
            }
            for index, item in enumerate(rubric["sections"][0]["criteria"], start=1)
        ]
        row = {
            "task_id": task_id,
            "sector": sector,
            "occupation": occupation,
            "motif": motif,
            "prompt": prompt,
            "reference_files": [f"reference_files/{path.name}" for path in sorted(refs.iterdir())],
            "deliverable_files": deliverable_files,
            "rubric": "\n".join(f"- [{item['score']} pts] {item['criterion']}" for item in rubric_items),
            "rubric_json": json.dumps(rubric_items, ensure_ascii=False),
            "extra": {
                "export_status": "candidate_ready_export",
                "not_final_training_data": False,
                "rw_task_export_ready": True,
                "rw_task_exporter": {"export_version": "v3.whole_task_materialized.1", "export_decision": "exported"},
                "experimental_motif": motif == "evidence_to_deliverable",
                "default_promotion_allowed": False,
            },
        }
        self._atomic_json(export_dir / "dataset_row.json", row)
        self._atomic_json(deliverables / "expected_deliverables.json", {"case_id": task_id, "deliverables": deliverable_files})
        self._atomic_json(artifacts / "rubric.json", rubric)
        self._atomic_json(export_dir / "rw_task_export_report.json", {
            "export_version": "v3.whole_task_materialized.1",
            "export_decision": "exported",
            "case_id": task_id,
            "reference_file_count": len(row["reference_files"]),
            "deliverable_file_count": len(deliverable_files),
        })

    @staticmethod
    def _atomic_json(path: Path, payload: Any) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(path.suffix + ".tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(path)


def _looks_like_heading(text: str) -> bool:
    prefix = text.split(maxsplit=1)[0] if text else ""
    return prefix.startswith(("REC-", "AP-", "POL-", "E2D-")) or text in {
        "Purpose", "Applicable Guidance", "Evidence Interpretation", "Decision Rules"
    }
