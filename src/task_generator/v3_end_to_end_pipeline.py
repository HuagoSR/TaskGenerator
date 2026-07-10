from __future__ import annotations

import json
import hashlib
import os
import platform
import shutil
import subprocess
import sys
import time
import zipfile
from importlib import metadata as importlib_metadata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from pydantic import BaseModel, Field

from task_generator.v3_source_schema import load_json_file
from task_generator.v3_skill_registry import SkillRegistryBuilder
from task_generator.v3_domain_profile import DEFAULT_DOMAIN_PROFILE_PATH, load_domain_profile


ROOT = Path(__file__).resolve().parents[2]
TEST_DIR = ROOT / "Test"
DEFAULT_OUTPUT_ROOT = Path(os.environ.get("TASKGEN_OUTPUT_ROOT", str(ROOT / "artifacts" / "end_to_end_runs")))
DEFAULT_REGISTRY_PATH = ROOT / "SkillRegistry" / "v3_skill_registry.json"
DEFAULT_MOTIF_GRAMMAR_PATH = ROOT / "SkillRegistry" / "v3_motif_graph_grammar.experimental.json"
DEFAULT_WORKFLOW_ASSET_PATH = ROOT / "SkillRegistry" / "v3_workflow_archetype_registry.experimental.json"
DEFAULT_ENV_PATH = Path(os.environ.get("TASKGEN_ENV_PATH", str(ROOT.parent / "rw-task" / ".env")))
DEFAULT_RW_TASK_ROOT = Path(os.environ.get("TASKGEN_RW_TASK_ROOT", str(ROOT.parent / "rw-task")))
DEFAULT_DEEPSEEK_KEY_PATH = Path(
    os.environ.get("TASKGEN_DEEPSEEK_KEY_PATH", str(ROOT / "deepseek-key.txt"))
)

PipelineStage = Literal[
    "source_to_skills",
    "registry_prepare",
    "task_generation",
    "production_review",
    "rw_task_eval",
]
SourceMode = Literal["web", "local", "existing", "public_fixture"]
EvalMode = Literal["dry-run", "prepare_only", "execute"]
RegistryMode = Literal["existing", "fresh_scratch", "snapshot_scratch"]
StageState = Literal["pending", "running", "skipped", "completed", "failed", "reused", "invalidated"]
CollectorBackend = Literal["direct", "stirrup"]
ExtractorMode = Literal["none", "mock", "llm"]
ExtractorOutputProfile = Literal["standard", "bounded_smoke", "bounded_production"]
RunAction = Literal["run", "resume", "status", "rerun"]
RunProfile = Literal[
    "custom",
    "public-smoke-offline",
    "public-smoke-llm",
    "local-existing",
    "local-source",
    "web-source",
    "finance-production",
]

STAGE_ORDER: List[PipelineStage] = [
    "source_to_skills",
    "registry_prepare",
    "task_generation",
    "production_review",
    "rw_task_eval",
]


class StageCommandResult(BaseModel):
    label: str
    command: List[str] = Field(default_factory=list)
    returncode: int = 0
    stdout_path: str = ""
    stderr_path: str = ""
    parsed_stdout: Dict[str, Any] = Field(default_factory=dict)
    duration_seconds: float = 0.0


class StageAttempt(BaseModel):
    attempt: int
    started_at: str
    finished_at: Optional[str] = None
    outcome: str = "running"
    failure_reason: Optional[str] = None


class StageStatus(BaseModel):
    stage: PipelineStage
    state: StageState = "pending"
    started_at: Optional[str] = None
    finished_at: Optional[str] = None
    artifact_paths: Dict[str, str] = Field(default_factory=dict)
    command_results: List[StageCommandResult] = Field(default_factory=list)
    warnings: List[str] = Field(default_factory=list)
    blocking_reasons: List[str] = Field(default_factory=list)
    summary: Dict[str, Any] = Field(default_factory=dict)
    dependencies: List[str] = Field(default_factory=list)
    attempt_count: int = 0
    attempts: List[StageAttempt] = Field(default_factory=list)
    input_fingerprint: str = ""
    output_fingerprint: str = ""
    artifact_checksums: Dict[str, str] = Field(default_factory=dict)
    logical_artifact_paths: Dict[str, str] = Field(default_factory=dict)


class StageFailure(RuntimeError):
    def __init__(self, message: str, status: StageStatus) -> None:
        super().__init__(message)
        self.status = status


class EndToEndRequest(BaseModel):
    run_id: str
    selected_stages: List[PipelineStage]
    action: RunAction = "run"
    profile: RunProfile = "custom"
    from_stage: Optional[PipelineStage] = None
    source_mode: SourceMode = "existing"
    registry_mode: Optional[RegistryMode] = None
    registry_path: Optional[str] = None
    topics: List[str] = Field(default_factory=list)
    collection_dirs: List[str] = Field(default_factory=list)
    local_source_dir: Optional[str] = None
    source_limit: int = 3
    collector_backend: CollectorBackend = "direct"
    collector_max_turns: int = 32
    topic_queries_path: Optional[str] = None
    max_cases: int = 4
    apply_registry_update: bool = False
    allow_web_collection: bool = False
    allow_external_upload: bool = False
    allow_external_source_upload: bool = False
    allow_external_eval: bool = False
    extractor_mode: ExtractorMode = "none"
    public_package_path: str = str(ROOT / "Test" / "v3_public_smoke_package" / "skill_extraction_prompt_package.json")
    provider: str = "deepseek"
    model: Optional[str] = None
    deepseek_model: str = "deepseek-v4-flash"
    max_candidates: int = 8
    extractor_max_tokens: int = 6000
    extractor_output_profile: ExtractorOutputProfile = "standard"
    reuse_existing: bool = False
    force_stage: bool = False
    review_spec_path: Optional[str] = None
    eval_mode: EvalMode = "dry-run"
    run_eval: bool = False
    allow_draft_eval: bool = False
    models: List[str] = Field(default_factory=list)
    rw_task_root: str = str(DEFAULT_RW_TASK_ROOT)
    env_path: str = str(DEFAULT_ENV_PATH)
    deepseek_key_path: str = str(DEFAULT_DEEPSEEK_KEY_PATH)
    python_exe: str = sys.executable
    timeout_seconds: int = 0
    generation_seed: int = 0
    domain_profile: str = "finance_audit"
    domain_profile_path: str = str(DEFAULT_DOMAIN_PROFILE_PATH)
    motifs: List[str] = Field(default_factory=list)
    case_index_offset: int = 0
    motif_occurrence_offsets: Dict[str, int] = Field(default_factory=dict)


class ExternalEffectsLedger(BaseModel):
    web_collection: bool = False
    external_source_upload: bool = False
    llm_extraction: bool = False
    eval_preparation: bool = False
    external_eval: bool = False


class EndToEndManifest(BaseModel):
    end_to_end_manifest_version: str = "v3.end_to_end_pipeline.2"
    run_id: str
    created_at: str
    updated_at: str
    request: EndToEndRequest
    run_dir: str
    stage_status_path: str
    summary_report_path: str
    registry_mode: RegistryMode = "existing"
    active_registry_path: str = str(DEFAULT_REGISTRY_PATH)
    canonical_registry_path: str = str(DEFAULT_REGISTRY_PATH)
    scratch_registry_initialized: bool = False
    registry_entry_count_initial: int = 0
    stages: Dict[str, StageStatus] = Field(default_factory=dict)
    notes: List[str] = Field(default_factory=list)
    run_status: str = "pending"
    request_sha256: str = ""
    git_commit: str = "unknown"
    runtime_summary: Dict[str, Any] = Field(default_factory=dict)
    input_fingerprints: Dict[str, str] = Field(default_factory=dict)
    external_effects: ExternalEffectsLedger = Field(default_factory=ExternalEffectsLedger)
    artifact_index: Dict[str, Dict[str, str]] = Field(default_factory=dict)
    content_fingerprint: str = ""
    acceptance_report_path: str = ""
    lifecycle_index_path: str = ""
    training_candidate_index_path: str = ""


class EndToEndPipeline:
    """Report-first central runner for the V3 source-to-eval workflow."""

    def __init__(self, repo_root: str | Path = ROOT) -> None:
        self.root = Path(repo_root)
        self.test_dir = self.root / "Test"

    def run(self, request: EndToEndRequest, output_root: str | Path = DEFAULT_OUTPUT_ROOT) -> EndToEndManifest:
        run_dir = Path(output_root) / request.run_id
        run_dir.mkdir(parents=True, exist_ok=True)
        manifest = self._load_or_create_manifest(request, run_dir)
        if manifest.end_to_end_manifest_version != "v3.end_to_end_pipeline.2":
            raise ValueError("Legacy V1 manifests are status-only; start a new V2 run instead of resuming in place.")
        if manifest.stages and request.action == "run" and not request.reuse_existing:
            raise ValueError("Run id already exists. Use --action resume or --action rerun instead of overwriting it.")
        if request.action in {"resume", "rerun"}:
            if not self._requests_resume_compatible(manifest.request, request):
                raise ValueError("Resume configuration conflicts with the frozen request; start a new run id.")
            request = manifest.request.model_copy(
                update={
                    "action": request.action,
                    "from_stage": request.from_stage,
                    "reuse_existing": True,
                    "force_stage": request.action == "rerun",
                    "timeout_seconds": request.timeout_seconds or manifest.request.timeout_seconds,
                }
            )
        manifest.run_status = "running"
        manifest = self._prepare_registry_state(request, run_dir, manifest)
        selected_stages = self._expanded_stages(request.selected_stages)

        if request.action == "rerun":
            if not request.from_stage:
                raise ValueError("--action rerun requires --from-stage.")
            self._invalidate_from(manifest, request.from_stage)
            selected_stages = STAGE_ORDER[STAGE_ORDER.index(request.from_stage) :]

        for stage in selected_stages:
            existing = manifest.stages.get(stage)
            if existing and existing.state in {"completed", "reused"} and request.reuse_existing and not request.force_stage:
                current_input_fingerprint = self._stage_input_fingerprint(stage, request, manifest)
                if self._stage_artifacts_valid(existing) and existing.input_fingerprint == current_input_fingerprint:
                    existing.state = "reused"
                    existing.warnings = sorted(set(existing.warnings + ["Stage reused after checksum validation."]))
                    self._write_manifest(manifest)
                    self._write_summary(manifest)
                    continue
                existing.state = "invalidated"
                existing.warnings = sorted(set(existing.warnings + ["Completed stage invalidated because input fingerprints or artifact checksums no longer match."]))

            prior_attempts = list(existing.attempts) if existing else []
            attempt_number = (existing.attempt_count if existing else 0) + 1
            running = StageStatus(
                stage=stage,
                state="running",
                started_at=self._now(),
                dependencies=STAGE_ORDER[: STAGE_ORDER.index(stage)],
                attempt_count=attempt_number,
                attempts=prior_attempts + [StageAttempt(attempt=attempt_number, started_at=self._now())],
                input_fingerprint=self._stage_input_fingerprint(stage, request, manifest),
            )
            manifest.stages[stage] = running
            self._write_manifest(manifest)
            self._write_summary(manifest)
            try:
                result = self._run_stage(stage, request, run_dir, manifest)
                result.dependencies = running.dependencies
                result.attempt_count = attempt_number
                result.attempts = running.attempts
                result.input_fingerprint = running.input_fingerprint
                result.artifact_checksums = self._artifact_checksums(result.artifact_paths)
                result.logical_artifact_paths = self._logical_paths(result.artifact_paths, run_dir)
                result.output_fingerprint = self._hash_json(
                    {"summary": result.summary, "checksums": result.artifact_checksums}
                )
                result.attempts[-1].finished_at = self._now()
                result.attempts[-1].outcome = "completed"
                manifest.stages[stage] = result
                if existing and existing.output_fingerprint and existing.output_fingerprint != result.output_fingerprint:
                    next_index = STAGE_ORDER.index(stage) + 1
                    if next_index < len(STAGE_ORDER):
                        self._invalidate_from(manifest, STAGE_ORDER[next_index])
            except StageFailure as exc:
                failed = exc.status
                failed.state = "failed"
                failed.finished_at = self._now()
                failed.dependencies = running.dependencies
                failed.attempt_count = attempt_number
                failed.attempts = running.attempts
                failed.attempts[-1].finished_at = failed.finished_at
                failed.attempts[-1].outcome = "failed"
                failed.attempts[-1].failure_reason = str(exc)
                failed.blocking_reasons = sorted(set(failed.blocking_reasons + [str(exc)]))
                manifest.stages[stage] = failed
                manifest.run_status = "failed"
                self._write_manifest(manifest)
                self._write_summary(manifest)
                raise RuntimeError(str(exc))
            except Exception as exc:
                failed = manifest.stages.get(stage) or running
                failure_report = run_dir / self._stage_dir_name(stage) / "extraction" / "skill_extraction_failure_report.json"
                if failure_report.is_file():
                    failed.artifact_paths["skill_extraction_failure_report"] = str(failure_report)
                    failed.artifact_checksums[str(failure_report)] = self._sha256_file(failure_report)
                    failed.logical_artifact_paths["skill_extraction_failure_report"] = self._logical_path(failure_report, run_dir)
                failed.state = "failed"
                failed.finished_at = self._now()
                failed.attempt_count = attempt_number
                failed.attempts = running.attempts
                failed.attempts[-1].finished_at = failed.finished_at
                failed.attempts[-1].outcome = "failed"
                failed.attempts[-1].failure_reason = str(exc)
                failed.blocking_reasons = sorted(set(failed.blocking_reasons + [str(exc)]))
                manifest.stages[stage] = failed
                manifest.run_status = "failed"
                self._write_manifest(manifest)
                self._write_summary(manifest)
                raise
            self._write_manifest(manifest)
            self._write_summary(manifest)

        manifest.run_status = "completed_with_warnings" if any(s.warnings for s in manifest.stages.values()) else "completed"
        if not self._write_closure_artifacts(manifest):
            manifest.run_status = "failed"
            self._write_manifest(manifest)
            self._write_summary(manifest)
            raise RuntimeError("End-to-end acceptance contract failed; inspect acceptance_report.json.")
        self._write_manifest(manifest)
        self._write_summary(manifest)
        return manifest

    def _run_stage(
        self,
        stage: PipelineStage,
        request: EndToEndRequest,
        run_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        status = StageStatus(stage=stage, state="pending", started_at=self._now())
        stage_dir = run_dir / self._stage_dir_name(stage)
        stage_dir.mkdir(parents=True, exist_ok=True)

        if stage == "source_to_skills":
            status = self._stage_source_to_skills(status, request, stage_dir, manifest)
        elif stage == "registry_prepare":
            status = self._stage_registry_prepare(status, request, stage_dir, manifest)
        elif stage == "task_generation":
            status = self._stage_task_generation(status, request, stage_dir, manifest)
        elif stage == "production_review":
            status = self._stage_production_review(status, request, stage_dir, manifest)
        elif stage == "rw_task_eval":
            status = self._stage_rw_task_eval(status, request, stage_dir, manifest)
        else:
            raise ValueError(f"Unsupported stage: {stage}")

        self._validate_stage_outcome(status, request, manifest)
        status.state = "completed"
        status.finished_at = self._now()
        return status

    def _stage_source_to_skills(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        manifest_registry_path = manifest.active_registry_path
        status.summary["active_registry_path"] = manifest_registry_path
        status.summary["registry_mode"] = manifest.registry_mode
        if request.source_mode == "public_fixture":
            package_path = Path(request.public_package_path)
            if not package_path.exists():
                raise ValueError(f"Public smoke prompt package does not exist: {package_path}")
            return self._run_prompt_package_pipeline(
                status,
                request,
                stage_dir,
                manifest,
                package_path,
                smoke_only=request.extractor_mode == "mock",
            )
        if request.source_mode == "existing":
            status.summary.update(
                {
                    "source_mode": "existing",
                    "note": "Existing SkillRegistry and source artifacts are reused.",
                }
            )
            return status

        if request.source_mode == "local":
            if not request.local_source_dir:
                raise ValueError("--local-source-dir is required when --source-mode local.")
            output_dir = stage_dir / "local_source_package"
            cmd = [
                request.python_exe,
                str(self.test_dir / "run_v3_local_source_to_skill.py"),
                "--input-dir",
                request.local_source_dir,
                "--output-dir",
                str(output_dir),
            ]
            domain_profile = load_domain_profile(request.domain_profile, request.domain_profile_path)
            for domain_tag in domain_profile.domain_tags:
                cmd.extend(["--domain-tag", domain_tag])
            result = self._run_command("local_source_to_skill", cmd, stage_dir, request.timeout_seconds)
            status.command_results.append(result)
            status.artifact_paths["local_prompt_package"] = str(output_dir / "skill_extraction_prompt_package.json")
            if request.extractor_mode in {"mock", "llm"}:
                return self._run_prompt_package_pipeline(
                    status,
                    request,
                    stage_dir,
                    manifest,
                    output_dir / "skill_extraction_prompt_package.json",
                    smoke_only=request.extractor_mode == "mock",
                )
            status.summary.update(
                {
                    "source_mode": "local",
                    "local_prompt_package_only": True,
                    "registry_entry_count_before": self._registry_entry_count(Path(manifest_registry_path)),
                    "registry_entry_count_after": self._registry_entry_count(Path(manifest_registry_path)),
                    "registry_entry_delta_this_run": 0,
                    "note": "Local mode builds a prompt package; it does not perform external LLM extraction by itself.",
                }
            )
            return status

        output_root = stage_dir / "web_source_collections"
        batch_report = stage_dir / "web_source_pipeline_a_batch_report.json"
        registry_path = manifest_registry_path
        cmd = [
            request.python_exe,
            str(self.test_dir / "run_v3_web_source_pipeline_a_batch.py"),
            "--output-root",
            str(output_root),
            "--batch-report-path",
            str(batch_report),
            "--registry-path",
            str(registry_path),
            "--env-path",
            request.env_path,
            "--limit",
            str(request.source_limit),
            "--collector-backend",
            request.collector_backend,
            "--collector-max-turns",
            str(request.collector_max_turns),
            "--build-transition-graph",
            "--provider",
            request.provider,
            "--deepseek-model",
            request.deepseek_model,
            "--max-candidates",
            str(request.max_candidates),
            "--max-tokens",
            str(request.extractor_max_tokens),
            "--output-profile",
            request.extractor_output_profile,
            "--deepseek-key-path",
            request.deepseek_key_path,
        ]
        if request.topic_queries_path:
            cmd.extend(["--topic-queries-path", request.topic_queries_path])
        if request.extractor_output_profile == "bounded_production":
            cmd.extend(["--max-pipeline-attempts", "2"])
        for topic in request.topics:
            cmd.extend(["--topic", topic])
        for collection_dir in request.collection_dirs:
            cmd.extend(["--collection-dir", collection_dir])
        if request.allow_web_collection:
            cmd.append("--allow-web-collection")
            manifest.external_effects.web_collection = True
        if request.allow_external_upload or request.allow_external_source_upload:
            cmd.append("--allow-external-upload")
            manifest.external_effects.external_source_upload = True
            manifest.external_effects.llm_extraction = True
        if request.reuse_existing:
            cmd.append("--reuse-existing")
        if not self._should_write_registry(request, manifest):
            cmd.append("--skip-registry-update")
        result = self._run_command("web_source_pipeline_a_batch", cmd, stage_dir, request.timeout_seconds)
        status.command_results.append(result)
        status.artifact_paths["web_source_batch_report"] = str(batch_report)
        status.artifact_paths["web_source_output_root"] = str(output_root)
        payload = self._safe_load(batch_report)
        totals = payload.get("totals") or {}
        status.summary.update(
            {
                "source_mode": "web",
                "collector_backend": request.collector_backend,
                "collection_dir_count": len(request.collection_dirs),
                "status": payload.get("status"),
                "batch_count": payload.get("batch_count"),
                "successful_batch_count": payload.get("successful_batch_count"),
                "failed_batch_count": payload.get("failed_batch_count"),
                "collection_failed_batch_count": payload.get("collection_failed_batch_count"),
                "candidate_count": totals.get("candidate_count"),
                "accepted_count": totals.get("accepted_count"),
                "revise_count": totals.get("revise_count"),
                "rejected_count": totals.get("rejected_count"),
                "registry_update_mode": payload.get("registry_update_mode"),
                "registry_entry_count_before": payload.get("registry_entry_count_before"),
                "registry_entry_count_after": payload.get("registry_entry_count_after"),
                "registry_entry_delta_this_run": payload.get("registry_entry_delta_this_run"),
            }
        )
        return status

    def _run_prompt_package_pipeline(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
        prompt_package: Path,
        *,
        smoke_only: bool,
    ) -> StageStatus:
        if request.extractor_mode not in {"mock", "llm"}:
            raise ValueError("Prompt-package source mode requires extractor_mode mock or llm.")
        extraction_dir = stage_dir / "extraction"
        review_dir = stage_dir / "review"
        extractor_runner = (
            "run_v3_mock_skill_extractor.py"
            if request.extractor_mode == "mock"
            else "run_v3_llm_skill_extractor.py"
        )
        cmd = [
            request.python_exe,
            str(self.test_dir / extractor_runner),
            "--prompt-package",
            str(prompt_package),
            "--output-dir",
            str(extraction_dir),
            "--max-candidates",
            str(request.max_candidates),
        ]
        if request.extractor_mode == "llm":
            if not (request.allow_external_source_upload or request.allow_external_upload):
                raise ValueError("LLM extraction requires --allow-external-source-upload.")
            cmd.extend(
                [
                    "--max-tokens",
                    str(request.extractor_max_tokens),
                    "--output-profile",
                    request.extractor_output_profile,
                    "--provider",
                    request.provider,
                    "--deepseek-model",
                    request.deepseek_model,
                    "--env-path",
                    request.env_path,
                    "--deepseek-key-path",
                    request.deepseek_key_path,
                    "--allow-external-upload",
                ]
            )
            if request.model:
                cmd.extend(["--model", request.model])
            manifest.external_effects.external_source_upload = True
            manifest.external_effects.llm_extraction = True
        status.command_results.append(
            self._run_command("skill_extractor", cmd, stage_dir, request.timeout_seconds)
        )
        candidate_path = extraction_dir / "extracted_skill_candidates.json"
        status.command_results.append(
            self._run_command(
                "skill_candidate_reviewer",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_skill_candidate_reviewer.py"),
                    "--candidates",
                    str(candidate_path),
                    "--output-dir",
                    str(review_dir),
                ],
                stage_dir,
                request.timeout_seconds,
            )
        )
        accepted_path = review_dir / "accepted_skill_candidates.json"
        registry_report = stage_dir / "scratch_registry_update_report.json"
        status.command_results.append(
            self._run_command(
                "scratch_registry_update",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_skill_registry_update.py"),
                    "--candidates",
                    str(accepted_path),
                    "--registry-path",
                    manifest.active_registry_path,
                    "--report-path",
                    str(registry_report),
                ],
                stage_dir,
                request.timeout_seconds,
            )
        )
        extraction = self._safe_load(extraction_dir / "skill_extraction_report.json")
        review = self._safe_load(review_dir / "skill_candidate_review_report.json")
        update = self._safe_load(registry_report)
        final_count = self._registry_entry_count(Path(manifest.active_registry_path))
        status.artifact_paths.update(
            {
                "prompt_package": str(prompt_package),
                "extraction_dir": str(extraction_dir),
                "review_dir": str(review_dir),
                "extracted_candidates": str(candidate_path),
                "accepted_candidates": str(accepted_path),
                "registry_update_report": str(registry_report),
            }
        )
        status.summary.update(
            {
                "source_mode": request.source_mode,
                "extractor_mode": request.extractor_mode,
                "provider_used": extraction.get("provider_used") or extraction.get("extractor") or request.extractor_mode,
                "smoke_only": smoke_only,
                "candidate_count": extraction.get("candidate_count"),
                "accepted_count": review.get("accepted_count"),
                "revise_count": review.get("revise_count"),
                "rejected_count": review.get("rejected_count"),
                "registry_update_mode": "updated",
                "registry_entry_count_before": manifest.registry_entry_count_initial,
                "registry_entry_count_after": final_count,
                "registry_entry_delta_this_run": final_count - manifest.registry_entry_count_initial,
                "registry_report": update,
            }
        )
        return status

    def _stage_registry_prepare(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        registry_path = manifest.active_registry_path
        audit_report = stage_dir / "v3_skill_registry_audit_report.json"
        readiness_report = stage_dir / "v3_registry_sampling_readiness_report.json"
        seed_report = stage_dir / "v3_pipeline_b_seed_set_report.json"
        transition_report = stage_dir / "v3_skill_transition_graph_report.json"
        composition_report = stage_dir / "v3_composition_readiness_report.json"

        commands = [
            (
                "skill_registry_audit",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_skill_registry_audit.py"),
                    "--registry-path",
                    str(registry_path),
                    "--audit-report-path",
                    str(audit_report),
                ],
            ),
            (
                "registry_sampling_readiness",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_registry_sampling_readiness.py"),
                    "--registry-path",
                    str(registry_path),
                    "--audit-report",
                    str(audit_report),
                    "--no-default-reports",
                    "--output-path",
                    str(readiness_report),
                ],
            ),
            (
                "pipeline_b_seed_set",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_pipeline_b_seed_set.py"),
                    "--registry-path",
                    str(registry_path),
                    "--readiness-report",
                    str(readiness_report),
                    "--output-path",
                    str(seed_report),
                    "--domain-profile",
                    request.domain_profile,
                    "--domain-profile-path",
                    request.domain_profile_path,
                ],
            ),
        ]
        for label, cmd in commands:
            status.command_results.append(self._run_command(label, cmd, stage_dir, request.timeout_seconds))

        candidates = self._find_latest_source_artifact(manifest, "extracted_skill_candidates.json")
        accepted = self._find_latest_source_artifact(manifest, "accepted_skill_candidates.json")
        if candidates:
            cmd = [
                request.python_exe,
                str(self.test_dir / "run_v3_skill_transition_graph.py"),
                "--candidates",
                str(candidates),
                "--registry-path",
                str(registry_path),
                "--readiness-report",
                str(readiness_report),
                "--transition-report-path",
                str(transition_report),
                "--composition-report-path",
                str(composition_report),
            ]
            if accepted:
                cmd.extend(["--accepted-candidates", str(accepted)])
            status.command_results.append(self._run_command("skill_transition_graph", cmd, stage_dir, request.timeout_seconds))
            status.artifact_paths["transition_report"] = str(transition_report)
            status.artifact_paths["composition_report"] = str(composition_report)
        else:
            status.warnings.append("No extracted_skill_candidates.json found from source_to_skills; transition graph was skipped.")

        status.artifact_paths.update(
            {
                "audit_report": str(audit_report),
                "readiness_report": str(readiness_report),
                "seed_report": str(seed_report),
            }
        )
        seed_payload = self._safe_load(seed_report)
        status.summary = {
            "active_registry_path": registry_path,
            "registry_mode": manifest.registry_mode,
            "selected_count": seed_payload.get("selected_count"),
            "selected_readiness_counts": seed_payload.get("selected_readiness_counts"),
            "selected_motif_counts": seed_payload.get("selected_motif_counts"),
        }
        return status

    def _stage_task_generation(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        registry_path = manifest.active_registry_path
        seed_report = self._artifact(manifest, "registry_prepare", "seed_report") or str(Path(registry_path).parent / "v3_pipeline_b_seed_set_report.json")
        cmd = [
            request.python_exe,
            str(self.test_dir / "run_v3_production_batch_runner.py"),
            "--production-batch-id",
            request.run_id,
            "--mode",
            "candidate_run",
            "--max-cases",
            str(request.max_cases),
            "--registry-path",
            registry_path,
            "--seed-report",
            seed_report,
            "--output-root",
            str(stage_dir),
            "--rw-task-root",
            request.rw_task_root,
            "--python-exe",
            request.python_exe,
            "--domain-scope",
            load_domain_profile(request.domain_profile, request.domain_profile_path).domain_scope,
            "--domain-profile",
            request.domain_profile,
            "--domain-profile-path",
            request.domain_profile_path,
            "--case-index-offset",
            str(request.case_index_offset),
        ]
        for motif, offset in sorted(request.motif_occurrence_offsets.items()):
            cmd.extend(["--motif-occurrence-offset", f"{motif}={offset}"])
        domain_profile = load_domain_profile(request.domain_profile, request.domain_profile_path)
        for file_type in domain_profile.allowed_input_file_types:
            cmd.extend(["--file-type", file_type])
        for motif in (request.motifs or domain_profile.allowed_motifs):
            cmd.extend(["--motif", motif])
        motif_grammar_path = self.root / "SkillRegistry" / "v3_motif_graph_grammar.experimental.json"
        workflow_asset_path = self.root / "SkillRegistry" / "v3_workflow_archetype_registry.experimental.json"
        if motif_grammar_path.exists():
            cmd.extend(["--motif-grammar-path", str(motif_grammar_path)])
        if workflow_asset_path.exists():
            cmd.extend(["--workflow-asset-path", str(workflow_asset_path)])
        result = self._run_command("production_batch_runner", cmd, stage_dir, request.timeout_seconds)
        status.command_results.append(result)
        batch_dir = stage_dir / request.run_id
        manifest_path = batch_dir / "production_batch_manifest.json"
        batch_report = batch_dir / "batch" / "pipeline_b_batch_report.json"
        status.artifact_paths.update(
            {
                "production_batch_manifest": str(manifest_path),
                "pipeline_b_batch_report": str(batch_report),
                "production_batch_dir": str(batch_dir),
            }
        )
        payload = self._safe_load(manifest_path)
        status.summary = payload.get("summary") or {}
        status.summary["active_registry_path"] = registry_path
        status.summary["registry_mode"] = manifest.registry_mode
        return status

    def _stage_production_review(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        manifest_path = self._artifact(manifest, "task_generation", "production_batch_manifest")
        if not manifest_path:
            raise ValueError("task_generation must run before production_review.")
        diversity_dir = stage_dir / "production_diversity"
        qa_dir = stage_dir / "production_qa_gate"
        diversity_report = diversity_dir / "production_batch_diversity_report.json"
        qa_report = qa_dir / "production_qa_gate_report.json"

        status.command_results.append(
            self._run_command(
                "production_batch_diversity",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_production_batch_diversity.py"),
                    "--manifest-path",
                    manifest_path,
                    "--output-dir",
                    str(diversity_dir),
                ],
                stage_dir,
                request.timeout_seconds,
            )
        )
        status.command_results.append(
            self._run_command(
                "production_qa_gate",
                [
                    request.python_exe,
                    str(self.test_dir / "run_v3_production_qa_gate.py"),
                    "--manifest-path",
                    manifest_path,
                    "--diversity-report-path",
                    str(diversity_report),
                    "--output-dir",
                    str(qa_dir),
                ],
                stage_dir,
                request.timeout_seconds,
            )
        )

        effective_qa_report = qa_report
        release_manifest: Optional[Path] = None
        if request.review_spec_path:
            review_dir = stage_dir / "production_review" / "applied"
            status.command_results.append(
                self._run_command(
                    "production_promotion_manager",
                    [
                        request.python_exe,
                        str(self.test_dir / "run_v3_production_promotion_manager.py"),
                        "--manifest-path",
                        manifest_path,
                        "--qa-gate-report-path",
                        str(qa_report),
                        "--review-spec-path",
                        request.review_spec_path,
                        "--output-dir",
                        str(review_dir),
                    ],
                    stage_dir,
                    request.timeout_seconds,
                )
            )
            effective_qa_report = review_dir / "reviewed_production_qa_gate_report.json"
            status.artifact_paths["reviewed_qa_report"] = str(effective_qa_report)

            release_id = f"{request.run_id}_release"
            status.command_results.append(
                self._run_command(
                    "dataset_release_packager",
                    [
                        request.python_exe,
                        str(self.test_dir / "run_v3_dataset_release_packager.py"),
                        "--release-id",
                        release_id,
                        "--manifest-path",
                        manifest_path,
                        "--qa-gate-report-path",
                        str(effective_qa_report),
                        "--diversity-report-path",
                        str(diversity_report),
                        "--output-root",
                        str(stage_dir / "releases"),
                    ],
                    stage_dir,
                    request.timeout_seconds,
                )
            )
            release_manifest = stage_dir / "releases" / release_id / "release_manifest.json"
            status.artifact_paths["release_manifest"] = str(release_manifest)
        else:
            status.warnings.append("No review spec was provided; reviewed promotion and strict release packaging were skipped.")

        dashboard_dir = stage_dir / "production_dashboard"
        dashboard_cmd = [
            request.python_exe,
            str(self.test_dir / "run_v3_production_dashboard.py"),
            "--manifest-path",
            manifest_path,
            "--diversity-report-path",
            str(diversity_report),
            "--qa-gate-report-path",
            str(effective_qa_report),
            "--output-dir",
            str(dashboard_dir),
        ]
        if release_manifest and release_manifest.exists():
            dashboard_cmd.extend(["--release-manifest-path", str(release_manifest)])
        status.command_results.append(self._run_command("production_dashboard", dashboard_cmd, stage_dir, request.timeout_seconds))

        status.artifact_paths.update(
            {
                "diversity_report": str(diversity_report),
                "qa_gate_report": str(qa_report),
                "dashboard_report": str(dashboard_dir / "production_dashboard_report.json"),
                "release_readiness_report": str(dashboard_dir / "release_readiness_report.json"),
            }
        )
        readiness = self._safe_load(dashboard_dir / "release_readiness_report.json")
        qa_payload = self._safe_load(effective_qa_report)
        diversity = self._safe_load(diversity_report)
        status.summary = {
            "release_readiness_status": readiness.get("readiness_status"),
            "approved_production_candidate_count": (qa_payload.get("summary") or {}).get("approved_production_candidate_count"),
            "review_required_count": (qa_payload.get("summary") or {}).get("review_required_count"),
            "blocked_count": (qa_payload.get("summary") or {}).get("blocked_count"),
            "diversity_warnings": ((diversity.get("diagnostics") or {}).get("warnings") or []),
        }
        return status

    def _stage_rw_task_eval(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        stage_dir: Path,
        manifest: EndToEndManifest,
    ) -> StageStatus:
        if request.eval_mode == "execute" and not (request.run_eval and request.allow_external_eval):
            raise ValueError("External eval requires eval_mode=execute, --run-eval, and --allow-external-eval.")
        manifest_path = self._artifact(manifest, "task_generation", "production_batch_manifest")
        if not manifest_path:
            raise ValueError("task_generation must run before rw_task_eval.")
        case_dir = self._first_eval_case_dir(Path(manifest_path))
        models = request.models or ["gpt-5.4-pro", "gpt-4o-mini"]
        cmd = [
            request.python_exe,
            str(self.test_dir / "run_v3_eval_orchestrator.py"),
            "--case-dir",
            str(case_dir),
            "--output-dir",
            str(stage_dir / "eval_orchestrator"),
            "--rw-task-root",
            request.rw_task_root,
            "--python-exe",
            request.python_exe,
            "--overwrite",
        ]
        for model in models:
            cmd.extend(["--model", model])
        manifest.external_effects.eval_preparation = True
        if request.eval_mode == "execute":
            cmd.append("--run-eval")
            manifest.external_effects.external_eval = True
            if request.allow_draft_eval:
                cmd.append("--allow-draft-eval")
        status.command_results.append(self._run_command("eval_orchestrator", cmd, stage_dir, request.timeout_seconds))
        report_path = stage_dir / "eval_orchestrator" / "evaluation_orchestration_report.json"
        status.artifact_paths["evaluation_orchestration_report"] = str(report_path)
        payload = self._safe_load(report_path)
        diagnostics = payload.get("diagnostics") or {}
        status.summary = {
            "execution_mode": payload.get("execution_mode"),
            "model_count": diagnostics.get("model_count"),
            "prepared_model_count": diagnostics.get("prepared_model_count"),
            "executed_model_count": diagnostics.get("executed_model_count"),
            "usable_summary_model_count": diagnostics.get("usable_summary_model_count"),
            "model_separation_profile_path": payload.get("model_separation_profile_path"),
        }
        return status

    def _run_command(
        self,
        label: str,
        cmd: List[str],
        stage_dir: Path,
        timeout_seconds: int,
    ) -> StageCommandResult:
        command_dir = stage_dir / "commands"
        command_dir.mkdir(parents=True, exist_ok=True)
        stdout_path = command_dir / f"{label}.stdout.txt"
        stderr_path = command_dir / f"{label}.stderr.txt"
        started = time.monotonic()
        result = subprocess.run(
            cmd,
            cwd=str(self.root),
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            timeout=timeout_seconds or None,
        )
        stdout_path.write_text(result.stdout or "", encoding="utf-8")
        stderr_path.write_text(result.stderr or "", encoding="utf-8")
        parsed = self._parse_json_stdout(result.stdout or "")
        command_result = StageCommandResult(
            label=label,
            command=self._redact_command(cmd),
            returncode=result.returncode,
            stdout_path=str(stdout_path),
            stderr_path=str(stderr_path),
            parsed_stdout=parsed,
            duration_seconds=round(time.monotonic() - started, 6),
        )
        if result.returncode != 0:
            tail = (result.stderr or result.stdout or "")[-1200:]
            raise RuntimeError(f"{label} failed with return code {result.returncode}: {tail}")
        return command_result

    def _write_manifest(self, manifest: EndToEndManifest) -> None:
        manifest.updated_at = self._now()
        manifest.artifact_index = {
            stage: dict(status.logical_artifact_paths)
            for stage, status in manifest.stages.items()
            if status.logical_artifact_paths
        }
        self._atomic_write_json(
            Path(manifest.stage_status_path),
            {stage: status.model_dump(mode="json") for stage, status in manifest.stages.items()},
        )
        manifest_path = Path(manifest.run_dir) / "end_to_end_run_manifest.json"
        self._atomic_write_json(manifest_path, manifest.model_dump(mode="json"))

    def _write_summary(self, manifest: EndToEndManifest) -> None:
        task_summary = self._task_summary(manifest)
        review_summary = self._summary_for(manifest, "production_review")
        eval_summary = self._summary_for(manifest, "rw_task_eval")
        source_summary = self._summary_for(manifest, "source_to_skills")
        payload = {
            "end_to_end_summary_version": "v3.end_to_end_summary.2",
            "run_id": manifest.run_id,
            "profile": manifest.request.profile,
            "run_status": manifest.run_status,
            "updated_at": self._now(),
            "stage_states": {stage: status.state for stage, status in manifest.stages.items()},
            "registry_summary": self._registry_summary(manifest),
            "questions": {
                "reusable_skills_extracted": self._skill_answer(source_summary),
                "tasks_generated": self._task_answer(task_summary),
                "task_diversity": self._diversity_answer(review_summary),
                "model_separation": self._model_separation_answer(eval_summary),
            },
            "source_summary": source_summary,
            "task_generation_summary": task_summary,
            "production_review_summary": review_summary,
            "rw_task_eval_summary": eval_summary,
            "artifact_paths": {
                stage: status.artifact_paths for stage, status in manifest.stages.items()
            },
            "warnings": {
                stage: status.warnings for stage, status in manifest.stages.items() if status.warnings
            },
            "blocking_reasons": {
                stage: status.blocking_reasons for stage, status in manifest.stages.items() if status.blocking_reasons
            },
            "external_effects": manifest.external_effects.model_dump(mode="json"),
            "content_fingerprint": manifest.content_fingerprint,
            "acceptance_report_path": manifest.acceptance_report_path,
            "lifecycle_index_path": manifest.lifecycle_index_path,
            "training_candidate_index_path": manifest.training_candidate_index_path,
            "notes": [
                "This is a report-first end-to-end summary; it does not mutate registry state unless apply_registry_update was explicitly enabled.",
                "External model evaluation remains dry-run unless eval_mode=execute and run_eval are both set.",
            ],
        }
        self._atomic_write_json(Path(manifest.summary_report_path), payload)

    def _load_or_create_manifest(self, request: EndToEndRequest, run_dir: Path) -> EndToEndManifest:
        now = self._now()
        manifest_path = run_dir / "end_to_end_run_manifest.json"
        if manifest_path.exists():
            payload = load_json_file(str(manifest_path))
            if isinstance(payload, dict):
                payload["updated_at"] = now
                payload["run_dir"] = str(run_dir)
                payload["stage_status_path"] = str(run_dir / "stage_status.json")
                payload["summary_report_path"] = str(run_dir / "end_to_end_summary_report.json")
                loaded = EndToEndManifest.model_validate(payload)
                loaded.input_fingerprints.update(self._collect_input_fingerprints(loaded.request))
                loaded.runtime_summary.update(self._runtime_summary(run_dir))
                return loaded
        frozen = request.model_copy(update={"action": "run", "from_stage": None, "reuse_existing": False, "force_stage": False})
        request_sha = self._request_sha(frozen)
        manifest = EndToEndManifest(
            run_id=request.run_id,
            created_at=now,
            updated_at=now,
            request=frozen,
            run_dir=str(run_dir),
            stage_status_path=str(run_dir / "stage_status.json"),
            summary_report_path=str(run_dir / "end_to_end_summary_report.json"),
            registry_mode="existing",
            active_registry_path=str(DEFAULT_REGISTRY_PATH),
            canonical_registry_path=str(DEFAULT_REGISTRY_PATH),
            scratch_registry_initialized=False,
            registry_entry_count_initial=0,
            notes=[
                "The central pipeline reuses existing report-first runners and records their original artifacts.",
                "Use --force-stage to rerun a completed stage, and --reuse-existing to avoid rerunning completed stages.",
            ],
            request_sha256=request_sha,
            git_commit=self._git_commit(),
            runtime_summary=self._runtime_summary(run_dir),
        )
        manifest.input_fingerprints.update(self._collect_input_fingerprints(request))
        return manifest

    def _expanded_stages(self, stages: List[PipelineStage]) -> List[PipelineStage]:
        if not stages:
            return STAGE_ORDER
        furthest = max(STAGE_ORDER.index(stage) for stage in stages)
        return STAGE_ORDER[: furthest + 1]

    def _artifact(self, manifest: EndToEndManifest, stage: str, key: str) -> Optional[str]:
        status = manifest.stages.get(stage)
        if not status:
            return None
        value = status.artifact_paths.get(key)
        return value if value else None

    def _summary_for(self, manifest: EndToEndManifest, stage: str) -> Dict[str, Any]:
        status = manifest.stages.get(stage)
        return dict(status.summary) if status else {}

    def _task_summary(self, manifest: EndToEndManifest) -> Dict[str, Any]:
        summary = self._summary_for(manifest, "task_generation")
        return summary if summary else {}

    def _skill_answer(self, source_summary: Dict[str, Any]) -> Dict[str, Any]:
        mode = source_summary.get("source_mode")
        registry_mode = source_summary.get("registry_mode")
        if registry_mode == "fresh_scratch" and mode == "existing":
            return {"status": "fresh_registry_not_populated", "evidence": source_summary}
        if mode == "existing":
            return {"status": "reused_existing_registry", "evidence": source_summary}
        if mode == "local":
            return {"status": "prompt_package_built_only", "evidence": source_summary}
        delta = source_summary.get("registry_entry_delta_this_run")
        update_mode = source_summary.get("registry_update_mode")
        if update_mode == "updated" and delta not in {None, 0}:
            return {"status": "registry_populated_for_this_run", "evidence": source_summary}
        return {
            "status": "observed" if delta not in {None, 0} else "no_registry_growth_or_update_skipped",
            "evidence": source_summary,
        }

    def _task_answer(self, task_summary: Dict[str, Any]) -> Dict[str, Any]:
        ready = int(task_summary.get("candidate_ready_count") or 0)
        states = task_summary.get("task_state_counts") or {}
        return {
            "status": "candidate_ready_tasks_present" if ready else "no_candidate_ready_tasks_observed",
            "candidate_ready_count": ready,
            "task_state_counts": states,
        }

    def _diversity_answer(self, review_summary: Dict[str, Any]) -> Dict[str, Any]:
        warnings = review_summary.get("diversity_warnings") or []
        return {
            "status": "no_diversity_warnings" if not warnings else "diversity_review_needed",
            "warnings": warnings,
        }

    def _model_separation_answer(self, eval_summary: Dict[str, Any]) -> Dict[str, Any]:
        usable = int(eval_summary.get("usable_summary_model_count") or 0)
        executed = int(eval_summary.get("executed_model_count") or 0)
        return {
            "status": "diagnostic_evidence_available" if usable >= 2 else "not_enough_executed_evidence",
            "executed_model_count": executed,
            "usable_summary_model_count": usable,
            "model_separation_profile_path": eval_summary.get("model_separation_profile_path"),
        }

    def _first_eval_case_dir(self, manifest_path: Path) -> Path:
        payload = self._safe_load(manifest_path)
        for case in payload.get("cases") or []:
            if not isinstance(case, dict):
                continue
            case_dir = Path(str(case.get("case_dir") or ""))
            if case_dir.exists() and (case_dir / "rw_task_export").exists():
                return case_dir
        raise ValueError(f"No eval-ready case_dir found in {manifest_path}")

    def _find_latest_source_artifact(self, manifest: EndToEndManifest, filename: str) -> Optional[Path]:
        source_status = manifest.stages.get("source_to_skills")
        if not source_status:
            return None
        roots = [
            Path(path)
            for key, path in source_status.artifact_paths.items()
            if key.endswith("root") or key.endswith("dir")
        ]
        for root in roots:
            if not root.exists():
                continue
            matches = sorted(root.rglob(filename), key=lambda path: path.stat().st_mtime, reverse=True)
            if matches:
                return matches[0]
        return None

    def _parse_json_stdout(self, stdout: str) -> Dict[str, Any]:
        text = stdout.strip()
        if not text:
            return {}
        try:
            payload = json.loads(text)
            return payload if isinstance(payload, dict) else {}
        except json.JSONDecodeError:
            return {}

    def _safe_load(self, path: str | Path) -> Dict[str, Any]:
        path = Path(path)
        if not path.exists():
            return {}
        payload = load_json_file(str(path))
        return payload if isinstance(payload, dict) else {}

    def _prepare_registry_state(
        self,
        request: EndToEndRequest,
        run_dir: Path,
        manifest: EndToEndManifest,
    ) -> EndToEndManifest:
        resumed_active_path = Path(manifest.active_registry_path) if manifest.active_registry_path else None
        resumed_mode = manifest.registry_mode if resumed_active_path else None
        registry_mode = request.registry_mode or resumed_mode or "existing"
        if registry_mode == "fresh_scratch" and request.registry_path:
            raise ValueError("--registry-path cannot be combined with --registry-mode fresh_scratch.")

        canonical_path = self.root / "SkillRegistry" / "v3_skill_registry.json"
        if canonical_path.exists():
            manifest.input_fingerprints["canonical_registry_before"] = self._sha256_file(canonical_path)

        if registry_mode == "snapshot_scratch":
            if request.registry_path:
                raise ValueError("--registry-path cannot be combined with --registry-mode snapshot_scratch.")
            active_path = run_dir / "00_run_state" / "v3_skill_registry.snapshot.json"
            if not active_path.exists():
                active_path.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(canonical_path, active_path)
            manifest.registry_mode = "snapshot_scratch"
            manifest.active_registry_path = str(active_path)
            manifest.canonical_registry_path = str(canonical_path)
            manifest.scratch_registry_initialized = True
            if not manifest.stages:
                manifest.registry_entry_count_initial = self._registry_entry_count(active_path)
            return manifest

        if registry_mode == "fresh_scratch":
            active_path = run_dir / "00_run_state" / "v3_skill_registry.scratch.json"
            scratch_initialized = self._initialize_scratch_registry(active_path)
            if resumed_mode != "fresh_scratch":
                manifest.registry_entry_count_initial = 0
            manifest.registry_mode = "fresh_scratch"
            manifest.active_registry_path = str(active_path)
            manifest.canonical_registry_path = str(canonical_path)
            manifest.scratch_registry_initialized = manifest.scratch_registry_initialized or scratch_initialized or active_path.exists()
            return manifest

        if resumed_mode == "fresh_scratch" and request.registry_mode is None and request.registry_path is None:
            active_path = Path(manifest.active_registry_path)
            registry_mode = "fresh_scratch"
            manifest.registry_mode = registry_mode
            manifest.active_registry_path = str(active_path)
            manifest.canonical_registry_path = str(canonical_path)
            manifest.scratch_registry_initialized = active_path.exists()
            if manifest.registry_entry_count_initial == 0 and not active_path.exists():
                manifest.registry_entry_count_initial = 0
            return manifest

        active_path = Path(request.registry_path) if request.registry_path else canonical_path
        initial_count = manifest.registry_entry_count_initial
        if active_path != resumed_active_path:
            initial_count = self._registry_entry_count(active_path)
        elif registry_mode == "existing" and initial_count == 0 and active_path.exists():
            initial_count = self._registry_entry_count(active_path)
        elif initial_count == 0 and not manifest.stages:
            initial_count = self._registry_entry_count(active_path)
        manifest.registry_mode = registry_mode
        manifest.active_registry_path = str(active_path)
        manifest.canonical_registry_path = str(canonical_path)
        manifest.scratch_registry_initialized = False
        manifest.registry_entry_count_initial = initial_count
        return manifest

    def _initialize_scratch_registry(self, path: Path) -> bool:
        if path.exists():
            return False
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            json.dumps(
                {
                    "registry_version": "v3.0",
                    "entry_count": 0,
                    "entries": [],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return True

    def _registry_entry_count(self, path: Path) -> int:
        return len(SkillRegistryBuilder().load_registry(path))

    def _registry_summary(self, manifest: EndToEndManifest) -> Dict[str, Any]:
        active_path = Path(manifest.active_registry_path)
        final_count = self._registry_entry_count(active_path) if active_path.exists() else 0
        return {
            "registry_mode": manifest.registry_mode,
            "active_registry_path": manifest.active_registry_path,
            "canonical_registry_path": manifest.canonical_registry_path,
            "scratch_registry_initialized": manifest.scratch_registry_initialized,
            "initial_entry_count": manifest.registry_entry_count_initial,
            "final_entry_count": final_count,
            "entry_delta": final_count - manifest.registry_entry_count_initial,
        }

    def _validate_stage_outcome(
        self,
        status: StageStatus,
        request: EndToEndRequest,
        manifest: EndToEndManifest,
    ) -> None:
        if status.stage == "source_to_skills":
            if manifest.registry_mode == "fresh_scratch" and request.source_mode == "existing":
                status.summary.setdefault("source_mode", "existing")
                status.summary["note"] = "Fresh scratch registry requires web or future local extraction that can populate skills."
                raise StageFailure(
                    "fresh_scratch requires source_to_skills to extract new skills; source-mode existing cannot populate an empty registry.",
                    status,
                )
            if (
                manifest.registry_mode == "fresh_scratch"
                and request.source_mode == "local"
                and request.extractor_mode == "none"
            ):
                raise StageFailure(
                    "fresh_scratch cannot continue from source-mode local yet because local mode only builds a prompt package and does not populate a registry.",
                    status,
                )
            if request.source_mode == "web":
                successful_batch_count = int(status.summary.get("successful_batch_count") or 0)
                batch_count = int(status.summary.get("batch_count") or 0)
                if batch_count > 0 and successful_batch_count == 0:
                    status.warnings.append(
                        "No topic collection finished successfully. Inspect the batch report before retrying."
                    )
                    raise StageFailure(
                        "source_to_skills did not produce any successful web collection or extraction batch, so downstream task generation was stopped.",
                        status,
                    )
                if batch_count > 0 and successful_batch_count < batch_count:
                    status.warnings.append(
                        "Only part of the requested web collection batch succeeded. Downstream stages will continue from the successful collections only."
                    )
            if manifest.registry_mode == "fresh_scratch":
                accepted_count = int(status.summary.get("accepted_count") or 0)
                final_count = self._registry_entry_count(Path(manifest.active_registry_path))
                status.summary["registry_entry_count_after"] = final_count
                if accepted_count == 0 or final_count == 0:
                    raise StageFailure(
                        "No accepted skills were written into the scratch registry, so Pipeline B cannot start from this fresh run.",
                        status,
                    )
        if status.stage == "registry_prepare":
            selected_count = int(status.summary.get("selected_count") or 0)
            if selected_count == 0:
                status.warnings.append(
                    "Seed set is empty. Increase topic/source coverage or inspect extraction and review reports before retrying."
                )
                raise StageFailure(
                    "registry_prepare produced an empty seed set; no skills are available for Pipeline B sampling.",
                    status,
                )

    def _should_write_registry(
        self,
        request: EndToEndRequest,
        manifest: EndToEndManifest,
    ) -> bool:
        if manifest.registry_mode in {"fresh_scratch", "snapshot_scratch"}:
            return True
        return request.apply_registry_update

    def _write_closure_artifacts(self, manifest: EndToEndManifest) -> bool:
        run_dir = Path(manifest.run_dir)
        production_manifest_path = self._artifact(manifest, "task_generation", "production_batch_manifest")
        production = self._safe_load(production_manifest_path) if production_manifest_path else {}
        cases = [case for case in (production.get("cases") or []) if isinstance(case, dict)]
        source = self._summary_for(manifest, "source_to_skills")
        registry = self._summary_for(manifest, "registry_prepare")
        tasks = self._summary_for(manifest, "task_generation")
        review = self._summary_for(manifest, "production_review")
        eval_summary = self._summary_for(manifest, "rw_task_eval")
        metrics = {
            "candidate_count": int(source.get("candidate_count") or 0),
            "accepted_count": int(source.get("accepted_count") or 0),
            "sample_ready_count": int(registry.get("selected_count") or 0),
            "generated_case_count": int(tasks.get("completed_case_count") or 0),
            "candidate_ready_count": int(tasks.get("candidate_ready_count") or 0),
            "verifier_pass_count": sum(case.get("verifier_status") == "pass" for case in cases),
            "export_compatible_count": sum("compatible" in str(case.get("validation_status") or "") for case in cases),
            "qa_blocked_count": int(review.get("blocked_count") or 0),
            "prepared_eval_count": int(eval_summary.get("prepared_model_count") or 0),
            "executed_eval_count": int(eval_summary.get("executed_model_count") or 0),
        }
        requirements: Dict[str, bool] = {"executed_eval_count_is_zero": metrics["executed_eval_count"] == 0}
        contract = self._safe_load(Path(manifest.request.public_package_path).parent / "acceptance_contract.json")
        if manifest.request.profile == "public-smoke-offline":
            expected = contract.get("offline") or {}
            requirements.update(
                {
                    "candidate_count_exact": metrics["candidate_count"] == int(expected.get("candidate_count", 4)),
                    "accepted_count_exact": metrics["accepted_count"] == int(expected.get("accepted_count", 4)),
                    "sample_ready_count_exact": metrics["sample_ready_count"] == int(expected.get("sample_ready_count", 3)),
                    "candidate_ready_exact": metrics["candidate_ready_count"] == int(expected.get("candidate_ready_count", 2)),
                    "verifier_exact": metrics["verifier_pass_count"] == int(expected.get("verifier_pass_count", 2)),
                    "export_exact": metrics["export_compatible_count"] == int(expected.get("export_compatible_count", 2)),
                    "qa_non_blocking": metrics["qa_blocked_count"] == 0,
                    "prepared_eval_exact": metrics["prepared_eval_count"] == 1,
                    "no_external_source_effect": not manifest.external_effects.external_source_upload,
                    "no_external_eval": not manifest.external_effects.external_eval,
                }
            )
        elif manifest.request.profile == "public-smoke-llm":
            expected = contract.get("llm_minimum") or {}
            requirements.update(
                {
                    "candidate_count_min": metrics["candidate_count"] >= int(expected.get("candidate_count", 3)),
                    "accepted_count_min": metrics["accepted_count"] >= int(expected.get("accepted_count", 2)),
                    "sample_ready_count_min": metrics["sample_ready_count"] >= int(expected.get("sample_ready_count", 2)),
                    "candidate_ready_min": metrics["candidate_ready_count"] >= int(expected.get("candidate_ready_count", 1)),
                    "verifier_min": metrics["verifier_pass_count"] >= int(expected.get("verifier_pass_count", 1)),
                    "export_min": metrics["export_compatible_count"] >= int(expected.get("export_compatible_count", 1)),
                    "qa_non_blocking": metrics["qa_blocked_count"] == 0,
                    "llm_extraction_recorded": manifest.external_effects.llm_extraction,
                    "no_external_eval": not manifest.external_effects.external_eval,
                }
            )

        canonical = Path(manifest.canonical_registry_path)
        canonical_after = self._sha256_file(canonical) if canonical.exists() else ""
        canonical_before = manifest.input_fingerprints.get("canonical_registry_before", "")
        requirements["canonical_registry_unchanged"] = bool(canonical_before) and canonical_before == canonical_after

        manifest.content_fingerprint = self._content_fingerprint(manifest, cases)
        acceptance_path = run_dir / "acceptance_report.json"
        acceptance = {
            "acceptance_version": "v3.end_to_end_acceptance.1",
            "run_id": manifest.run_id,
            "profile": manifest.request.profile,
            "passed": all(requirements.values()),
            "metrics": metrics,
            "requirements": requirements,
            "content_fingerprint": manifest.content_fingerprint,
            "external_effects": manifest.external_effects.model_dump(mode="json"),
            "canonical_registry_before_sha256": canonical_before,
            "canonical_registry_after_sha256": canonical_after,
        }
        self._atomic_write_json(acceptance_path, acceptance)
        manifest.acceptance_report_path = str(acceptance_path)

        training_index_path = run_dir / "training_candidate_index.json"
        training_rows = []
        for case in cases:
            if not case.get("training_pool_candidate_eligible"):
                continue
            case_dir = Path(str(case.get("case_dir") or ""))
            training_rows.append(
                {
                    "case_id": case.get("case_id"),
                    "motif": case.get("motif"),
                    "case_dir": self._logical_path(case_dir, run_dir),
                    "training_annotation": self._logical_path(case_dir / "training_annotation" / "training_annotation.json", run_dir),
                    "dataset_row": self._logical_path(case_dir / "rw_task_export" / "dataset_row.json", run_dir),
                    "formal_training_export_ready": False,
                }
            )
        self._atomic_write_json(
            training_index_path,
            {"version": "v3.training_candidate_index.1", "count": len(training_rows), "candidates": training_rows},
        )
        manifest.training_candidate_index_path = str(training_index_path)

        lifecycle_path = run_dir / "lifecycle_index.json"
        lifecycle = self._lifecycle_index(manifest, cases[0] if cases else None)
        self._atomic_write_json(lifecycle_path, lifecycle)
        manifest.lifecycle_index_path = str(lifecycle_path)
        return bool(acceptance["passed"])

    def _lifecycle_index(self, manifest: EndToEndManifest, case: Optional[Dict[str, Any]]) -> Dict[str, Any]:
        run_dir = Path(manifest.run_dir)
        if not case:
            return {"version": "v3.lifecycle_index.1", "run_id": manifest.run_id, "status": "no_task"}
        case_dir = Path(str(case.get("case_dir") or ""))
        paths = {
            "source_evidence": self._artifact(manifest, "source_to_skills", "prompt_package") or "",
            "accepted_skills": self._artifact(manifest, "source_to_skills", "accepted_candidates") or "",
            "subgraph": str(case_dir / "subgraph_sampler" / "pipeline_b_subgraph_report.json"),
            "blueprint": str(case_dir / "prototype" / "draft_task_blueprint.json"),
            "reference_files": str(case_dir / "reference_file_generation" / "reference_files"),
            "golden_run": str(case_dir / "teacher_runner" / "golden_run.json"),
            "training_annotation": str(case_dir / "training_annotation" / "training_annotation.json"),
            "rubric": str(case_dir / "rubric" / "rubric.json"),
            "quality_report": str(case_dir / "quality_gate" / "pipeline_b_quality_report.json"),
            "rw_task_export": str(case_dir / "rw_task_export"),
        }
        return {
            "version": "v3.lifecycle_index.1",
            "run_id": manifest.run_id,
            "case_id": case.get("case_id"),
            "task_state": case.get("task_state"),
            "verifier_status": case.get("verifier_status"),
            "validation_status": case.get("validation_status"),
            "paths": {key: self._logical_path(Path(value), run_dir) if value else "" for key, value in paths.items()},
        }

    def _content_fingerprint(self, manifest: EndToEndManifest, cases: List[Dict[str, Any]]) -> str:
        payload: Dict[str, Any] = {
            "source": {
                key: value
                for key, value in self._summary_for(manifest, "source_to_skills").items()
                if key in {"extractor_mode", "candidate_count", "accepted_count", "revise_count", "rejected_count"}
            },
            "registry": {
                key: value
                for key, value in self._summary_for(manifest, "registry_prepare").items()
                if key in {"selected_count", "selected_readiness_counts", "selected_motif_counts"}
            },
            "cases": [],
        }
        for case in cases:
            case_dir = Path(str(case.get("case_dir") or ""))
            row = self._safe_load(case_dir / "rw_task_export" / "dataset_row.json")
            normalized_row = {
                "prompt": row.get("prompt"),
                "rubric_json": row.get("rubric_json"),
                # Keep the Milestone C canonical representation stable across Windows and Linux.
                "deliverable_files": [
                    str(item).replace("/", "\\") for item in (row.get("deliverable_files") or [])
                ],
                "reference_file_names": sorted(Path(str(item)).name for item in (row.get("reference_files") or [])),
            }
            reference_hashes = []
            for path in sorted((case_dir / "rw_task_export" / "reference_files").glob("*")):
                if path.is_file():
                    reference_hashes.append({"name": path.name, "sha256": self._semantic_file_hash(path)})
            payload["cases"].append(
                {
                    "motif": case.get("motif"),
                    "selected_skill_ids": sorted(case.get("selected_skill_ids") or []),
                    "row": normalized_row,
                    "reference_hashes": reference_hashes,
                }
            )
        return self._hash_json(payload)

    def _invalidate_from(self, manifest: EndToEndManifest, stage: PipelineStage) -> None:
        start = STAGE_ORDER.index(stage)
        for name in STAGE_ORDER[start:]:
            if name in manifest.stages:
                manifest.stages[name].state = "invalidated"
                manifest.stages[name].warnings = sorted(
                    set(manifest.stages[name].warnings + [f"Invalidated by rerun from {stage}."])
                )

    def _stage_artifacts_valid(self, status: StageStatus) -> bool:
        if not status.artifact_checksums:
            return False
        for value, expected in status.artifact_checksums.items():
            path = Path(value)
            if not path.is_file() or self._sha256_file(path) != expected:
                return False
        return True

    def _artifact_checksums(self, paths: Dict[str, str]) -> Dict[str, str]:
        return {
            value: self._sha256_file(Path(value))
            for value in paths.values()
            if value and Path(value).is_file()
        }

    def _logical_paths(self, paths: Dict[str, str], run_dir: Path) -> Dict[str, str]:
        return {key: self._logical_path(Path(value), run_dir) for key, value in paths.items() if value}

    def _logical_path(self, path: Path, run_dir: Path) -> str:
        try:
            return str(path.resolve().relative_to(run_dir.resolve())).replace("\\", "/")
        except (ValueError, OSError):
            try:
                return "repo://" + str(path.resolve().relative_to(self.root.resolve())).replace("\\", "/")
            except (ValueError, OSError):
                return "runtime://" + path.name

    def _stage_input_fingerprint(
        self, stage: PipelineStage, request: EndToEndRequest, manifest: EndToEndManifest
    ) -> str:
        dependencies = {
            name: manifest.stages[name].output_fingerprint
            for name in STAGE_ORDER[: STAGE_ORDER.index(stage)]
            if name in manifest.stages
        }
        return self._hash_json(
            {
                "stage": stage,
                "request_sha256": manifest.request_sha256,
                "dependencies": dependencies,
                "input_fingerprints": manifest.input_fingerprints,
            }
        )

    def _request_sha(self, request: EndToEndRequest) -> str:
        return self._hash_json(request.model_dump(mode="json"))

    def _requests_resume_compatible(self, frozen: EndToEndRequest, supplied: EndToEndRequest) -> bool:
        ignored = {"action", "from_stage", "reuse_existing", "force_stage", "timeout_seconds"}
        left = frozen.model_dump(mode="json")
        right = supplied.model_dump(mode="json")
        for key in ignored:
            left.pop(key, None)
            right.pop(key, None)
        return left == right

    def _collect_input_fingerprints(self, request: EndToEndRequest) -> Dict[str, str]:
        inputs = {
            "public_package": Path(request.public_package_path),
            "public_acceptance_contract": Path(request.public_package_path).parent / "acceptance_contract.json",
            "motif_grammar": self.root / "SkillRegistry" / "v3_motif_graph_grammar.experimental.json",
            "workflow_asset": self.root / "SkillRegistry" / "v3_workflow_archetype_registry.experimental.json",
            "domain_profiles": Path(request.domain_profile_path),
        }
        if request.review_spec_path:
            inputs["review_spec"] = Path(request.review_spec_path)
        return {name: self._sha256_file(path) for name, path in inputs.items() if path.is_file()}

    def _runtime_summary(self, run_dir: Path) -> Dict[str, Any]:
        packages = {}
        for name in ("pydantic", "pandas", "openpyxl", "httpx", "trafilatura", "stirrup", "e2b"):
            try:
                packages[name] = importlib_metadata.version(name)
            except importlib_metadata.PackageNotFoundError:
                packages[name] = "not_installed"
        return {
            "python_version": platform.python_version(),
            "platform": platform.platform(),
            "python_executable": sys.executable,
            "repo_root": str(self.root),
            "run_root": str(run_dir),
            "dependencies": packages,
            "container": {
                "image_id": os.environ.get("TASKGEN_IMAGE_ID", "not_set"),
                "release_id": os.environ.get("TASKGEN_RELEASE_ID", "not_set"),
                "cpu_limit": os.environ.get("TASKGEN_CPU_LIMIT", "not_set"),
                "memory_limit": os.environ.get("TASKGEN_MEMORY_LIMIT", "not_set"),
                "execution_platform": os.environ.get("TASKGEN_EXECUTION_PLATFORM", "host"),
                "rw_task_snapshot": os.environ.get("TASKGEN_RW_TASK_SNAPSHOT", "not_set"),
            },
        }

    def _hash_json(self, payload: Any) -> str:
        text = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return hashlib.sha256(text.encode("utf-8")).hexdigest()

    def _sha256_file(self, path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _semantic_file_hash(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".md", ".txt", ".csv"}:
            normalized = path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n").replace("\r", "\n")
            return hashlib.sha256(normalized.replace("\n", "\r\n").encode("utf-8")).hexdigest()
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=False, read_only=True)
            payload = []
            for sheet in workbook.worksheets:
                rows = []
                for row in sheet.iter_rows():
                    cells = [
                        {"coordinate": cell.coordinate, "value": cell.value, "data_type": cell.data_type}
                        for cell in row
                        if cell.value is not None
                    ]
                    if cells:
                        rows.append(cells)
                payload.append({"title": sheet.title, "rows": rows})
            workbook.close()
            return self._hash_json(payload)
        if suffix == ".docx":
            with zipfile.ZipFile(path) as archive:
                members = [
                    name
                    for name in archive.namelist()
                    if name.startswith("word/") and name.endswith(".xml")
                ]
                payload = {name: archive.read(name).decode("utf-8", errors="replace") for name in sorted(members)}
            return self._hash_json(payload)
        return self._sha256_file(path)

    def _atomic_write_json(self, path: Path, payload: Any) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temporary = path.with_name(path.name + ".tmp")
        temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temporary, path)

    def _git_commit(self) -> str:
        injected = os.environ.get("TASKGEN_SOURCE_COMMIT")
        if injected:
            return injected
        result = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=str(self.root),
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
        )
        return result.stdout.strip() if result.returncode == 0 else "unknown"

    def _redact_command(self, command: List[str]) -> List[str]:
        redacted: List[str] = []
        hide_next = False
        secret_flags = {"--api-key", "--token", "--secret", "--bearer-token"}
        for value in command:
            if hide_next:
                redacted.append("<redacted>")
                hide_next = False
            else:
                redacted.append(value)
                hide_next = value.lower() in secret_flags
        return redacted

    def _stage_dir_name(self, stage: PipelineStage) -> str:
        index = STAGE_ORDER.index(stage) + 1
        return f"{index:02d}_{stage}"

    def _now(self) -> str:
        return datetime.now(timezone.utc).isoformat()
