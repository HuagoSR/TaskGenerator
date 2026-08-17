from __future__ import annotations

from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from task_generator.v3_behavioral_authorization import (
    BehavioralAuthorizationManager,
    BehavioralPreflightPanelMemberV1,
    BehavioralRetainedPanelEvidenceV1,
)
from task_generator.v3_behavioral_preflight_evidence import (
    BehavioralEvidenceIntegrityRecordV1,
    BehavioralEvidenceIntegrityReportV1,
)
from task_generator.v3_behavioral_preflight_execution import BehavioralPreflightExecutor
from task_generator.v3_behavioral_validation import SolverToolPreflight
from task_generator.v3_evaluation_calibration import (
    EvaluationCalibrationCompiler,
    SolverPanelMemberV1,
)
from task_generator.v3_solver_execution_budget import SolverExecutionBudgetV1
from task_generator.v3_solver_panel_admission import SolverPanelAdmissionCompiler
from task_generator.v3_source_fingerprint import governed_source_fingerprint


def _write(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class BehavioralReplacementTest(unittest.TestCase):
    def _passing_member(
        self, root: Path, model: str, stratum: str
    ) -> SolverPanelMemberV1:
        preflight = SolverToolPreflight()
        preflight.create_fixture(root)
        created = root / "deliverable_files" / preflight.EXPECTED_CREATED
        workbook = Workbook()
        workbook.active["A1"] = "created_marker"
        workbook.save(created)
        edited = root / "deliverable_files" / preflight.EXPECTED_EDITED
        shutil.copy2(root / "reference_files" / "copy_template.xlsx", edited)
        workbook = load_workbook(edited)
        workbook["Template"]["B1"] = "edited_marker"
        workbook.save(edited)
        workbook.close()
        report = preflight.inspect(root, model, "replacement-environment")
        report_path = root / "preflight_report.json"
        report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
        return SolverPanelMemberV1(
            solver_model=model,
            stratum=stratum,
            preflight_report_path=str(report_path),
            preflight_report=report,
            provider="test-provider",
            maximum_cost_per_task_usd=1,
        )

    def _prepared(self, root: Path) -> tuple[BehavioralAuthorizationManager, Path, Path]:
        _write(
            root / "campaign_manifest.json",
            {
                "comparison_id": "replacement-comparison",
                "status": "evaluation_ready",
                "assignments": [
                    {"blind_task_id": f"blind-{index}"} for index in range(12)
                ],
            },
        )
        _write(
            root / "governance" / "route_comparison_manifest.json",
            {"environment_contract_id": "replacement-environment"},
        )
        _write(
            root / "governance" / "provider_screening_outcome.json",
            {"decision": "proceed_to_behavioral_evaluation"},
        )
        _write(
            root / "blind_staging" / "governance" / "route_blind_staging_report.json",
            {"decision": "pass", "package_count": 12, "teacher_artifacts_included": False},
        )

        panel_members = []
        integrity_records = []
        retained_report_path = None
        for model, stratum, passing in (
            ("failed-weak", "weak", False),
            ("retained-medium", "medium", True),
            ("failed-strong", "strong", False),
        ):
            case = root / "prior" / model
            preflight = SolverToolPreflight()
            preflight.create_fixture(case)
            if passing:
                created = case / "deliverable_files" / preflight.EXPECTED_CREATED
                created.parent.mkdir(parents=True, exist_ok=True)
                workbook = Workbook()
                workbook.active["A1"] = "created_marker"
                workbook.save(created)
                edited = case / "deliverable_files" / preflight.EXPECTED_EDITED
                shutil.copy2(case / "reference_files" / "copy_template.xlsx", edited)
                workbook = load_workbook(edited)
                workbook["Template"]["B1"] = "edited_marker"
                workbook.save(edited)
                workbook.close()
            report = preflight.inspect(case, model, "replacement-environment")
            report_path = case / "preflight_report.json"
            report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
            if passing:
                retained_report_path = report_path
            panel_members.append(
                SolverPanelMemberV1(
                    solver_model=model,
                    stratum=stratum,
                    preflight_report_path=str(report_path),
                    preflight_report=report,
                    provider="test-provider",
                    maximum_cost_per_task_usd=1,
                )
            )
            integrity_records.append(
                BehavioralEvidenceIntegrityRecordV1(
                    solver_model=model,
                    decision="pass",
                    budget_sha256_valid=True,
                    outcome_sha256_valid=True,
                    stdout_sha256_valid=True,
                    stderr_sha256_valid=True,
                    output_tree_sha256_valid=True,
                    identity_valid=True,
                )
            )
        panel = EvaluationCalibrationCompiler().freeze_solver_panel(
            panel_id="prior-panel", members=panel_members
        )
        panel_path = root / "governance" / "prior_panel.json"
        panel_path.write_text(panel.model_dump_json(indent=2), encoding="utf-8")
        integrity = BehavioralEvidenceIntegrityReportV1(
            comparison_id="replacement-comparison",
            authorization_request_sha256="a" * 64,
            execution_manifest_path=str(root / "prior_manifest.json"),
            execution_manifest_sha256="b" * 64,
            record_count=3,
            records=integrity_records,
            decision="pass",
        )
        integrity_path = root / "governance" / "prior_integrity.json"
        integrity_path.write_text(integrity.model_dump_json(indent=2), encoding="utf-8")
        admission = SolverPanelAdmissionCompiler().compile(
            frozen_panel_path=panel_path,
            evidence_integrity_path=integrity_path,
        )
        admission_path = root / "governance" / "admission.json"
        admission_path.write_text(admission.model_dump_json(indent=2), encoding="utf-8")
        boundary = SolverPanelAdmissionCompiler().compile_replacement_boundary(
            admission_plan_path=admission_path
        )
        boundary_path = root / "governance" / "boundary.json"
        boundary_path.write_text(boundary.model_dump_json(indent=2), encoding="utf-8")
        parity_path = root / "governance" / "parity.json"
        _write(
            parity_path,
            {
                "passed": True,
                "returncode": 0,
                "network_mode": "none",
                "read_only_root": True,
                "provider_credentials_mounted": False,
                "cleanup_returncode": 0,
                "external_model_execution_authorized": False,
                "source_fingerprint": governed_source_fingerprint(ROOT),
            },
        )
        candidates = [
            BehavioralPreflightPanelMemberV1(
                solver_model=f"new-{stratum}",
                stratum=stratum,
                provider="test-provider",
                maximum_cost_usd=1,
            )
            for stratum in ("weak", "strong")
        ]
        budgets = [
            SolverExecutionBudgetV1(
                solver_model=item.solver_model,
                maximum_provider_calls=2,
                maximum_agent_turns=2,
                maximum_request_bytes_per_call=65536,
                maximum_completion_tokens_per_call=2048,
                maximum_provider_input_tokens=131072,
                maximum_provider_output_tokens=4096,
                input_cost_usd_per_million_tokens=1,
                output_cost_usd_per_million_tokens=2,
                maximum_contract_cost_usd=1,
            )
            for item in candidates
        ]
        manager = BehavioralAuthorizationManager(root)
        assert retained_report_path is not None
        manager.write_replacement_request_v3(
            panel_members=candidates,
            retained_evidence=[
                BehavioralRetainedPanelEvidenceV1(
                    solver_model="retained-medium",
                    stratum="medium",
                    preflight_report_path=str(retained_report_path),
                    preflight_report_sha256=_sha(retained_report_path),
                )
            ],
            solver_execution_budgets=budgets,
            pricing_schedule_source="test conservative schedule",
            admission_plan_path=admission_path,
            replacement_boundary_path=boundary_path,
            retained_frozen_panel_path=panel_path,
            retained_evidence_integrity_path=integrity_path,
            container_parity_report_path=parity_path,
            maximum_total_cost_usd=2,
        )
        request_sha = _sha(manager.current_request_path)
        receipt_path = root / "governance" / "replacement_receipt.json"
        manager.compile_replacement_receipt_v3(
            authorization_request_path=(
                manager.immutable_request_root / f"{request_sha}.json"
            ),
            authorization_id="replacement-test",
            authorization_statement="User authorized this exact replacement test request.",
            expires_at=(datetime.now(timezone.utc) + timedelta(days=1)).isoformat(),
            output_path=receipt_path,
        )
        input_root = root / "public_input"
        case = input_root / "solver_tool_preflight"
        fixture = SolverToolPreflight().create_fixture(case)
        _write(
            case / "dataset_row.json",
            {
                "task_id": "solver_tool_preflight",
                "prompt": Path(fixture["prompt_path"]).read_text(encoding="utf-8"),
                "reference_files": ["reference_files/copy_template.xlsx"],
                "deliverable_files": [
                    "deliverable_files/created_workbook.xlsx",
                    "deliverable_files/edited_template.xlsx",
                ],
                "extra": {
                    "tool_only_preflight": True,
                    "business_task": False,
                    "grader_authorized": False,
                },
            },
        )
        return manager, receipt_path, input_root

    def test_dry_run_contains_only_two_replacement_models(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manager, receipt, input_root = self._prepared(root)
            result = BehavioralPreflightExecutor(root).execute(
                authorization_receipt_path=receipt,
                input_root=input_root,
                output_root=root / "outputs",
                real_world_python=sys.executable,
                rw_task_root=root,
            )
            self.assertEqual(result.status, "dry_run_ready")
            self.assertEqual(
                {item.solver_model for item in result.records},
                {"new-weak", "new-strong"},
            )
            self.assertNotIn("retained-medium", [item.solver_model for item in result.records])
            self.assertEqual(manager.check_replacement_v3(receipt).decision, "ready")

    def test_retained_report_mutation_blocks_request_and_execution(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manager, receipt, input_root = self._prepared(root)
            request = json.loads(manager.current_request_path.read_text(encoding="utf-8"))
            Path(request["retained_evidence"][0]["preflight_report_path"]).write_text(
                "{}", encoding="utf-8"
            )
            check = manager.check_replacement_v3(receipt)
            self.assertEqual(check.decision, "blocked")
            with self.assertRaisesRegex(PermissionError, "authorization_not_ready"):
                BehavioralPreflightExecutor(root).execute(
                    authorization_receipt_path=receipt,
                    input_root=input_root,
                    output_root=root / "outputs",
                    real_world_python=sys.executable,
                    rw_task_root=root,
                )

    def test_failed_model_cannot_be_selected_as_replacement(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manager, _, _ = self._prepared(root)
            request = json.loads(manager.current_request_path.read_text(encoding="utf-8"))
            request["panel_members"][0]["solver_model"] = "failed-weak"
            request["solver_execution_budgets"][0]["solver_model"] = "failed-weak"
            manager.current_request_path.write_text(
                json.dumps(request, indent=2) + "\n", encoding="utf-8"
            )
            self.assertEqual(manager.check_replacement_v3().decision, "blocked")

    def test_hash_bound_parity_may_use_repository_level_artifact_root(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            base = Path(temporary)
            manager, _, _ = self._prepared(base / "campaign")
            request = json.loads(manager.current_request_path.read_text(encoding="utf-8"))
            original = Path(request["container_parity_report_path"])
            external = base / "docker_parity" / "parity_report.json"
            external.parent.mkdir(parents=True)
            shutil.copy2(original, external)
            request["container_parity_report_path"] = str(external)
            manager.current_request_path.write_text(
                json.dumps(request, indent=2) + "\n", encoding="utf-8"
            )
            self.assertEqual(
                manager.check_replacement_v3().decision, "authorization_required"
            )

    def test_merge_produces_hash_bound_passing_three_stratum_panel(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            manager, _, _ = self._prepared(root)
            request_path = (
                manager.immutable_request_root
                / f"{_sha(manager.current_request_path)}.json"
            )
            request = json.loads(request_path.read_text(encoding="utf-8"))
            replacements = [
                self._passing_member(root / "replacement" / "new-weak", "new-weak", "weak"),
                self._passing_member(root / "replacement" / "new-strong", "new-strong", "strong"),
            ]
            replacement_integrity = BehavioralEvidenceIntegrityReportV1(
                comparison_id="replacement-comparison",
                authorization_request_sha256=_sha(request_path),
                execution_manifest_path=str(root / "replacement_manifest.json"),
                execution_manifest_sha256="c" * 64,
                record_count=2,
                records=[
                    BehavioralEvidenceIntegrityRecordV1(
                        solver_model=item.solver_model,
                        decision="pass",
                        budget_sha256_valid=True,
                        outcome_sha256_valid=True,
                        stdout_sha256_valid=True,
                        stderr_sha256_valid=True,
                        output_tree_sha256_valid=True,
                        identity_valid=True,
                    )
                    for item in replacements
                ],
                decision="pass",
            )
            replacement_integrity_path = root / "governance" / "replacement_integrity.json"
            replacement_integrity_path.write_text(
                replacement_integrity.model_dump_json(indent=2), encoding="utf-8"
            )
            final_panel_path = root / "governance" / "final_panel.json"
            report = SolverPanelAdmissionCompiler().merge_replacement_evidence(
                admission_plan_path=request["admission_plan_path"],
                replacement_boundary_path=request["replacement_boundary_path"],
                replacement_authorization_request_path=request_path,
                retained_frozen_panel_path=request["retained_frozen_panel_path"],
                retained_evidence_integrity_path=request[
                    "retained_evidence_integrity_path"
                ],
                replacement_evidence_integrity_path=replacement_integrity_path,
                replacement_members=replacements,
                final_panel_id="replacement-final-panel",
                output_panel_path=final_panel_path,
            )
            self.assertEqual(report.decision, "pass")
            self.assertEqual(report.final_panel.decision, "pass")
            self.assertEqual(_sha(final_panel_path), report.final_panel_sha256)
            self.assertEqual(report.retained_models, ["retained-medium"])
            self.assertEqual(
                set(report.replacement_models), {"new-weak", "new-strong"}
            )


if __name__ == "__main__":
    unittest.main()
