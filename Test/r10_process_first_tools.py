"""Concrete no-provider preparation and file checks inside the fixed Linux image."""
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import subprocess
import sys
import zipfile
from email.parser import BytesParser

PINS = ["python-docx==1.2.0", "pypdf==6.10.0", "reportlab==4.4.9", "Pillow==12.3.0"]


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def files(root):
    result = {}
    for p in sorted(root.rglob("*")):
        if p.is_symlink() or (not p.is_dir() and not p.is_file()):
            raise ValueError("unsafe_output_path")
        if p.is_file():
            if p.stat().st_nlink != 1:
                raise ValueError("hardlinked_output_path")
            result[p.relative_to(root).as_posix()] = digest(p)
    return result


def dependencies(root):
    wheels = root / "wheels"
    wheels.mkdir()
    subprocess.run([sys.executable, "-m", "pip", "--isolated", "download", "--disable-pip-version-check", "--index-url", "https://pypi.org/simple", "--only-binary=:all:", "--dest", str(wheels), *PINS], check=True)
    lock, missing = [], []
    for wheel in sorted(wheels.glob("*.whl")):
        with zipfile.ZipFile(wheel) as archive:
            metadata = BytesParser().parsebytes(archive.read(next(n for n in archive.namelist() if n.endswith(".dist-info/METADATA"))))
        name, version = metadata["Name"], metadata["Version"]
        try:
            installed = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            installed = None
        if installed is not None and installed != version:
            raise ValueError("dependency_would_override_image_package:" + name)
        line = f"{name}=={version} --hash=sha256:{digest(wheel)}"
        lock.append({"name": name, "version": version, "wheel": wheel.name, "sha256": digest(wheel), "from_image": installed is not None})
        if installed is None:
            missing.append(line)
    (root / "requirements.lock").write_text("\n".join(missing) + "\n")
    subprocess.run([sys.executable, "-m", "pip", "--isolated", "install", "--disable-pip-version-check", "--no-index", "--no-deps", "--no-compile", "--require-hashes", "--find-links", str(wheels), "--target", str(root / "site"), "-r", str(root / "requirements.lock")], check=True)
    (root / "lock.json").write_text(json.dumps({"python": sys.version, "platform": sys.platform, "pins": PINS, "packages": lock, "site_hashes": files(root / "site")}, indent=2))


def verify(root):
    lock = json.loads((root / "lock.json").read_text())
    assert lock["pins"] == PINS and lock["site_hashes"] == files(root / "site")
    for package in lock["packages"]:
        assert digest(root / "wheels" / package["wheel"]) == package["sha256"]
    print("DEPENDENCIES_VERIFIED")


def smoke(root):
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfReader
    from reportlab.pdfgen.canvas import Canvas
    from PIL import Image
    # This fixture exercises the same Docker read-only bind semantics as sessions.
    before = files(Path("/inputs"))
    try:
        Path("/inputs/should_not_write.txt").write_text("test")
    except OSError as error:
        assert error.errno in (13, 30)
    else:
        raise AssertionError("input_mount_is_writable")
    assert files(Path("/inputs")) == before
    assert not Path("/run/codex-home/auth.json").exists()
    assert not Path("/run/secrets/deepseek_api_key").exists()
    root.mkdir()
    doc = Document()
    doc.add_paragraph("Procurement document smoke check")
    doc.save(root / "memo.docx")
    assert Document(root / "memo.docx").paragraphs[0].text
    book = Workbook()
    book.active.append([2, 3, "=A1+B1"])
    book.save(root / "table.xlsx")
    book = load_workbook(root / "table.xlsx")
    assert book.active["C1"].value == "=A1+B1"
    book.close()
    pdf = Canvas(str(root / "source.pdf"))
    pdf.drawString(50, 750, "Procurement PDF smoke check")
    pdf.save()
    assert "Procurement" in PdfReader(root / "source.pdf").pages[0].extract_text()
    for name in ("memo.docx", "table.xlsx"):
        subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/pilot-lo", "--headless", "--convert-to", "pdf", "--outdir", str(root), str(root / name)], check=True, timeout=60)
        assert (root / Path(name).with_suffix(".pdf")).is_file()
    for name in ("memo", "table", "source"):
        subprocess.run(["pdftoppm", "-f", "1", "-singlefile", "-scale-to", "800", "-png", str(root / f"{name}.pdf"), str(root / name)], check=True, timeout=30)
        with Image.open(root / f"{name}.png") as im:
            im.verify()
    print("OFFLINE_DOCX_XLSX_PDF_CREATE_READ_RENDER_OK")


def validate(root):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    files(root)
    for p in root.rglob("*"):
        if not p.is_file():
            continue
        if p.suffix.lower() not in {".docx", ".xlsx", ".pdf", ".csv", ".txt", ".md", ".json", ".sha256"} or p.name in {"auth.json", ".env", "deepseek-key.txt"}:
            raise ValueError("unexpected_output_type")
        if not p.stat().st_size:
            raise ValueError("empty_required_file")
        if p.suffix == ".docx":
            Document(p)
        elif p.suffix == ".xlsx":
            load_workbook(p).close()
        elif p.suffix == ".pdf":
            assert len(PdfReader(p).pages)
        else:
            p.read_text(encoding="utf-8")
    print("OUTPUT_PATHS_AND_OPENABILITY_OK")


def freeze(root):
    if (root / "candidate").exists() and any((root / "candidate").rglob("*")):
        raise ValueError("candidate_precedes_process_freeze")
    record = root / "hidden/process.md"
    assert record.stat().st_size
    value = digest(record)
    with (root / "hidden/process.sha256").open("x") as handle:
        handle.write(value + "\n")
    print("PROCESS_FROZEN " + value)


if __name__ == "__main__":
    {"dependencies": dependencies, "verify-dependencies": verify, "smoke": smoke,
     "validate-output": validate, "freeze-process": freeze}[sys.argv[1]](Path(sys.argv[2]))
