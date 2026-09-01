"""Run the bounded R10.8B-2 same-scenario compiler-revision experiment.

This is deliberately a small experiment runner, not a new production framework.
It derives two new packages from frozen Bibles and Skills, never edits the four
pilot packages, and records all semantic inputs before a local Codex session.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "Test"))

from r10_local_codex_judge import _redact, _terminate_process_tree, local_codex_version
from task_generator.core.deliverable_contract import DeliverableContractCompiler
from task_generator.core.scenario_first import ProfessionalRuleSetV1, ScenarioBibleV1
from task_generator.planning.scenario_evidence_experiment import ScenarioEvidenceExperiment, ScenarioEvidenceSessionV1
from task_generator.planning.scenario_task_compiler import (
    ScenarioTaskAdmissionValidator,
    ScenarioTaskCompilationPlanV1,
    ScenarioTaskSpecV1,
    TaskCompilationOutputV1,
    compiler_prompt,
    productive_workload_rubric_errors,
    sha256_json,
    tree_sha256,
)
from task_generator.substrate.professional_skills import ProfessionalSkillLoader


SCENARIOS: dict[str, dict[str, str]] = {
    "r10r_audit_revenue_evidence_reliability": {
        "scenario_id": "scenario_audit_company_information_reliability",
        "domain": "audit_compliance",
        "skill_id": "r10.audit-evidence-reliability",
        "file_name": "audit_evidence_workpaper.xlsx",
        "format": "xlsx",
    },
    "r10r_procurement_price_reasonableness": {
        "scenario_id": "scenario_procurement_price_reasonableness",
        "domain": "procurement_operations",
        "skill_id": "r10.procurement-price-reasonableness",
        "file_name": "price_reasonableness_memo.docx",
        "format": "docx",
    },
}
BIBLES = ROOT / "artifacts/r10/r10_2_scenario_bible_batch_20260829_distribution_repair/bibles"
RULE_SETS = ROOT / "data/r10/work_seeds/professional_rule_sets.json"
CATALOG = ROOT / "data/r10/professional_skills/catalog.json"
SKILLS_ROOT = ROOT / ".agents/skills/r10"


def _now() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def _write(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load() -> tuple[dict[str, ScenarioBibleV1], dict[str, ProfessionalRuleSetV1], dict[str, Any]]:
    bibles = {path.stem: ScenarioBibleV1.model_validate_json(path.read_text(encoding="utf-8")) for path in BIBLES.glob("*.json")}
    rules = {item.rule_set_id: item for item in (ProfessionalRuleSetV1.model_validate(value) for value in json.loads(RULE_SETS.read_text(encoding="utf-8")))}
    loader = ProfessionalSkillLoader()
    catalog = loader.load_catalog(CATALOG)
    skills = {
        config["skill_id"]: loader.load_skill(
            next(entry for entry in catalog.entries if entry.skill_id == config["skill_id"] and entry.status == "curated"),
            skills_root=SKILLS_ROOT,
        )
        for config in SCENARIOS.values()
    }
    return bibles, rules, skills


def _source_commit() -> str:
    return subprocess.run(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True, encoding="utf-8", capture_output=True, check=True).stdout.strip()


def _scope(*, bibles: dict[str, ScenarioBibleV1], rules: dict[str, ProfessionalRuleSetV1], skills: dict[str, Any], cli_version: str) -> dict[str, Any]:
    assignments = []
    for task_id, config in SCENARIOS.items():
        bible, skill = bibles[config["scenario_id"]], skills[config["skill_id"]]
        assignments.append({
            "task_id": task_id,
            "scenario_id": bible.scenario_id,
            "domain": config["domain"],
            "bible_sha256": bible.canonical_sha256(),
            "rule_set_sha256": rules[bible.rule_set_id].canonical_sha256(),
            "skill_id": skill.entry.skill_id,
            "skill_sha256": ScenarioEvidenceExperiment.sha256_json(skill.model_dump(mode="json")),
            "deliverable": {"file_name": config["file_name"], "format": config["format"]},
        })
    return {
        "scope_version": "r10.compiler_revision_scope.1",
        "campaign_id": "r10_8b2_compiler_revision",
        "source_commit": _source_commit(),
        "calibration_report_sha256": _sha256_path(ROOT / "artifacts/r10/r10_8b2_gdpval_calibration/report.json"),
        "assignments": assignments,
        "generator": {"model": "gpt-5.6-sol", "codex_cli": cli_version, "public_probe_calls": 1, "private_evidence_calls": 2, "private_compiler_calls": 2, "timeout_seconds": 1800, "project_retry_count": 0},
        "allowed_private_uploads": ["frozen Scenario Bible", "professional rules", "curated factory-side Skill", "candidate bundle derived in this same campaign"],
        "excluded_actions": ["frozen_candidate_mutation", "solver", "judge", "grader", "release", "training", "promotion", "GDPval_content_upload"],
    }


def _codex_command(*, workspace: Path) -> list[str]:
    return [
        "codex", "exec", "--approve-for-me", "--model", "gpt-5.6-sol",
        "-c", "project_doc_max_bytes=0", "-c", "agents.enabled=false",
        "--disable", "plugins", "--disable", "apps", "--disable", "multi_agent", "--disable", "skill_search",
        "--json", "--ephemeral", "--ignore-user-config", "--ignore-rules", "--skip-git-repo-check",
        "-C", str(workspace), "-",
    ]


def _run_codex(*, workspace: Path, prompt: str, timeout_seconds: int = 1800) -> dict[str, Any]:
    # ``codex -C`` resolves relative paths against the child process cwd.  Use
    # one absolute workspace for both, rather than passing the same relative
    # path twice and accidentally nesting it.
    workspace = workspace.resolve()
    workspace.mkdir(parents=True, exist_ok=True)
    started = datetime.now(UTC)
    process = subprocess.Popen(
        _codex_command(workspace=workspace), cwd=workspace, text=True, encoding="utf-8", errors="replace",
        stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        creationflags=getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0),
    )
    timed_out = False
    try:
        stdout, stderr = process.communicate(prompt, timeout=timeout_seconds)
    except subprocess.TimeoutExpired:
        timed_out = True
        _terminate_process_tree(process)
        stdout, stderr = process.communicate(timeout=30)
    (workspace / "codex.jsonl").write_text(_redact(stdout), encoding="utf-8")
    (workspace / "stderr.txt").write_text(_redact(stderr), encoding="utf-8")
    return {
        "returncode": 124 if timed_out else process.returncode,
        "timed_out": timed_out,
        "started_at": started.isoformat(),
        "finished_at": _now(),
        "cli_version": local_codex_version(),
    }


def _stage_evidence(*, workspace: Path, session: ScenarioEvidenceSessionV1, bible: ScenarioBibleV1, rules: ProfessionalRuleSetV1, skill: Any) -> None:
    ScenarioEvidenceExperiment().stage_session(
        workspace=workspace, session=session, bible=bible, rules=rules, skill=skill, productive_workload=True,
    )


def _stage_compiler(*, task_root: Path, bundle: Path, spec: ScenarioTaskSpecV1, bible: ScenarioBibleV1, rules: ProfessionalRuleSetV1, skill: Any) -> None:
    shutil.copytree(bundle / "candidate", task_root / "reference_files")
    shutil.copytree(bundle / "candidate", task_root / "_frozen_candidate")
    teacher = task_root / "teacher"
    teacher.mkdir()
    _write(teacher / "scenario_bible.json", bible.model_dump(mode="json"))
    _write(teacher / "professional_rules.json", rules.model_dump(mode="json"))
    for name in ("evidence_map.json", "scenario_extension.json"):
        shutil.copy2(bundle / "teacher" / name, teacher / name)
    package = teacher / "professional_skill"
    package.mkdir()
    (package / "SKILL.md").write_text(skill.skill_markdown, encoding="utf-8")
    for relative, content in skill.reference_markdown.items():
        target = package / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8")
    _write(task_root / "deliverable_contract.json", spec.deliverable_contract.model_dump(mode="json"))
    (task_root / "TASK.md").write_text(compiler_prompt(spec=spec, productive_workload=True), encoding="utf-8")


def _materialize_task(*, task_root: Path, spec: ScenarioTaskSpecV1, output: TaskCompilationOutputV1) -> None:
    prompt = DeliverableContractCompiler().compile_prompt(output.base_prompt, spec.deliverable_contract)
    (task_root / "TASK.md").write_text(prompt, encoding="utf-8")
    _write(task_root / "teacher" / "teacher_truth.json", {"truth": [item.model_dump(mode="json") for item in output.teacher_truth]})
    _write(task_root / "teacher" / "decision_matrix.json", output.decision_matrix.model_dump(mode="json"))
    _write(task_root / "teacher" / "task_specific_rubric.json", output.task_specific_rubric.model_dump(mode="json"))
    _write(task_root / "teacher" / "task_compilation.json", output.model_dump(mode="json"))
    _write(task_root / "task_manifest.json", {
        "manifest_version": "r10.compiler_revision_task_manifest.1", "task_id": spec.task_id,
        "scenario_id": spec.scenario_id, "parent_bundle_sha256": spec.parent_bundle_sha256,
        "candidate_tree_sha256": spec.candidate_tree_sha256,
        "task_compilation_sha256": output.canonical_sha256(),
    })


def _review_pack(*, task_root: Path, target: Path, output: TaskCompilationOutputV1) -> None:
    target.mkdir(parents=True)
    candidate = target / "candidate_task"
    candidate.mkdir()
    for name in ("TASK.md", "deliverable_contract.json"):
        shutil.copy2(task_root / name, candidate / name)
    shutil.copytree(task_root / "reference_files", candidate / "reference_files")
    text = ["# Teacher Truth 摘要", "", "本文件只用于审题，不属于 candidate package。", ""]
    for point in output.teacher_truth:
        text.extend([f"## {point.decision_id}", point.conclusion, "", "后续行动：", *(f"- {value}" for value in point.required_follow_up_actions), ""])
    (target / "teacher_truth_summary.md").write_text("\n".join(text), encoding="utf-8")


def _xlsx_has_visible_content(path: Path) -> bool:
    """Accept a workbook only when at least one visible cell has content.

    ``max_row`` and ``max_column`` are unreliable for some valid workbooks in
    openpyxl read-only mode, so probe admission reads the cells instead.
    """
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                if any(value is not None and str(value).strip() for value in row):
                    return True
    finally:
        book.close()
    return False


def _probe(root: Path) -> bool:
    workspace = root / "public_probe"
    prompt = "Create a non-empty candidate/public_probe.xlsx and candidate/public_probe.docx. Do not read or write teacher/ files."
    outcome = _run_codex(workspace=workspace, prompt=prompt, timeout_seconds=300)
    _write(workspace / "diagnostics.json", outcome)
    xlsx, docx = workspace / "candidate/public_probe.xlsx", workspace / "candidate/public_probe.docx"
    try:
        xlsx_ok = _xlsx_has_visible_content(xlsx)
        with zipfile.ZipFile(docx) as archive:
            docx_ok = "[Content_Types].xml" in archive.namelist() and "word/document.xml" in archive.namelist()
    except Exception:
        xlsx_ok, docx_ok = False, False
    return outcome["returncode"] == 0 and xlsx_ok and docx_ok


def run(*, output_root: Path, authorized_scope_sha256: str) -> dict[str, Any]:
    bibles, rules, skills = _load()
    scope = _scope(bibles=bibles, rules=rules, skills=skills, cli_version=local_codex_version())
    if authorized_scope_sha256 != sha256_json(scope):
        raise PermissionError("r10_compiler_revision_scope_authorization_mismatch")
    output_root.mkdir(parents=True, exist_ok=True)
    _write(output_root / "campaign_scope.json", scope)
    _write(output_root / "scope_receipt.json", {"scope_sha256": sha256_json(scope), "status": "consumed", "consumed_at": _now()})
    if not _probe(output_root):
        result = {"decision": "incomplete", "first_failure": "public_local_codex_probe_failed"}
        _write(output_root / "result.json", result)
        return result
    experiment = ScenarioEvidenceExperiment()
    reports: list[dict[str, Any]] = []
    staged: dict[str, tuple[Path, ScenarioTaskSpecV1, ScenarioBibleV1, ProfessionalRuleSetV1, Any]] = {}
    for task_id, config in SCENARIOS.items():
        bible, rule_set, skill = bibles[config["scenario_id"]], rules[bibles[config["scenario_id"]].rule_set_id], skills[config["skill_id"]]
        session = ScenarioEvidenceSessionV1(
            session_id=f"{task_id}__productive_evidence", scenario_id=bible.scenario_id, domain=config["domain"], condition="with_skill",
            bible_sha256=bible.canonical_sha256(), rule_set_sha256=rule_set.canonical_sha256(), skill_id=skill.entry.skill_id,
            skill_sha256=experiment.sha256_json(skill.model_dump(mode="json")), image="local-codex", image_sha256=hashlib.sha256(local_codex_version().encode()).hexdigest(),
        )
        evidence_root = output_root / "evidence" / task_id
        _stage_evidence(workspace=evidence_root, session=session, bible=bible, rules=rule_set, skill=skill)
        outcome = _run_codex(workspace=evidence_root, prompt=(evidence_root / "TASK.md").read_text(encoding="utf-8"))
        report = experiment.admit(session=session, workspace=evidence_root, bible=bible)
        if outcome["returncode"] != 0:
            report = report.model_copy(update={"decision": "incomplete", "first_failure": f"local_codex_exit:{outcome['returncode']}"})
        _write(evidence_root / "admission_report.json", report.model_dump(mode="json"))
        if report.decision != "pass":
            result = {"decision": "incomplete", "first_failure": f"evidence_admission_failed:{task_id}", "reports": reports + [report.model_dump(mode="json")]}
            _write(output_root / "result.json", result)
            return result
        contract = DeliverableContractCompiler().build(case_id=task_id, deliverable_specs=[{"file_name": config["file_name"], "format": config["format"], "creation_mode": "create"}])
        spec = ScenarioTaskSpecV1(task_id=task_id, scenario_id=bible.scenario_id, domain=config["domain"], parent_bundle_sha256=tree_sha256(evidence_root), candidate_tree_sha256=tree_sha256(evidence_root / "candidate"), skill_id=skill.entry.skill_id, deliverable_contract=contract)
        staged[task_id] = (evidence_root, spec, bible, rule_set, skill)
        reports.append(report.model_dump(mode="json"))
    plan = ScenarioTaskCompilationPlanV1(tasks=[item[1] for item in staged.values()])
    _write(output_root / "compilation_plan.json", plan.model_dump(mode="json"))
    task_reports = []
    for task_id, (bundle, spec, bible, rule_set, skill) in staged.items():
        task_root = output_root / "tasks" / task_id
        _stage_compiler(task_root=task_root, bundle=bundle, spec=spec, bible=bible, rules=rule_set, skill=skill)
        outcome = _run_codex(workspace=task_root, prompt=(task_root / "TASK.md").read_text(encoding="utf-8"))
        try:
            output = TaskCompilationOutputV1.model_validate_json((task_root / "teacher/task_compilation.json").read_text(encoding="utf-8"))
        except Exception as exc:
            task_reports.append({"task_id": task_id, "decision": "incomplete", "first_failure": f"task_compilation_json_or_schema_invalid:{type(exc).__name__}"})
            continue
        boundary_errors = productive_workload_rubric_errors(output)
        _materialize_task(task_root=task_root, spec=spec, output=output)
        report = ScenarioTaskAdmissionValidator().validate(spec=spec, package_root=task_root, bible=bible, rules=rule_set, output=output).model_dump(mode="json")
        if outcome["returncode"] != 0:
            report.update({"decision": "incomplete", "first_failure": f"local_codex_exit:{outcome['returncode']}"})
        if boundary_errors:
            report.update({"decision": "blocked", "first_failure": boundary_errors[0], "productive_workload_rubric_errors": boundary_errors})
        _write(task_root / "admission_report.json", report)
        task_reports.append(report)
        if report["decision"] == "pass":
            _review_pack(task_root=task_root, target=output_root / "user_review" / task_id, output=output)
    decision = "compiled_for_user_review" if len(task_reports) == 2 and all(item["decision"] == "pass" for item in task_reports) else "incomplete"
    result = {"result_version": "r10.compiler_revision_result.1", "decision": decision, "evidence_reports": reports, "task_reports": task_reports}
    _write(output_root / "result.json", result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, default=ROOT / "artifacts/r10/r10_8b2_compiler_revision")
    parser.add_argument("--scope-only", action="store_true")
    parser.add_argument("--authorized-scope-sha256")
    args = parser.parse_args()
    bibles, rules, skills = _load()
    scope = _scope(bibles=bibles, rules=rules, skills=skills, cli_version=local_codex_version())
    if args.scope_only:
        if args.output_root.exists():
            raise FileExistsError("r10_compiler_revision_output_already_exists")
        args.output_root.mkdir(parents=True)
        _write(args.output_root / "campaign_scope.json", scope)
        print(json.dumps({"decision": "awaiting_private_upload_authorization", "scope_sha256": sha256_json(scope), "output_root": str(args.output_root)}, ensure_ascii=False))
        return
    if not args.authorized_scope_sha256:
        parser.error("--authorized-scope-sha256 is required unless --scope-only")
    if args.output_root.exists():
        allowed = {"campaign_scope.json"}
        entries = {path.name for path in args.output_root.iterdir()}
        if entries != allowed:
            raise FileExistsError("r10_compiler_revision_output_already_exists")
        persisted_scope = json.loads((args.output_root / "campaign_scope.json").read_text(encoding="utf-8"))
        if sha256_json(persisted_scope) != sha256_json(scope):
            raise PermissionError("r10_compiler_revision_persisted_scope_drift")
    print(json.dumps(run(output_root=args.output_root, authorized_scope_sha256=args.authorized_scope_sha256), ensure_ascii=False))


if __name__ == "__main__":
    main()
