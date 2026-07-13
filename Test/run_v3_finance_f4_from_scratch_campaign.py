from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from task_generator.v3_f4_from_scratch_campaign import (  # noqa: E402
    AssistantReviewRecord,
    F4FromScratchCampaign,
    utc_now,
)
from task_generator.v3_holistic_editorial import HolisticCostLedger, HolisticEditorialExecutor  # noqa: E402


def main() -> None:
    parser = argparse.ArgumentParser(description="Govern the sequential F4 from-scratch validation campaign.")
    parser.add_argument("--action", choices=["init", "status", "begin-source", "source-run", "freeze-source", "begin-slot", "generate-next", "mark-generated", "holistic", "mark-holistic", "review"], required=True)
    parser.add_argument("--campaign-root", type=Path, required=True)
    parser.add_argument("--spec", type=Path, default=ROOT / "SkillRegistry" / "v3_finance_f4_from_scratch_validation.experimental.json")
    parser.add_argument("--canonical-registry", type=Path, default=ROOT / "SkillRegistry" / "v3_skill_registry.json")
    parser.add_argument("--slot", type=int)
    parser.add_argument("--run-manifest", type=Path)
    parser.add_argument("--source-run-manifest", type=Path)
    parser.add_argument("--source-root", type=Path)
    parser.add_argument("--evidence-dir", type=Path)
    parser.add_argument("--review-json", type=Path)
    parser.add_argument("--env-path", default="/run/secrets/provider_env")
    parser.add_argument("--deepseek-key-path", default="/run/secrets/deepseek_api_key")
    parser.add_argument("--tuzi-env-path", default="/run/secrets/eval_tuzi_env")
    parser.add_argument("--rw-task-root", default="/opt/rw-task")
    parser.add_argument("--timeout-seconds", type=int, default=7200)
    args = parser.parse_args()
    campaign = F4FromScratchCampaign(args.campaign_root, args.spec, args.canonical_registry)
    if args.action == "init": result = campaign.initialize()
    elif args.action == "status": result = campaign.read()
    elif args.action == "begin-source": result = campaign.begin_source_collection()
    elif args.action == "source-run": result = _source_run(campaign, args)
    elif args.action == "freeze-source":
        if not args.source_root or not args.source_run_manifest: parser.error("freeze-source requires --source-root and --source-run-manifest")
        result = campaign.freeze_sources([p for p in args.source_root.rglob("*") if p.is_file()], args.source_run_manifest)
    elif args.action == "begin-slot": result = campaign.begin_slot(_required(args.slot, parser, "--slot"))
    elif args.action == "generate-next": result = _generate_next(campaign, args, _required(args.slot, parser, "--slot"))
    elif args.action == "mark-generated": result = campaign.mark_generated(_required(args.slot, parser, "--slot"), _required(args.run_manifest, parser, "--run-manifest"))
    elif args.action == "holistic": result = _holistic(campaign, args, _required(args.slot, parser, "--slot"))
    elif args.action == "mark-holistic": result = campaign.mark_holistic_complete(_required(args.slot, parser, "--slot"), _required(args.evidence_dir, parser, "--evidence-dir"))
    else:
        if not args.review_json: parser.error("review requires --review-json")
        payload = json.loads(args.review_json.read_text(encoding="utf-8")); payload.setdefault("reviewed_at", utc_now())
        result = campaign.apply_review(AssistantReviewRecord.model_validate(payload))
    print(json.dumps(result, ensure_ascii=False, indent=2))


def _required(value, parser, flag):
    if value is None: parser.error(f"{flag} is required")
    return value


def _source_run(campaign, args):
    manifest = campaign.initialize()
    if manifest["state"] == "new": campaign.begin_source_collection()
    elif manifest["state"] != "collecting_sources": raise RuntimeError("source_run_not_allowed")
    query_path = campaign.root / "state" / "topic_queries.json"
    _atomic(query_path, {"topics": {item["topic"]: item["queries"] for item in campaign.spec["topics"]}})
    command = [sys.executable, str(ROOT / "Test" / "run_v3_end_to_end_pipeline.py"),
        "--run-id", "source_freeze", "--action", "run", "--profile", "custom", "--stage", "source_to_skills",
        "--source-mode", "web", "--registry-mode", "existing", "--registry-path", str(campaign.registry_path),
        "--apply-registry-update", "--allow-web-collection", "--allow-external-source-upload",
        "--collector-backend", "direct", "--source-limit", str(campaign.spec["source_limit_per_topic"]),
        "--provider", campaign.spec["source_provider"], "--deepseek-model", campaign.spec["source_model"],
        "--max-candidates", str(campaign.spec["max_candidates_per_topic"]), "--extractor-max-tokens", str(campaign.spec["max_tokens"]),
        "--extractor-output-profile", campaign.spec["extractor_output_profile"], "--topic-queries-path", str(query_path),
        "--env-path", args.env_path, "--deepseek-key-path", args.deepseek_key_path, "--output-root", str(campaign.root / "pipeline_runs")]
    for item in campaign.spec["topics"]: command += ["--topic", item["topic"]]
    _run(command, args.timeout_seconds)
    run_root = campaign.root / "pipeline_runs" / "source_freeze"
    campaign.freeze_sources([p for p in run_root.rglob("*") if p.is_file()], run_root / "end_to_end_run_manifest.json")
    return campaign.read()


def _generate_next(campaign, args, slot):
    campaign.begin_slot(slot); item = campaign.spec["slots"][slot - 1]
    attempt = campaign.read()["slots"][slot - 1]["attempts"]
    run_id = f"slot_{slot:02d}_{item['motif']}_r{attempt:02d}"
    command = [sys.executable, str(ROOT / "Test" / "run_v3_end_to_end_pipeline.py"),
        "--run-id", run_id, "--action", "run", "--profile", "custom", "--stage", "production_review",
        "--source-mode", "existing", "--registry-mode", "existing", "--registry-path", str(campaign.registry_path),
        "--domain-profile", "finance_audit", "--max-cases", "1", "--case-index-offset", str(slot - 1),
        "--target-difficulty-profile", "finance_semantic_contract_v2" if item["motif"] != "evidence_to_deliverable" else "finance_production_v1",
        "--motif", item["motif"], "--motif-occurrence-offset", f"{item['motif']}={item['occurrence']}",
        "--rw-task-root", args.rw_task_root, "--python-exe", sys.executable,
        "--output-root", str(campaign.root / "pipeline_runs"), "--timeout-seconds", str(args.timeout_seconds)]
    _run(command, args.timeout_seconds)
    run_manifest = campaign.root / "pipeline_runs" / run_id / "end_to_end_run_manifest.json"
    return campaign.mark_generated(slot, run_manifest)


def _holistic(campaign, args, slot):
    manifest = campaign.read(); record = manifest["slots"][slot - 1]
    if record["state"] != "awaiting_holistic_review": raise RuntimeError("slot_not_ready_for_holistic")
    run_manifest = Path(record["task_run_manifest"]); run_root = run_manifest.parent
    lifecycle = json.loads((run_root / "lifecycle_index.json").read_text(encoding="utf-8"))
    paths = lifecycle["paths"]
    blueprint = _json(run_root / paths["blueprint"]); golden = _json(run_root / paths["golden_run"])
    annotation = _json(run_root / paths["training_annotation"]); rubric = _json(run_root / paths["rubric"])
    refs = run_root / paths["reference_files"]
    prompt = "\n".join(blueprint.get("prompt_spec", {}).get("visible_requirements", []))
    out = campaign.root / "slots" / f"slot_{slot:02d}" / f"revision_{record['attempts']:02d}"
    out.mkdir(parents=True, exist_ok=True)
    ledger = HolisticCostLedger(campaign.root / "holistic_cost_ledger.json", campaign.spec["tuzi_budget_rmb"], campaign.spec["tuzi_request_cap"])
    executor = HolisticEditorialExecutor(args.tuzi_env_path, ledger)
    context = {"task_id": lifecycle["case_id"], "motif": record["motif"], "required_business_goal": blueprint.get("task_metadata", {}).get("task_goal"),
        "prompt": prompt, "candidate_files": _contents(refs), "teacher_artifacts": {"golden_run": golden, "training_annotation": annotation, "rubric": rubric},
        "minimum_repairs": ["candidate-visible answerability", "independently reproducible truth", "fact-centered rubric >= 60%"],
        "experimental_evidence_to_deliverable": bool(record.get("experimental_motif"))}
    revision = executor.revise(context); _atomic(out / "whole_task_revision_bundle.json", revision.model_dump(mode="json")); _atomic(out / "terra_provider_diagnostics.json", executor.last_diagnostics)
    deterministic = {"decision": "pass" if revision.overall_decision == "already_valid" and not revision.candidate_file_changes and not revision.unresolved_questions else "revise_system",
        "original_verifier_pass": lifecycle.get("verifier_status") == "pass", "candidate_file_changes_requested": len(revision.candidate_file_changes),
        "unresolved_questions": revision.unresolved_questions, "truth_accepted_from_llm": False}
    _atomic(out / "deterministic_validation.json", deterministic)
    solve = executor.solve({"task_id": lifecycle["case_id"], "prompt": revision.revised_prompt or prompt, "candidate_files": _contents(refs)})
    _atomic(out / "luna_candidate_solve.json", solve.model_dump(mode="json")); _atomic(out / "luna_provider_diagnostics.json", executor.last_diagnostics)
    _atomic(out / "holistic_summary.json", {"slot": slot, "motif": record["motif"], "deterministic_decision": deterministic["decision"], "luna_answerable": solve.answerable,
        "unresolved_material_ambiguity": bool(solve.ambiguity_or_hidden_assumptions), "raw_provider_response_included": False})
    return campaign.mark_holistic_complete(slot, out)


def _run(command, timeout):
    result = subprocess.run(command, cwd=ROOT, timeout=timeout, check=False)
    if result.returncode: raise RuntimeError(f"subprocess_failed:{result.returncode}")


def _json(path): return json.loads(Path(path).read_text(encoding="utf-8"))
def _atomic(path, payload):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"); tmp.replace(path)


def _contents(root):
    result = {}
    for path in sorted(Path(root).iterdir()):
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=False); result[path.name] = {ws.title: [list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}; wb.close()
        elif path.suffix.lower() == ".docx":
            with zipfile.ZipFile(path) as archive: xml = ElementTree.fromstring(archive.read("word/document.xml"))
            result[path.name] = "\n".join(node.text or "" for node in xml.iter() if node.tag.endswith("}t"))
        else: result[path.name] = path.read_text(encoding="utf-8", errors="replace")[:30000]
    return result


if __name__ == "__main__": main()
