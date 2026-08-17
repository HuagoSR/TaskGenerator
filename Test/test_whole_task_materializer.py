from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from pydantic import ValidationError

from src.task_generator.v3_whole_task_materializer import (
    CandidateFileRevisionV2,
    WholeTaskMaterializer,
    WholeTaskRevisionBundleV2,
)


class WholeTaskMaterializerTests(unittest.TestCase):
    def test_rejects_path_traversal_and_unsupported_file(self):
        with self.assertRaises(ValidationError):
            CandidateFileRevisionV2(file_name="../secret.json", change_type="add", content={})
        with self.assertRaises(ValidationError):
            CandidateFileRevisionV2(file_name="macro.xlsm", change_type="add", content={})

    def test_cash_revision_materializes_recomputes_and_exports(self):
        from docx import Document
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            refs = root / "refs"; refs.mkdir()
            self._xlsx(refs / "bank_statement.xlsx", "Bank_Activity", [
                ["Bank_ID", "Reference", "Amount"], ["B-1", "R-1", 100], ["B-2", "R-2", -25],
            ])
            self._xlsx(refs / "cash_ledger.xlsx", "Cash_Ledger", [
                ["Ledger_ID", "Reference", "Amount"], ["L-1", "R-1", 100], ["L-2", "R-X", -25],
            ])
            doc = Document(); doc.add_paragraph("REC-001 Exact reference and amount match"); doc.save(refs / "reconciliation_rules.docx")
            bundle = WholeTaskRevisionBundleV2(
                task_id="cash_test", motif="fan_in_reconciliation", overall_decision="revise",
                revised_prompt="Reconcile the supplied bank and ledger activity and save cash_reconciliation.xlsx.",
                candidate_files=[
                    {"file_name": "bank_statement.xlsx", "change_type": "retain"},
                    {"file_name": "cash_ledger.xlsx", "change_type": "retain"},
                    {"file_name": "reconciliation_rules.docx", "change_type": "retain"},
                ],
            )
            blueprint = {
                "template_family": "finance_cash_reconciliation_v1",
                "data_spec": {"reference_files": []},
                "deliverable_spec": [{"file_name": "cash_reconciliation.xlsx", "requirements": ["reconcile activity"]}],
                "prompt_spec": {"visible_requirements": ["reconcile activity"]},
                "golden_plan": {"required_final_checks": ["reconcile activity"]},
            }
            report = WholeTaskMaterializer().materialize(bundle, refs, root / "revision", blueprint)
            self.assertEqual(report.decision, "pass")
            self.assertTrue(report.deterministic_recomputation_pass)
            self.assertTrue(report.rw_task_export_compatible)
            self.assertTrue(report.deliverable_contract_valid)
            prompt = (root / "revision" / "prompt.md").read_text(encoding="utf-8")
            self.assertIn("deliverable_files/cash_reconciliation.xlsx", prompt)
            self.assertTrue((root / "revision" / "deliverable_contract.json").is_file())
            self.assertTrue(
                (root / "revision" / "rw_task_export" / "deliverable_contract.json").is_file()
            )
            from openpyxl import load_workbook
            rendered = load_workbook(root / "revision" / "reference_files" / "bank_statement.xlsx")
            self.assertEqual(rendered.active.page_setup.fitToWidth, 1)
            self.assertEqual(rendered.active.page_setup.orientation, "landscape")
            rendered.close()
            truth = json.loads((root / "revision" / "teacher" / "deterministic_answer_key.json").read_text(encoding="utf-8"))
            self.assertEqual(truth["claim_001"]["matched_count"], 1)
            self.assertEqual(truth["claim_001"]["bank_only_count"], 1)
            self.assertEqual(truth["claim_001"]["activity_difference"], 0.0)
            self.assertEqual(truth["claim_001"]["bank_only_items"][0]["Bank_ID"], "B-2")
            self.assertEqual(truth["claim_001"]["ledger_only_items"][0]["Ledger_ID"], "L-2")

    def test_xlsx_shape_limit_is_enforced(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); refs = root / "refs"; refs.mkdir()
            bundle = WholeTaskRevisionBundleV2(
                task_id="t", motif="fan_in_reconciliation", overall_decision="revise",
                revised_prompt="A sufficiently long prompt for validation and materialization.",
                candidate_files=[{"file_name": "bad.xlsx", "change_type": "add", "content": {"S": [[str(i) for i in range(41)]]}}],
            )
            with self.assertRaises(ValueError):
                WholeTaskMaterializer().materialize(bundle, refs, root / "out", {
                    "data_spec": {"reference_files": []}, "deliverable_spec": [{"file_name": "out.xlsx"}],
                    "prompt_spec": {}, "golden_plan": {},
                })

    def test_candidate_text_mojibake_is_repaired_or_rejected(self):
        materializer = WholeTaskMaterializer()
        self.assertEqual(materializer._normalize_candidate_text("invoice line鈥檚 key"), "invoice line's key")
        with self.assertRaisesRegex(ValueError, "candidate_text_encoding_corruption"):
            materializer._normalize_candidate_text("broken 鈥 text")

    def test_wide_workbook_renders_two_pages_wide(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "wide.xlsx"
            self._xlsx(path, "Validation", [[f"C{i}" for i in range(19)]])
            WholeTaskMaterializer._normalize_xlsx_print_layout(path)
            workbook = load_workbook(path)
            self.assertEqual(workbook.active.page_setup.fitToWidth, 2)
            workbook.close()

    @staticmethod
    def _xlsx(path: Path, sheet_name: str, rows):
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active; sheet.title = sheet_name
        for row in rows: sheet.append(row)
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
