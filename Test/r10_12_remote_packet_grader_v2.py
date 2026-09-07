"""Build a coordinate-preserving XLSX evidence packet and call the fixed judge."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import urllib.request
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries


ROOT = Path("/workspace")


def sha(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def recalc(path: Path) -> tuple[Path | None, str]:
    source = ROOT / ".recalc_source"
    output = ROOT / ".recalc_output"
    shutil.rmtree(source, ignore_errors=True)
    shutil.rmtree(output, ignore_errors=True)
    source.mkdir(); output.mkdir()
    copied = source / path.name
    shutil.copy2(path, copied)
    try:
        result = subprocess.run(
            ["libreoffice", "-env:UserInstallation=file:///tmp/lo-profile", "--headless",
             "--convert-to", "xlsx", "--outdir", str(output), str(copied)],
            text=True, capture_output=True, timeout=180, check=False,
        )
    except (OSError, subprocess.TimeoutExpired):
        return None, "recalculation_failed"
    generated = output / path.name
    return (generated, "recalculated") if result.returncode == 0 and generated.is_file() else (None, "recalculation_failed")


def refs(formula: str, sheet: str) -> set[tuple[str, str]]:
    found = set()
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return found
    for token in tokens:
        if token.subtype != "RANGE" or "[" in token.value:
            continue
        value = token.value.replace("$", "")
        target_sheet, target = sheet, value
        if "!" in value:
            prefix, target = value.rsplit("!", 1)
            target_sheet = prefix.strip("'")
        try:
            min_col, min_row, max_col, max_row = range_boundaries(target)
        except ValueError:
            continue
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 10000:
            continue
        from openpyxl.utils import get_column_letter
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                found.add((target_sheet, f"{get_column_letter(col)}{row}"))
    return found


def spreadsheet(path: Path, relative: str) -> dict:
    recalculated, status = recalc(path)
    formula_book = load_workbook(path, data_only=False)
    value_book = load_workbook(path, data_only=True)
    recalc_book = load_workbook(recalculated, data_only=True) if recalculated else None
    sheets, edges, diagnostics = [], {}, []
    try:
        for ws in formula_book.worksheets:
            value_ws = value_book[ws.title]
            recalc_ws = recalc_book[ws.title] if recalc_book and ws.title in recalc_book.sheetnames else None
            cells, headers, styles = [], [], {}
            meaningful = sorted((cell for cell in ws._cells.values() if cell.value is not None), key=lambda cell: (cell.row, cell.column))
            min_row = min((cell.row for cell in meaningful), default=1); max_row = max((cell.row for cell in meaningful), default=1)
            min_col = min((cell.column for cell in meaningful), default=1); max_col = max((cell.column for cell in meaningful), default=1)
            for cell in meaningful:
                formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                if formula:
                    node = (ws.title, cell.coordinate)
                    edges[node] = refs(formula, ws.title)
                    if node in edges[node]:
                        diagnostics.append({"kind": "formula_cycle", "locations": [f"{ws.title}!{cell.coordinate}"], "detail": "formula directly or range-references itself"})
                value = value_ws[cell.coordinate].value
                recalculated_value = recalc_ws[cell.coordinate].value if recalc_ws else None
                if isinstance(value, str) and value.startswith("#"):
                    diagnostics.append({"kind": "formula_error", "locations": [f"{ws.title}!{cell.coordinate}"], "detail": value})
                def safe(item): return item if item is None or isinstance(item, (str, int, float, bool)) else str(item)
                cells.append({"address": cell.coordinate, "value": safe(value), "formula": formula,
                              "recalculated_value": safe(recalculated_value), "number_format": cell.number_format,
                              "style_id": cell.style_id})
                styles[str(cell.style_id)] = styles.get(str(cell.style_id), 0) + 1
                if cell.row <= 10 and isinstance(value, str) and value.strip(): headers.append(value.strip())
            sheets.append({"name": ws.title, "used_range": f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}", "max_row": max_row,
                           "max_column": max_col, "headers": headers[:40], "cells": cells,
                           "merged_ranges": sorted(str(row) for row in ws.merged_cells.ranges),
                           "hidden_rows": [key for key, row in ws.row_dimensions.items() if row.hidden],
                           "hidden_columns": [key for key, row in ws.column_dimensions.items() if row.hidden],
                           "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                           "column_widths": {key: row.width for key, row in ws.column_dimensions.items() if row.width is not None},
                           "style_counts": styles})
        # Detect indirect cycles among formula cells.
        active, visited, stack, cycles = set(), set(), [], set()
        def visit(node):
            if node in active:
                cycle = tuple(sorted(stack[stack.index(node):] + [node]))
                cycles.add(cycle); return
            if node in visited: return
            visited.add(node); active.add(node); stack.append(node)
            for target in edges.get(node, set()):
                if target in edges: visit(target)
            stack.pop(); active.remove(node)
        for node in edges: visit(node)
        for cycle in cycles:
            diagnostics.append({"kind": "formula_cycle", "locations": [f"{sheet}!{cell}" for sheet, cell in cycle], "detail": "formula dependency cycle"})
    finally:
        formula_book.close(); value_book.close()
        if recalc_book: recalc_book.close()
    return {"evidence_version": "r10.spreadsheet_evidence.1", "relative_path": relative,
            "original_sha256": sha(path), "recalculated_sha256": sha(recalculated) if recalculated else None,
            "recalculation_status": status, "sheets": sheets, "diagnostics": diagnostics}


def packet() -> dict:
    materials = (ROOT / "materials_text.txt").read_text(encoding="utf-8")
    workbook_evidence = {}
    for root in (ROOT / "reference_files", ROOT / "anonymous_submission"):
        for path in sorted(root.rglob("*.xlsx")):
            relative = path.relative_to(ROOT).as_posix()
            workbook_evidence[relative] = spreadsheet(path, relative)
    # Repeated JSON field names made an otherwise small workbook exceed the
    # provider context budget.  This is lossless, structured columnar data:
    # ``cells[i][j]`` maps to ``cell_columns[j]``.  It is expanded and checked
    # by the controller before a score is accepted.
    compact = {}
    columns = ["address", "value", "formula", "recalculated_value", "number_format", "style_id"]
    for relative, evidence in workbook_evidence.items():
        compact[relative] = dict(evidence)
        compact_sheets = []
        for sheet in evidence["sheets"]:
            encoded = dict(sheet)
            encoded["cell_columns"] = columns
            encoded["cells"] = [[cell[key] for key in columns] for cell in sheet["cells"]]
            compact_sheets.append(encoded)
        compact[relative]["sheets"] = compact_sheets
    return {"packet_version": "r10.strict_grade_packet.2", "task_id": json.loads((ROOT / "task.json").read_text())["task_id"],
            "strict_system": (ROOT / "strict_system.txt").read_text(encoding="utf-8"),
            "rubric_items": json.loads((ROOT / "rubric_items.json").read_text(encoding="utf-8")),
            "materials_text": materials, "workbook_evidence": compact,
            "output_schema": json.loads((ROOT / "grade_schema.json").read_text(encoding="utf-8"))}


def main() -> None:
    value = packet()
    (ROOT / "strict_grade_packet_v2.json").write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    key = os.environ["DEEPSEEK_API_KEY"].strip()
    body = {"model": "deepseek-v4-pro", "messages": [
        {"role": "system", "content": value["strict_system"]},
        {"role": "user", "content": json.dumps({key: value[key] for key in ("rubric_items", "materials_text", "workbook_evidence", "output_schema")}, ensure_ascii=False)},
    ], "response_format": {"type": "json_object"}, "temperature": 0, "max_tokens": 32768, "stream": False,
            "thinking": {"type": "disabled"}}
    request = urllib.request.Request("https://api.deepseek.com/chat/completions", data=json.dumps(body, ensure_ascii=False).encode(),
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {key}"}, method="POST")
    with urllib.request.urlopen(request, timeout=900) as response:
        result = json.loads(response.read().decode("utf-8"))
        (ROOT / "provider_status.json").write_text(json.dumps({"http_status": response.status, "response_model": result.get("model"),
            "response_object": result.get("object"), "request_id": response.headers.get("x-request-id"), "usage": result.get("usage")}, ensure_ascii=False, indent=2), encoding="utf-8")
    content = result["choices"][0]["message"]["content"]
    # Kept only in ignored campaign artifacts.  This distinguishes a provider
    # JSON truncation from a controller/schema failure without exposing it to Git.
    (ROOT / "provider_response.json").write_text(content, encoding="utf-8")
    parsed = json.loads(content)
    (ROOT / "grade.raw.json").write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
