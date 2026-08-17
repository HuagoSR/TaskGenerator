from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from src.task_generator.v3_behavioral_validation import SolverToolPreflight
from src.task_generator.v3_deliverable_contract import DeliverableContractCompiler
from src.task_generator.v3_rw_task_eval_prep import (
    RwTaskEvalPrepReport,
    RwTaskEvalPrepRequest,
)
from src.task_generator.v3_rw_task_eval_runner import RwTaskEvalRunner


class BehavioralEvalRunnerIntegrationTests(unittest.TestCase):
    def test_hybrid_route_blocks_execution_without_preflight(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            prep_path = self._write_prep(root, hybrid=True)
            with patch(
                "src.task_generator.v3_rw_task_eval_runner.subprocess.run"
            ) as mocked:
                report = RwTaskEvalRunner().run(
                    prep_report_path=prep_path,
                    output_dir=root / "run_report",
                    run_eval=True,
                )
            mocked.assert_not_called()
            self.assertEqual(report.run_status, "blocked")
            self.assertIn("solver_preflight_required", report.blocking_reasons)
            self.assertEqual(report.behavioral_failure_category, "tool_failure")
            self.assertFalse(report.grader_eligible)
            self.assertTrue(Path(report.behavioral_execution_report_path).is_file())

    def test_hybrid_route_blocks_failed_or_mismatched_preflight(self):
        for mode in ("failed", "mismatched"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                prep_path = self._write_prep(root, hybrid=True)
                preflight_path = self._write_preflight(
                    root / "preflight",
                    solver_model="other" if mode == "mismatched" else "solver",
                    passing=mode == "mismatched",
                )
                with patch(
                    "src.task_generator.v3_rw_task_eval_runner.subprocess.run"
                ) as mocked:
                    report = RwTaskEvalRunner().run(
                        prep_report_path=prep_path,
                        output_dir=root / "run_report",
                        run_eval=True,
                        solver_preflight_report_path=preflight_path,
                    )
                mocked.assert_not_called()
                expected = (
                    "solver_preflight_status:fail"
                    if mode == "failed"
                    else "solver_preflight_model_mismatch"
                )
                self.assertIn(expected, report.blocking_reasons)
                self.assertFalse(report.grader_eligible)

    def test_passing_preflight_and_valid_delivery_reaches_grader(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            prep_path = self._write_prep(root, hybrid=True)
            prep = RwTaskEvalPrepReport.model_validate_json(
                prep_path.read_text(encoding="utf-8")
            )
            solver_output = Path(prep.would_run_commands[0][5])
            preflight_path = self._write_preflight(
                root / "preflight",
                solver_model="solver",
                passing=True,
            )

            def run_side_effect(*args, **kwargs):
                if run_side_effect.calls == 0:
                    delivery = solver_output / "run_case" / "deliverable_files"
                    delivery.mkdir(parents=True)
                    workbook = Workbook()
                    workbook.active.append(["value"])
                    workbook.save(delivery / "output.xlsx")
                run_side_effect.calls += 1
                return subprocess.CompletedProcess([], 0, "", "")

            run_side_effect.calls = 0
            with patch(
                "src.task_generator.v3_rw_task_eval_runner.subprocess.run",
                side_effect=run_side_effect,
            ) as mocked:
                report = RwTaskEvalRunner().run(
                    prep_report_path=prep_path,
                    output_dir=root / "run_report",
                    run_eval=True,
                    solver_preflight_report_path=preflight_path,
                )
            self.assertEqual(mocked.call_count, 2)
            self.assertEqual(report.run_status, "completed")
            self.assertEqual(report.delivery_status, "valid")
            self.assertTrue(report.grader_eligible)
            self.assertEqual(report.behavioral_failure_category, "none")

    def test_legacy_route_remains_compatible_without_preflight(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            prep_path = self._write_prep(root, hybrid=False)
            with patch(
                "src.task_generator.v3_rw_task_eval_runner.subprocess.run",
                return_value=subprocess.CompletedProcess([], 1, "", "tool failed"),
            ) as mocked:
                report = RwTaskEvalRunner().run(
                    prep_report_path=prep_path,
                    output_dir=root / "run_report",
                    run_eval=True,
                )
            self.assertEqual(mocked.call_count, 1)
            self.assertNotIn("solver_preflight_required", report.blocking_reasons)
            self.assertEqual(report.run_status, "failed")

    def _write_prep(self, root: Path, *, hybrid: bool) -> Path:
        case_dir = root / "eval_input" / "case"
        case_dir.mkdir(parents=True)
        contract = DeliverableContractCompiler().build(
            "case",
            [{"file_name": "output.xlsx"}],
        )
        (case_dir / "deliverable_contract.json").write_text(
            contract.model_dump_json(indent=2),
            encoding="utf-8",
        )
        extra = {"deliverable_contract": contract.model_dump(mode="json")}
        if hybrid:
            extra["materialization_route"] = "hybrid_deterministic_fixture_v1"
        (case_dir / "dataset_row.json").write_text(
            json.dumps(
                {
                    "task_id": "case",
                    "reference_files": [],
                    "deliverable_files": ["deliverable_files/output.xlsx"],
                    "extra": extra,
                },
                indent=2,
            ),
            encoding="utf-8",
        )
        solver_output = root / "solver_output"
        grade_output = root / "grade_output"
        commands = [
            [
                sys.executable,
                "-m",
                "task_generator.v3_rw_task_eval_stirrup_wrapper",
                str(root / "eval_input"),
                "--output",
                str(solver_output),
            ],
            [
                sys.executable,
                "-m",
                "bench_standalone.grade_deliverables",
                str(solver_output),
                "--out-dir",
                str(grade_output),
            ],
        ]
        prep = RwTaskEvalPrepReport(
            request=RwTaskEvalPrepRequest(
                case_dir=str(case_dir),
                validation_report_path=str(root / "validation.json"),
                eval_input_dir=str(root / "eval_input"),
                model="solver",
                rw_task_root=str(root),
                python_exe=sys.executable,
            ),
            case_id="case",
            batch_case_id="case",
            blueprint_id="blueprint",
            rw_task_task_id="case",
            evaluated_model_name="solver",
            prep_status="prepared",
            evaluation_mode="candidate_ready_eval_candidate",
            eval_input_case_dir=str(case_dir),
            would_run_commands=commands,
        )
        path = root / "prep_report.json"
        path.write_text(prep.model_dump_json(indent=2), encoding="utf-8")
        return path

    def _write_preflight(
        self,
        root: Path,
        *,
        solver_model: str,
        passing: bool,
    ) -> Path:
        preflight = SolverToolPreflight()
        fixture = preflight.create_fixture(root)
        if passing:
            created = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_CREATED
            workbook = Workbook()
            workbook.active["A1"] = "created_marker"
            workbook.save(created)
            edited = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_EDITED
            shutil.copy2(fixture["template_path"], edited)
            workbook = load_workbook(edited)
            workbook["Template"]["B1"] = "edited_marker"
            workbook.save(edited)
            workbook.close()
        report = preflight.inspect(root, solver_model, "local-test")
        path = root / "solver_tool_preflight_report.json"
        path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
        return path


if __name__ == "__main__":
    unittest.main()
