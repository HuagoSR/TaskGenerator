from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from task_generator.v3_behavioral_validation import SolverToolPreflight  # noqa: E402


DEFAULT_OUTPUT = (
    ROOT
    / "artifacts"
    / "pipeline_reconstruction"
    / "solver_preflight_smoke"
    / "local_reference"
)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Create and inspect a deterministic local reference execution of the "
            "solver file-tool preflight. This does not call a model or provider."
        )
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    preflight = SolverToolPreflight()
    fixture = preflight.create_fixture(args.output_dir)
    deliverable_dir = Path(fixture["deliverable_dir"])

    created = deliverable_dir / preflight.EXPECTED_CREATED
    workbook = Workbook()
    workbook.active["A1"] = "created_marker"
    workbook.save(created)

    edited = deliverable_dir / preflight.EXPECTED_EDITED
    shutil.copy2(fixture["template_path"], edited)
    workbook = load_workbook(edited)
    workbook["Template"]["B1"] = "edited_marker"
    workbook.save(edited)
    workbook.close()

    report = preflight.inspect(
        args.output_dir,
        solver_model="deterministic-local-reference",
        environment_id="local-taskgenerator",
    )
    report_path = args.output_dir / "solver_tool_preflight_report.json"
    report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "status": report.status,
                "eligible_for_business_eval": report.eligible_for_business_eval,
                "operations": {
                    item.operation: item.passed for item in report.operations
                },
                "report_path": str(report_path),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    if report.status != "pass":
        raise SystemExit(1)


if __name__ == "__main__":
    main()
