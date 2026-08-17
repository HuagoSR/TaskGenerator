from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.task_generator.v3_behavioral_validation import (
    BehavioralExecutionBuilder,
    SolverToolPreflight,
)
from src.task_generator.v3_deliverable_contract import (
    DeliverableContractCompiler,
    inspect_delivery,
)


class BehavioralValidationTests(unittest.TestCase):
    def test_solver_preflight_covers_create_copy_edit_save_submit(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            preflight = SolverToolPreflight()
            fixture = preflight.create_fixture(root)
            created = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_CREATED
            workbook = Workbook()
            workbook.active["A1"] = "created_marker"
            workbook.save(created)
            edited = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_EDITED
            shutil.copy2(fixture["template_path"], edited)
            workbook = load_workbook(edited)
            workbook["Template"]["B1"] = "edited_marker"
            workbook.save(edited)
            workbook.close()

            report = preflight.inspect(root, "fixture-solver", "local-test")
            self.assertEqual(report.status, "pass")
            self.assertTrue(report.eligible_for_business_eval)
            self.assertTrue(all(item.passed for item in report.operations))
            self.assertEqual(
                [item.operation for item in report.operations],
                ["create", "copy", "edit", "save", "submit"],
            )
            self.assertEqual(len(report.output_fingerprints), 2)

    def test_solver_preflight_failure_preserves_first_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            preflight = SolverToolPreflight()
            preflight.create_fixture(root)
            report = preflight.inspect(root, "fixture-solver", "local-test")
            self.assertEqual(report.status, "fail")
            self.assertFalse(report.eligible_for_business_eval)
            self.assertEqual(report.first_failure, "preflight_operation:create")
            self.assertEqual(report.attempts[0].failure_reason, "create")

    def test_behavioral_report_separates_process_delivery_and_grader(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            delivery = self._valid_delivery(root)
            preflight = self._passing_preflight(root / "preflight")
            report = BehavioralExecutionBuilder().build(
                case_id="case",
                solver_model="solver",
                environment_id="local",
                command_records=[
                    {
                        "command_name": "solver",
                        "status": "succeeded",
                        "stdout_excerpt": "done",
                        "stderr_excerpt": "",
                    },
                    {
                        "command_name": "bench_standalone.grade_deliverables",
                        "status": "succeeded",
                    },
                ],
                delivery_inspection=delivery,
                output_root=root,
                preflight_report=preflight,
                preflight_required=True,
                business_validity_status="pass",
                professional_quality_status="provisional",
                business_score=0.72,
            )
            self.assertEqual(report.process_status, "succeeded")
            self.assertEqual(report.delivery_status, "valid")
            self.assertEqual(report.file_validity_status, "pass")
            self.assertTrue(report.grader_eligible)
            self.assertTrue(report.grader_executed)
            self.assertEqual(report.failure_category, "none")
            self.assertEqual(report.business_score, 0.72)

    def test_failure_taxonomy(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            builder = BehavioralExecutionBuilder()
            preflight = self._passing_preflight(root / "preflight")
            provider = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[
                    {
                        "command_name": "solver",
                        "status": "failed",
                        "failure_stage": "solver",
                        "stderr_excerpt": "HTTP 429 rate limit",
                    }
                ],
                delivery_inspection=None,
                output_root=root,
                preflight_report=preflight,
            )
            self.assertEqual(provider.failure_category, "provider_failure")

            missing_contract = DeliverableContractCompiler().build(
                "c", [{"file_name": "output.xlsx"}]
            )
            missing = inspect_delivery(root / "missing", missing_contract)
            non_delivery = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=missing,
                output_root=root / "missing",
                preflight_report=preflight,
            )
            self.assertEqual(non_delivery.failure_category, "non_delivery")

            wrong_root = root / "wrong"
            wrong_dir = wrong_root / "run" / "deliverable_files"
            wrong_dir.mkdir(parents=True)
            workbook = Workbook()
            workbook.save(wrong_dir / "wrong.xlsx")
            wrong = inspect_delivery(wrong_root, missing_contract)
            wrong_report = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=wrong,
                output_root=wrong_root,
                preflight_report=preflight,
            )
            self.assertEqual(wrong_report.failure_category, "wrong_path")

            invalid_root = root / "invalid"
            invalid_dir = invalid_root / "deliverable_files"
            invalid_dir.mkdir(parents=True)
            (invalid_dir / "output.xlsx").write_bytes(b"not-a-workbook")
            invalid = inspect_delivery(invalid_root, missing_contract)
            invalid_report = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=invalid,
                output_root=invalid_root,
                preflight_report=preflight,
            )
            self.assertEqual(invalid_report.failure_category, "invalid_file")

            valid = self._valid_delivery(root / "business")
            business = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=valid,
                output_root=root / "business",
                preflight_report=preflight,
                business_validity_status="fail",
            )
            self.assertEqual(business.failure_category, "business_error")

            professional = builder.build(
                case_id="c",
                solver_model="m",
                environment_id="e",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=valid,
                output_root=root / "business",
                preflight_report=preflight,
                business_validity_status="pass",
                professional_quality_status="fail",
            )
            self.assertEqual(
                professional.failure_category,
                "professional_quality_failure",
            )

    def test_failed_or_missing_preflight_is_never_grader_eligible(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            delivery = self._valid_delivery(root / "delivery")
            report = BehavioralExecutionBuilder().build(
                case_id="case",
                solver_model="solver",
                environment_id="local",
                command_records=[{"command_name": "solver", "status": "succeeded"}],
                delivery_inspection=delivery,
                output_root=root,
                preflight_report=None,
                preflight_required=True,
            )
            self.assertFalse(report.grader_eligible)
            self.assertEqual(report.failure_category, "tool_failure")

    def _passing_preflight(self, root: Path):
        preflight = SolverToolPreflight()
        fixture = preflight.create_fixture(root)
        created = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_CREATED
        workbook = Workbook()
        workbook.active["A1"] = "created_marker"
        workbook.save(created)
        edited = Path(fixture["deliverable_dir"]) / preflight.EXPECTED_EDITED
        shutil.copy2(fixture["template_path"], edited)
        workbook = load_workbook(edited)
        workbook["Template"]["B1"] = "edited_marker"
        workbook.save(edited)
        workbook.close()
        return preflight.inspect(root, "solver", "local")

    def _valid_delivery(self, root: Path):
        contract = DeliverableContractCompiler().build(
            "case", [{"file_name": "output.xlsx"}]
        )
        delivery = root / "deliverable_files"
        delivery.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        workbook.active.append(["value"])
        workbook.save(delivery / "output.xlsx")
        return inspect_delivery(root, contract)


if __name__ == "__main__":
    unittest.main()
