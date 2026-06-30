import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TASK_ID = "TASK_11C0BEA2"

GOLDEN_PATH = ROOT / "Test" / "v2_outputs" / "finance_batch_01" / TASK_ID / "golden_run" / "golden_intermediate_values.json"
CANDIDATE_DIRS = {
    "claude-3-7-sonnet-latest": ROOT / "Test" / "v2_outputs" / "rw_task_eval_claude37_top2_results" / f"run_{TASK_ID}",
    "deepseek-v3.2": ROOT / "Test" / "v2_outputs" / "rw_task_eval_deepseek_results" / f"run_{TASK_ID}",
}
OUTPUT_DIR = ROOT / "Test" / "v2_outputs" / "analysis_reports"


@dataclass
class CandidateSnapshot:
    model_name: str
    workbook_path: Path
    sheet_names: list[str]
    summary_totals: dict[str, float]
    show_rows: list[dict[str, Any]]
    diagnosis: list[str]


def _normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return None


def _find_header_map(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_idx, row in enumerate(rows):
        texts = [_normalize_text(v).lower() for v in row]
        if not texts:
            continue
        if "show" in texts and "city" in texts and "country" in texts:
            header_map: dict[str, int] = {}
            for idx, text in enumerate(texts):
                if "show" == text:
                    header_map["show"] = idx
                elif "city" == text:
                    header_map["city"] = idx
                elif "country" == text:
                    header_map["country"] = idx
                elif "gross revenue" in text:
                    header_map["gross"] = idx
                elif "tax rate" in text or "wht rate" in text:
                    header_map["tax_rate"] = idx
                elif "withholding tax" in text or "wht amount" in text:
                    header_map["tax_amount"] = idx
                elif "net revenue" in text:
                    header_map["net"] = idx
            if {"show", "city", "country", "gross"} <= set(header_map):
                return row_idx, header_map
    return None


def _extract_show_rows(ws) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    header = _find_header_map(rows)
    if not header:
        return []

    header_idx, header_map = header
    show_rows: list[dict[str, Any]] = []
    for row in rows[header_idx + 1 :]:
        show_name = _normalize_text(row[header_map["show"]])
        if not show_name.startswith("Show "):
            if show_rows:
                break
            continue
        show_rows.append(
            {
                "show": show_name,
                "city": _normalize_text(row[header_map["city"]]),
                "country": _normalize_text(row[header_map["country"]]).replace(" *", ""),
                "gross_revenue": _to_float(row[header_map["gross"]]),
                "tax_rate": _to_float(row[header_map.get("tax_rate", -1)]) if "tax_rate" in header_map else None,
                "tax_amount": _to_float(row[header_map.get("tax_amount", -1)]) if "tax_amount" in header_map else None,
                "net_revenue": _to_float(row[header_map.get("net", -1)]) if "net" in header_map else None,
                "sheet_name": ws.title,
            }
        )
    return show_rows


def _extract_summary_totals(wb) -> dict[str, float]:
    totals: dict[str, float] = {}
    wanted = {
        "gross revenue": "gross_revenue",
        "total gross revenue": "gross_revenue",
        "less: withholding taxes": "withholding_tax",
        "withholding taxes": "withholding_tax",
        "total withholding taxes": "withholding_tax",
        "production costs": "production_costs",
        "total production costs": "production_costs",
        "net income": "net_income",
    }
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            texts = [_normalize_text(v) for v in row]
            lowered = [t.lower() for t in texts]
            nums = [_to_float(v) for v in row]
            numeric_values = [n for n in nums if n is not None]
            if not numeric_values:
                continue
            for cell_text in lowered:
                if cell_text in wanted:
                    key = wanted[cell_text]
                    candidate = numeric_values[-1]
                    if key == "withholding_tax":
                        candidate = abs(candidate)
                    totals.setdefault(key, candidate)
    return totals


def _build_diagnosis(show_rows: list[dict[str, Any]], golden_rows: list[dict[str, Any]]) -> list[str]:
    diagnosis: list[str] = []
    if not show_rows:
        return ["Could not locate a structured show-level revenue table in the workbook."]

    copied_raw_count = 0
    for candidate, golden in zip(show_rows, golden_rows):
        gross = candidate.get("gross_revenue")
        if gross is None:
            continue
        if abs(gross - golden["gross_revenue_usd"]) > 1e-6 and gross < golden["gross_revenue_usd"]:
            copied_raw_count += 1
    if copied_raw_count == len(golden_rows):
        diagnosis.append("All show-level gross revenue values are below the golden USD values, which indicates the workbook copied local-currency source amounts as if they were already USD.")

    germany_rates = [row["tax_rate"] for row in show_rows if row["country"] == "Germany" and row.get("tax_rate") is not None]
    if germany_rates:
        unique_rates = sorted({round(rate, 6) for rate in germany_rates})
        diagnosis.append(f"Germany withholding tax was applied at {unique_rates}, while the golden run resolves Germany to 0.15825.")

    return diagnosis


def analyze_candidate(model_name: str, run_dir: Path, golden_rows: list[dict[str, Any]]) -> CandidateSnapshot:
    workbook_path = next((run_dir / "deliverable_files").rglob("*.xlsx"))
    wb = load_workbook(workbook_path, data_only=True)

    show_rows: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        extracted = _extract_show_rows(ws)
        if extracted:
            show_rows = extracted
            break

    summary_totals = _extract_summary_totals(wb)
    diagnosis = _build_diagnosis(show_rows, golden_rows)
    return CandidateSnapshot(
        model_name=model_name,
        workbook_path=workbook_path,
        sheet_names=wb.sheetnames,
        summary_totals=summary_totals,
        show_rows=show_rows,
        diagnosis=diagnosis,
    )


def make_report() -> dict[str, Any]:
    golden = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    golden_totals = golden["net_income_totals"]
    golden_rows = golden["revenue_line_items"]

    snapshots = [analyze_candidate(name, path, golden_rows) for name, path in CANDIDATE_DIRS.items()]

    report: dict[str, Any] = {
        "task_id": TASK_ID,
        "golden_totals": golden_totals,
        "golden_revenue_line_items": golden_rows,
        "candidates": [],
    }
    for snap in snapshots:
        report["candidates"].append(
            {
                "model_name": snap.model_name,
                "workbook_path": str(snap.workbook_path),
                "sheet_names": snap.sheet_names,
                "summary_totals": snap.summary_totals,
                "show_rows": snap.show_rows,
                "diagnosis": snap.diagnosis,
            }
        )
    return report


def render_markdown(report: dict[str, Any]) -> str:
    lines: list[str] = []
    golden_totals = report["golden_totals"]
    lines.append(f"# Finance Failure Analysis: {report['task_id']}")
    lines.append("")
    lines.append("## Golden Totals")
    lines.append("")
    lines.append(f"- Revenue USD: {golden_totals['revenue_usd']:.2f}")
    lines.append(f"- Withholding Tax USD: {golden_totals['tax_usd']:.2f}")
    lines.append(f"- Production Cost USD: {golden_totals['expense_usd']:.2f}")
    lines.append(f"- Net Income USD: {golden_totals['net_income_usd']:.2f}")
    lines.append("")

    for candidate in report["candidates"]:
        lines.append(f"## {candidate['model_name']}")
        lines.append("")
        lines.append(f"- Workbook: `{candidate['workbook_path']}`")
        lines.append(f"- Sheets: {', '.join(candidate['sheet_names'])}")
        summary = candidate["summary_totals"]
        if summary:
            lines.append(f"- Reported revenue total: {summary.get('gross_revenue', float('nan')):.2f}")
            lines.append(f"- Reported withholding total: {summary.get('withholding_tax', float('nan')):.2f}")
            lines.append(f"- Reported production cost total: {summary.get('production_costs', float('nan')):.2f}")
            lines.append(f"- Reported net income total: {summary.get('net_income', float('nan')):.2f}")
        lines.append("")
        lines.append("### Diagnosis")
        lines.append("")
        for item in candidate["diagnosis"]:
            lines.append(f"- {item}")
        lines.append("")
        lines.append("### Show-Level Gross Revenue Comparison")
        lines.append("")
        lines.append("| Show | City | Candidate Gross | Golden Gross | Delta |")
        lines.append("|---|---|---:|---:|---:|")
        for cand_row, golden_row in zip(candidate["show_rows"], report["golden_revenue_line_items"]):
            cand_gross = cand_row.get("gross_revenue") or 0.0
            golden_gross = golden_row["gross_revenue_usd"]
            delta = cand_gross - golden_gross
            lines.append(
                f"| {cand_row['show']} | {cand_row['city']} | {cand_gross:.2f} | {golden_gross:.2f} | {delta:.2f} |"
            )
        lines.append("")

    return "\n".join(lines)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    report = make_report()
    json_path = OUTPUT_DIR / f"{TASK_ID}_finance_failure_analysis.json"
    md_path = OUTPUT_DIR / f"{TASK_ID}_finance_failure_analysis.md"
    json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    md_path.write_text(render_markdown(report), encoding="utf-8")
    print(json_path)
    print(md_path)


if __name__ == "__main__":
    main()



