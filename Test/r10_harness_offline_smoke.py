"""No-provider container check of the real socket service and CLI."""
import json
from pathlib import Path
import subprocess
import sys
import time

from task_generator.production import agent_factory_tools as tools
from task_generator.production import agent_factory as f


def main():
    blocked = []
    for folder in ('/workspace/inputs', '/control/workspace/inputs'):
        try:
            (Path(folder) / 'forbidden-write.txt').write_text('must fail')
        except OSError:
            blocked.append(folder)
    if len(blocked) != 2:
        raise AssertionError('input_alias_writable')
    assert not any(p.is_file() for folder in ('/run/secrets', '/run/codex-home')
                   for p in Path(folder).rglob('*'))
    server = subprocess.Popen([sys.executable, '-m', 'task_generator.production.agent_factory_tools',
                               'serve', '/control/broker_config.json'])
    try:
        for _ in range(100):
            if Path('/run/factory/tool.sock').exists():
                break
            if server.poll() is not None:
                raise AssertionError('broker_failed')
            time.sleep(.02)
        Path('/tmp/factory-requests').mkdir(exist_ok=True)
        Path('/tmp/factory-requests/check.json').write_text('{"mode":"draft"}')
        result = subprocess.run([sys.executable, '-m', 'task_generator.production.agent_factory_tools',
                                 'check', '--input-file', '/tmp/factory-requests/check.json'],
                                stdin=subprocess.DEVNULL, capture_output=True, text=True, timeout=20)
        assert result.returncode == 0, result.stdout + result.stderr
        assert tools.client('schema', {'artifact': 'basis'})['artifact'] == 'basis'
        inspection = tools.client('inspect', {'area': 'inputs', 'path': 'record.txt', 'max_chars': 20})
        assert inspection['truncated']
        from openpyxl import Workbook, load_workbook
        workbook = Workbook()
        workbook.active['A1'] = '=SUM(1,2,3)'
        original = Path('/draft/formula.xlsx')
        workbook.save(original)
        original_sha = f.digest(original)
        cached = load_workbook(original, data_only=True)
        assert cached.active['A1'].value is None
        cached.close()
        converted = tools.client('render', {'area': 'draft', 'path': 'formula.xlsx', 'recalculate': True})
        cached = load_workbook(converted['output'], data_only=True)
        assert cached.active['A1'].value == 6
        cached.close()
        assert f.digest(original) == original_sha
        try:
            tools.client('check', {})
        except (ValueError, FileNotFoundError):
            pass
        else:
            raise AssertionError('incomplete_draft_ready')
        tools.client('consult', {'reason': 'Synthetic incomplete draft help'})
        status = tools.client('status', {})
        assert status['pending_action'] == 'consult'
        assert not any('check.json' in p for p in status['current_draft_hashes'])
        result = {'status': 'passed', 'network': 'none', 'credentials': 'not_mounted',
                  'readonly_aliases': blocked, 'tool_handoff': 'consult_queued',
                  'incomplete_ready_rejected': True, 'scratch_excluded': True,
                  'workbook': {'formula_count': 1, 'original_cached_values': 0,
                               'copy_cached_values': 1, 'copy_value_A1': 6, 'original_unchanged': True,
                               'scope': 'Synthetic A1 SUM only; no professional calculation claim'}}
        Path('/control/offline_tool_result.json').write_text(json.dumps(result))
        print(json.dumps(result))
    finally:
        server.terminate()
        server.wait(timeout=10)


if __name__ == '__main__':
    main()
