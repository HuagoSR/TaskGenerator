"""A single authorized Agent factory world, with bounded cooperative handoffs."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import shlex
import shutil
import sys
import subprocess
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'src'), str(ROOT / 'Test')]
import run_r10_process_first_pilot as base
import run_r10_task_method_pilot as previous
from task_generator.production import agent_factory as f
from task_generator.production import agent_factory_tools as ft
from task_generator.production import task_factory_harness as harness

HOST, IMAGE, IMAGE_SHA = base.HOST, base.IMAGE, base.IMAGE_SHA
DEFAULT_DEPS = '/home/huagosr/taskgenerator-data/r10-process-first/r10_process_first_pilot_20260905'
DEFAULT_LOCK = ROOT / 'artifacts/r10/r10_process_first_pilot_20260905/dependency_lock.json'
REMOTE_BASE = '/home/huagosr/taskgenerator-data/r10-agent-factory'
RUNTIME_FILES = ['src/task_generator/' + p for p in (
    '__init__.py', 'core/__init__.py', 'core/deliverable_contract.py', 'core/scenario_first.py',
    'planning/__init__.py', 'planning/rubric_compiler_v2.py', 'production/__init__.py',
    'production/task_method_pilot.py', 'production/agent_factory.py', 'production/agent_factory_tools.py'
)] + ['Test/r10_process_first_tools.py']
HARNESS_RUNTIME_FILES = [
    'src/task_generator/production/task_factory_harness.py',
    'src/task_generator/production/obligation_trace.py',
    'src/task_generator/production/quality_diagnostics.py',
    'Test/r10_calculation_replay.py',
]

COLLECTION_RETRY_DELAYS = (0, 2, 5)
LOCAL_ADMISSION_RETRY_DELAYS = (0, 0.1, 0.5, 1.0)


def _transport_retry(operation, *, deadline_epoch, label):
    """Retry only bounded transport failures; callers still validate content."""
    failures = []
    for index, delay in enumerate(COLLECTION_RETRY_DELAYS):
        if delay:
            time.sleep(delay)
        remaining = int(deadline_epoch - time.time())
        if remaining < 1:
            raise ValueError(f'collection_time_exhausted:{label}')
        try:
            result = operation(min(120, remaining))
        except subprocess.TimeoutExpired as error:
            failures.append({'attempt': index + 1, 'kind': 'timeout', 'detail': str(error)})
            continue
        if result.returncode == 0:
            return result, failures
        failures.append({'attempt': index + 1, 'kind': 'returncode',
                         'returncode': result.returncode, 'detail': result.stderr[-800:]})
        transient = (result.returncode == 255 or any(token in result.stderr.lower() for token in (
            'connection reset', 'connection timed out', 'operation timed out', 'broken pipe',
            'connection closed')))
        if not transient:
            raise ValueError(f'collection_command_failed:{label}:{result.returncode}')
    raise ValueError(f'collection_transport_failed:{label}:{json.dumps(failures, ensure_ascii=False)}')


def _admit_collected_tree(staging, raw, *, deadline_epoch):
    """Retry only transient local rename failures while preserving one atomic admission."""
    failures = []
    for delay in LOCAL_ADMISSION_RETRY_DELAYS:
        if delay:
            time.sleep(delay)
        if time.time() >= deadline_epoch:
            raise ValueError('collection_time_exhausted:local_admission')
        try:
            staging.replace(raw)
            return failures
        except PermissionError as error:
            failures.append({'kind': 'permission', 'detail': str(error)})
    raise ValueError('collection_local_admission_failed:'
                     + json.dumps(failures, ensure_ascii=False))


def _collect_remote_outputs(root, state, attempt, remote, runtime_remote, local, deps, role,
                            deadline_epoch):
    """Validate one remote inventory and atomically admit its downloaded copy."""
    deadline = min(deadline_epoch, time.time() + 360)
    inventory_code = (
        "import json,pathlib; "
        "from task_generator.production.agent_factory import files,digest; "
        "r=pathlib.Path('/raw'); out={}; "
        "names=('draft','snapshots','processes','broker.json'); "
        "[(out.__setitem__(n, {'kind':'directory','hashes':files(p)}) if p.is_dir() "
        "else out.__setitem__(n, {'kind':'file','sha256':digest(p)})) "
        "for n in names for p in (r/n,) if p.exists()]; "
        "print(json.dumps(out,sort_keys=True))"
    )
    safe_cmd = base.docker_base(deps) + [
        '-e', 'PYTHONPATH=/code/src:/code/Test:/deps/site',
        '-v', f'{runtime_remote}/runtime:/code:ro',
        '-v', f'{remote}:/raw:ro', '--entrypoint', 'python', IMAGE, '-c', inventory_code]
    result, retries = _transport_retry(
        lambda timeout: base._ssh(HOST, shlex.join(safe_cmd), timeout=timeout, check=False),
        deadline_epoch=deadline, label='inventory')
    try:
        inventory = json.loads(result.stdout)
    except json.JSONDecodeError as error:
        raise ValueError('collection_inventory_invalid_json') from error
    required = {'draft', 'broker.json'} | ({'snapshots'} if role in f.AUTHOR_ROLES else set())
    missing = sorted(required - set(inventory))
    if missing:
        raise ValueError('collection_required_output_missing:' + ','.join(missing))
    collection = attempt.setdefault('collection', {})
    collection.update(status='inventory_validated', inventory=inventory)
    collection.setdefault('transport_retries', []).extend(retries)
    collection.setdefault('confirmed', {})
    f.write(root / 'receipt.json', state)
    staging = local / 'collection'
    staging.mkdir()
    for name, expected in inventory.items():
        destination = staging / name
        result, item_retries = _transport_retry(
            lambda timeout, n=name, d=destination: base._run(
                base._scp_command(f'{HOST}:{remote}/{n}', str(d)), timeout=timeout, check=False),
            deadline_epoch=deadline, label='download_' + name)
        if expected['kind'] == 'directory':
            actual = {'kind': 'directory', 'hashes': f.files(destination)}
        else:
            actual = {'kind': 'file', 'sha256': f.digest(destination)}
        if actual != expected:
            raise ValueError('collection_download_hash_mismatch:' + name)
        collection['confirmed'][name] = actual
        collection['transport_retries'].extend(item_retries)
        f.write(root / 'receipt.json', state)
    raw = local / 'raw'
    local_retries = _admit_collected_tree(staging, raw, deadline_epoch=deadline)
    collection.setdefault('local_admission_retries', []).extend(local_retries)
    collection.update(status='collected', raw_path=raw.relative_to(root).as_posix())
    attempt.update(raw_path=raw.relative_to(root).as_posix(), raw_hashes=f.files(raw))
    f.write(root / 'receipt.json', state)
    return raw


def _verify_remote_tree_with_retry(remote_path, expected, deadline_epoch):
    code = ("import hashlib,json,pathlib,sys; r=pathlib.Path(sys.argv[1]); "
            "print(json.dumps({p.relative_to(r).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() "
            "for p in sorted(r.rglob('*')) if p.is_file()}))")
    result, retries = _transport_retry(
        lambda timeout: base._ssh(
            HOST, shlex.join(['python3', '-c', code, remote_path]), timeout=timeout, check=False),
        deadline_epoch=deadline_epoch, label='verify_remote_inputs')
    try:
        actual = json.loads(result.stdout)
    except json.JSONDecodeError as error:
        raise ValueError('remote_input_fingerprint_invalid_json') from error
    if actual != expected:
        raise ValueError('remote_input_fingerprint_changed')
    return retries


def now():
    return datetime.now(timezone.utc).isoformat()


def safe_root(run_id):
    if not re.fullmatch(r'[A-Za-z0-9_-]+', run_id):
        raise ValueError('unsafe_run_id')
    return ROOT / 'artifacts/r10' / run_id


def prepare(root, dependency_lock=DEFAULT_LOCK, dependency_remote=DEFAULT_DEPS, *, spec=None, source_bundle=None,
            batch_id=None, protocol=None, harness_options=None, synthetic_public=None):
    if root.exists():
        raise FileExistsError('run_exists')
    if not re.fullmatch(r'/home/huagosr/taskgenerator-data/[A-Za-z0-9_/-]+', dependency_remote) or '..' in dependency_remote:
        raise ValueError('unsafe_dependency_path')
    if f.read(dependency_lock)['pins'] != base.PINS:
        raise ValueError('dependency_pins_changed')
    root.mkdir()
    spec = spec or previous.method.CASES[0]
    if synthetic_public is not None:
        if ((harness_options or {}).get('purpose') not in ('tool_microtest', 'quality_diagnostic_microtest')
                or spec['seed'] != 'synthetic_tool_fixture'):
            raise ValueError('synthetic_public_requires_microtest_scope')
        shutil.copytree(synthetic_public, root / 'public')
        sources = {}
    elif source_bundle is not None:
        from task_generator.production import agent_factory_batch as batch
        sources = batch.public_inputs(sys.modules[__name__], root / 'public', spec, source_bundle)
    else:
        sources = previous.public_inputs(root / 'public', spec)
    hashes = base.code_hashes()
    for path in hashes:
        target = root / 'code' / path
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / path, target)
    shutil.copyfile(dependency_lock, root / 'dependency_lock.json')
    factory = f
    runtime_files = list(RUNTIME_FILES)
    if protocol == 'task_factory_harness_v1':
        from task_generator.production import task_factory_harness as factory
        runtime_files += HARNESS_RUNTIME_FILES
    for path in runtime_files:
        target = root / 'runtime' / path
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / path, target)
    prompt_options = harness_options or {}
    for role in factory.ROLES:
        p = root / 'prompts' / (role + '.md')
        p.parent.mkdir(exist_ok=True)
        p.write_text(factory.prompt(role, prompt_options), encoding='utf-8')
    scope = {'created_at': now(), 'seed': spec['seed'], 'host': HOST, 'image': IMAGE, 'image_sha': IMAGE_SHA,
             'remote': REMOTE_BASE + '/' + root.name, 'dependency_remote': dependency_remote,
             'dependency_lock_sha': f.digest(root / 'dependency_lock.json'), 'public_hashes': f.files(root / 'public'),
             'source_hashes': sources, 'code_hashes': hashes, 'runtime_hashes': f.files(root / 'runtime'), 'prompt_hashes': f.files(root / 'prompts'),
             'models': factory.MODELS, 'max_launches': 16 if protocol else 12,
             'production_launches': 14 if protocol else 10, 'seconds': 14400 if protocol else 21600,
             'per_launch_seconds': 1800, 'terminal_reserve_seconds': 3600, 'world_count': 1,
             'authorization': 'User approved one new procurement Agent development world, autonomous versioned revision, independent consultation, final review once and admitted trial once. No grading.',
             'cost_cap': None, 'no_automatic_expansion': True, 'domain': spec['domain']}
    if protocol:
        smoke = root / 'replay_smoke'
        (smoke / 'scripts').mkdir(parents=True)
        (smoke / 'scripts/value.py').write_text(
            "import argparse,json\np=argparse.ArgumentParser();p.add_argument('--inputs');p.parse_args()\nprint(json.dumps({'value': 3}))\n",
            encoding='utf-8')
        f.write(smoke / 'execution_manifest.json', {'version': 2, 'executions': [{
            'execution_id': 'smoke', 'script': 'value.py'}]})
        (smoke / 'output').mkdir()
        scope.update(protocol=protocol, max_upstream_revisions=1, calculation_contract_version=2,
                     obligation_trace_version=2,
                     replay_smoke_hashes=f.files(smoke),
                     authorization='User approved one new development world in this scope with isolated development solving, reproducible calculation evidence, one upstream revision, one final review and one admitted final solve. No grading.')
        scope.update(harness_options or {})
    if batch_id:
        if protocol == 'task_factory_harness_v1':
            scope.update(batch_id=batch_id, case_id=spec['id'], seconds=14400,
                         authorization='One fixed position within the user-approved task factory harness batch. No independent execution or grading.')
        else:
            scope.update(batch_id=batch_id, case_id=spec['id'], seconds=10800,
                         authorization='One fixed position within the user-approved 60-launch/15-hour Agent factory batch. No independent execution.')
    f.write(root / 'scope.json', scope)
    f.write(root / 'initial_git.json', {'index_sha256': f.digest(ROOT / '.git/index'),
            'status': base._run(['git', 'status', '--porcelain']).stdout})
    f.write(root / 'receipt.json', {'status': 'prepared', 'scope_sha': f.digest(root / 'scope.json'),
            'attempts': [], 'current': {}, 'sessions': {}, 'checks': [], 'consultations': [], 'dispositions': [],
            'next': {'role': 'world'}, 'submitted': False, 'recoveries': 0,
            'task_manual_edits': 0, 'controller_interventions': [], 'first_failure': None,
            **({'compile_stage': 'basis', 'upstream_revisions': 0, 'calculation_replays': [],
                'candidate_edit_version': scope.get('candidate_edit_version'),
                'atomic_rubric_version': scope.get('atomic_rubric_version'),
                'quality_diagnostics_version': scope.get('quality_diagnostics_version')}
               if protocol else {})})
    return scope


def verify(root):
    scope, state = f.read(root / 'scope.json'), f.read(root / 'receipt.json')
    assert state['scope_sha'] == f.digest(root / 'scope.json'), 'scope_changed'
    assert scope['code_hashes'] == base.code_hashes(), 'code_changed'
    assert scope['code_hashes'] == f.files(root / 'code'), 'code_snapshot_changed'
    assert scope['runtime_hashes'] == f.files(root / 'runtime'), 'runtime_changed'
    assert scope['public_hashes'] == f.files(root / 'public'), 'public_changed'
    assert scope['prompt_hashes'] == f.files(root / 'prompts'), 'prompts_changed'
    assert scope['dependency_lock_sha'] == f.digest(root / 'dependency_lock.json'), 'dependency_changed'
    if scope.get('protocol'):
        assert scope['replay_smoke_hashes'] == f.files(root / 'replay_smoke'), 'replay_smoke_changed'
    for relative, value in scope['source_hashes'].items():
        assert f.digest(ROOT / relative) == value, 'source_changed'
    for a in state['attempts']:
        if a.get('raw_hashes'):
            assert f.files(root / a['raw_path']) == a['raw_hashes'], 'raw_changed'
    return scope, state


def environment(root, scope):
    remote, deps = scope['remote'], scope['dependency_remote']
    identity = base._ssh(HOST, f"docker image inspect --format '{{{{.Id}}}}' {IMAGE}", timeout=60).stdout.strip().removeprefix('sha256:')
    if identity != IMAGE_SHA:
        raise ValueError('image_changed')
    lock = base._ssh(HOST, f'sha256sum {deps}/deps/lock.json', timeout=60).stdout.split()[0]
    if lock != scope['dependency_lock_sha']:
        raise ValueError('remote_dependencies_changed')
    if not (root / 'environment.json').exists():
        base._ssh(HOST, f'mkdir -p {REMOTE_BASE} && mkdir {remote}', timeout=60)
        names = ['public', 'runtime'] + (['replay_smoke'] if scope.get('protocol') else [])
        for name in names:
            base._run(base._scp_command(str(root / name), f'{HOST}:{remote}/{name}'), timeout=180)
        for action, target in [('verify-dependencies', '/deps'), ('smoke', '/tmp/factory-smoke')]:
            cmd = base.docker_base(deps) + ['-e', 'PYTHONPATH=/deps/site', '-v', f'{deps}/deps:/deps:ro',
                    '-v', f'{remote}/runtime/Test/r10_process_first_tools.py:/tools.py:ro', '-v', f'{remote}/public:/inputs:ro',
                    '--entrypoint', 'python', IMAGE, '/tools.py', action, target]
            result = base._ssh(HOST, shlex.join(cmd), timeout=180)
            (root / (action + '.log')).write_text(result.stdout + result.stderr, encoding='utf-8')
        if scope.get('protocol'):
            cmd = base.docker_base(deps) + ['-e', 'PYTHONPATH=/deps/site',
                    '-v', f'{remote}/runtime:/code:ro', '-v', f'{deps}/deps:/deps:ro',
                    '-v', f'{remote}/replay_smoke/scripts:/scripts:ro',
                    '-v', f'{remote}/replay_smoke/execution_manifest.json:/control/execution_manifest.json:ro',
                    '-v', f'{remote}/public:/inputs:ro', '-v', f'{remote}/replay_smoke/output:/result:rw',
                    '--entrypoint', 'python', IMAGE, '/code/Test/r10_calculation_replay.py',
                    '--scripts', '/scripts', '--manifest', '/control/execution_manifest.json',
                    '--inputs', '/inputs', '--output', '/result/result.json']
            result = base._ssh(HOST, shlex.join(cmd), timeout=180)
            (root / 'replay-smoke.log').write_text(result.stdout + result.stderr, encoding='utf-8')
        f.write(root / 'environment.json', {'checked_at': now(), 'image': identity, 'dependency_lock': lock,
                'offline_smoke': 'passed', 'credentials': 'not_read_or_copied'})
    base.verify_remote_tree(remote + '/runtime', scope['runtime_hashes'], time.monotonic() + 120)


def copy_file(source, target):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)


def version_path(root, entry):
    p = root / entry['path']
    if f.files(p) != entry['hashes']:
        raise ValueError('version_changed')
    return p


def build_inputs(root, state, role, target, consult_role=None, protocol=None):
    target.mkdir(parents=True)
    public = root / 'public'
    if role == 'world':
        shutil.copytree(public, target, dirs_exist_ok=True)
        return
    world = version_path(root, state['current']['world'])
    shutil.copytree(world / 'candidate', target / 'reference_files')
    effective = consult_role if role == 'consult' else role
    if role != 'solve' or protocol:
        copy_file(public / 'public_context.json', target / 'public_context.json')
    if effective in ('edit', 'compile', 'review') or role == 'consult':
        for name in ('professional_rules.json', 'sources.json'):
            copy_file(public / name, target / name)
        if state.get('candidate_edit_version') and effective in ('edit', 'compile') and (public / 'SKILL.md').is_file():
            copy_file(public / 'SKILL.md', target / 'SKILL.md')
    if effective == 'world' or (effective == 'mine' and role != 'consult'):
        return
    mine = state['current']['mine']
    use_edited = (state.get('candidate_edit_version') and 'edit' in state.get('current', {})
                  and not (effective == 'edit' and role != 'consult'))
    task_owner = state['current']['edit'] if use_edited else mine
    task = version_path(root, task_owner) / 'task.json'
    parsed = previous.method.task_result(task.parent, root / task_owner['inputs'])
    (target / 'candidate_task.md').write_text(parsed['candidate_task'], encoding='utf-8')
    f.write(target / 'deliverable_contract.json', parsed['contract'])
    if role not in ('devsolve', 'solve'):
        copy_file(task, target / 'task.json')
    if role == 'consult' and consult_role == 'edit':
        copy_file(version_path(root, state['current']['edit']) / 'edit_record.json', target / 'edit_record.json')
    if protocol == 'task_factory_harness_v1' and role == 'compile' and state.get('development_trial'):
        from task_generator.production import task_factory_harness as harness
        if not harness.trial_is_current(state):
            raise ValueError('development_trial_outside_current_candidate')
        trial = root / state['development_trial']['path']
        if f.files(trial) != state['development_trial']['hashes']:
            raise ValueError('development_trial_changed')
        shutil.copytree(trial / 'deliverable_files', target / 'development_trial/deliverable_files')
        copy_file(trial / 'diagnostic.json', target / 'development_trial/diagnostic.json')
        copy_file(version_path(root, mine) / 'design_intent.json', target / 'design_intent.json')
        initial = state.get('initial_basis_snapshot', {}).get('entry')
        if initial is not None:
            source = version_path(root, initial)
            for name in ('basis_draft.json', 'new_rubric.json', 'calculation_evidence.json'):
                copy_file(source / name, target / 'initial_basis' / name)
    if role == 'review' or (role == 'consult' and consult_role == 'compile'):
        compiler = version_path(root, state['current']['compile'])
        names = ['supervision.json', 'new_rubric.json']
        if protocol == 'task_factory_harness_v1':
            names += ['basis_draft.json', 'calculation_evidence.json']
            shutil.copytree(compiler / 'calculation_scripts', target / 'calculation_scripts')
            replay = next((row for row in reversed(state['calculation_replays'])
                           if row['compile_snapshot'] == state['current']['compile']['id']), None)
            if replay is not None:
                f.write(target / 'replay_results.json', replay['result'])
            elif role == 'review':
                raise ValueError('current_calculation_replay_required')
        for name in names:
            copy_file(compiler / name, target / name)


def agent_script(role, session_id=None):
    if session_id and not re.fullmatch(r'[A-Za-z0-9_-]+', session_id):
        raise ValueError('unsafe_session_id')
    lines = ['#!/bin/sh', 'set -eu', 'export HOME=/tmp',
             'export PYTHONPATH=/code/src:/code/Test:/deps/site', 'export PYTHONDONTWRITEBYTECODE=1',
             'export PATH=/workspace/bin:$PATH', 'export PIP_NO_INDEX=1',
             'mkdir -p /tmp/factory-requests']
    if role in ('world', 'edit', 'compile', 'devsolve', 'solve'):
        lines += ['export CODEX_HOME=/state/codex', 'mkdir -p "$CODEX_HOME"',
                  'ln -sfn /run/codex-home/auth.json "$CODEX_HOME/auth.json"']
        cmd = ['codex', 'exec'] + (['resume', session_id] if session_id else [])
        cmd += ['--dangerously-bypass-approvals-and-sandbox', '--model', 'gpt-5.6-sol', '-c', 'model_reasoning_effort="high"',
                '-c', 'project_doc_max_bytes=0', '--disable', 'plugins', '--disable', 'apps', '--disable', 'multi_agent',
                '--disable', 'skill_search', '--ignore-user-config', '--ignore-rules', '--skip-git-repo-check', '--json', '-']
        lines.append(shlex.join(cmd) + ' < /workspace/TASK.md')
    else:
        lines += ['export XDG_DATA_HOME=/state/data', 'export XDG_CONFIG_HOME=/tmp/opencode-config',
                  'export DEEPSEEK_API_KEY="$(cat /run/secrets/deepseek_api_key)"']
        cmd = ['opencode', 'run', '--pure', '--format', 'json', '--model', 'deepseek/deepseek-v4-pro', '--variant', 'max', '--auto', '--dir', '/workspace']
        if session_id:
            cmd += ['--session', session_id]
        lines.append(shlex.join(cmd) + ' "$(cat /workspace/TASK.md)"')
    return '\n'.join(lines) + '\n'


def session_identity(path, role):
    ids = set()
    for line in path.read_text(encoding='utf-8').splitlines():
        e = json.loads(line)
        if e.get('type') == 'thread.started' and e.get('thread_id'):
            ids.add(e['thread_id'])
        if e.get('sessionID'):
            ids.add(e['sessionID'])
    if len(ids) != 1:
        raise ValueError('session_identity_not_unique')
    return ids.pop()


def run_calculation_replay(root, scope, state, attempt, entry, remote):
    """Run compiler-authored calculations without network or credential mounts."""
    ordinal = attempt['ordinal']
    local = root / 'replays' / f'{ordinal:02d}'
    request = local / 'request'
    request.mkdir(parents=True)
    snapshot = version_path(root, entry)
    manifest = f.read(snapshot / 'calculation_evidence.json')
    shutil.copytree(snapshot / 'calculation_scripts', request / 'scripts')
    version = str(manifest.get('version'))
    if version == '2':
        execution_manifest = {'version': 2, 'executions': [
            {'execution_id': row['execution_id'], 'script': row['script'],
             'script_sha256': f.digest(request / 'scripts' / row['script']), 'sources': row['sources']}
            for row in manifest['executions']
        ]}
    else:
        execution_manifest = {'version': 1, 'calculations': [
            {'calculation_id': row['calculation_id'], 'script': row['script'],
             'script_sha256': f.digest(request / 'scripts' / row['script']), 'sources': row['sources']}
            for row in manifest['calculations']
        ]}
    f.write(request / 'execution_manifest.json', execution_manifest)
    output_remote = f'{remote}/replay_output/result.json'
    base._run(base._scp_command(str(request), f'{HOST}:{remote}/replay_request'), timeout=90)
    base._ssh(HOST, f'mkdir -p {remote}/replay_output', timeout=30)
    reserve = scope.get('terminal_reserve_seconds', 3600)
    available = int(state['deadline_epoch'] - time.time() - reserve)
    if available < 1:
        raise ValueError('time_budget_exhausted_before_replay')
    replay_limit = min(120, available)
    command = base.docker_base(scope['dependency_remote']) + [
        '-e', 'PYTHONPATH=/code/src:/code/Test:/deps/site',
        '-v', f'{scope["remote"]}/runtime:/code:ro',
        '-v', f'{scope["dependency_remote"]}/deps:/deps:ro',
        '-v', f'{remote}/replay_request/scripts:/scripts:ro',
        '-v', f'{remote}/replay_request/execution_manifest.json:/control/execution_manifest.json:ro',
        '-v', f'{remote}/workspace/inputs/reference_files:/inputs/reference_files:ro',
        '-v', f'{remote}/replay_output:/result:rw', '--entrypoint', 'python', IMAGE,
        '/code/Test/r10_calculation_replay.py', '--scripts', '/scripts',
        '--manifest', '/control/execution_manifest.json', '--inputs', '/inputs',
        '--output', '/result/result.json', '--timeout-seconds', str(replay_limit)]
    replay_name = f'{root.name}-replay-{ordinal:02d}'
    command[2:2] = ['--name', replay_name, '--stop-timeout', '1']
    try:
        result = base._ssh(
            HOST,
            f'timeout --signal=TERM --kill-after=2s {replay_limit + 5}s {shlex.join(command)}',
            timeout=replay_limit + 30, check=False)
    finally:
        base._ssh(HOST, f'docker rm -f {replay_name}', timeout=30, check=False)
    version_path(root, entry)
    base.verify_remote_tree(remote + '/replay_request', f.files(request), time.monotonic() + 60)
    if attempt.get('inputs'):
        base.verify_remote_tree(remote + '/workspace/inputs/reference_files',
                                f.files(root / attempt['inputs'] / 'reference_files'), time.monotonic() + 60)
    exists = base._ssh(HOST, f'test -e {output_remote}', timeout=30, check=False).returncode == 0
    if exists:
        base._run(base._scp_command(f'{HOST}:{output_remote}', str(local / 'execution_results.json')), timeout=60)
        execution = f.read(local / 'execution_results.json')
    else:
        execution = {'status': 'error', 'calculations': [], 'network': 'none',
                     'credentials': 'not_mounted', 'diagnostic': (result.stdout + result.stderr)[-2000:]}
        f.write(local / 'execution_results.json', execution)
    payload = compare_calculation_results(manifest, execution)
    payload.update(container_returncode=result.returncode, execution_result_sha256=f.digest(
        local / 'execution_results.json'), request_hashes=f.files(request))
    f.write(local / 'replay_results.json', payload)
    record = {'compile_snapshot': entry['id'], 'status': payload['status'], 'result': payload,
              'result_sha256': f.digest(local / 'replay_results.json'), 'elapsed_in_case_clock': True}
    state.setdefault('calculation_replays', []).append(record)
    return record


def compare_calculation_results(manifest, execution):
    if execution.get('status') != 'executed':
        rows = execution.get('executions', execution.get('calculations', []))
        return {'status': execution.get('status', 'error'), 'calculations': rows,
                'network': 'none', 'credentials': 'not_mounted'}
    if str(manifest.get('version')) == '2':
        from task_generator.production import task_factory_harness as factory
        actual_rows = execution.get('executions', [])
        actual = {row.get('execution_id'): row for row in actual_rows}
        expected_ids = [row['execution_id'] for row in manifest['executions']]
        if len(actual) != len(actual_rows) or set(actual) != set(expected_ids):
            return {'status': 'error', 'calculations': actual_rows,
                    'diagnostic': 'execution_result_ids_do_not_match_manifest',
                    'network': 'none', 'credentials': 'not_mounted'}
        rows = []
        for item in manifest['calculations']:
            execution_row = actual[item['execution_id']]
            try:
                value = factory.resolve_json_pointer(execution_row.get('output'), item['result_pointer'])
            except ValueError as error:
                rows.append({'calculation_id': item['calculation_id'], 'execution_id': item['execution_id'],
                             'result_pointer': item['result_pointer'], 'status': 'error',
                             'diagnostic': str(error), 'matched': False})
                continue
            target, tolerance = item['expected']['value'], item['expected']['tolerance']
            if (isinstance(value, (int, float)) and not isinstance(value, bool)
                    and isinstance(target, (int, float)) and not isinstance(target, bool)):
                matched = abs(value - target) <= tolerance
            else:
                matched = type(value) is type(target) and value == target and tolerance == 0
            rows.append({'calculation_id': item['calculation_id'], 'execution_id': item['execution_id'],
                         'result_pointer': item['result_pointer'], 'script_sha256': execution_row.get('script_sha256'),
                         'value': value, 'expected': target, 'tolerance': tolerance, 'matched': matched})
        return {'status': 'passed' if rows and all(row['matched'] for row in rows) else 'mismatch',
                'calculations': rows, 'executions': actual_rows,
                'network': 'none', 'credentials': 'not_mounted'}
    actual_rows = execution.get('calculations', [])
    actual = {row.get('calculation_id'): row for row in actual_rows}
    expected_ids = [row['calculation_id'] for row in manifest['calculations']]
    if len(actual) != len(actual_rows) or set(actual) != set(expected_ids):
        return {'status': 'error', 'calculations': actual_rows,
                'diagnostic': 'execution_result_ids_do_not_match_manifest',
                'network': 'none', 'credentials': 'not_mounted'}
    rows = []
    for item in manifest['calculations']:
        row = actual[item['calculation_id']]
        value, expected = row.get('value'), item['expected']
        target, tolerance = expected['value'], expected['tolerance']
        if (isinstance(value, (int, float)) and not isinstance(value, bool)
                and isinstance(target, (int, float)) and not isinstance(target, bool)):
            matched = abs(value - target) <= tolerance
        else:
            matched = type(value) is type(target) and value == target and tolerance == 0
        rows.append({**row, 'expected': target, 'tolerance': tolerance, 'matched': matched})
    return {'status': 'passed' if all(row['matched'] for row in rows) else 'mismatch',
            'calculations': rows, 'network': 'none', 'credentials': 'not_mounted'}


def checked_output(attempt, *args):
    try:
        return ft.check(*args)
    except ValueError:
        attempt['failure_kind'] = 'output_contract'
        raise


def run_turn(root, scope, state, request, batch_root=None):
    factory = f
    if scope.get('protocol') == 'task_factory_harness_v1':
        from task_generator.production import task_factory_harness as factory
    if scope.get('purpose') == 'tool_microtest':
        if not state.get('microtest_execution'):
            raise ValueError('microtest_entry_required')
        state['budget_limits'] = {key: scope[key] for key in (
            'max_launches', 'production_launches', 'seconds', 'per_launch_seconds', 'terminal_reserve_seconds')}
    role = request['role']
    from task_generator.production.quality_diagnostics import micro_enabled
    diagnostic_micro = micro_enabled(scope)
    if diagnostic_micro:
        from run_r10_quality_diagnostic_micro import check_batch_launch
        deadline = check_batch_launch(root, scope)
        if state.get('deadline_epoch') != deadline:
            raise ValueError('diagnostic_batch_deadline_changed')
    if diagnostic_micro and role not in ('compile', 'review'):
        raise ValueError('diagnostic_micro_role_not_allowed')
    feedback = factory.role_feedback(state, request)
    session = factory.validate_session(state, role, root.name) if role in factory.AUTHOR_ROLES else None
    if session and not re.fullmatch(re.escape(scope['remote']) + r'/state_\d{2}', session['storage']):
        raise ValueError('session_storage_outside_scope')
    if role == 'consult' and (request.get('snapshot') != state['current'][request['consult_role']]['id']
                             or request.get('parents') != factory.expected_parents(state, request['consult_role'])):
        raise ValueError('consultation_outside_current_version')
    limit = factory.cap(state, role)
    if batch_root is not None:
        if scope.get('protocol') == 'task_factory_harness_v1':
            limit = min(limit, factory.batch_cap(batch_root, root, role))
        else:
            from task_generator.production import agent_factory_batch as batch
            limit = min(limit, batch.cap(sys.modules[__name__], batch_root, root, role))
    ordinal = len(state['attempts']) + 1
    name = f'{root.name}-{ordinal:02d}'
    remote = f'{scope["remote"]}/turn_{ordinal:02d}'
    local = root / 'turns' / f'{ordinal:02d}'
    workspace = local / 'workspace'
    workspace.mkdir(parents=True)
    build_inputs(root, state, role, workspace / 'inputs', request.get('consult_role'), scope.get('protocol'))
    if diagnostic_micro and role == 'compile':
        shutil.copytree(root / 'fixture/diagnostic_context', workspace / 'inputs/diagnostic_context')
    if feedback is not None:
        f.write(workspace / 'feedback.json', feedback)
    task = (root / 'prompts' / (role + '.md')).read_text(encoding='utf-8')
    if feedback is not None:
        task += '\nRead /workspace/feedback.json as feedback, investigate it rather than blindly accepting it.\n'
    if role == 'consult':
        task += '\nQuestion: ' + request['question']
    production_cap = scope.get('production_launches', 10)
    if scope.get('protocol') == 'task_factory_harness_v1':
        workflow = factory.workflow_status(state, role)
        task += ('\nRead factory-tools status before choosing the next action. '
                 f'Production launches remaining after this launch: '
                 f'{workflow["budget"]["production_remaining_after_current"]}. '
                 f'Total launches remaining after this launch: '
                 f'{workflow["budget"]["launches_remaining_after_current"]}.\n')
    else:
        workflow = None
        task += f'\nProduction launches remaining including this one: {production_cap - sum(a["role"] not in ("review", "solve") for a in state["attempts"])}.\n'
    (workspace / 'TASK.md').write_text(task, encoding='utf-8')
    parents = factory.expected_parents(state, role) if role in factory.AUTHOR_ROLES else {}
    consultations = factory.relevant_consultations(state, role) if role in factory.AUTHOR_ROLES else []
    current_id = state.get('current', {}).get(role, {}).get('id')
    replay_passed = any(row.get('compile_snapshot') == current_id and row.get('status') == 'passed'
                        for row in state.get('calculation_replays', []))
    passed_replay_snapshots = sorted({row['compile_snapshot'] for row in state.get('calculation_replays', [])
                                      if row.get('status') == 'passed'})
    f.write(workspace / 'role.json', {'role': role, 'parents': parents,
                                    'diagnostic_micro_version': scope.get('diagnostic_micro_version'),
                                    'diagnostic_subjects': scope.get('diagnostic_subjects'),
                                    'event_version': 2,
                                    'purpose': scope.get('purpose'),
                                    'remaining_production_launches': (
                                        workflow['budget']['production_remaining_after_current'] if workflow
                                        else production_cap - len(state['attempts'])),
                                    'consultations': consultations, 'protocol': scope.get('protocol'),
                                    'development_trial_current': (
                                        factory.trial_is_current(state) if workflow and 'compile' in state['current'] else False),
                                    'passed_replay_snapshots': passed_replay_snapshots,
                                    'calculation_contract_version': scope.get('calculation_contract_version'),
                                    'obligation_trace_version': scope.get('obligation_trace_version'),
                                    'candidate_edit_version': scope.get('candidate_edit_version'),
                                    'atomic_rubric_version': scope.get('atomic_rubric_version'),
                                    'quality_diagnostics_version': scope.get('quality_diagnostics_version'),
                                    'initial_basis_snapshot': state.get('initial_basis_snapshot', {}).get('compile_snapshot'),
                                    'deadline_epoch': state.get('deadline_epoch'),
                                    'workflow_status': workflow})
    session_id = session['id'] if session else None
    storage = session['storage'] if session else f'{scope["remote"]}/state_{ordinal:02d}'
    (workspace / 'agent.sh').write_text(agent_script(role, session_id), encoding='utf-8', newline='\n')
    (workspace / 'bin').mkdir()
    (workspace / 'bin/factory-tools').write_text('#!/bin/sh\nexec python -m task_generator.production.agent_factory_tools "$@"\n', encoding='utf-8', newline='\n')
    draft = local / 'draft'
    if role in state['current'] and role in factory.AUTHOR_ROLES:
        shutil.copytree(version_path(root, state['current'][role]), draft)
    elif (scope.get('purpose') == 'quality_diagnostic_microtest' and role == 'compile'
          and scope.get('initial_compile_draft')):
        initial = root / scope['initial_compile_draft']
        if f.files(initial) != scope.get('initial_compile_draft_hashes'):
            raise ValueError('quality_microtest_initial_draft_changed')
        shutil.copytree(initial, draft)
    else:
        draft.mkdir()
    process = state['current'].get('world', {}).get('process') if role == 'world' else None
    dispositions = [d for d in state.get('dispositions', []) if d['request_id'] in {c['request_id'] for c in consultations}]
    f.write(local / 'broker_config.json', {'root': remote, 'role': role, 'parents': parents, 'process': process,
                                         'diagnostic_micro_version': scope.get('diagnostic_micro_version'),
                                         'diagnostic_subjects': scope.get('diagnostic_subjects'),
                                         'purpose': scope.get('purpose'),
                                         'starting_snapshot': state.get('current', {}).get(role),
                                         'consultations': consultations, 'dispositions': dispositions,
                                         'protocol': scope.get('protocol'),
                                         'development_trial_current': (
                                             factory.trial_is_current(state) if workflow and 'compile' in state['current'] else False),
                                         'passed_replay_snapshots': passed_replay_snapshots,
                                         'calculation_contract_version': scope.get('calculation_contract_version'),
                                         'obligation_trace_version': scope.get('obligation_trace_version'),
                                         'candidate_edit_version': scope.get('candidate_edit_version'),
                                         'atomic_rubric_version': scope.get('atomic_rubric_version'),
                                         'quality_diagnostics_version': scope.get('quality_diagnostics_version'),
                                         'initial_basis_snapshot': state.get('initial_basis_snapshot', {}).get('compile_snapshot'),
                                         'deadline_epoch': state.get('deadline_epoch'),
                                         'workflow_status': workflow})
    attempt = {'ordinal': ordinal, 'role': role, 'status': 'staging', 'started_at': now(), 'remote': remote,
               'container': name, 'timeout_seconds': limit, 'resumed_session': session_id,
               'input_hashes': f.files(workspace), 'inputs': (workspace / 'inputs').relative_to(root).as_posix()}
    state['attempts'].append(attempt)
    f.write(root / 'receipt.json', state)
    base._ssh(HOST, f'mkdir {remote} && mkdir -p {storage}', timeout=60)
    for p in ('workspace', 'draft', 'broker_config.json'):
        base._run(base._scp_command(str(local / p), f'{HOST}:{remote}/{p}'), timeout=180)
    base._ssh(HOST, f'chmod +x {remote}/workspace/bin/factory-tools', timeout=30)
    base.verify_remote_tree(remote + '/workspace', attempt['input_hashes'], time.monotonic() + 60)
    if (root / 'STOP').exists():
        raise ValueError('user_stopped_before_semantic_start')
    if state.get('deadline_epoch') is None:
        state.update(started_at=now(), deadline_epoch=time.time() + scope.get('seconds', 21600))
    if batch_root is not None:
        if scope.get('protocol') == 'task_factory_harness_v1':
            limit = min(limit, factory.batch_cap(
                batch_root, root, role, started=True, appended=False))
        else:
            limit = min(limit, batch.cap(
                sys.modules[__name__], batch_root, root, role, started=True, appended=True))
    reserve = scope.get('terminal_reserve_seconds', 3600) if role not in ('review', 'solve') else (1800 if role == 'review' else 0)
    limit = min(limit, int(state['deadline_epoch'] - time.time() - reserve))
    if limit < 1:
        raise ValueError('time_budget_exhausted_before_semantic_start')
    # Clocks start after staging, then the exact launch-time identity is frozen.
    # Re-upload only controller-owned configuration before either process starts.
    for config_path in (workspace / 'role.json', local / 'broker_config.json'):
        config = f.read(config_path)
        config['deadline_epoch'] = state['deadline_epoch']
        config['launch_deadline_epoch'] = time.time() + limit
        if workflow:
            runtime_state = {**state, 'attempts': [*state['attempts'][:-1],
                {**state['attempts'][-1], 'status': 'started', 'phase': 'native'}]}
            config['workflow_status'] = factory.workflow_status(runtime_state, role, current_launch_included=True)
        f.write(config_path, config)
        destination = remote + ('/workspace/role.json' if config_path.name == 'role.json' else '/broker_config.json')
        base._run(base._scp_command(str(config_path), f'{HOST}:{destination}'), timeout=60)
    attempt['input_hashes'] = f.files(workspace)
    attempt['broker_config_sha256'] = f.digest(local / 'broker_config.json')
    base.verify_remote_tree(remote + '/workspace', attempt['input_hashes'], time.monotonic() + 60)
    limit = min(limit, int(state['deadline_epoch'] - time.time() - reserve))
    if limit < 1:
        raise ValueError('time_budget_exhausted_before_semantic_start')
    if diagnostic_micro:
        check_batch_launch(root, scope)
    attempt.update(status='started', started_at=now(), timeout_seconds=limit, phase='native', semantic_started=True)
    f.write(root / 'receipt.json', state)
    if diagnostic_micro:
        from run_r10_quality_diagnostic_micro import persist_batch_launches
        persist_batch_launches(root)
    deps = scope['dependency_remote']
    cmd = base.docker_base(deps, network=True)
    cmd[2:2] = ['--name', name, '--stop-timeout', '1']
    cmd += ['-v', f'{scope["remote"]}/runtime:/code:ro', '-v', f'{deps}/deps:/deps:ro',
            '-v', f'{remote}/workspace:/workspace:ro', '-v', f'{remote}/draft:/draft:rw',
            '-v', f'{remote}/socket:/run/factory:ro', '-v', f'{storage}:/state:rw', '-w', '/workspace']
    cmd += (['-v', '/home/huagosr/taskgenerator-secrets/codex-auth-current:/run/codex-home:rw'] if role in ('world', 'edit', 'compile', 'devsolve', 'solve')
            else ['-v', '/home/huagosr/taskgenerator-secrets/deepseek_api_key:/run/secrets/deepseek_api_key:ro'])
    cmd += ['--entrypoint', '/bin/sh', IMAGE, '/workspace/agent.sh']
    command = f'''set -eu
export PYTHONDONTWRITEBYTECODE=1
export PYTHONPATH={scope['remote']}/runtime/src
python3 -m task_generator.production.agent_factory_tools serve {remote}/broker_config.json > {remote}/broker.log 2>&1 &
factory_broker_pid=$!
trap 'timeout 5s docker stop -t 1 {name} >/dev/null 2>&1 || true; kill "$factory_broker_pid" 2>/dev/null || true' EXIT
for factory_wait in $(seq 1 100); do [ -S {remote}/socket/tool.sock ] && break; sleep 0.1; done
test -S {remote}/socket/tool.sock
timeout --signal=TERM --kill-after=2s {max(1, limit-5)}s {shlex.join(cmd)} > {remote}/agent.jsonl 2> {remote}/stderr.txt
'''
    print(f'{role}: launch {ordinal}/{scope.get("max_launches", 12)}' + (' resume' if session_id else ' fresh'), flush=True)
    start = time.monotonic()
    result = base._ssh(HOST, command, timeout=limit, check=False)
    attempt.update(returncode=result.returncode, elapsed_seconds=time.monotonic() - start)
    attempt['phase'] = 'collection_logs'
    collection_deadline = state['deadline_epoch']
    if batch_root is not None:
        batch_deadline = f.read(batch_root / 'receipt.json').get('deadline_epoch')
        if batch_deadline is not None:
            collection_deadline = min(collection_deadline, batch_deadline)
    attempt['collection'] = {'status': 'collecting_logs', 'confirmed': {}, 'transport_retries': []}
    f.write(root / 'receipt.json', state)
    for p in ('agent.jsonl', 'stderr.txt', 'broker.log'):
        transferred, retries = _transport_retry(
            lambda timeout, name=p: base._run(
                base._scp_command(f'{HOST}:{remote}/{name}', str(local / name)),
                timeout=timeout, check=False),
            deadline_epoch=collection_deadline, label='download_' + p)
        attempt['collection']['transport_retries'].extend(retries)
        attempt['collection']['confirmed'][p] = {'sha256': f.digest(local / p)}
        f.write(root / 'receipt.json', state)
    attempt['usage'] = base.usage(local / 'agent.jsonl')
    native, stderr = (local / 'agent.jsonl').read_text(encoding='utf-8'), (local / 'stderr.txt').read_text(encoding='utf-8')
    authentication_failed = previous.authentication_failed(result.returncode, attempt['usage'], stderr, native)
    native_completed = (not authentication_failed and not result.returncode
                        and attempt['usage']['completed'] and not attempt['usage']['error_events'])
    if native_completed:
        actual_id = session_identity(local / 'agent.jsonl', role)
        if session_id and actual_id != session_id:
            raise ValueError('resumed_identity_changed')
        attempt.update(status='native_completed', native_status='completed', native_completed_at=now(),
                       acceptance_status='pending', session_id=actual_id, phase='collection_outputs')
        if role in factory.AUTHOR_ROLES:
            state['sessions'][role] = {'id': actual_id, 'storage': storage, 'scope': root.name, 'role': role,
                                       'parents': parents, 'normal_end': True}
    else:
        attempt.update(native_status='failed', acceptance_status='not_applicable', phase='collection_outputs')
    f.write(root / 'receipt.json', state)
    raw = _collect_remote_outputs(
        root, state, attempt, remote, scope['remote'], local, deps, role, collection_deadline)
    attempt['collection']['input_verification_retries'] = _verify_remote_tree_with_retry(
        remote + '/workspace', attempt['input_hashes'], collection_deadline)
    f.write(root / 'receipt.json', state)
    if authentication_failed:
        raise ValueError('authentication_failed')
    if not native_completed:
        raise ValueError('started_session_failed_no_redraw')
    journal = f.read(raw / 'broker.json') if (raw / 'broker.json').exists() else {'events': [], 'pending': None}
    for event in journal['events']:
        if event['action'] == 'log-check':
            state['checks'].append({'launch': ordinal, **event['result']['recorded']})
    for disposition in journal.get('dispositions', []):
        if disposition not in state.setdefault('dispositions', []):
            state['dispositions'].append(disposition)
    attempt['phase'] = 'output_check'
    if diagnostic_micro and role == 'compile':
        from task_generator.production.quality_diagnostics import micro_result
        pending = journal.get('pending') or {}
        if pending.get('action') != 'finish' or pending.get('hashes') != f.files(raw / 'draft'):
            raise ValueError('diagnostic_output_not_finished_or_changed')
        outcome = micro_result(raw / 'draft', workspace / 'inputs', scope)
        entry = pending['snapshot_entry']
        entry.update(path=(raw / 'snapshots' / entry['id']).relative_to(root).as_posix(), inputs=attempt['inputs'])
        if f.files(version_path(root, entry)) != pending['hashes']:
            raise ValueError('diagnostic_snapshot_changed')
        factory.accept_snapshot(state, role, entry)
        state['diagnostic_result'] = outcome
        if outcome['status'] == 'upstream_issue':
            replay = {'status': 'not_run', 'reason': 'diagnostic_upstream_issue'}
            state.update(status='diagnostic_upstream_issue', next=None)
        else:
            replay = run_calculation_replay(root, scope, state, attempt, entry, remote)
            state.update(status='running' if replay['status'] == 'passed' else 'diagnostic_replay_failed',
                         next={'role': 'review'} if replay['status'] == 'passed' else None)
        attempt['replay'] = replay
        attempt.update(outcome=outcome, status='completed', acceptance_status='accepted', completed_at=now())
        f.write(root / 'receipt.json', state)
        return
    if role in factory.AUTHOR_ROLES:
        pending = journal['pending']
        if not pending:
            raise ValueError('normal_turn_without_explicit_next_action')
        if pending['action'] == 'handoff' and pending.get('target') == 'stop':
            state.update(status='incomplete', stop_reason=pending['reason'], next=None)
            attempt.update(status='completed', acceptance_status='accepted', completed_at=now())
            f.write(root / 'receipt.json', state)
            return
        entry = pending['snapshot_entry']
        entry.update(path=(raw / 'snapshots' / entry['id']).relative_to(root).as_posix(), inputs=attempt['inputs'])
        if f.files(raw / 'draft') != entry['hashes']:
            raise ValueError('draft_changed_after_yield')
        action = pending['action']
        outcome = checked_output(attempt, role, version_path(root, entry), workspace / 'inputs',
                           scope.get('protocol'), 'draft' if action == 'consult' else 'ready',
                           scope.get('calculation_contract_version'), scope.get('obligation_trace_version'),
                           state.get('initial_basis_snapshot', {}).get('compile_snapshot'),
                           scope.get('atomic_rubric_version'), scope.get('quality_diagnostics_version'))
        attempt['outcome'] = outcome
        factory.accept_snapshot(state, role, entry)
        if scope.get('protocol') == 'task_factory_harness_v1':
            factory.validate_next_action(
                factory.workflow_status(state, role, current_launch_included=True),
                role, action, pending)
        if action == 'consult':
            question = ('Independently read these business records. Locate material contradictions, provenance and knowledge-time problems, useful work and concrete information gaps. Ask clarifying questions when evidence is insufficient.' if role == 'world' else pending['reason'])
            state['next'] = {'role': 'consult', 'consult_role': role, 'question': question,
                             'request_id': pending['request_id'], 'snapshot': entry['id'], 'parents': parents}
        elif action == 'replay':
            replay = run_calculation_replay(root, scope, state, attempt, entry, remote)
            next_action = pending.get('next_action', 'return')
            if next_action == 'stop' and scope.get('purpose') == 'tool_microtest':
                state.update(status='microtest_completed' if replay['status'] == 'passed' else 'microtest_failed', next=None)
            elif replay['status'] != 'passed' or next_action == 'return':
                state['next'] = {'role': 'compile', 'feedback': replay['result'],
                                 'feedback_origin': 'replay', 'feedback_role': 'compile',
                                 'feedback_snapshot': entry['id']}
            elif next_action == 'development_trial':
                if state.get('initial_basis_snapshot') is None:
                    state['initial_basis_snapshot'] = {
                        'compile_snapshot': entry['id'],
                        'entry': dict(entry),
                        'candidate_parents': factory.expected_parents(state, 'compile'),
                        'basis_sha256': f.digest(version_path(root, entry) / 'basis_draft.json'),
                        'rubric_sha256': f.digest(version_path(root, entry) / 'new_rubric.json'),
                    }
                state['compile_stage'] = 'development_trial'
                state['next'] = {'role': 'devsolve', 'compile_snapshot': entry['id'],
                                 'candidate_parents': factory.expected_parents(state, 'compile')}
            else:
                if outcome['status'] != 'completed':
                    raise ValueError('cannot_submit_upstream_issue')
                factory.assert_submission(state)
                state.update(submitted=True, submitted_at=now(), next={'role': 'review'})
        elif action == 'request-development-trial':
            if not any(row.get('compile_snapshot') == entry['id'] and row.get('status') == 'passed'
                       for row in state.get('calculation_replays', [])):
                raise ValueError('current_calculation_replay_required')
            if state.get('initial_basis_snapshot') is None:
                state['initial_basis_snapshot'] = {
                    'compile_snapshot': entry['id'],
                    'entry': dict(entry),
                    'candidate_parents': factory.expected_parents(state, 'compile'),
                    'basis_sha256': f.digest(version_path(root, entry) / 'basis_draft.json'),
                    'rubric_sha256': f.digest(version_path(root, entry) / 'new_rubric.json'),
                }
            state['compile_stage'] = 'development_trial'
            state['next'] = {'role': 'devsolve', 'compile_snapshot': entry['id'],
                             'candidate_parents': factory.expected_parents(state, 'compile')}
        elif action == 'submit':
            if outcome['status'] != 'completed':
                raise ValueError('cannot_submit_upstream_issue')
            factory.assert_submission(state)
            state.update(submitted=True, submitted_at=now(), next={'role': 'review'})
        elif pending['target'] == 'stop':
            state.update(status='incomplete', stop_reason=pending['reason'], next=None)
        else:
            state['next'] = factory.handoff_request(state, role, pending['target'], pending['reason'])
    else:
        pending = journal.get('pending') or {}
        if pending.get('action') != 'finish' or pending.get('hashes') != f.files(raw / 'draft'):
            raise ValueError('readonly_output_not_finished_or_changed')
        outcome = checked_output(attempt, role, raw / 'draft', workspace / 'inputs', scope.get('protocol'),
                           'ready', scope.get('calculation_contract_version'),
                           scope.get('obligation_trace_version'), None,
                           scope.get('atomic_rubric_version'), scope.get('quality_diagnostics_version'))
        attempt['outcome'] = outcome
        if role == 'devsolve':
            destination = root / 'development_trials' / f'{ordinal:02d}'
            shutil.copytree(raw / 'draft', destination)
            trial = {'compile_snapshot': request['compile_snapshot'],
                     'candidate_parents': request['candidate_parents'],
                     'path': destination.relative_to(root).as_posix(), 'hashes': f.files(destination),
                     'outcome': outcome}
            state['development_trial'] = trial
            current_compile = state['current']['compile']['id']
            state['compile_stage'] = 'refine'
            state['next'] = {'role': 'compile', 'feedback': {
                                'status': 'development_trial_complete',
                                'diagnostic': f.read(destination / 'diagnostic.json')},
                             'feedback_origin': 'development_trial', 'feedback_role': 'compile',
                             'feedback_snapshot': current_compile}
        elif role == 'consult':
            feedback = f.read(raw / 'draft/consultation.json')
            if request['snapshot'] != state['current'][request['consult_role']]['id']:
                raise ValueError('stale_consultation_result')
            state['consultations'].append({'launch': ordinal, 'role': request['consult_role'], 'snapshot': request['snapshot'],
                                          'parents': request['parents'], 'request_id': request['request_id'],
                                          'finding_ids': [str(i) for i in range(1, len(feedback['findings']) + 1)]})
            feedback = {**feedback, 'request_id': request['request_id'],
                        'findings': [{'finding_id': str(i), **item} for i, item in enumerate(feedback['findings'], 1)]}
            owner = request['consult_role']
            state['next'] = {'role': owner, 'feedback': feedback, 'feedback_origin': 'consultation',
                             'feedback_role': owner, 'feedback_snapshot': state['current'][owner]['id']}
        elif role == 'review':
            state['review'] = outcome
            if diagnostic_micro:
                state.update(status='diagnostic_review_complete', next=None)
            elif outcome['quality'] == 'pass':
                export_package(root, state, workspace / 'inputs')
                state['next'] = {'role': 'solve'}
            else:
                state.update(status='quality_not_passed', next=None)
        else:
            state.update(trial=outcome, status='completed', next=None)
    attempt.update(status='completed', acceptance_status='accepted', completed_at=now())
    f.write(root / 'receipt.json', state)


def export_package(root, state, inputs):
    package = root / 'package'
    package.mkdir()
    shutil.copytree(inputs / 'reference_files', package / 'reference_files')
    atomic = f.read(root / 'scope.json').get('atomic_rubric_version') if (root / 'scope.json').exists() else None
    if atomic and (inputs / 'public_context.json').is_file():
        copy_file(inputs / 'public_context.json', package / 'reference_files/public_context.json')
    (package / 'deliverable_files').mkdir()
    rubric = f.read(inputs / 'new_rubric.json')
    task = f.read(inputs / 'task.json')
    formal_rubric = json.dumps({
        'rubric_version': rubric.get('rubric_version'),
        'scoring': rubric.get('scoring', 'integer_boundaries'),
        'criteria': rubric['criteria'],
    }, ensure_ascii=False, sort_keys=True)
    f.write(package / 'dataset_row.json', {'task_id': 'anonymous_task', 'title': task['title'],
            'sector': f.read(root / 'scope.json').get('domain', 'procurement_operations') if (root / 'scope.json').exists() else 'procurement_operations',
            'occupation': f.read(root / 'public/public_context.json')['role'],
            'prompt': (inputs / 'candidate_task.md').read_text(encoding='utf-8'),
            'reference_files': ['reference_files/' + p for p in f.files(package / 'reference_files')],
            'deliverable_files': [r['relative_path'] for r in f.read(inputs / 'deliverable_contract.json')['deliverables']],
            'rubric': formal_rubric if atomic else 'Original Rubric V2 criteria retained; provisional LLM-proxy.',
            'rubric_json': json.dumps(rubric['criteria'], ensure_ascii=False),
            'extra': {'internal_research_only': True, 'professional_status': 'provisional/LLM-proxy'}})
    copy_file(inputs / 'new_rubric.json', package / ('atomic_rubric_v1.json' if atomic else 'rubric_v2.json'))
    copy_file(inputs / 'deliverable_contract.json', package / 'deliverable_contract.json')
    if (inputs / 'calculation_evidence.json').is_file():
        copy_file(inputs / 'calculation_evidence.json', package / 'calculation_evidence.json')
        copy_file(inputs / 'replay_results.json', package / 'replay_results.json')
        shutil.copytree(inputs / 'calculation_scripts', package / 'calculation_scripts')
    state['package_hashes'] = f.files(package)


def record_execution_failure(root, state, error):
    reason = type(error).__name__ + ':' + str(error)
    if state['attempts'] and state['attempts'][-1]['status'] in ('started', 'staging'):
        state['attempts'][-1].update(status='incomplete', reason=reason)
    elif state['attempts'] and state['attempts'][-1]['status'] == 'native_completed':
        last_attempt = state['attempts'][-1]
        if str(last_attempt.get('phase', '')).startswith('collection'):
            last_attempt.setdefault('collection', {}).update(status='collection_failed', reason=reason)
            last_attempt.update(acceptance_status='pending', acceptance_reason=reason)
        else:
            last_attempt.update(acceptance_status='failed', acceptance_reason=reason)
    state.update(status='incomplete', stop_reason=reason, stopped_at=now())
    critical = ('authentication_failed', 'batch_', 'feedback_outside', 'session_outside', 'session_storage',
                'resumed_identity', 'consultation_outside', 'stale_consultation', 'source_changed', 'code_changed',
                'raw_changed', 'runtime_changed', 'scope_changed', 'version_changed', 'public_changed',
                'prompts_changed', 'dependencies_changed', 'path_escape', 'linked_', 'hardlinked_',
                'session_identity_not_unique', 'already_submitted', 'draft_changed',
                'readonly_output_not_finished_or_changed')
    last = state['attempts'][-1] if state['attempts'] else {}
    if any(token in reason for token in critical):
        category = 'global'
    elif ('normal_turn_without_explicit_next_action' in reason
          and state.get('attempts')
          and (root / 'turns' / f'{state["attempts"][-1]["ordinal"]:02d}' / 'agent.jsonl').is_file()
          and 'factory-tools: not found' in (root / 'turns' / f'{state["attempts"][-1]["ordinal"]:02d}' / 'agent.jsonl').read_text(encoding='utf-8')):
        category = 'global'
    elif ('budget_exhausted' in reason or 'started_session_failed' in reason or 'user_stopped' in reason
          or last.get('failure_kind') == 'output_contract'):
        category = 'case'
    else:
        category = 'controller'
    state['failure_category'] = category
    state['first_failure'] = state.get('first_failure') or {
        'launch': len(state['attempts']), 'reason': reason}
    return reason


def execute(root, *, batch_root=None):
    scope, state = verify(root)
    if scope.get('batch_id'):
        if batch_root is None or batch_root.name != scope['batch_id'] or root.parent.parent.resolve() != batch_root.resolve():
            raise ValueError('batch_execution_required')
    elif batch_root is not None:
        raise ValueError('not_a_batch_scope')
    if state['status'] not in ('prepared', 'paused'):
        raise ValueError('scope_not_executable')
    if state['status'] == 'paused':
        (root / 'STOP').unlink(missing_ok=True)
    ready = f.read(root / 'readiness.json')
    if not ready.get('tests_passed') or ready.get('code_hashes') != scope['code_hashes']:
        raise ValueError('readiness_required')
    lock = root / 'execution.lock'
    with lock.open('x') as handle:
        handle.write(str(__import__('os').getpid()))
    try:
        environment(root, scope)
        state['status'] = 'running'
        f.write(root / 'receipt.json', state)
        while state.get('next'):
            if (root / 'STOP').exists():
                state.update(status='paused', paused_at=now())
                break
            verify(root)
            try:
                if batch_root is None:
                    run_turn(root, scope, state, state['next'])
                else:
                    run_turn(root, scope, state, state['next'], batch_root=batch_root)
            except (subprocess.TimeoutExpired, subprocess.CalledProcessError) as error:
                attempt = state['attempts'][-1]
                recoverable = (attempt['status'] == 'staging' and attempt['role'] not in ('review', 'solve')
                               and state['recoveries'] == 0
                               and (isinstance(error, subprocess.TimeoutExpired) or error.returncode == 255))
                if not recoverable:
                    raise
                attempt.update(status='presemantic_transport_failure', reason=type(error).__name__)
                state['first_failure'] = state.get('first_failure') or {'launch': attempt['ordinal'], 'reason': attempt['reason']}
                state['recoveries'] += 1
                f.write(root / 'receipt.json', state)
        state['stopped_at'] = now()
    except (Exception, KeyboardInterrupt) as error:
        record_execution_failure(root, state, error)
    finally:
        f.write(root / 'receipt.json', state)
        lock.unlink(missing_ok=True)
    if batch_root is not None:
        persisted = f.read(root / 'receipt.json')
        return {key: persisted.get(key) for key in (
            'status', 'submitted', 'review', 'trial', 'first_failure', 'stop_reason')}
    return report(root)


def emit_json(value, stream=None):
    """Write one JSON document without coupling persisted execution state to console encoding."""
    stream = stream or sys.stdout
    if stream is sys.stdout and hasattr(stream, 'reconfigure'):
        try:
            stream.reconfigure(encoding='utf-8', errors='strict')
        except (AttributeError, OSError, ValueError):
            pass
    payload = json.dumps(value, ensure_ascii=False, indent=2)
    encoding = getattr(stream, 'encoding', None) or 'utf-8'
    try:
        payload.encode(encoding)
    except (LookupError, UnicodeEncodeError):
        payload = json.dumps(value, ensure_ascii=True, indent=2)
    stream.write(payload + '\n')
    stream.flush()


def stop(root):
    scope, state = f.read(root / 'scope.json'), f.read(root / 'receipt.json')
    (root / 'STOP').touch()
    if state['attempts'] and state['attempts'][-1]['status'] == 'started':
        name = state['attempts'][-1]['container']
        if not re.fullmatch(re.escape(root.name) + r'-\d{2}', name):
            raise ValueError('container_identity_invalid')
        base._ssh(HOST, f'docker stop -t 1 {name}', timeout=15, check=False)
    return {'stop_requested': True, 'scope': root.name}


def report(root):
    if not root.exists():
        raise FileNotFoundError(root)
    try:
        return _report(root)
    except (OSError, ValueError, KeyError, TypeError) as error:
        try:
            state = f.read(root / 'receipt.json')
        except (OSError, ValueError):
            state = {}
        return {'status': state.get('status', 'unreadable'), 'submitted': state.get('submitted'),
                'first_failure': state.get('first_failure'), 'report_incomplete': True,
                'report_errors': [type(error).__name__ + ':' + str(error)],
                'professional_status': 'provisional/LLM-proxy'}


def _report(root):
    state = f.read(root / 'receipt.json')
    consultations = state.get('consultations', [])
    finding_keys = {(row['request_id'], finding)
                    for row in consultations for finding in row.get('finding_ids', [])}
    current_ids = {entry['id'] for entry in state.get('current', {}).values()}
    dispositions = state.get('dispositions', [])
    successful_tool_calls = 0
    disposition_calls = 0
    call_records = []
    report_errors = []
    for attempt in state.get('attempts', []):
        raw = root / attempt.get('raw_path', '') / 'broker.json' if attempt.get('raw_path') else None
        if raw and raw.is_file():
            try:
                events = f.read(raw).get('events', [])
            except (OSError, ValueError) as error:
                report_errors.append({'path': raw.relative_to(root).as_posix(), 'error': str(error)})
                continue
            call_records.extend(event['result']['recorded'] for event in events if event.get('action') == 'log-call')
            successful_tool_calls += len(events)
            disposition_calls += sum(event.get('action') in ('record-disposition', 'record-dispositions')
                                     for event in events)
    replay_execution_counts = []
    duplicate_executions = 0
    for replay in state.get('calculation_replays', []):
        rows = replay.get('result', {}).get('executions', replay.get('result', {}).get('calculations', []))
        identities = [row.get('execution_id', row.get('calculation_id')) for row in rows]
        replay_execution_counts.append(len(identities))
        duplicate_executions += len(identities) - len(set(identities))
    requirement_changes = None
    compile_entry = state.get('current', {}).get('compile')
    if compile_entry:
        try:
            comparison_path = version_path(root, compile_entry) / 'comparison.json'
            if comparison_path.is_file():
                requirement_changes = f.read(comparison_path).get('requirement_changes')
        except (OSError, ValueError) as error:
            report_errors.append({'path': compile_entry['path'], 'error': str(error)})
    workbook_evidence = []
    workbook_roots = []
    if state.get('development_trial'):
        workbook_roots.append(('development_trial', root / state['development_trial']['path']))
    solve_attempt = next((attempt for attempt in reversed(state.get('attempts', []))
                          if attempt.get('role') == 'solve' and attempt.get('raw_path')), None)
    if solve_attempt:
        workbook_roots.append(('blind_solve', root / solve_attempt['raw_path'] / 'draft'))
    for stage, workbook_root in workbook_roots:
        for path in workbook_root.rglob('*.xlsx'):
            from openpyxl import load_workbook
            try:
                formulas = load_workbook(path, read_only=True, data_only=False)
                cached = load_workbook(path, read_only=True, data_only=True)
                formula_count = cached_count = 0
                for formula_sheet, cached_sheet in zip(formulas.worksheets, cached.worksheets):
                    for formula_row, cached_row in zip(formula_sheet.iter_rows(), cached_sheet.iter_rows()):
                        for formula_cell, cached_cell in zip(formula_row, cached_row):
                            if formula_cell.data_type == 'f':
                                formula_count += 1
                                cached_count += cached_cell.value is not None
                formulas.close()
                cached.close()
                evidence = {'formula_cells': formula_count, 'formula_cached_values': cached_count}
            except Exception as error:
                evidence = {'read_error': type(error).__name__ + ':' + str(error)}
            workbook_evidence.append({
                'stage': stage, 'path': path.relative_to(root).as_posix(),
                'sha256': f.digest(path), **evidence, 'copy_recalculation': None,
                'claim_limit': 'Openability and cached-value coverage do not prove calculation correctness.',
            })
    return {k: state.get(k) for k in ('status', 'submitted', 'review', 'trial', 'first_failure', 'stop_reason')} | {
        'launches_used': state.get('prior_launches', 0) + sum(
            harness.attempt_consumes_launch(a) for a in state['attempts']),
        'operations_recorded': len(state['attempts']), 'consultations': len(consultations),
        'distinct_findings': len(finding_keys), 'disposition_records': len(dispositions),
        'dispositions': len(dispositions),
        'final_version_valid_dispositions': sum(
            any(row.get('snapshot') == state['current'].get(c['role'], {}).get('id')
                and row.get('request_id') == c['request_id'] and row.get('finding_id') in c['finding_ids']
                and c.get('parents') == f.expected_parents(state, c['role']) for c in consultations
                if c['role'] in state['current']) for row in dispositions),
        'disposition_tool_calls': disposition_calls,
        'successful_broker_tool_calls': successful_tool_calls,
        'broker_events_not_explicit_calls': successful_tool_calls,
        'explicit_tool_calls': sum(r['origin'] == 'explicit' for r in call_records) if call_records else None,
        'explicit_tool_failures': sum(r['origin'] == 'explicit' and not r['ok'] for r in call_records) if call_records else None,
        'internal_tool_calls': sum(r['origin'] == 'internal' for r in call_records) if call_records else None,
        'tool_call_measurement': 'recorded calls only; CLI parse failures before dispatch and historical uninstrumented calls unknown',
        'report_errors': report_errors,
        'failed_tool_calls': None,
        'explicit_vs_internal_checks': {
            'explicit_successful': sum(a.get('tool') == 'check' for a in state.get('checks', [])),
            'internal_stage_checks': sum(a.get('outcome') is not None for a in state.get('attempts', [])),
        },
        'calculation_replay_count': len(state.get('calculation_replays', [])),
        'replay_execution_counts': replay_execution_counts,
        'duplicate_executions_within_replays': duplicate_executions,
        'initial_basis_snapshot': state.get('initial_basis_snapshot'),
        'compiler_requirement_changes': requirement_changes,
        'workbook_evidence': workbook_evidence,
        'final_versions': {role: item['id'] for role, item in state['current'].items()},
        'structural_results': [{'launch': a['ordinal'], 'role': a['role'], 'outcome': a.get('outcome')}
                               for a in state['attempts']],
        'native_elapsed_seconds': state.get('prior_native_elapsed_seconds', 0) +
                                  sum(a.get('elapsed_seconds', 0) for a in state['attempts']),
        'native_completed_operations': sum(a.get('native_status') == 'completed' for a in state['attempts']),
        'collection_failed_operations': sum(
            a.get('collection', {}).get('status') == 'collection_failed' for a in state['attempts']),
        'pending_controller_acceptance': sum(
            a.get('acceptance_status') == 'pending' for a in state['attempts']),
        'started_at': state.get('started_at'), 'stopped_at': state.get('stopped_at'),
        'model_request_count': None, 'total_cost': None,
        'professional_status': 'provisional/LLM-proxy', 'task_manual_edits': state['task_manual_edits']}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['prepare', 'execute', 'status', 'stop', 'report'])
    identity = parser.add_mutually_exclusive_group(required=True)
    identity.add_argument('--run-id')
    identity.add_argument('--batch-id')
    parser.add_argument('--source-bundle', type=Path)
    parser.add_argument('--dependency-lock', type=Path, default=DEFAULT_LOCK)
    parser.add_argument('--dependency-remote', default=DEFAULT_DEPS)
    args = parser.parse_args()
    root = safe_root(args.batch_id or args.run_id)
    if args.batch_id:
        from task_generator.production import agent_factory_batch as batch
        if args.action == 'prepare':
            if args.source_bundle is None:
                parser.error('--source-bundle is required to prepare a batch')
            batch.prepare(sys.modules[__name__], root, args.source_bundle, args.dependency_lock, args.dependency_remote)
            result = batch.report(root)
        elif args.action == 'execute':
            result = batch.execute(sys.modules[__name__], root)
        elif args.action == 'stop':
            result = batch.stop(sys.modules[__name__], root)
        else:
            result = batch.report(root)
        emit_json(result)
    elif args.action == 'prepare':
        prepare(root, args.dependency_lock, args.dependency_remote)
        emit_json(report(root))
    elif args.action == 'execute':
        emit_json(execute(root))
    else:
        emit_json(stop(root) if args.action == 'stop' else report(root))
