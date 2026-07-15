from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def atomic_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def inspect_xlsx(path: Path) -> dict:
    workbook = load_workbook(path, read_only=False, data_only=False)
    formula_count = 0
    error_cells: list[str] = []
    sheets = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                if cell.data_type == "e" or value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                    error_cells.append(f"{sheet.title}!{cell.coordinate}")
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column})
    workbook.close()
    return {"sheets": sheets, "formula_count": formula_count, "formula_error_cells": error_cells}


def render(source: Path, output_root: Path) -> dict:
    if source.suffix.lower() not in {".xlsx", ".docx"}:
        raise ValueError("unsupported_visual_qa_file")
    item_root = output_root / source.stem
    item_root.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="taskgen-lo-") as profile:
        command = [
            "soffice", "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
            f"-env:UserInstallation=file://{Path(profile).as_posix()}",
            "--convert-to", "pdf", "--outdir", str(item_root), str(source),
        ]
        result = subprocess.run(command, capture_output=True, text=True, timeout=180, check=False)
    pdf = item_root / f"{source.stem}.pdf"
    if result.returncode != 0 or not pdf.exists():
        raise RuntimeError(f"libreoffice_render_failed:{source.name}:{result.returncode}")
    png_prefix = item_root / "page"
    png_result = subprocess.run(
        ["pdftoppm", "-png", "-r", "110", str(pdf), str(png_prefix)],
        capture_output=True, text=True, timeout=180, check=False,
    )
    if png_result.returncode != 0:
        raise RuntimeError(f"pdf_png_render_failed:{source.name}:{png_result.returncode}")
    pages = sorted(item_root.glob("page-*.png"))
    page_count = len(PdfReader(str(pdf)).pages)
    if page_count < 1 or len(pages) != page_count:
        raise RuntimeError(f"visual_page_count_mismatch:{source.name}")
    payload = {
        "source": source.name,
        "source_sha256": sha256_file(source),
        "pdf": str(pdf.relative_to(output_root)).replace("\\", "/"),
        "pdf_sha256": sha256_file(pdf),
        "page_count": page_count,
        "page_images": [str(item.relative_to(output_root)).replace("\\", "/") for item in pages],
        "structural": inspect_xlsx(source) if source.suffix.lower() == ".xlsx" else {},
    }
    if payload["structural"].get("formula_error_cells"):
        raise RuntimeError(f"xlsx_formula_error:{source.name}")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser(description="Render every candidate-visible XLSX/DOCX for visual review.")
    parser.add_argument("--reference-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if not args.reference_dir.is_dir():
        parser.error("reference directory does not exist")
    if args.output_dir.exists():
        raise FileExistsError("visual_qa_output_already_exists")
    args.output_dir.mkdir(parents=True)
    files = sorted(
        item for item in args.reference_dir.iterdir()
        if item.is_file() and item.suffix.lower() in {".xlsx", ".docx"}
    )
    if not files:
        raise RuntimeError("no_renderable_candidate_files")
    report = {
        "report_version": "v3.task_package_visual_qa.1",
        "reference_dir": str(args.reference_dir),
        "files": [render(item, args.output_dir) for item in files],
        "rendered_file_count": len(files),
        "decision": "pass",
    }
    atomic_json(args.output_dir / "visual_qa_report.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
