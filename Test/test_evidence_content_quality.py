from __future__ import annotations

import sys
import tempfile
import unittest
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from task_generator.v3_evidence_content_quality import (  # noqa: E402
    EvidenceContentQualityValidator,
)
from task_generator.v3_task_design_frontend import (  # noqa: E402
    CapabilityBriefV1,
    EvidenceArtifactSpecV1,
    TaskDesignProposalV1,
)
from task_generator.v3_hybrid_task_materializer import (  # noqa: E402
    HybridEvidenceDossierV1,
    HybridTaskMaterializer,
)


class EvidenceContentQualityTest(unittest.TestCase):
    def _semantic_v2_proposal(self) -> TaskDesignProposalV1:
        proposal_path = (
            ROOT
            / "Test"
            / "fixtures"
            / "pipeline_reconstruction"
            / "task_design"
            / "valid_proposal.json"
        )
        payload = json.loads(proposal_path.read_text(encoding="utf-8"))
        payload["proposal_version"] = "v3.task_design_proposal.2"
        node_specs = {
            "node_primary": {
                "transaction_amount": ("Transaction Amount", "currency"),
                "transaction_status": ("Transaction Status", "enum"),
                "transaction_date": ("Transaction Date", "date"),
                "transaction_note": ("Transaction Note", "text"),
            },
            "node_control": {
                "expected_amount": ("Expected Amount", "currency"),
                "policy_condition": ("Policy Condition", "text"),
                "control_status": ("Control Status", "enum"),
            },
        }
        for node in payload["evidence_nodes"]:
            specific = node_specs[node["node_id"]]
            fields = [
                {
                    "field_name": "record_id",
                    "display_name": "Record ID",
                    "data_type": "identifier",
                    "description": "Stable candidate-visible record identifier.",
                },
                {
                    "field_name": "case_id",
                    "display_name": "Case ID",
                    "data_type": "identifier",
                    "description": "Shared key for governed cross-file analysis.",
                },
            ]
            fields.extend(
                {
                    "field_name": field_name,
                    "display_name": display_name,
                    "data_type": data_type,
                    "description": f"Candidate-visible {display_name.lower()} for analysis.",
                }
                for field_name, (display_name, data_type) in specific.items()
            )
            records = []
            for index in range(1, 4):
                record_id = f"{node['node_id'].upper()}-{index:03d}"
                values = {
                    "record_id": record_id,
                    "case_id": f"CASE-{index:03d}",
                }
                if node["node_id"] == "node_primary":
                    values.update(
                        {
                            "transaction_amount": 1000.0 * index,
                            "transaction_status": "posted",
                            "transaction_date": f"2026-0{index}-28",
                            "transaction_note": (
                                "Candidate-visible transaction support includes the approved "
                                "source document, reviewer evidence, and a detailed explanation "
                                "of the exception that requires professional follow-up."
                            ),
                        }
                    )
                else:
                    values.update(
                        {
                            "expected_amount": 1000.0 * index,
                            "policy_condition": "Manager approval required",
                            "control_status": "documented",
                        }
                    )
                records.append({"record_id": record_id, "values": values})
            node["artifact_spec"] = {
                "fact_origin": "governed_scenario_fact",
                "record_type": node["artifact_role"],
                "primary_key_field": "record_id",
                "fields": fields,
                "records": records,
                "methodological_source_ref_ids": node["source_ref_ids"],
            }
        relation = payload["evidence_relations"][0]
        relation["from_join_field"] = "case_id"
        relation["to_join_field"] = "case_id"
        relation["comparison_fields"] = [
            {
                "from_field": "transaction_amount",
                "to_field": "expected_amount",
                "operator": "equal",
            }
        ]
        return TaskDesignProposalV1.model_validate(payload)

    def _write_workbook(
        self,
        root: Path,
        file_name: str,
        headers: list[str],
        rows: list[list[object]],
        *,
        width: float = 24.0,
    ) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Evidence"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = width
        workbook.save(root / file_name)

    def test_blocks_legacy_placeholder_content_and_clipping(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "materiality_context.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Evidence_ID",
                    "Record_Key",
                    "Observed_Value",
                    "Observation_Status",
                    "Source_Ref",
                ],
                [["RESOURCE-LONG-ID-001", "REC-001", 100, "observed", "source_candidate_1"]],
                width=8.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_placeholder",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_materiality",
                        artifact_role="documented materiality threshold",
                        intended_contents="Threshold amount, currency, period, scope and approval source.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_materiality", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "blocked")
            self.assertIn("generic_placeholder_schema", report.blocking_reason_codes)
            self.assertIn("insufficient_business_fields", report.blocking_reason_codes)
            self.assertIn("visible_column_clipping_risk", report.blocking_reason_codes)

    def test_passes_semantic_business_fields_with_readable_widths(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "ownership_frequency.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Control_ID",
                    "Control_Owner",
                    "Expected_Frequency",
                    "Observed_Execution_Date",
                    "Frequency_Discrepancy",
                ],
                [
                    ["CTRL-001", "Revenue Controller", "Monthly", "2026-03-31", "None"],
                    ["CTRL-002", "AP Manager", "Quarterly", "2026-02-15", "Out of cycle"],
                    ["CTRL-003", "Treasury Director", "Monthly", "", "Missing evidence"],
                ],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_semantic",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_ownership",
                        artifact_role="control ownership and frequency evidence",
                        intended_contents="Control owners, expected frequency, execution dates and discrepancies.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_ownership", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertEqual(report.blocking_reason_codes, [])

    def test_role_header_overlap_accepts_conservative_plural_forms(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "materiality_thresholds.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Threshold ID",
                    "Metric ID",
                    "Review Threshold Amount",
                    "Escalation Threshold Amount",
                    "Threshold Basis",
                ],
                [["TH-001", "MET-001", 50000, 150000, "Current-period policy"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_plural_role",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_thresholds",
                        artifact_role="policy source for materiality thresholds",
                        intended_contents="Scenario thresholds paired to exposure metrics.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_thresholds", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertGreater(report.checks[0].role_header_overlap_count, 0)

    def test_corroborating_role_accepts_support_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "independent_corroborating_evidence.xlsx"
            self._write_workbook(
                root,
                file_name,
                ["Corroboration ID", "Population Item ID", "Support Type", "Support Status", "Support Note"],
                [["COR-1", "POP-1", "Approval", "Present", "Independent support"]],
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_corroborating_support",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_corroborating",
                        artifact_role="independent corroborating evidence",
                        intended_contents="Corroboration used to support operating-test observations.",
                        artifact_spec=SimpleNamespace(record_type="corroborating_evidence_record"),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_corroborating", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertGreater(report.checks[0].role_header_overlap_count, 0)

    def test_plural_normalization_does_not_match_unrelated_roles(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "materiality_thresholds.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Threshold ID",
                    "Metric ID",
                    "Review Threshold Amount",
                    "Escalation Threshold Amount",
                    "Threshold Basis",
                ],
                [["TH-001", "MET-001", 50000, 150000, "Current-period policy"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_unrelated_role",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_invoices",
                        artifact_role="vendor invoices and payment exceptions",
                        intended_contents="Invoice dates and disputed supplier payments.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_invoices", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "blocked")
            self.assertIn(
                "artifact_role_not_represented_in_headers",
                report.blocking_reason_codes,
            )

    def test_role_header_overlap_accepts_criteria_and_criterion(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "audit_evidence_assessment_criteria.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Criterion ID",
                    "Criterion Name",
                    "Evidence Expectation",
                    "Assessment Focus",
                ],
                [["CRIT-001", "Relevance", "Address the assertion", "Appropriateness"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_irregular_plural_role",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_audit_evidence_criteria",
                        artifact_role="criteria for audit evidence assessment",
                        intended_contents="Assessment criteria for sufficiency and appropriateness.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_audit_evidence_criteria",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertGreater(report.checks[0].role_header_overlap_count, 0)

    def test_irregular_plural_normalization_remains_role_specific(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "audit_evidence_assessment_criteria.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Criterion ID",
                    "Criterion Name",
                    "Evidence Expectation",
                    "Assessment Focus",
                ],
                [["CRIT-001", "Relevance", "Address the assertion", "Appropriateness"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_irregular_plural_unrelated_role",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_bank_balances",
                        artifact_role="bank balances and cash transactions",
                        intended_contents="Amounts and posting dates for cash activity.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_bank_balances",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "blocked")
            self.assertIn(
                "artifact_role_not_represented_in_headers",
                report.blocking_reason_codes,
            )

    def test_record_type_can_ground_role_header_overlap(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "audit_evidence_assessment_context.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Criterion ID",
                    "Criterion Name",
                    "Application Question",
                    "Expected Documentation",
                    "Response Required",
                ],
                [["CRIT-01", "Relevance", "Does it apply?", "Document linkage", "Yes"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_record_type_role",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_evidence_assessment_criteria",
                        artifact_role="audit evidence assessment context",
                        intended_contents="Methodological prompts for professional evaluation.",
                        artifact_spec=SimpleNamespace(
                            record_type="evidence_assessment_criterion"
                        ),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_evidence_assessment_criteria",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertGreater(report.checks[0].role_header_overlap_count, 0)

    def test_unrelated_record_type_does_not_bypass_role_header_gate(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "audit_evidence_assessment_context.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Criterion ID",
                    "Criterion Name",
                    "Application Question",
                    "Expected Documentation",
                    "Response Required",
                ],
                [["CRIT-01", "Relevance", "Does it apply?", "Document linkage", "Yes"]],
                width=32.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_unrelated_record_type",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_bank_balances",
                        artifact_role="bank balances and cash transactions",
                        intended_contents="Amounts and posting dates for cash activity.",
                        artifact_spec=SimpleNamespace(record_type="cash_balance_record"),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_bank_balances",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "blocked")
            self.assertIn(
                "artifact_role_not_represented_in_headers",
                report.blocking_reason_codes,
            )

    def test_blocks_candidate_output_modeled_as_reference_input(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "candidate_synthesis.xlsx"
            self._write_workbook(
                root,
                file_name,
                ["Finding_ID", "Control_ID", "Assessment", "Rationale", "Evidence_Reference"],
                [["F-001", "CTRL-001", "Effective", "Supported", "EV-001"]],
                width=30.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_output_leak",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_synthesis",
                        artifact_role="candidate-authored assessment synthesis",
                        intended_contents="Candidate-produced review conclusions and prioritization rationale.",
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[SimpleNamespace(node_id="node_synthesis", file_name=file_name)],
                reference_root=root,
            )
            self.assertEqual(report.decision, "blocked")
            self.assertIn(
                "candidate_output_modeled_as_reference_input",
                report.blocking_reason_codes,
            )

    def test_allows_methodological_input_supporting_candidate_authored_mapping(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "coso_component_reference.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Component ID",
                    "Component Name",
                    "Component Focus",
                    "Mapping Evidence",
                ],
                [["COSO-01", "Control Environment", "Governance", "EV-001"]],
                width=30.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_methodological_input",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_coso_components",
                        artifact_role="Methodological COSO component reference",
                        intended_contents=(
                            "Methodological component descriptions used to support "
                            "candidate-authored mapping of control objectives to COSO components."
                        ),
                        artifact_spec=SimpleNamespace(
                            record_type="coso_component_reference",
                            fact_origin="methodological_context",
                        ),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_coso_components",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertNotIn(
                "candidate_output_modeled_as_reference_input",
                report.blocking_reason_codes,
            )

    def test_allows_governed_findings_requiring_candidate_classification(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "nonconformance_evidence.xlsx"
            self._write_workbook(
                root,
                file_name,
                [
                    "Nonconformance ID",
                    "Inspection ID",
                    "Finding Code",
                    "Finding Description",
                    "Safety Impact Confirmed",
                ],
                [["NC-001", "INSP-001", "F-01", "Fastener defect", "No"]],
                width=30.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_governed_findings",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_nonconformance_log",
                        artifact_role="Quality finding and nonconformance evidence",
                        intended_contents=(
                            "Governed fictional quality findings with impact indicators "
                            "that require candidate-authored severity classifications."
                        ),
                        artifact_spec=SimpleNamespace(
                            record_type="nonconformance finding",
                            fact_origin="governed_scenario_fact",
                        ),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_nonconformance_log",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertNotIn(
                "candidate_output_modeled_as_reference_input",
                report.blocking_reason_codes,
            )

    def test_allows_governed_prior_draft_supplied_for_validation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            file_name = "prior_draft_report_section_register.xlsx"
            self._write_workbook(
                root,
                file_name,
                ["Section ID", "Control ID", "Requirement ID", "Draft Position"],
                [["DS-001", "C-101", "R-001", "No conclusion recorded"]],
                width=30.0,
            )
            proposal = SimpleNamespace(
                proposal_id="proposal_prior_draft_input",
                evidence_nodes=[
                    SimpleNamespace(
                        node_id="node_prior_draft",
                        artifact_role="prior draft report section register",
                        intended_contents=(
                            "Governed prior-draft section entries supplied for validation, "
                            "not candidate-authored final report content."
                        ),
                    )
                ],
            )
            report = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=[
                    SimpleNamespace(
                        node_id="node_prior_draft",
                        file_name=file_name,
                    )
                ],
                reference_root=root,
            )
            self.assertEqual(report.decision, "pass")
            self.assertNotIn(
                "candidate_output_modeled_as_reference_input",
                report.blocking_reason_codes,
            )

    def test_v2_proposal_requires_artifact_spec_for_every_candidate_node(self) -> None:
        proposal_path = (
            ROOT
            / "Test"
            / "fixtures"
            / "pipeline_reconstruction"
            / "task_design"
            / "valid_proposal.json"
        )
        payload = json.loads(proposal_path.read_text(encoding="utf-8"))
        payload["proposal_version"] = "v3.task_design_proposal.2"
        with self.assertRaisesRegex(
            ValidationError,
            "missing_candidate_evidence_artifact_spec",
        ):
            TaskDesignProposalV1.model_validate(payload)

    def test_artifact_spec_requires_exact_primary_keys_and_required_fields(self) -> None:
        valid = {
            "fact_origin": "governed_scenario_fact",
            "record_type": "control ownership record",
            "primary_key_field": "control_id",
            "fields": [
                {
                    "field_name": "control_id",
                    "display_name": "Control ID",
                    "data_type": "identifier",
                    "description": "Stable identifier for the control.",
                },
                {
                    "field_name": "control_owner",
                    "display_name": "Control Owner",
                    "data_type": "text",
                    "description": "Documented owner of the control.",
                },
                {
                    "field_name": "expected_frequency",
                    "display_name": "Expected Frequency",
                    "data_type": "enum",
                    "description": "Policy-required performance frequency.",
                },
                {
                    "field_name": "observed_date",
                    "display_name": "Observed Date",
                    "data_type": "date",
                    "description": "Candidate-visible execution date.",
                },
            ],
            "records": [
                {
                    "record_id": f"CTRL-{index:03d}",
                    "values": {
                        "control_id": f"CTRL-{index:03d}",
                        "control_owner": "Revenue Controller",
                        "expected_frequency": "Monthly",
                        "observed_date": f"2026-0{index}-28",
                    },
                }
                for index in range(1, 4)
            ],
            "methodological_source_ref_ids": ["source_candidate_1"],
        }
        spec = EvidenceArtifactSpecV1.model_validate(valid)
        self.assertEqual(spec.primary_key_field, "control_id")
        invalid = json.loads(json.dumps(valid))
        invalid["records"][0]["values"]["control_id"] = "CTRL-WRONG"
        with self.assertRaisesRegex(
            ValidationError,
            "evidence_artifact_primary_key_record_id_mismatch",
        ):
            EvidenceArtifactSpecV1.model_validate(invalid)

    def test_v2_proposal_accepts_structured_artifacts_and_join_contracts(self) -> None:
        proposal_path = (
            ROOT
            / "Test"
            / "fixtures"
            / "pipeline_reconstruction"
            / "task_design"
            / "valid_proposal.json"
        )
        payload = json.loads(proposal_path.read_text(encoding="utf-8"))
        payload["proposal_version"] = "v3.task_design_proposal.2"
        for node in payload["evidence_nodes"]:
            node["artifact_spec"] = {
                "fact_origin": "governed_scenario_fact",
                "record_type": node["artifact_role"],
                "primary_key_field": "record_id",
                "fields": [
                    {
                        "field_name": "record_id",
                        "display_name": "Record ID",
                        "data_type": "identifier",
                        "description": "Stable candidate-visible record identifier.",
                    },
                    {
                        "field_name": "case_id",
                        "display_name": "Case ID",
                        "data_type": "identifier",
                        "description": "Shared key for governed cross-file analysis.",
                    },
                    {
                        "field_name": "period",
                        "display_name": "Period",
                        "data_type": "text",
                        "description": "Applicable business reporting period.",
                    },
                    {
                        "field_name": "record_status",
                        "display_name": "Record Status",
                        "data_type": "enum",
                        "description": "Candidate-visible business record status.",
                    },
                ],
                "records": [
                    {
                        "record_id": f"REC-{index:03d}",
                        "values": {
                            "record_id": f"REC-{index:03d}",
                            "case_id": f"CASE-{index:03d}",
                            "period": "2026-Q1",
                            "record_status": "open" if index == 3 else "complete",
                        },
                    }
                    for index in range(1, 4)
                ],
                "methodological_source_ref_ids": node["source_ref_ids"],
            }
        for relation in payload["evidence_relations"]:
            relation["from_join_field"] = "case_id"
            relation["to_join_field"] = "case_id"
            relation["comparison_fields"] = []
        proposal = TaskDesignProposalV1.model_validate(payload)
        self.assertEqual(proposal.proposal_version, "v3.task_design_proposal.2")
        self.assertTrue(
            all(node.artifact_spec is not None for node in proposal.evidence_nodes)
        )

    def test_v2_materialization_writes_semantic_content_and_provenance(self) -> None:
        brief_path = (
            ROOT
            / "Test"
            / "fixtures"
            / "pipeline_reconstruction"
            / "task_design"
            / "capability_brief.json"
        )
        brief = CapabilityBriefV1.model_validate_json(
            brief_path.read_text(encoding="utf-8")
        )
        proposal = self._semantic_v2_proposal()
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary) / "semantic_v2"
            materialization = HybridTaskMaterializer().materialize(
                brief,
                proposal,
                root,
            )
            self.assertEqual(materialization.decision, "pass")
            self.assertEqual(
                materialization.evidence_content_quality_decision,
                "pass",
            )
            self.assertEqual(
                materialization.materialization_backend,
                "hybrid_semantic_artifact_v2",
            )
            validity = json.loads(
                (root / "governance" / "validity_vector.json").read_text(
                    encoding="utf-8"
                )
            )
            factual = next(
                item
                for item in validity["dimensions"]
                if item["dimension"] == "factual_validity"
            )
            self.assertEqual(factual["status"], "provisional")
            self.assertIn("business grounding", factual["rationale"])
            utility = json.loads(
                (root / "governance" / "utility_profile.json").read_text(
                    encoding="utf-8"
                )
            )
            missing_fields = next(
                item
                for item in utility["accidental_difficulty_findings"]
                if item["finding_code"] == "missing_candidate_field"
            )
            self.assertEqual(missing_fields["status"], "clear")
            dossier = HybridEvidenceDossierV1.model_validate_json(
                (root / "governance" / "evidence_dossier_plan.json").read_text(
                    encoding="utf-8"
                )
            )
            content = EvidenceContentQualityValidator().validate(
                proposal=proposal,
                files=dossier.files,
                reference_root=root / "candidate" / "reference_files",
            )
            self.assertEqual(content.decision, "pass")
            workbook_path = (
                root
                / "candidate"
                / "reference_files"
                / dossier.files[0].file_name
            )
            workbook = load_workbook(workbook_path, read_only=False, data_only=True)
            self.assertIn("Provenance", workbook.sheetnames)
            self.assertIn("Transaction Amount", [cell.value for cell in workbook["Evidence"][1]])
            note_column = next(
                cell.column
                for cell in workbook["Evidence"][1]
                if cell.value == "Transaction Note"
            )
            self.assertTrue(
                workbook["Evidence"].cell(
                    row=2,
                    column=note_column,
                ).alignment.wrap_text
            )
            self.assertGreaterEqual(
                workbook["Evidence"].row_dimensions[2].height,
                36.0,
            )
            self.assertEqual(
                workbook["Provenance"]["B2"].value,
                "Governed fictional scenario facts",
            )
            workbook.close()
            anchors = json.loads(
                (root / "teacher" / "deterministic_fact_anchors.json").read_text(
                    encoding="utf-8"
                )
            )
            relation_anchor = next(
                item
                for item in anchors["anchors"]
                if item["anchor_type"] == "relation_crosscheck"
            )
            self.assertEqual(relation_anchor["expected_value"]["shared_key_count"], 3)
            self.assertEqual(relation_anchor["expected_value"]["exact_match_count"], 3)

    def test_v2_materialization_blocks_generic_display_schema(self) -> None:
        brief_path = (
            ROOT
            / "Test"
            / "fixtures"
            / "pipeline_reconstruction"
            / "task_design"
            / "capability_brief.json"
        )
        brief = CapabilityBriefV1.model_validate_json(
            brief_path.read_text(encoding="utf-8")
        )
        payload = self._semantic_v2_proposal().model_dump(mode="json")
        generic_headers = [
            "Evidence ID",
            "Record Key",
            "Observed Value",
            "Observation Status",
            "Source Ref",
        ]
        for node in payload["evidence_nodes"]:
            for field, display_name in zip(
                node["artifact_spec"]["fields"],
                generic_headers,
            ):
                field["display_name"] = display_name
        proposal = TaskDesignProposalV1.model_validate(payload)
        with tempfile.TemporaryDirectory() as temporary:
            report = HybridTaskMaterializer().materialize(
                brief,
                proposal,
                Path(temporary) / "generic_v2",
            )
            self.assertEqual(report.decision, "blocked")
            self.assertEqual(report.evidence_content_quality_decision, "blocked")
            self.assertIn(
                "candidate_evidence_content_quality_failed",
                report.reason_codes,
            )


if __name__ == "__main__":
    unittest.main()
