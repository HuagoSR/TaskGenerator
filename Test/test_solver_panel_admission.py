from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest
import sys

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from task_generator.v3_behavioral_preflight_evidence import (
    BehavioralEvidenceIntegrityRecordV1,
    BehavioralEvidenceIntegrityReportV1,
)
from task_generator.v3_behavioral_validation import SolverToolPreflight
from task_generator.v3_evaluation_calibration import (
    EvaluationCalibrationCompiler,
    SolverPanelMemberV1,
)
from task_generator.v3_solver_panel_admission import (
    SolverPanelAdmissionCompiler,
)


class SolverPanelAdmissionTest(unittest.TestCase):
    def test_retains_only_passing_integrity_complete_member(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            members = []
            integrity_records = []
            for index, (model, stratum, passing) in enumerate(
                [
                    ("weak-model", "weak", False),
                    ("medium-model", "medium", True),
                    ("strong-model", "strong", False),
                ]
            ):
                fixture = root / model
                preflight = SolverToolPreflight()
                preflight.create_fixture(fixture)
                if passing:
                    from openpyxl import Workbook, load_workbook
                    import shutil

                    created = fixture / "deliverable_files" / preflight.EXPECTED_CREATED
                    workbook = Workbook()
                    workbook.active["A1"] = "created_marker"
                    workbook.save(created)
                    edited = fixture / "deliverable_files" / preflight.EXPECTED_EDITED
                    shutil.copy2(fixture / "reference_files" / "copy_template.xlsx", edited)
                    workbook = load_workbook(edited)
                    workbook["Template"]["B1"] = "edited_marker"
                    workbook.save(edited)
                    workbook.close()
                report = preflight.inspect(
                    fixture, solver_model=model, environment_id="test-env"
                )
                report_path = fixture / "report.json"
                report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
                members.append(
                    SolverPanelMemberV1(
                        solver_model=model,
                        stratum=stratum,
                        preflight_report_path=str(report_path),
                        preflight_report=report,
                        provider="test-provider",
                        maximum_cost_per_task_usd=1,
                    )
                )
                integrity_records.append(
                    BehavioralEvidenceIntegrityRecordV1(
                        solver_model=model,
                        decision="pass",
                        budget_sha256_valid=True,
                        outcome_sha256_valid=True,
                        stdout_sha256_valid=True,
                        stderr_sha256_valid=True,
                        output_tree_sha256_valid=True,
                        identity_valid=True,
                    )
                )
            panel = EvaluationCalibrationCompiler().freeze_solver_panel(
                panel_id="test-panel", members=members
            )
            panel_path = root / "panel.json"
            panel_path.write_text(panel.model_dump_json(indent=2), encoding="utf-8")
            integrity = BehavioralEvidenceIntegrityReportV1(
                comparison_id="comparison",
                authorization_request_sha256="a" * 64,
                execution_manifest_path=str(root / "manifest.json"),
                execution_manifest_sha256="b" * 64,
                record_count=3,
                records=integrity_records,
                decision="pass",
            )
            integrity_path = root / "integrity.json"
            integrity_path.write_text(integrity.model_dump_json(indent=2), encoding="utf-8")
            plan = SolverPanelAdmissionCompiler().compile(
                frozen_panel_path=panel_path,
                evidence_integrity_path=integrity_path,
            )
            self.assertEqual(plan.decision, "retain_passing_replace_failed")
            self.assertEqual(plan.retained_models, ["medium-model"])
            self.assertEqual(plan.replacement_required_strata, ["weak", "strong"])
            plan_path = root / "admission.json"
            plan_path.write_text(plan.model_dump_json(indent=2), encoding="utf-8")
            boundary = SolverPanelAdmissionCompiler().compile_replacement_boundary(
                admission_plan_path=plan_path
            )
            self.assertEqual(boundary.maximum_external_models, 2)
            self.assertEqual(boundary.retained_models, ["medium-model"])
            self.assertEqual(
                {item.failed_model for item in boundary.slots},
                {"weak-model", "strong-model"},
            )
            self.assertIn("claude-sonnet-4-6", boundary.prohibited_models)

    def test_integrity_failure_prevents_reuse(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            fixture = root / "model"
            preflight = SolverToolPreflight()
            preflight.create_fixture(fixture)
            report = preflight.inspect(fixture, "model", "env")
            report_path = fixture / "report.json"
            report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
            members = [
                SolverPanelMemberV1(
                    solver_model=f"{stratum}-model",
                    stratum=stratum,
                    preflight_report_path=str(report_path),
                    preflight_report=report.model_copy(
                        update={"solver_model": f"{stratum}-model"}
                    ),
                    provider="provider",
                    maximum_cost_per_task_usd=1,
                )
                for stratum in ("weak", "medium", "strong")
            ]
            for member in members:
                path = root / f"{member.solver_model}.json"
                path.write_text(member.preflight_report.model_dump_json(indent=2), encoding="utf-8")
                member.preflight_report_path = str(path)
            panel = EvaluationCalibrationCompiler().freeze_solver_panel(
                panel_id="panel", members=members
            )
            panel_path = root / "panel.json"
            panel_path.write_text(panel.model_dump_json(indent=2), encoding="utf-8")
            integrity_records = []
            for member in members:
                integrity_records.append(
                    BehavioralEvidenceIntegrityRecordV1(
                        solver_model=member.solver_model,
                        decision="blocked" if member.stratum == "weak" else "pass",
                        budget_sha256_valid=member.stratum != "weak",
                        outcome_sha256_valid=True,
                        stdout_sha256_valid=True,
                        stderr_sha256_valid=True,
                        output_tree_sha256_valid=True,
                        identity_valid=True,
                        blocking_reasons=(
                            ["budget_sha256_mismatch"] if member.stratum == "weak" else []
                        ),
                    )
                )
            integrity = BehavioralEvidenceIntegrityReportV1(
                comparison_id="comparison",
                authorization_request_sha256="a" * 64,
                execution_manifest_path="manifest.json",
                execution_manifest_sha256="b" * 64,
                record_count=3,
                records=integrity_records,
                decision="blocked",
                blocking_reasons=["weak-model:budget_sha256_mismatch"],
            )
            integrity_path = root / "integrity.json"
            integrity_path.write_text(integrity.model_dump_json(indent=2), encoding="utf-8")
            plan = SolverPanelAdmissionCompiler().compile(
                frozen_panel_path=panel_path,
                evidence_integrity_path=integrity_path,
            )
            self.assertEqual(plan.decision, "evidence_redesign_required")
            weak = next(item for item in plan.members if item.stratum == "weak")
            self.assertEqual(weak.action, "evidence_blocked")
            self.assertFalse(weak.prior_evidence_reusable)


if __name__ == "__main__":
    unittest.main()
