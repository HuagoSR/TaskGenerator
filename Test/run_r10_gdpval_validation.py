"""Run the fixed public GDPval solver and pairwise-evaluator calibration.

This is an evaluation-only harness.  It does not expose GDPval content to any
TaskGenerator generation component, and it persists each provider session
before continuing so completed work is never silently redrawn.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
from datetime import UTC, datetime
from itertools import combinations
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "Test"))

from openpyxl import load_workbook

from run_r10_behavioral_pilot import (
    CHATGPT_CODEX_STACK,
    DEEPSEEK_OPENCODE_STACK,
    _codex_turn_completed,
    _extract_opencode_json,
    _remote_docx_opened,
    _run_remote,
)
from run_r10_evaluator_v2_calibration import _extract_fenced_json_from_opencode
from task_generator.evaluation.codex_grader import _strict_output_schema
from task_generator.evaluation.r10_gdpval_validation import (
    GDPvalPairReviewV1,
    GDPvalRankingSnapshotV1,
    GDPvalTaskBindingV1,
    GDPvalSingleReviewV1,
    GDPvalMaterialIncomplete,
    balanced_order,
    bootstrap_model_order,
    canonical_sha256,
    normalize_pair_review,
    ordinal_direction_agreement,
    score_gdpval_bundle,
    score_gdpval_single,
    summarize_single_scores,
)
from task_generator.planning.scenario_task_compiler import tree_sha256
from task_generator.production.campaign import atomic_json


DEFAULT_DATA = ROOT / "artifacts/r10/r10_10_gdpval_public_20260903"
DEFAULT_RUN = ROOT / "artifacts/r10/r10_10_gdpval_validation_20260903"
IMAGE = "taskgenerator-eval:milestone-r9-huago-cone-a2d8a5e35099"
IMAGE_SHA256 = "3c6c76324780a911ff7b47bf52629280b4f2a49f26248540de049a61b9129363"
CODEX_AUTH_DIR = "/home/huagosr/taskgenerator-secrets/codex-auth-current"
REMOTE_HOME = "/home/huagosr/taskgenerator-data/r10-gdpval-validation"
PRIMARY_JUDGE = "gpt-5.6-terra@chatgpt_codex"
SECONDARY_JUDGE = "deepseek-v4-pro@official_opencode"
SOLVERS = {
    "gpt-5.6-sol@chatgpt_codex": {
        "stack": CHATGPT_CODEX_STACK, "model": "gpt-5.6-sol", "reasoning": "none",
    },
    "deepseek-v4-pro@official_opencode": {
        "stack": DEEPSEEK_OPENCODE_STACK, "model": "deepseek-v4-pro", "variant": "max",
    },
    "deepseek-v4-flash@official_opencode": {
        "stack": DEEPSEEK_OPENCODE_STACK, "model": "deepseek-v4-flash", "variant": "high",
    },
}
LUNA = {
    "gpt-5.6-luna@chatgpt_codex": {
        "stack": CHATGPT_CODEX_STACK, "model": "gpt-5.6-luna", "reasoning": "medium",
    }
}
JUDGES = {
    PRIMARY_JUDGE: {"stack": CHATGPT_CODEX_STACK, "model": "gpt-5.6-terra", "reasoning": "medium"},
    SECONDARY_JUDGE: {"stack": DEEPSEEK_OPENCODE_STACK, "model": "deepseek-v4-pro", "variant": "max"},
}
PAIR_SEED = 1010
SENTINEL_COUNT = 6

SINGLE_TASKS = (
    "7d7fc9a7-21a7-4b83-906f-416dea5ad04f",
    "1b1ade2d-f9f6-4a04-baa5-aa15012b53be",
    "36d567ba-e205-4313-9756-931c6e4691fe",
)


def _single_assignments() -> list[dict[str, str]]:
    models = list(SOLVERS)
    specs = [(task, model, SECONDARY_JUDGE, "primary") for task in SINGLE_TASKS for model in models]
    specs += [(task, models[index], PRIMARY_JUDGE, "check") for index, task in enumerate(SINGLE_TASKS)]
    return [{"assignment_id": "single_" + canonical_sha256(spec)[:24], "task_id": spec[0],
             "solver_id": spec[1], "judge_id": spec[2], "role": spec[3]} for spec in specs]


def _single_scope(run_id: str, data_root: Path, bindings: list[GDPvalTaskBindingV1], solver_root: Path) -> dict[str, Any]:
    scope = _judge_scope(run_id, data_root, bindings, solver_root)
    chosen = [b for b in bindings if b.task_id in SINGLE_TASKS]
    if len(chosen) != 3 or any(b.split != "development" for b in chosen) or len({b.occupation for b in chosen}) != 3:
        raise ValueError("gdpval_single_development_subset_invalid")
    assignments = _single_assignments()
    outputs = scope["solver_output_sha256"]
    required = {f"{a['solver_id']}/{a['task_id']}" for a in assignments}
    if not required <= outputs.keys():
        raise ValueError("gdpval_single_missing_valid_delivery")
    for key in ("ranking_snapshot_sha256", "sentinel_count", "sentinel_design", "pair_seed", "optional_luna"):
        scope.pop(key, None)
    scope.update({"scope_version": "r10.gdpval_single_scope.1", "protocol": "balanced_integer_rubric_only",
                  "task_binding_sha256": {b.task_id: b.canonical_sha256() for b in chosen},
                  "solver_output_sha256": {key: outputs[key] for key in sorted(required)},
                  "evaluator_sha256": _sha(ROOT / "src/task_generator/evaluation/r10_gdpval_validation.py"),
                  "assignments": assignments, "max_attempts": 14, "global_retry_limit": 2,
                  "format_retry_limit": 1, "solver_execution_allowed": False})
    return scope


def _stage_single(target: Path, data_root: Path, run_root: Path, binding: GDPvalTaskBindingV1, solver_id: str) -> None:
    target.mkdir(parents=True)
    task = _task_root(data_root, binding.task_id)
    shutil.copy2(task / "prompt.txt", target / "candidate_task.md")
    shutil.copytree(task / "reference_files", target / "reference_files")
    shutil.copytree(_delivery_root(run_root, data_root, binding.task_id, solver_id), target / "anonymous_submission")
    _write(target / "human_rubric.json", [row.model_dump(mode="json") for row in binding.rubric_items])
    _write(target / "grade_schema.json", _strict_output_schema(GDPvalSingleReviewV1.model_json_schema()))
    (target / "TASK.md").write_text("""Grade ONE anonymous submission independently against human_rubric.json.
Read candidate_task.md, every reference_files file and the whole anonymous_submission bundle.
Use a balanced policy: full credit when satisfied, partial integer credit when partially satisfied,
zero when unmet. The rubric score is the maximum for that row; use its actual value, not a fixed weight list.
Include each rubric_item_id exactly once with integer awarded, applicability, concise rationale and actual evidence_paths.
evidence_paths must contain only existing input-relative filenames (anonymous_submission/..., reference_files/...
or candidate_task.md). Put sheet/cell/page locations in rationale, not in the filename. Do not cite temporary renders as paths.
Do not invent criteria, add a holistic veto, issue a preference, infer authorship, or penalize high scores.
Conditional items: if the reference-supported prerequisite is absent, use not_triggered and full row credit,
explaining the absent prerequisite. If necessary materials cannot be read or applicability cannot be established,
use unresolved and material_status incomplete; never guess or manufacture evidence.
Deduct only under a rubric item, not merely because you prefer different wording. If the rubric explicitly
requires a format/title, apply that requirement; do not add one where absent.
Inspect XLSX cell contents/formulas and document contents using the installed tools. Use python3 (openpyxl),
LibreOffice and pdftotext as appropriate. Keep each file's converted output in a separate directory to avoid
same-name overwrites. Try OOXML text extraction for readable DOCX contents if conversion fails; report any
remaining inspection limits. Layout claims or deductions need rendered-file evidence, not guesses from plain text.
Return one JSON object matching grade_schema.json. Do not return totals, preference, model, task or judge identity.
The controller alone computes scores. For Codex, return the JSON as your final answer; do not change any inputs.
""", encoding="utf-8")


def _parse_single(workspace: Path, binding: GDPvalTaskBindingV1, judge_id: str, stdout: str = "") -> GDPvalSingleReviewV1:
    raw = workspace / "grade.raw.json"
    if not raw.is_file() and judge_id.startswith("deepseek"):
        text = (workspace / "agent.jsonl").read_text(encoding="utf-8") if (workspace / "agent.jsonl").is_file() else stdout
        # JSONL is LF-delimited. Legal Unicode separators inside tool text are not record boundaries.
        text = "\n".join(json.dumps(json.loads(line), ensure_ascii=True) for line in text.split("\n") if line.strip())
        extracted = _extract_opencode_json(text) or _extract_fenced_json_from_opencode(text)
        if extracted:
            raw.write_text(extracted, encoding="utf-8")
    review = GDPvalSingleReviewV1.model_validate_json(raw.read_text(encoding="utf-8"))
    score_gdpval_single(binding, review)
    review, normalizations = _normalize_single_references(workspace, review)
    if normalizations:
        _write(workspace / "evidence_reference_normalization.json", {
            "raw_sha256": _sha(raw), "changes": normalizations,
            "scores_and_rationales_unchanged": True})
    for item in review.assessments:
        for name in item.evidence_paths:
            path = Path(name)
            if (path.is_absolute() or ".." in path.parts or "\\" in name or
                not (name == "candidate_task.md" or name.startswith(("reference_files/", "anonymous_submission/"))) or
                not (workspace / path).is_file() or not (workspace / path).resolve().is_relative_to(workspace.resolve())):
                raise ValueError("gdpval_single_invalid_evidence_path")
    return review


def _normalize_single_references(workspace: Path, review: GDPvalSingleReviewV1) -> tuple[GDPvalSingleReviewV1, list]:
    """Resolve filename + location citations without editing a grade or inventing evidence."""
    files = {p.relative_to(workspace).as_posix() for folder in ("anonymous_submission", "reference_files")
             for p in (workspace / folder).rglob("*") if p.is_file() and p.resolve().is_relative_to(workspace.resolve())}
    if (workspace / "candidate_task.md").is_file():
        files.add("candidate_task.md")
    items, changes = [], []
    for item in review.assessments:
        normalized = []
        for citation in item.evidence_paths:
            paths, auxiliary = [], []
            for segment in re.split(r";\s+(?=/(?:workspace|tmp)/)", citation):
                value = segment.removeprefix("/workspace/")
                matches = [name for name in files if value == name or value.startswith(name + " (")]
                if len(matches) == 1:
                    paths.append(matches[0])
                elif segment.startswith("/tmp/opencode/") and " (rendered layout)" in segment:
                    auxiliary.append(segment.split(" (", 1)[0])
                else:
                    raise ValueError("gdpval_single_invalid_evidence_path")
            # A temporary rendering is commentary, not an accepted input file. It must accompany
            # its uniquely named original; preserve the verbatim citation in the audit record.
            if not paths or any(not any(Path(a).stem == Path(p).stem for p in paths) for a in auxiliary):
                raise ValueError("gdpval_single_unbound_render_reference")
            normalized.extend(paths)
            if paths != [citation]:
                changes.append({"rubric_item_id": item.rubric_item_id, "original": citation,
                                "input_paths": paths, "auxiliary_render_not_file_verified": auxiliary})
        items.append(item.model_copy(update={"evidence_paths": list(dict.fromkeys(normalized))}))
    return review.model_copy(update={"assessments": items}), changes


def _single_recovery_source(parent: Path, scope: dict[str, Any]) -> dict[str, Any]:
    """Bind the one stopped canary response. Never restart its semantic session."""
    prior = json.loads((parent / "scope.json").read_text(encoding="utf-8"))
    receipt = json.loads((parent / "receipt.json").read_text(encoding="utf-8"))
    if receipt["scope_sha256"] != canonical_sha256(prior):
        raise ValueError("single_recovery_receipt_drift")
    for key in ("scope_version", "protocol", "image_sha256", "judges", "assignments", "task_binding_sha256", "solver_output_sha256"):
        if prior[key] != scope[key]:
            raise ValueError("single_recovery_contract_drift")
    first = scope["assignments"][0]
    states = list((parent / "singles").glob("*/state.json"))
    root = parent / "singles" / first["assignment_id"]
    state = _state(root)
    if len(states) != 1 or state is None or state["status"] != "incomplete" or state["attempt"] != 1:
        raise ValueError("single_recovery_requires_one_stopped_canary")
    workspace = root / "attempt_1/workspace"
    if not _agent_completed(workspace, JUDGES[first["judge_id"]]["stack"]):
        raise ValueError("single_recovery_response_not_terminal")
    return {"parent_root": str(parent.resolve()), "parent_scope_sha256": canonical_sha256(prior),
            "parent_receipt_sha256": _sha(parent / "receipt.json"), "assignment_id": first["assignment_id"],
            "source_tree_sha256": tree_sha256(root), "prior_attempts": 1}


def _import_single_canary(parent: Path, data_root: Path, run_root: Path, scope: dict[str, Any]) -> None:
    recovery = _single_recovery_source(parent, scope)
    if recovery != scope["recovery"]:
        raise ValueError("single_recovery_source_drift")
    spec = scope["assignments"][0]
    target = run_root / "singles" / spec["assignment_id"]
    if target.exists():
        if (_state(target) or {}).get("recovered_from") != recovery:
            raise ValueError("single_recovery_target_already_started")
        return
    source = parent / "singles" / spec["assignment_id"]
    binding = next(b for b in _bindings(data_root) if b.task_id == spec["task_id"])
    original = source / "attempt_1/workspace"
    task = _task_root(data_root, binding.task_id)
    if (tree_sha256(original / "anonymous_submission") != scope["solver_output_sha256"][f"{spec['solver_id']}/{binding.task_id}"]
            or tree_sha256(original / "reference_files") != binding.reference_tree_sha256
            or (original / "candidate_task.md").read_text(encoding="utf-8") != (task / "prompt.txt").read_text(encoding="utf-8")
            or json.loads((original / "human_rubric.json").read_text(encoding="utf-8")) != [r.model_dump(mode="json") for r in binding.rubric_items]):
        raise ValueError("single_recovery_input_drift")
    # Copy into a new scope; parent receipt, failure, response and grading remain immutable.
    shutil.copytree(source, target)
    review = _parse_single(target / "attempt_1/workspace", binding, spec["judge_id"])
    record = {**spec, "score": score_gdpval_single(binding, review), "review_sha256": canonical_sha256(review)}
    _write(target / "review.json", review)
    _write(target / "result.json", record)
    _write(target / "state.json", {"status": "completed", "attempt": 1, "completed_at": _now(),
        "first_failure": _state(source)["first_failure"], "recovered_from": recovery,
        "recovery_mode": "offline_parse_only_no_provider_call", "result_sha256": canonical_sha256(record)})


def _single_budget(run_root: Path, retry: bool) -> None:
    states = [_state(path.parent) for path in (run_root / "singles").glob("*/state.json")]
    used = sum(s.get("attempt", 0) for s in states)
    retries = sum(max(0, s.get("attempt", 0) - 1) for s in states)
    if used >= 14 or (retry and retries >= 2):
        raise RuntimeError("gdpval_single_global_attempt_limit")


def _run_single(host: str, data_root: Path, run_root: Path, run_id: str,
                binding: GDPvalTaskBindingV1, spec: dict[str, str]) -> dict[str, Any]:
    root = run_root / "singles" / spec["assignment_id"]
    state = _state(root)
    if state:
        if state["status"] != "completed":
            raise RuntimeError("gdpval_single_started_assignment_not_rerunnable")
        record = json.loads((root / "result.json").read_text(encoding="utf-8"))
        review = GDPvalSingleReviewV1.model_validate_json((root / "review.json").read_text(encoding="utf-8"))
        if (canonical_sha256(record) != state["result_sha256"] or canonical_sha256(review) != record["review_sha256"]
                or any(record[k] != v for k, v in spec.items()) or score_gdpval_single(binding, review) != record["score"]):
            raise ValueError("gdpval_single_completed_record_drift")
        return record
    config = JUDGES[spec["judge_id"]]
    first_failure = None
    frozen = None
    started = time.monotonic()
    for attempt in (1, 2):
        _single_budget(run_root, retry=attempt == 2)
        workspace = root / f"attempt_{attempt}" / "workspace"
        _stage_single(workspace, data_root, run_root, binding, spec["solver_id"])
        digest = tree_sha256(workspace)
        if frozen is not None and digest != frozen:
            raise ValueError("gdpval_single_retry_input_drift")
        frozen = digest
        _write(root / "state.json", {"status": "running", "attempt": attempt, "started_at": _now(),
                                     "input_sha256": digest, "first_failure": first_failure})
        retryable = False
        try:
            code, stdout, stderr = _run_remote(
                host=host, remote=f"{REMOTE_HOME}/{run_id}/singles/{spec['assignment_id']}/attempt_{attempt}",
                local=workspace, stack=config["stack"], grade=True, image=IMAGE, codex_auth_dir=CODEX_AUTH_DIR,
                timeout_seconds=1800, model_override=config["model"],
                codex_reasoning_effort=config.get("reasoning"), opencode_variant=config.get("variant"))
            if code != 0 or not _agent_completed(workspace, config["stack"]):
                raise RuntimeError(f"gdpval_single_agent_not_terminal:{code}")
            retryable = True
            review = _parse_single(workspace, binding, spec["judge_id"], stdout)
            record = {**spec, "score": score_gdpval_single(binding, review), "review_sha256": canonical_sha256(review)}
            _write(root / "review.json", review)
            _write(root / "result.json", record)
            _write(root / "state.json", {"status": "completed", "attempt": attempt, "completed_at": _now(),
                "first_failure": first_failure, "input_sha256": digest,
                "duration_seconds": round(time.monotonic() - started, 3), "result_sha256": canonical_sha256(record)})
            return record
        except Exception as exc:
            # Never persist raw provider exception text that may contain credentials.
            failure = f"{type(exc).__name__}:single_review_failed"
            first_failure = first_failure or failure
            if isinstance(exc, GDPvalMaterialIncomplete):
                retryable = False
            details = ([{"loc": list(e["loc"]), "type": e["type"]} for e in exc.errors(include_input=False, include_context=False)]
                       if hasattr(exc, "errors") else [])
            _write(root / f"attempt_{attempt}" / "failure.json", {"failure": failure, "retryable": retryable,
                "schema_errors": details,
                "agent_jsonl_sha256": _sha(workspace / "agent.jsonl") if (workspace / "agent.jsonl").is_file() else None})
            _write(root / "state.json", {"status": "incomplete", "attempt": attempt,
                                         "first_failure": first_failure, "completed_at": _now()})
            if not retryable or attempt == 2:
                raise RuntimeError(failure) from None
    raise AssertionError("unreachable")


def run_single_scoring(host: str, data_root: Path, run_root: Path, run_id: str,
                       scope: dict[str, Any]) -> dict[str, Any]:
    bindings = {b.task_id: b for b in _bindings(data_root)}
    rows = []
    failure = None
    for spec in scope["assignments"]:
        try:
            if (run_root / "STOP_REQUESTED").exists():
                raise RuntimeError("user_requested_stop_before_next_assignment")
            binding = bindings[spec["task_id"]]
            if binding.canonical_sha256() != scope["task_binding_sha256"][binding.task_id]:
                raise ValueError("single_binding_drift")
            _verify_public_files(data_root, [binding])
            if tree_sha256(_delivery_root(run_root, data_root, binding.task_id, spec["solver_id"])) != scope["solver_output_sha256"][f"{spec['solver_id']}/{binding.task_id}"]:
                raise ValueError("single_delivery_drift")
            rows.append(_run_single(host, data_root, run_root, run_id, binding, spec))
        except Exception as exc:
            failure = {"assignment_id": spec["assignment_id"], "error_type": type(exc).__name__}
        report = summarize_single_scores(rows, scope["assignments"])
        report.update({"updated_at": _now(), "first_failure": failure})
        # Retain concrete cross-judge disagreements without averaging judges.
        for check in report["cross_judge_checks"]:
            matching = [r for r in rows if (r["task_id"], r["solver_id"]) == (check["task_id"], check["solver_id"])]
            assessments = {}
            for r in matching:
                review = GDPvalSingleReviewV1.model_validate_json((run_root / "singles" / r["assignment_id"] / "review.json").read_text(encoding="utf-8"))
                assessments[r["role"]] = {a.rubric_item_id: a.awarded for a in review.assessments}
            check["differing_items"] = [{"rubric_item_id": k, "primary": v, "check": assessments["check"][k]}
                                        for k, v in assessments["primary"].items() if v != assessments["check"][k]]
        _write(run_root / "single_result.json", report)
        print(json.dumps({"completed": len(rows), "planned": 12, "first_failure": failure}), flush=True)
        if failure:
            break
    return report


def _now() -> str:
    return datetime.now(UTC).isoformat().replace("+00:00", "Z")


def _write(path: Path, value: Any) -> None:
    atomic_json(path, value.model_dump(mode="json") if hasattr(value, "model_dump") else value)


def _sha(path: Path) -> str:
    import hashlib
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _bindings(data_root: Path) -> list[GDPvalTaskBindingV1]:
    paths = sorted((data_root / "tasks").glob("*/binding.json"))
    values = [GDPvalTaskBindingV1.model_validate_json(path.read_text(encoding="utf-8")) for path in paths]
    if len(values) != 12 or {item.occupation for item in values} != {
        "Accountants and Auditors", "Buyers and Purchasing Agents", "Compliance Officers"
    }:
        raise ValueError("gdpval_fixed_subset_binding_invalid")
    return values


def _verify_public_files(data_root: Path, bindings: list[GDPvalTaskBindingV1]) -> None:
    for binding in bindings:
        task = _task_root(data_root, binding.task_id)
        # The preparation contract fingerprints Unicode prompt text before
        # platform newline conversion, unlike the byte-based artifact trees.
        import hashlib
        prompt_sha = hashlib.sha256((task / "prompt.txt").read_text(encoding="utf-8").encode()).hexdigest()
        if (prompt_sha != binding.prompt_sha256
            or tree_sha256(task / "reference_files") != binding.reference_tree_sha256
            or tree_sha256(task / "human_gold") != binding.gold_tree_sha256):
            raise ValueError(f"gdpval_public_input_drift:{binding.task_id}")


def _scope(run_id: str, data_root: Path, bindings: list[GDPvalTaskBindingV1]) -> dict[str, Any]:
    source_commit = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True, capture_output=True, check=True
    ).stdout.strip()
    snapshot = GDPvalRankingSnapshotV1.model_validate_json(
        (data_root / "ranking_snapshot.json").read_text(encoding="utf-8")
    )
    return {
        "scope_version": "r10.gdpval_validation_scope.1", "campaign_id": run_id,
        "purpose": "eval_calibration_only", "source_commit": source_commit,
        "image": IMAGE, "image_sha256": IMAGE_SHA256,
        "dataset_manifest_sha256": _sha(data_root / "manifest.json"),
        "task_binding_sha256": {item.task_id: item.canonical_sha256() for item in bindings},
        "ranking_snapshot_sha256": snapshot.canonical_sha256(),
        "solvers": SOLVERS,
        "optional_luna": LUNA,
        "judges": JUDGES,
        "pair_seed": PAIR_SEED, "sentinel_count": SENTINEL_COUNT,
        "timeout_seconds": 1800, "semantic_retry_limit": 0,
        "excluded_actions": ["task_generation", "skill_extraction", "training", "task_mutation", "rerun_low_score"],
    }


def _state(root: Path) -> dict[str, Any] | None:
    path = root / "state.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None


def _task_root(data_root: Path, task_id: str) -> Path:
    root = data_root / "tasks" / task_id
    if not root.is_dir():
        raise FileNotFoundError(f"gdpval_task_missing:{task_id}")
    return root


def _stage_solver(task: Path, binding: GDPvalTaskBindingV1, target: Path) -> None:
    target.mkdir(parents=True)
    shutil.copytree(task / "reference_files", target / "reference_files")
    (target / "deliverable_files").mkdir()
    deliverables = "\n".join(f"- deliverable_files/{name}" for name in binding.expected_deliverables)
    prompt = (task / "prompt.txt").read_text(encoding="utf-8")
    (target / "TASK.md").write_text(
        prompt + "\n\nSubmission contract\nCreate exactly these final files at these paths:\n" + deliverables
        + "\nWork autonomously with local tools. Verify every final file can be reopened.\n",
        encoding="utf-8",
    )


def _input_hashes(task: Path) -> set[str]:
    return {_sha(path) for path in (task / "reference_files").rglob("*") if path.is_file()}


def _validate_file(path: Path, *, input_hashes: set[str], office_opened: bool | None = None) -> str | None:
    if not path.is_file() or path.stat().st_size == 0:
        return "missing_or_empty"
    if _sha(path) in input_hashes:
        return "input_copy"
    suffix = path.suffix.casefold()
    try:
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            if not workbook.sheetnames:
                return "xlsx_no_sheets"
            workbook.close()
        elif suffix == ".docx":
            if not zipfile.is_zipfile(path):
                return "docx_not_ooxml"
            with zipfile.ZipFile(path) as archive:
                if "word/document.xml" not in archive.namelist():
                    return "docx_document_xml_missing"
            if office_opened is False:
                return "docx_libreoffice_failed"
        elif suffix == ".pdf":
            if path.read_bytes()[:5] != b"%PDF-":
                return "pdf_invalid"
        elif suffix in {".txt", ".csv"}:
            path.read_text(encoding="utf-8-sig")
        else:
            return f"unsupported_extension:{suffix}"
    except Exception as exc:
        return f"unopenable:{type(exc).__name__}"
    return None


def _validate_delivery(workspace: Path, binding: GDPvalTaskBindingV1, task: Path) -> dict[str, Any]:
    checks = {}
    inputs = _input_hashes(task)
    for name in binding.expected_deliverables:
        relative = f"deliverable_files/{name}"
        office = _remote_docx_opened(workspace, relative) if name.casefold().endswith(".docx") else None
        failure = _validate_file(workspace / relative, input_hashes=inputs, office_opened=office)
        checks[name] = {"valid": failure is None, "failure": failure}
    unexpected = sorted(
        str(path.relative_to(workspace / "deliverable_files")).replace("\\", "/")
        for path in (workspace / "deliverable_files").rglob("*") if path.is_file()
        and path.name not in set(binding.expected_deliverables)
    )
    return {"valid": all(item["valid"] for item in checks.values()) and not unexpected,
            "files": checks, "unexpected": unexpected}


def _run_solver(
    *, host: str, remote_root: str, data_root: Path, run_root: Path,
    binding: GDPvalTaskBindingV1, solver_id: str, config: dict[str, str],
) -> dict[str, Any]:
    root = run_root / "solvers" / solver_id / binding.task_id
    existing = _state(root)
    if existing:
        if existing["status"] == "completed":
            return existing
        raise RuntimeError(f"gdpval_started_solver_not_rerunnable:{solver_id}:{binding.task_id}")
    workspace = root / "workspace"
    task = _task_root(data_root, binding.task_id)
    _stage_solver(task, binding, workspace)
    _write(root / "state.json", {"status": "running", "started_at": _now(), "input_sha256": tree_sha256(workspace)})
    started = time.monotonic()
    code, stdout, stderr = _run_remote(
        host=host, remote=f"{remote_root}/solvers/{solver_id}/{binding.task_id}",
        local=workspace, stack=config["stack"], image=IMAGE, codex_auth_dir=CODEX_AUTH_DIR,
        timeout_seconds=1800, model_override=config["model"],
        codex_reasoning_effort=config.get("reasoning"), opencode_variant=config.get("variant"),
    )
    complete = _codex_turn_completed(workspace / "agent.jsonl") if config["stack"] == CHATGPT_CODEX_STACK else True
    delivery = _validate_delivery(workspace, binding, task)
    status = "completed" if code == 0 and complete and delivery["valid"] else (
        "infrastructure_failed" if code == 124 or not complete or any(
            marker in (stderr or "").casefold() for marker in ("auth", "connect", "rate limit", "service unavailable", "timeout")
        ) else "task_failed"
    )
    result = {
        "status": status, "task_id": binding.task_id, "solver_id": solver_id,
        "returncode": code, "duration_seconds": round(time.monotonic() - started, 3),
        "delivery": delivery, "output_tree_sha256": tree_sha256(workspace / "deliverable_files"),
        "agent_jsonl_sha256": _sha(workspace / "agent.jsonl") if (workspace / "agent.jsonl").is_file() else None,
        "first_failure": None if status == "completed" else (stderr or stdout)[-1000:],
        "completed_at": _now(),
    }
    _write(root / "state.json", result)
    return result


def recover_solver(data_root: Path, run_root: Path, solver_id: str, task_id: str,
                   recovered_root: Path) -> dict[str, Any]:
    """Import a terminal allowlisted return; never launch another agent."""
    root = run_root / "solvers" / solver_id / task_id
    original = _state(root)
    if solver_id not in SOLVERS or not original or original["status"] != "running":
        raise ValueError("gdpval_recovery_requires_running_assignment")
    binding = next(item for item in _bindings(data_root) if item.task_id == task_id)
    workspace = root / "workspace"
    # The interrupted controller may leave its generated shell script behind;
    # it was written after the semantic input fingerprint was persisted.
    with tempfile.TemporaryDirectory(prefix="gdpval-recovery-input-") as temporary:
        copied = Path(temporary) / "input"
        shutil.copytree(workspace, copied, ignore=shutil.ignore_patterns(".r10_agent.sh"))
        if tree_sha256(copied) != original["input_sha256"]:
            raise ValueError("gdpval_recovery_input_drift")
    allowed = {"agent.jsonl", "stderr.txt", "docker_stdout.txt", "docker_stderr.txt",
               "docx_office_opened.txt", "deliverable_files", ".docx_office_check"}
    if any(p.is_symlink() for p in recovered_root.rglob("*")) or any(
        p.name not in allowed for p in recovered_root.iterdir()
    ):
        raise ValueError("gdpval_recovery_non_allowlisted_output")
    events = [json.loads(line) for line in (recovered_root / "agent.jsonl").read_text(
        encoding="utf-8").splitlines() if line.strip()]
    terminal = bool(events) and (events[-1].get("type") == "turn.completed" or (
        events[-1].get("type") == "step_finish" and events[-1].get("part", {}).get("reason") == "stop"))
    if not terminal or any(e.get("type") in {"error", "turn.failed"} for e in events):
        raise ValueError("gdpval_recovery_no_clean_terminal")
    delivery = _validate_delivery(recovered_root, binding, _task_root(data_root, task_id))
    if not delivery["valid"]:
        raise ValueError("gdpval_recovery_delivery_invalid")
    _write(root / "recovery.json", {"original_state": original,
                                    "returned_tree_sha256": tree_sha256(recovered_root),
                                    "recovered_at": _now(), "provider_calls": 0})
    shutil.copytree(recovered_root, workspace, dirs_exist_ok=True)
    result = {"status": "completed", "task_id": task_id, "solver_id": solver_id,
              "returncode": None, "duration_seconds": None,
              "recovered_without_provider_call": True, "delivery": delivery,
              "output_tree_sha256": tree_sha256(workspace / "deliverable_files"),
              "agent_jsonl_sha256": _sha(workspace / "agent.jsonl"),
              "first_failure": "local_controller_interrupted_after_remote_completion",
              "completed_at": _now()}
    _write(root / "state.json", result)
    return result


def _delivery_root(run_root: Path, data_root: Path, task_id: str, identity: str) -> Path:
    if identity == "human_gold":
        return _task_root(data_root, task_id) / "human_gold"
    return _solver_source(run_root) / "solvers" / identity / task_id / "workspace" / "deliverable_files"


def _solver_source(run_root: Path) -> Path:
    scope_file = run_root / "scope.json"
    scope = json.loads(scope_file.read_text(encoding="utf-8")) if scope_file.is_file() else {}
    return Path(scope.get("solver_run_root", run_root))


def _judge_scope(run_id: str, data_root: Path, bindings: list[GDPvalTaskBindingV1],
                 solver_root: Path) -> dict[str, Any]:
    _verify_public_files(data_root, bindings)
    scope = _scope(run_id, data_root, bindings)
    parent = json.loads((solver_root / "scope.json").read_text(encoding="utf-8"))
    receipt = json.loads((solver_root / "receipt.json").read_text(encoding="utf-8"))
    if receipt["scope_sha256"] != canonical_sha256(parent):
        raise ValueError("gdpval_solver_receipt_drift")
    for key in ("task_binding_sha256", "dataset_manifest_sha256", "image_sha256", "solvers"):
        if parent[key] != scope[key]:
            raise ValueError(f"gdpval_solver_parent_drift:{key}")
    result = json.loads((solver_root / "solver_result.json").read_text(encoding="utf-8"))
    outputs = {}
    for identity, tasks in result["results"].items():
        for task_id, state in tasks.items():
            if state["status"] != "completed":
                continue
            path = solver_root / "solvers" / identity / task_id
            if _state(path) != state:
                raise ValueError("gdpval_solver_state_drift")
            digest = tree_sha256(path / "workspace" / "deliverable_files")
            if digest != state["output_tree_sha256"]:
                raise ValueError("gdpval_solver_delivery_drift")
            outputs[f"{identity}/{task_id}"] = digest
    # These are controller/protocol corrections, frozen before any GDPval
    # judging; they do not tune R10 profiles using held-out outcomes.
    scope.update({"scope_version": "r10.gdpval_judge_only_scope.2",
                  "solver_run_root": str(solver_root.resolve()),
                  "solver_scope_sha256": canonical_sha256(parent),
                  "solver_result_sha256": _sha(solver_root / "solver_result.json"),
                  "solver_output_sha256": outputs,
                  "runner_sha256": _sha(Path(__file__)),
                  "protocol": "anonymous_controller_metadata_independent_order_and_judge_checks",
                  "sentinel_design": "two_tasks_per_occupation_primary_reverse_secondary_same_order",
                  "format_retry_limit": 1})
    return scope


def _pair_schema() -> dict[str, Any]:
    schema = GDPvalPairReviewV1.model_json_schema()
    for name in ("task_id", "judge_id", "pair_id", "order_id"):
        schema["properties"].pop(name)
        schema["required"].remove(name)
    # The removed judge field is the only reference to this enum. Do not
    # advertise panel model names through otherwise-unused schema definitions.
    return _strict_output_schema(schema)


def _stage_pair(
    *, target: Path, data_root: Path, run_root: Path, binding: GDPvalTaskBindingV1,
    candidate_a: str, candidate_b: str, order: str, pair_id: str, judge_id: str,
) -> None:
    target.mkdir(parents=True)
    task = _task_root(data_root, binding.task_id)
    shutil.copy2(task / "prompt.txt", target / "candidate_task.md")
    shutil.copytree(task / "reference_files", target / "reference_files")
    _write(target / "human_rubric.json", [item.model_dump(mode="json") for item in binding.rubric_items])
    _write(target / "grade_schema.json", _pair_schema())
    identities = (candidate_a, candidate_b) if order == "a_b" else (candidate_b, candidate_a)
    for index, identity in enumerate(identities, start=1):
        shutil.copytree(_delivery_root(run_root, data_root, binding.task_id, identity), target / f"anonymous_slot_{index}")
    (target / "TASK.md").write_text(
        f"""Evaluate two anonymous work products for public GDPval task `{binding.task_id}`.

Read candidate_task.md, reference_files, human_rubric.json, anonymous_slot_1 and anonymous_slot_2. Evaluate each work product independently against every human rubric item before comparing them. Include every rubric_item_id exactly once for each slot. Use met, partial, or not_met; keep rationales concise and cite actual files. Then select slot_1, slot_2, or tie based on overall professional quality and rubric coverage. Do not infer model identity or whether either slot is human-authored.

Return exactly one JSON object matching grade_schema.json. Do not include task, model, provider, judge, pair or display-order identifiers; the controller supplies administrative metadata after review.
""", encoding="utf-8"
    )


def _parse_pair(workspace: Path, judge_id: str, stdout: str, *, binding: GDPvalTaskBindingV1,
                pair_id: str, order: str) -> GDPvalPairReviewV1:
    raw = workspace / "grade.raw.json"
    if not raw.is_file() and judge_id.startswith("deepseek"):
        text = (workspace / "agent.jsonl").read_text(encoding="utf-8", errors="replace") if (workspace / "agent.jsonl").is_file() else stdout
        extracted = _extract_opencode_json(text) or _extract_fenced_json_from_opencode(text)
        if extracted:
            raw.write_text(extracted, encoding="utf-8")
    if not raw.is_file():
        raise ValueError("gdpval_pair_output_missing")
    value = json.loads(raw.read_text(encoding="utf-8"))
    if any(key in value for key in ("task_id", "judge_id", "pair_id", "order_id")):
        raise ValueError("gdpval_model_supplied_controller_metadata")
    review = GDPvalPairReviewV1.model_validate({**value, "task_id": binding.task_id,
        "judge_id": judge_id, "pair_id": pair_id, "order_id": order})
    for bundle in review.bundles:
        score_gdpval_bundle(binding, bundle)
    return review


def _agent_completed(workspace: Path, stack: str) -> bool:
    try:
        events = [json.loads(line) for line in (workspace / "agent.jsonl").read_text(
            encoding="utf-8").split("\n") if line.strip()]
    except (OSError, ValueError):
        return False
    if any(e.get("type") in {"error", "turn.failed"} for e in events):
        return False
    if stack == CHATGPT_CODEX_STACK:
        return any(e.get("type") == "turn.completed" for e in events)
    return bool(events) and events[-1].get("type") == "step_finish" and events[-1].get("part", {}).get("reason") == "stop"


def _run_pair(
    *, host: str, remote_root: str, data_root: Path, run_root: Path,
    binding: GDPvalTaskBindingV1, candidate_a: str, candidate_b: str,
    pair_id: str, order: str, judge_id: str,
) -> GDPvalPairReviewV1:
    assignment = "review_" + canonical_sha256([binding.task_id, pair_id, judge_id, order])[:24]
    root = run_root / "judges" / assignment
    existing = _state(root)
    if existing:
        if existing["status"] == "completed":
            review = GDPvalPairReviewV1.model_validate_json((root / "review.json").read_text(encoding="utf-8"))
            if canonical_sha256(review) != existing["review_sha256"] or (
                review.task_id, review.judge_id, review.pair_id, review.order_id
            ) != (binding.task_id, judge_id, pair_id, order):
                raise ValueError("gdpval_completed_review_drift")
            for bundle in review.bundles:
                score_gdpval_bundle(binding, bundle)
            return review
        workspace = root / f"attempt_{existing.get('attempt', 1)}" / "workspace"
        try:
            review = _parse_pair(workspace, judge_id, "", binding=binding, pair_id=pair_id, order=order)
            if not _agent_completed(workspace, JUDGES[judge_id]["stack"]):
                raise ValueError("gdpval_judge_recovery_no_terminal")
        except Exception as exc:
            raise RuntimeError(f"gdpval_started_judge_not_rerunnable:{assignment}") from exc
        _write(root / "review.json", review)
        _write(root / "state.json", {"status": "completed", "recovered_without_provider_call": True,
                                      "first_failure": existing.get("first_failure"),
                                      "review_sha256": canonical_sha256(review), "completed_at": _now()})
        return review
    config = JUDGES[judge_id]
    started = time.monotonic()
    first_failure = None
    frozen_input = None
    for attempt in (1, 2):
        workspace = root / f"attempt_{attempt}" / "workspace"
        _stage_pair(target=workspace, data_root=data_root, run_root=run_root, binding=binding,
                    candidate_a=candidate_a, candidate_b=candidate_b, order=order,
                    pair_id=pair_id, judge_id=judge_id)
        digest = tree_sha256(workspace)
        if frozen_input is not None and digest != frozen_input:
            raise ValueError("gdpval_format_retry_input_drift")
        frozen_input = digest
        _write(root / "state.json", {"status": "running", "attempt": attempt,
                                    "started_at": _now(), "input_sha256": digest,
                                    "first_failure": first_failure})
        code, stdout, stderr = _run_remote(
            host=host, remote=f"{remote_root}/judges/{assignment}/attempt_{attempt}", local=workspace,
            stack=config["stack"], grade=True, image=IMAGE, codex_auth_dir=CODEX_AUTH_DIR,
            timeout_seconds=1800, model_override=config["model"],
            codex_reasoning_effort=config.get("reasoning"), opencode_variant=config.get("variant"),
        )
        # A started but nonterminal agent is not a formatting failure. Never
        # draw a replacement while a remote semantic call may still be alive.
        retryable = False
        try:
            if code != 0 or not _agent_completed(workspace, config["stack"]):
                raise RuntimeError(f"gdpval_judge_not_completed:{code}:{stderr[-600:]}")
            retryable = True
            review = _parse_pair(workspace, judge_id, stdout, binding=binding, pair_id=pair_id, order=order)
            _write(root / "review.json", review)
            _write(root / "state.json", {"status": "completed", "attempt": attempt,
                "completed_at": _now(), "first_failure": first_failure,
                "duration_seconds": round(time.monotonic() - started, 3),
                "review_sha256": canonical_sha256(review)})
            return review
        except Exception as exc:
            failure = f"{type(exc).__name__}:{exc}"
            first_failure = first_failure or failure
            _write(root / f"attempt_{attempt}" / "failure.json", {
                "failure": failure, "returncode": code, "format_retryable": retryable,
                "agent_jsonl_sha256": _sha(workspace / "agent.jsonl") if (workspace / "agent.jsonl").is_file() else None})
            _write(root / "state.json", {"status": "incomplete", "attempt": attempt,
                "completed_at": _now(), "first_failure": first_failure})
            if not retryable or attempt == 2:
                raise
    raise AssertionError("unreachable")


def run_solvers(host: str, data_root: Path, run_root: Path, run_id: str) -> dict[str, Any]:
    bindings = _bindings(data_root)
    remote_root = f"{REMOTE_HOME}/{run_id}"
    results: dict[str, dict[str, Any]] = {}
    configs = dict(SOLVERS)
    for solver_id, config in configs.items():
        by_occupation = {}
        for binding in bindings:
            by_occupation.setdefault(binding.occupation, binding)
        canaries = list(by_occupation.values())
        solver_results = {}
        infrastructure_failed = False
        for binding in canaries:
            result = _run_solver(host=host, remote_root=remote_root, data_root=data_root,
                                 run_root=run_root, binding=binding, solver_id=solver_id, config=config)
            solver_results[binding.task_id] = result
            infrastructure_failed |= result["status"] == "infrastructure_failed"
        if not infrastructure_failed:
            for binding in bindings:
                if binding.task_id in solver_results:
                    continue
                result = _run_solver(host=host, remote_root=remote_root, data_root=data_root,
                                     run_root=run_root, binding=binding, solver_id=solver_id, config=config)
                solver_results[binding.task_id] = result
        results[solver_id] = solver_results
    valid_counts = {solver: sum(item["status"] == "completed" for item in values.values())
                    for solver, values in results.items()}
    if any(valid_counts.get(solver, 0) < 10 for solver in SOLVERS):
        for solver_id, config in LUNA.items():
            solver_results = {}
            for binding in bindings:
                solver_results[binding.task_id] = _run_solver(
                    host=host, remote_root=remote_root, data_root=data_root, run_root=run_root,
                    binding=binding, solver_id=solver_id, config=config,
                )
            results[solver_id] = solver_results
    summary = {"completed_at": _now(), "valid_counts": {
        solver: sum(item["status"] == "completed" for item in values.values())
        for solver, values in results.items()
    }, "results": results}
    _write(run_root / "solver_result.json", summary)
    return summary


def _eligible_solvers(run_root: Path) -> list[str]:
    result = json.loads((_solver_source(run_root) / "solver_result.json").read_text(encoding="utf-8"))
    return [solver for solver, count in result["valid_counts"].items() if count >= 10]


def run_judges(host: str, data_root: Path, run_root: Path, run_id: str) -> dict[str, Any]:
    bindings = _bindings(data_root)
    solvers = _eligible_solvers(run_root)
    if len(solvers) < 2:
        raise RuntimeError("gdpval_fewer_than_two_eligible_solvers")
    remote_root = f"{REMOTE_HOME}/{run_id}"
    assignments = []
    model_pairs = list(combinations(solvers, 2))
    sentinel_candidates = []
    for task_index, binding in enumerate(bindings):
        for first, second in model_pairs:
            state_first = _state(_solver_source(run_root) / "solvers" / first / binding.task_id)
            state_second = _state(_solver_source(run_root) / "solvers" / second / binding.task_id)
            if not state_first or not state_second or state_first["status"] != "completed" or state_second["status"] != "completed":
                continue
            pair_id = "pair_" + canonical_sha256([binding.task_id, first, second])[:24]
            order = balanced_order(task_id=binding.task_id, pair_id=pair_id, seed=PAIR_SEED)
            review = _run_pair(host=host, remote_root=remote_root, data_root=data_root, run_root=run_root,
                               binding=binding, candidate_a=first, candidate_b=second,
                               pair_id=pair_id, order=order, judge_id=PRIMARY_JUDGE)
            normalized = normalize_pair_review(review)
            row = _pair_result(binding, review, normalized, first, second, "model_pair")
            assignments.append(row)
            sentinel_candidates.append((binding, first, second, pair_id, order, row))
        anchor = solvers[task_index % len(solvers)]
        state = _state(_solver_source(run_root) / "solvers" / anchor / binding.task_id)
        if state and state["status"] == "completed":
            pair_id = "pair_" + canonical_sha256([binding.task_id, anchor, "human_gold"])[:24]
            order = balanced_order(task_id=binding.task_id, pair_id=pair_id, seed=PAIR_SEED)
            review = _run_pair(host=host, remote_root=remote_root, data_root=data_root, run_root=run_root,
                               binding=binding, candidate_a=anchor, candidate_b="human_gold",
                               pair_id=pair_id, order=order, judge_id=PRIMARY_JUDGE)
            assignments.append(_pair_result(binding, review, normalize_pair_review(review), anchor, "human_gold", "gold_anchor"))
    sentinel_results = []
    for binding, first, second, pair_id, original_order, primary in _select_sentinels(sentinel_candidates):
        reverse = "b_a" if original_order == "a_b" else "a_b"
        reversed_review = _run_pair(host=host, remote_root=remote_root, data_root=data_root, run_root=run_root,
                           binding=binding, candidate_a=first, candidate_b=second,
                           pair_id=pair_id, order=reverse, judge_id=PRIMARY_JUDGE)
        reversed_row = _pair_result(binding, reversed_review, normalize_pair_review(reversed_review), first, second, "sentinel_order")
        review = _run_pair(host=host, remote_root=remote_root, data_root=data_root, run_root=run_root,
                           binding=binding, candidate_a=first, candidate_b=second,
                           pair_id=pair_id, order=original_order, judge_id=SECONDARY_JUDGE)
        secondary = _pair_result(binding, review, normalize_pair_review(review), first, second, "sentinel")
        secondary["judge_agrees_with_primary"] = secondary["preference"] == primary["preference"]
        secondary["position_consistent"] = reversed_row["preference"] == primary["preference"]
        secondary["primary_reverse"] = reversed_row
        sentinel_results.append(secondary)
    result = {"completed_at": _now(), "eligible_solvers": solvers,
              "pair_results": assignments, "sentinel_results": sentinel_results}
    _write(run_root / "judge_result.json", result)
    return result


def _select_sentinels(candidates: list[tuple]) -> list[tuple]:
    """Six prespecified coverage strata, independent of scores/preferences."""
    selected = []
    for occupation in sorted({row[0].occupation for row in candidates}):
        for split in ("development", "holdout"):
            rows = [row for row in candidates if row[0].occupation == occupation and row[0].split == split]
            if rows:
                selected.append(min(rows, key=lambda row: canonical_sha256([PAIR_SEED, row[0].task_id, row[3]])))
    return selected


def _pair_result(binding, review, normalized, first, second, kind):
    return {
        "task_id": binding.task_id, "occupation": binding.occupation, "kind": kind,
        "judge_id": review.judge_id, "candidate_a": first, "candidate_b": second,
        "preference": normalized["preference"].replace("candidate_a", first).replace("candidate_b", second),
        "score_a": score_gdpval_bundle(binding, normalized["candidate_a"]),
        "score_b": score_gdpval_bundle(binding, normalized["candidate_b"]),
        "order_id": review.order_id,
    }


def aggregate(data_root: Path, run_root: Path) -> dict[str, Any]:
    solver = json.loads((_solver_source(run_root) / "solver_result.json").read_text(encoding="utf-8"))
    judges = json.loads((run_root / "judge_result.json").read_text(encoding="utf-8"))
    snapshot = GDPvalRankingSnapshotV1.model_validate_json((data_root / "ranking_snapshot.json").read_text(encoding="utf-8"))
    model_rows = [row for row in judges["pair_results"] if row["kind"] == "model_pair"]
    preferences = [{key: row[key] for key in ("task_id", "candidate_a", "candidate_b", "preference")} for row in model_rows]
    ranking = bootstrap_model_order(preferences, models=judges["eligible_solvers"], seed=PAIR_SEED)
    reference = [model for model in snapshot.reference_order if model in judges["eligible_solvers"]]
    ordinal = ordinal_direction_agreement(ranking["observed_order"], reference)
    sentinels = judges["sentinel_results"]
    position_rate = sum(row["position_consistent"] for row in sentinels) / len(sentinels) if sentinels else 0.0
    judge_rate = sum(row["judge_agrees_with_primary"] for row in sentinels) / len(sentinels) if sentinels else 0.0
    # Direction consistency is evaluated on the deliberately duplicated
    # sentinel pairs; the rest of the public set is judged once to bound cost.
    result = {
        "result_version": "r10.gdpval_validation_result.1", "completed_at": _now(),
        "valid_solver_counts": solver["valid_counts"], "ranking": ranking,
        "reference_order": reference, "ordinal_direction": ordinal,
        "sentinel_position_consistency": round(position_rate, 6),
        "judge_direction_consistency": round(judge_rate, 6),
        "sentinel_count": len(sentinels),
        "legacy_absolute_baseline": "not_evaluated",
        "reference_alias_match": "unverified_not_exact_model_snapshot_identity",
        "evidence_ceiling": "protocol_diagnostic_not_full_evaluator_validation",
        "gold_anchor_count": sum(row["kind"] == "gold_anchor" for row in judges["pair_results"]),
        "decision": (
            "gdpval_protocol_checks_passed" if len(judges["eligible_solvers"]) >= 3
            and all(solver["valid_counts"].get(model, 0) >= 10 for model in judges["eligible_solvers"][:3])
            and len(sentinels) == SENTINEL_COUNT and position_rate >= .9 and judge_rate >= .7
            and ordinal["rate"] is not None and ordinal["rate"] >= 2 / 3
            else "gdpval_protocol_partial"
        ),
    }
    _write(run_root / "result.json", result)
    return result


def status(run_root: Path) -> dict[str, Any]:
    return {
        "scope": (run_root / "scope.json").is_file(),
        "solver_result": json.loads((run_root / "solver_result.json").read_text(encoding="utf-8")) if (run_root / "solver_result.json").is_file() else None,
        "judge_result": (run_root / "judge_result.json").is_file(),
        "result": json.loads((run_root / "result.json").read_text(encoding="utf-8")) if (run_root / "result.json").is_file() else None,
        "solver_states": {str(path.parent.relative_to(run_root)): _state(path.parent)["status"] for path in run_root.glob("solvers/*/*/state.json")},
        "judge_states": {path.parent.name: _state(path.parent)["status"] for path in run_root.glob("judges/*/state.json")},
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=["scope", "solve", "judge", "aggregate", "all", "status", "recover-solver", "single-scope", "single-judge"])
    parser.add_argument("--host", default="huago-cone")
    parser.add_argument("--run-id", default="r10_10_gdpval_validation_20260903")
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--run-root", type=Path, default=DEFAULT_RUN)
    parser.add_argument("--solver-run-root", type=Path,
                        help="Frozen solver campaign; required for separate judge-only execution")
    parser.add_argument("--solver-id")
    parser.add_argument("--task-id")
    parser.add_argument("--recovered-root", type=Path)
    parser.add_argument("--resume-single-root", type=Path, help="Explicit offline recovery of one stopped, terminal canary into a new scope")
    args = parser.parse_args()
    if args.action == "status":
        print(json.dumps(status(args.run_root), ensure_ascii=False, indent=2)); return
    if args.action == "all":
        parser.error("run solve first, then create a separate judge-only scope; legacy all is read-only history")
    bindings = _bindings(args.data_root)
    if args.action in {"single-scope", "single-judge"}:
        if args.solver_run_root is None or args.solver_run_root.resolve() == args.run_root.resolve():
            parser.error("single scoring requires a separate frozen solver root")
        scope = _single_scope(args.run_id, args.data_root, bindings, args.solver_run_root)
        if args.resume_single_root is not None:
            if args.resume_single_root.resolve() == args.run_root.resolve():
                parser.error("single recovery requires a new run root")
            scope["recovery"] = _single_recovery_source(args.resume_single_root, scope)
        if (args.run_root / "scope.json").exists():
            if json.loads((args.run_root / "scope.json").read_text(encoding="utf-8")) != scope:
                raise RuntimeError("gdpval_single_scope_drift")
            receipt = json.loads((args.run_root / "receipt.json").read_text(encoding="utf-8"))
            if receipt["scope_sha256"] != canonical_sha256(scope):
                raise RuntimeError("gdpval_single_receipt_drift")
        else:
            if args.run_root.exists() and any(args.run_root.iterdir()):
                raise RuntimeError("gdpval_single_new_root_must_be_empty")
            _write(args.run_root / "scope.json", scope)
            _write(args.run_root / "receipt.json", {"scope_sha256": canonical_sha256(scope), "consumed_at": _now()})
        if args.action == "single-scope":
            print(json.dumps({"scope_sha256": canonical_sha256(scope)}, indent=2)); return
        if args.resume_single_root is not None:
            _import_single_canary(args.resume_single_root, args.data_root, args.run_root, scope)
        result = run_single_scoring(args.host, args.data_root, args.run_root, args.run_id, scope)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if result["status"] != "preliminary_scoring_complete":
            raise SystemExit(2)
        return
    if args.action == "judge" and (args.run_root / "scope_reduction.json").exists():
        raise RuntimeError("gdpval_full_queue_disabled_by_user")
    if args.action in {"judge", "aggregate"} and args.solver_run_root is None:
        parser.error("judge/aggregate require --solver-run-root and a new judge-only run root")
    if args.solver_run_root is not None:
        if args.action not in {"scope", "judge", "aggregate"} or args.solver_run_root.resolve() == args.run_root.resolve():
            parser.error("judge-only scope must be separate and cannot execute solvers")
        scope = _judge_scope(args.run_id, args.data_root, bindings, args.solver_run_root)
    else:
        scope = _scope(args.run_id, args.data_root, bindings)
    args.run_root.mkdir(parents=True, exist_ok=True)
    if not (args.run_root / "scope.json").exists():
        _write(args.run_root / "scope.json", scope)
        _write(args.run_root / "receipt.json", {"scope_sha256": canonical_sha256(scope), "consumed_at": _now()})
    elif json.loads((args.run_root / "scope.json").read_text(encoding="utf-8")) != scope:
        raise RuntimeError("gdpval_campaign_scope_drift")
    if args.action == "scope":
        print(json.dumps({"scope_sha256": canonical_sha256(scope)}, indent=2)); return
    if args.action == "recover-solver":
        if not all((args.solver_id, args.task_id, args.recovered_root)):
            parser.error("recovery requires solver-id, task-id and recovered-root")
        print(json.dumps(recover_solver(args.data_root, args.run_root, args.solver_id,
                                       args.task_id, args.recovered_root), indent=2)); return
    if args.action in {"solve", "all"}:
        print(json.dumps(run_solvers(args.host, args.data_root, args.run_root, args.run_id), ensure_ascii=False, indent=2))
    if args.action in {"judge", "all"}:
        print(json.dumps(run_judges(args.host, args.data_root, args.run_root, args.run_id), ensure_ascii=False, indent=2))
    if args.action in {"aggregate", "all"}:
        print(json.dumps(aggregate(args.data_root, args.run_root), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
