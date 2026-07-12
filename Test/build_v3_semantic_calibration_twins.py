from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path
from typing import Any, Dict

from openpyxl import load_workbook


CASH_CASES = {2, 6}
CROSS_CASES = {13, 31, 47}
POLICY_CASES = {18, 35, 38}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build repaired semantic-calibration twins without modifying source evidence.")
    parser.add_argument("--review-root", required=True)
    parser.add_argument("--output-dir", required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    source = Path(args.review_root)
    output = Path(args.output_dir)
    output.mkdir(parents=True, exist_ok=True)
    records = []
    for number in sorted(CASH_CASES | CROSS_CASES | POLICY_CASES):
        key = f"{number:02d}"
        candidate_target = output / "candidate_view" / key
        teacher_target = output / "teacher_view" / key
        if candidate_target.exists() or teacher_target.exists():
            raise FileExistsError(f"Twin already exists and will not be overwritten: {key}")
        shutil.copytree(source / "candidate_view" / key, candidate_target)
        shutil.copytree(source / "teacher_view" / key, teacher_target)
        dataset_path = candidate_target / "dataset_row.json"
        dataset = _read_json(dataset_path)
        if number in CASH_CASES:
            dataset["prompt"] = _replace_cash_requirement(str(dataset.get("prompt") or ""))
            repair = "replace unsupported adjusted-balance requirement with exact period-activity reconciliation"
        elif number in CROSS_CASES:
            rule_path = candidate_target / "reference_files" / "disposition_rules.md"
            rule_path.write_text(
                "# Invoice disposition rules\n\n"
                "- `clear`: PO, receipt, invoice quantity, and invoice unit price agree, and no duplicate invoice exists.\n"
                "- `hold`: the invoice belongs to a duplicate invoice pair; place every member of the duplicate pair on hold.\n"
                "- `investigate`: any quantity variance, unit-price variance, or missing receipt exists and the line is not already on duplicate hold.\n",
                encoding="utf-8",
            )
            dataset.setdefault("reference_files", []).append("reference_files\\disposition_rules.md")
            repair = "add candidate-visible clear/hold/investigate decision rules"
        else:
            workbook_path = candidate_target / "reference_files" / "expense_transactions.xlsx"
            _add_attendee_count(workbook_path)
            repair = "add attendee count so per-person meal thresholds are computable"
        _write_json(dataset_path, dataset)
        _patch_teacher_text(teacher_target)
        _write_json(teacher_target / "rubric.json", _fact_centered_rubric(number, teacher_target / "deterministic_answer_key.json"))
        records.append({"case_number": number, "repair": repair, "source_preserved": True})
    _write_json(
        output / "twin_manifest.json",
        {
            "manifest_version": "v3.semantic_calibration_twins.1",
            "case_count": len(records),
            "records": records,
            "notes": [
                "Original candidate and teacher artifacts remain unchanged.",
                "Twins repair semantic dependencies and replace the rubric with fact-centered criteria.",
            ],
        },
    )


def _replace_cash_requirement(prompt: str) -> str:
    replacements = {
        "calculate the adjusted balances": "calculate the bank and ledger period-activity totals",
        "calculate adjusted bank and ledger balances and explain whether they agree": "calculate bank and ledger period-activity totals and explain whether they agree",
    }
    for old, new in replacements.items():
        prompt = prompt.replace(old, new)
    return prompt


def _add_attendee_count(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = [cell.value for cell in sheet[1]]
    if "Attendee_Count" in headers:
        raise ValueError(f"Attendee_Count already exists: {path}")
    category_index = headers.index("Category") + 1
    attendee_column = len(headers) + 1
    sheet.cell(row=1, column=attendee_column, value="Attendee_Count")
    for row in range(2, sheet.max_row + 1):
        category = str(sheet.cell(row=row, column=category_index).value or "").lower()
        sheet.cell(row=row, column=attendee_column, value=1 if "meal" in category else None)
    workbook.save(path)


def _patch_teacher_text(root: Path) -> None:
    for name in ("golden_run.json", "training_annotation.json"):
        path = root / name
        payload = _read_json(path)
        text = json.dumps(payload, ensure_ascii=False)
        text = text.replace("adjusted bank and ledger balances", "bank and ledger period-activity totals")
        text = text.replace("adjusted balances", "period-activity totals")
        _write_json(path, json.loads(text))


def _fact_centered_rubric(number: int, answer_key_path: Path) -> Dict[str, Any]:
    answer = _read_json(answer_key_path)
    if number in CASH_CASES:
        facts = [
            ("Correctly identify matched, bank-only, and ledger-only records.", 30),
            ("Correctly calculate bank and ledger period-activity totals and matched difference.", 30),
            ("Explain whether the period-activity totals agree without claiming unsupported account balances.", 15),
        ]
    elif number in CROSS_CASES:
        facts = [
            ("Correctly calculate quantity and unit-price variances for every invoice line.", 30),
            ("Correctly identify the complete duplicate pair and all exception invoice IDs.", 25),
            ("Apply the candidate-visible disposition rules consistently to every line.", 20),
        ]
    else:
        facts = [
            ("Correctly apply receipt, approval, and per-person meal rules to every transaction.", 30),
            ("Correctly identify exception transactions, clause IDs, counts, and amounts.", 30),
            ("Compute per-person meal amounts using candidate-visible Attendee_Count.", 15),
        ]
    support = [
        ("Cite the source row or transaction and applicable rule for every exception.", 15),
        ("Produce the requested file in a readable and complete structure.", 10),
    ]
    criteria = []
    for index, (description, weight) in enumerate(facts + support, start=1):
        criteria.append(
            {
                "criterion_id": f"R_FACT_{index:03d}",
                "section": "fact_checks" if index <= 3 else "compliance_checks",
                "criterion_type": "fact" if index <= 3 else "compliance",
                "audience": "candidate",
                "description": description,
                "pass_condition": description,
                "weight": weight,
                "source_ids": ["deterministic_answer_key"],
            }
        )
    return {
        "rubric_version": "v3.semantic_twin_rubric.1",
        "fact_weight_ratio": 0.75,
        "answer_key_field_names": sorted(answer.keys()),
        "sections": [
            {"section_name": "fact_checks", "criteria": criteria[:3]},
            {"section_name": "compliance_checks", "criteria": criteria[3:]},
        ],
        "unresolved_gaps": [],
        "warning_reason_codes": [],
    }


def _read_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, payload: Dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
